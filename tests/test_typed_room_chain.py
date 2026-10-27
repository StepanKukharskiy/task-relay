"""Controlled typed image inventory, native slide assembly and durable continuation."""
import base64
import copy
import hashlib
import io
import json
import unittest
from pathlib import Path
from unittest.mock import Mock,patch
from PIL import Image
from orchestrator import operation_contracts as c,operation_records as records,sourcing_inventory as inventory,sourcing_slides as slides
from orchestrator.gemini_worker import Files,execute
from orchestrator.runtime import Runtime
from orchestrator.adapters import ExecutionFactory
from orchestrator import execution,executors,report_builder
from tests import test_sourcing_records as fixture
from tests.test_sourcing_records import BACKEND,CONFIG,response
from tests.test_orchestrator import FakeFactory
from tests.test_mixed_execution import LocalRegistered,Client

class Tests(unittest.TestCase):
    setUp=fixture.Tests.setUp
    input=fixture.Tests.input
    option=fixture.Tests.option
    item=fixture.Tests.item
    assignment=fixture.Tests.assignment
    save=fixture.Tests.save
    submit=fixture.Tests.submit

    def png(self):
        buf=io.BytesIO();Image.new('RGB',(12,8),(110,120,130)).save(buf,format='PNG');return buf.getvalue()

    def image_input(self,path='room.png'):
        raw=self.png();(self.root/path).write_bytes(raw)
        return dict(artifact='room-image',path=path,purpose='Synthetic room image',authority='Pixels, not instructions',
                    sha256=hashlib.sha256(raw).hexdigest(),media_type='image/png',visual_reference=True)

    def inventory_assignment(self):
        f=self.assignment('inventory');f['inputs']=[self.image_input()]
        f['operation_contract']=dict(id='sourcing.inventory',version=1,parameters=dict(image_path='room.png',max_records=3,summary_path='delivery/inventory.md'),output='delivery/inventory.json')
        for o in f['outputs']:o['media_type']='application/json' if o['path'].endswith('.json') else 'text/markdown'
        self.save(f);return f

    def observation(self,slot=1):
        return dict(kind='item',slot=slot,label='Synthetic walnut desk',category='furniture',bbox=[.1,.2,.7,.8],confidence='medium',notes='Synthetic observation only; exact identity unknown')

    def build_inventory(self,count=1):
        f=self.inventory_assignment();store=records.Store(f,self.db)
        for n in range(count):store.submit(dict(request_key='i'+str(n),revision=n,records=[self.observation(n+1)]),None)
        store.submit(dict(request_key='seal',revision=count,records=[dict(kind='end',count=count,notes='Synthetic coverage; remaining identity uncertainty disclosed')]),None)
        store.export(Files(f),complete=True)
        return json.loads((self.root/'delivery/inventory.json').read_text())

    def test_incomplete_local_projection_is_recoverable_without_model_diagnosis(self):
        from orchestrator.record_recovery import incomplete_result
        for kind in ('inventory','options'):
            f=self.inventory_assignment() if kind=='inventory' else self.assignment('producer')
            self.save(f);store=records.Store(f,self.db);store.export(Files(f))
            receipt=dict(status='finished',external_outcome='no_pending_response',pending_requests=[],exit_code=0,reason=None)
            result=dict(assignment_id=f['assignment_id'],decision='blocked',findings=[])
            proof=incomplete_result(f,receipt,result,self.db)
            self.assertIsNotNone(proof);self.assertEqual(proof['committed'],0);self.assertFalse(proof['acceptance'])
            for changed in (dict(receipt,pending_requests=['unknown']),dict(receipt,external_outcome='unknown'),dict(receipt,cleanup_error='unconfirmed')):
                self.assertIsNone(incomplete_result(f,changed,result,self.db))
            for changed in (dict(result,decision='delivered'),dict(result,assignment_id='other'),dict(result,findings=[dict(category='permission')])):
                self.assertIsNone(incomplete_result(f,receipt,changed,self.db))
            output=self.root/f['operation_contract']['output'];original=output.read_bytes();output.write_text('{}')
            self.assertIsNone(incomplete_result(f,receipt,result,self.db));output.write_bytes(original)
            if kind=='inventory':
                store.submit(dict(request_key='seal-empty',revision=0,records=[dict(kind='end',count=0,notes='No items observed')]),None)
                store.export(Files(f),complete=True)
                self.assertIsNone(incomplete_result(f,receipt,result,self.db))

    def assembler_fixture(self,count=5,photos=None):
        inv=self.build_inventory();self.inventory=inv;self.inputs[0]=self.input('inventory.json',inv)
        f=self.assignment('producer');self.save(f);store=records.Store(f,self.db)
        store.submit(dict(request_key='catalog',revision=0,records=[self.item('VIS-001',count)]),None);store.export(Files(f),complete=True)
        inputs=[self.inputs[0],self.inputs[1],self.input('catalog.json',json.loads((self.root/'delivery/producer.json').read_text())),self.image_input()]
        for i in inputs:
            if i['path'].endswith('.json'):i['media_type']='application/json'
        f=dict(inputs=inputs,execution=dict(capability='sourcing.slides',version=1,parameters=dict(inventory_path='inventory.json',catalog_path='catalog.json',source_pack_path='sources.json',image_path='room.png',photos_path=photos,title='Synthetic sourcing',layout='standard')))
        return f

    def test_generic_visual_subject_tool_loop_flows_to_editable_pptx(self):
        from orchestrator import subject_contract, sourcing_contract, pptx_document
        f=self.inventory_assignment()
        f['operation_contract'].update(id='collection.subjects',parameters=dict(source_path='room.png',summary_path='delivery/inventory.md',max_records=3,fields=['material']))
        self.save(f)
        observation=dict(kind='item',slot=1,label='Observed subject',attributes=dict(material=None),quote='',bbox=[.1,.2,.7,.8],notes='Identity and material remain uncertain')
        model=Mock();model.request.side_effect=[response('operation_submit',dict(request_key='one',revision=0,records=[observation])),response('operation_submit',dict(request_key='seal',revision=1,records=[dict(kind='end',count=1,notes='Observed coverage')])),response('finish',dict(summary='Subject observations',decision='delivered',instruction='',findings=[],checks={'c1':dict(passed=True,evidence='Controlled pixel observation')}))]
        control=self.root/'generic-vision';control.mkdir();(control/'launch.json').write_text(c.encoded(dict(credential_fingerprint=executors.fingerprint(CONFIG,BACKEND))))
        with patch('orchestrator.gemini_worker.emit'):execute(f,control,client=model,config_reader=lambda:(CONFIG,BACKEND))
        self.inventory=json.loads((self.root/'delivery/inventory.json').read_bytes())
        subject_contract.validate_document(self.inventory)
        self.assertEqual(self.inventory['items'][0]['id'],'SUB-001')
        self.inputs[0]=self.input('inventory.json',self.inventory)
        producer=self.assignment('generic-catalog');producer['operation_contract']['id']='collection.options'
        producer['operation_contract']['parameters'].update(field_definitions={k:'price' if k=='price' else 'text' for k in sourcing_contract.FIELDS},identity_field='name_variant')
        stage=self.root/'generic-catalog-worker';stage.mkdir();(stage/'.relay').mkdir()
        for item in producer['inputs']:(stage/item['path']).write_bytes((self.root/item['path']).read_bytes())
        producer['workspace']=str(stage)
        self.save(producer);store=records.Store(producer,self.db)
        store.submit(dict(request_key='catalog',revision=0,records=[self.item('SUB-001',1)]),{})
        store.export(Files(producer),complete=True)
        catalog=json.loads((stage/'delivery/generic-catalog.json').read_bytes())
        inputs=[self.inputs[0],self.inputs[1],self.input('catalog.json',catalog),self.image_input()]
        for item in inputs:
            if item['path'].endswith('.json'):item['media_type']='application/json'
        assembly=dict(inputs=inputs,execution=dict(capability='collection.slides',version=1,parameters=dict(inventory_path='inventory.json',catalog_path='catalog.json',source_pack_path='sources.json',image_path='room.png',photos_path=None,title='Generic subjects',layout='standard')))
        raw,proof=slides.build(assembly,lambda p:(self.root/p).read_bytes())
        specification=json.loads(raw)
        self.assertTrue(any(e.get('name')=='SUB-001' for page in specification['slides'] for e in page['elements']))
        output,_=pptx_document.create(specification,{'room.png':self.png()},5000000)
        from pptx import Presentation
        self.assertEqual(len(Presentation(io.BytesIO(output)).slides),proof['slides'])

    def test_inventory_seals_stable_ids_and_rejects_invalid_or_late_fragments(self):
        f=self.inventory_assignment();store=records.Store(f,self.db)
        bad=self.observation();bad['bbox']=[.5,.2,.1,.9]
        with self.assertRaisesRegex(ValueError,'bbox'):store.submit(dict(request_key='bad',revision=0,records=[bad]),None)
        invalid=self.observation();invalid['bbox']=[100,200,700,800]
        with self.assertRaisesRegex(ValueError,'bbox'):inventory.record(invalid,f['operation_contract']['parameters'])
        inventory.record(self.observation(),f['operation_contract']['parameters'])
        store.submit(dict(request_key='one',revision=0,records=[self.observation()]),None)
        with self.assertRaisesRegex(ValueError,'completion count'):store.submit(dict(request_key='end',revision=1,records=[dict(kind='end',count=2,notes='Invalid count')]),None)
        store.submit(dict(request_key='end',revision=1,records=[dict(kind='end',count=1,notes='Coverage notes')]),None)
        with self.assertRaisesRegex(ValueError,'follow completion'):store.submit(dict(request_key='late',revision=2,records=[self.observation(2)]),None)
        store.export(Files(f),complete=True);doc=json.loads((self.root/'delivery/inventory.json').read_text())
        self.assertEqual(doc['items'][0]['id'],'VIS-001');self.assertTrue(doc['complete'])
        self.assertEqual(doc['image']['sha256'],f['inputs'][0]['sha256'])

    def test_empty_inventory_is_explicit_and_incomplete_inventory_cannot_export(self):
        f=self.inventory_assignment();store=records.Store(f,self.db)
        with self.assertRaisesRegex(ValueError,'every inventory'):store.export(Files(f),complete=True)
        doc=self.build_inventory(0);self.assertEqual(doc['items'],[]);self.assertTrue(doc['complete'])

    def test_sealed_empty_inventory_has_honest_zero_coverage_not_legacy_empty_acceptance(self):
        from task_relay.delivery_completion import coverage_slots
        inv=self.build_inventory(0)
        self.assertEqual(coverage_slots(inv,{'items':[]},3,'target')['required_items'],0)
        with self.assertRaises(ValueError):coverage_slots({'items':[]},{'items':[]},3,'target')

    def test_malformed_inventory_image_identity_is_a_contract_error(self):
        data=self.build_inventory()
        for bad in (None,{'path':'not-the-image'},dict(data['image'],sha256='invented'),dict(data['image'],width=True)):
            candidate=copy.deepcopy(data);candidate['image']=bad
            with self.assertRaises(ValueError):inventory.validate_document(candidate)

    def test_worker_attaches_exact_png_pixels_and_exports_only_committed_fragments(self):
        f=self.inventory_assignment();model=Mock();model.request.side_effect=[response('operation_submit',dict(request_key='one',revision=0,records=[self.observation()])),response('operation_submit',dict(request_key='end',revision=1,records=[dict(kind='end',count=1,notes='Synthetic coverage')])),response('finish',dict(summary='Saved visual observations',decision='delivered',instruction='',findings=[],checks={'c1':dict(passed=True,evidence='Synthetic visual seam')}))]
        control=self.root/'vision-control';control.mkdir();(control/'launch.json').write_text(c.encoded(dict(credential_fingerprint=executors.fingerprint(CONFIG,BACKEND))))
        with patch('orchestrator.gemini_worker.emit'):execute(f,control,client=model,config_reader=lambda:(CONFIG,BACKEND))
        parts=model.request.call_args_list[0].args[1]['contents'][0]['parts']
        image=next(p['inlineData'] for p in parts if 'inlineData' in p)
        self.assertEqual(base64.b64decode(image['data']),self.png())
        self.assertTrue(records.Store(f,self.db).progress()['sealed'])
        with self.assertRaisesRegex(ValueError,'owned'):Files(f).call('file_write',dict(path='delivery/inventory.md',text='Invented inventory'))

    def test_inventory_worker_cannot_claim_vision_for_an_unsupported_profile(self):
        f=self.inventory_assignment();f['backend']={'type':'deepseek-agent','model':'fixture'}
        with self.assertRaisesRegex(ValueError,'vision API'):records.Store(f,self.db)
        f=self.inventory_assignment();f['inputs'][0].pop('visual_reference')
        with self.assertRaisesRegex(ValueError,'visual PNG'):records.Store(f,self.db)

    def test_slide_assembly_preserves_five_prices_urls_and_editable_map(self):
        f=self.assembler_fixture();raw,verification=slides.build(f,lambda p:(self.root/p).read_bytes());spec=json.loads(raw)
        self.assertEqual(verification['options'],5);self.assertEqual(len(verification['missing_photos']),5)
        urls=[e['hyperlink'] for s in spec['slides'] for e in s['elements'] if 'hyperlink' in e]
        self.assertEqual(len(set(urls)),5)
        self.assertTrue(any(e.get('name')=='VIS-001' for s in spec['slides'] for e in s['elements']))
        from orchestrator.pptx_document import create
        pptx,checked=create(spec,{'room.png':self.png()},5000000)
        from pptx import Presentation
        deck=Presentation(io.BytesIO(pptx))
        text='\n'.join(shape.text for s in deck.slides for shape in s.shapes if shape.has_text_frame)
        self.assertEqual(text.count('$349.99'),5);self.assertIn('Product photo unavailable',text)
        self.assertEqual(len(deck.slides),verification['slides'])

    def test_recovered_photo_zip_version_reaches_native_assembly_and_deck(self):
        f=self.assembler_fixture(1)
        catalog=json.loads((self.root/'catalog.json').read_text());option=catalog['items'][0]['options'][0]
        from orchestrator import browser_images,image_sources,pptx_document
        from tests.test_image_sources import photo
        manifest=dict(schema=browser_images.SCHEMA,job='synthetic-photos',subjects=[dict(id=option['option_id'],label='Synthetic desk',query='Synthetic desk')],candidates=[])
        base,_=browser_images.fetch_bundle(manifest,fetch=lambda *_:None)
        manifest['candidates']=[dict(subject=option['option_id'],download_url='https://images.example/desk.jpg',source_url=option['product_url'],observed_url=option['product_url'],observation='page-0-0',action_id='saved-source',title='Synthetic product photo')]
        recovered,_=browser_images.fetch_bundle(manifest,fetch=lambda *_:(photo(),'image/jpeg','https://images.example/desk.jpg'))
        merged,_=image_sources.merge_bundles([base,recovered]);path=self.root/'recovered.zip';path.write_bytes(merged)
        f['inputs'].append(dict(artifact='recovered',path='recovered.zip',purpose='Exact recovered photo version',authority='Synthetic reviewed photos',media_type='application/zip',sha256=hashlib.sha256(merged).hexdigest()))
        f['execution']['parameters']['photos_path']='recovered.zip'
        raw,verification=slides.build(f,lambda p:(self.root/p).read_bytes())
        self.assertEqual(verification['photos_sha256'],hashlib.sha256(merged).hexdigest());self.assertEqual(verification['missing_photos'],[])
        pptx,_=pptx_document.create(json.loads(raw),{'room.png':self.png()},5000000,bundles={'recovered.zip':merged})
        from pptx import Presentation
        images=[shape.image.blob for page in Presentation(io.BytesIO(pptx)).slides for shape in page.shapes if shape.shape_type==13]
        self.assertIn(photo(),images)

        # A saved section must survive a cross-worker bundle path alias without
        # re-authoring its claims or substituting a different product photo.
        from orchestrator import document_compose
        data={'section.json':raw.encode(),'assets/room.png':self.png(),'assets/photos.zip':merged}
        composed=dict(inputs=[dict(path=p,sha256=hashlib.sha256(b).hexdigest(),media_type=m)
            for p,b,m in ((p,b,'application/json' if p.endswith('.json') else 'application/zip' if p.endswith('.zip') else 'image/png') for p,b in data.items())],
            execution=dict(parameters=dict(document_paths=['section.json'],image_paths=[],image_titles=[],title='Composed sourcing')))
        result,receipt=document_compose.build(composed,data.__getitem__)
        saved=json.loads(raw);combined=json.loads(result)
        self.assertEqual([s['notes'] for s in combined['slides']],[s['notes'] for s in saved['slides']])
        self.assertEqual(receipt['model_calls'],0)
        deck,_=pptx_document.create(combined,{'assets/room.png':self.png()},5000000,bundles={'assets/photos.zip':merged})
        images=[shape.image.blob for page in Presentation(io.BytesIO(deck)).slides for shape in page.shapes if shape.shape_type==13]
        self.assertIn(photo(),images)

    def test_slide_assembly_rejects_changed_room_or_authored_price(self):
        f=self.assembler_fixture(1);(self.root/'room.png').write_bytes(b'wrong image')
        with self.assertRaisesRegex(ValueError,'pixel version'):slides.build(f,lambda p:(self.root/p).read_bytes())
        (self.root/'room.png').write_bytes(self.png());catalog=json.loads((self.root/'catalog.json').read_text());catalog['items'][0]['options'][0]['fields']['price']['amount']='999'
        replacement=self.input('catalog.json',catalog);replacement['media_type']='application/json';f['inputs'][2]=replacement
        with self.assertRaisesRegex(ValueError,'source projection'):slides.build(f,lambda p:(self.root/p).read_bytes())

    def runtime_inventory(self,max_attempts=2,extra=None):
        f=self.inventory_assignment();factory=FakeFactory();rt=Runtime(self.root/'native',factory);self.addCleanup(rt.db.close)
        original=f['inputs'][0];aid=rt.register(self.root/'room.png','Synthetic room image');image={**original,'artifact':aid}
        for k in ('assignment_id','workspace','operation_store','backend','report_contract'):f.pop(k,None)
        f.update(id='inventory',role='observer',objective='Inspect synthetic image',max_attempts=max_attempts,dependencies=[],inputs=[image])
        review=dict(id='inventory-review',role='reviewer',objective='Inspect exact pixels',review_of='inventory',max_attempts=max_attempts,dependencies=['inventory'],instruction='Synthetic independent visual seam',tools=['files'],limits=f['limits'].copy(),criteria=f['criteria']+[inventory.CRITERION],inputs=[image]+[dict(from_task='inventory',output=o['path'],path=o['path'],purpose='Exact inventory',authority='Unaccepted candidate') for o in f['outputs']],outputs=[dict(path='review.md',purpose='Independent review')])
        tasks=[f,review]+(extra(rt,image,f) if extra else [])
        rt.create(dict(id='room',brief='Synthetic typed image',backend=BACKEND,tasks=tasks));rt.tick('room');return rt,factory

    def test_confirmed_timeout_resumes_committed_record_without_attempt_reset(self):
        rt,factory=self.runtime_inventory();old=rt.task('room','inventory')['latest'];f=factory.sessions[old]['frozen'];store=records.Store(f,rt.db)
        store.submit(dict(request_key='one',revision=0,records=[self.observation()]),None);store.export(Files(f))
        factory.sessions[old]['status']=dict(status='finished',exit_code=1,reason='time_limit',external_outcome='no_pending_response',pending_requests=[],usage=[dict(input_tokens=15,output_tokens=5)])
        rt.tick('room');new=rt.task('room','inventory')['latest'];self.assertNotEqual(new,old)
        newf=factory.sessions[new]['frozen'];current=records.Store(newf,rt.db);progress=current.progress()
        self.assertEqual(progress['next_slot'],2);self.assertEqual(progress['committed'],1)
        self.assertEqual(rt.task('room','inventory')['attempts'],2)
        current.submit(dict(request_key='seal',revision=progress['revision'],records=[dict(kind='end',count=1,notes='Completed remaining inventory work')]),None);current.export(Files(newf),complete=True)
        self.assertEqual(rt.db.execute('select count(*) from operation_records where attempt=?',(old,)).fetchone()[0],1)
        self.assertEqual(rt.db.execute('select state from production_attempts where id=?',(old,)).fetchone()[0],'blocked')
        self.assertEqual(len(factory.calls),2)

    def test_changed_input_version_is_refused_before_a_resume_dispatch(self):
        rt,factory=self.runtime_inventory();old=rt.task('room','inventory')['latest'];f=factory.sessions[old]['frozen']
        records.Store(f,rt.db).submit(dict(request_key='one',revision=0,records=[self.observation()]),None)
        factory.sessions[old]['status']=dict(status='finished',exit_code=1,reason='time_limit',external_outcome='no_pending_response',pending_requests=[])
        rt.tick('room',dispatch=False)
        other=self.root/'different.png';buf=io.BytesIO();Image.new('RGB',(12,8),'red').save(buf,format='PNG');other.write_bytes(buf.getvalue())
        aid=rt.register(other,'Different pixels');task=rt.task('room','inventory');spec=rt.spec(task);spec['inputs'][0]['artifact']=aid
        assignment=rt.new_assignment('room',spec);rt.db.execute('update production_tasks set assignment=? where run=? and id=?',(assignment,'room','inventory'))
        rt.tick('room');self.assertEqual(rt.task('room','inventory')['status'],'blocked');self.assertEqual(rt.task('room','inventory')['attempts'],1)
        self.assertEqual(factory.calls,[old])
        self.assertIn('assignment changed after commit',rt.db.execute("select data from production_events where kind='continuation_preflight_blocked'").fetchone()[0])

    def test_catalog_and_review_resume_keep_original_price_and_uncertain_verdict(self):
        rt=Runtime(self.root/'catalog-resume',FakeFactory());self.addCleanup(rt.db.close);factory=rt.factory
        inputs=[]
        for i in self.inputs:inputs.append({**i,'artifact':rt.register(self.root/i['path'],'Synthetic exact source')})
        source=rt.artifact(inputs[1]['artifact']);receipt=dict(status='finished',operation=dict(output_sha256=source['sha256'],validation=dict(fetched=5)))
        rt.db.execute('INSERT INTO production_attempts VALUES (?,?,?,?,?,?,?,?,?,?)',('read','prior','sources','prior','completed',None,c.encoded(dict(execution=dict(capability='web.sources'))),'{}',c.encoded(receipt),None))
        rt.db.execute('UPDATE production_artifacts SET attempt=? WHERE id=?',('read',source['id']))
        producer=self.assignment('producer');reviewer=self.assignment('reviewer',True)
        for t in (producer,reviewer):
            for k in ('assignment_id','workspace','operation_store','backend','report_contract'):t.pop(k,None)
            t.update(id='reviewer' if t.get('review_of') else 'producer',role='agent',objective='Synthetic checkpoint fixture',max_attempts=2,inputs=copy.deepcopy(inputs),dependencies=['producer'] if t.get('review_of') else [])
        reviewer['inputs'] += [dict(from_task='producer',output='delivery/producer.json',path='catalog.json',purpose='Exact candidate',authority='Unaccepted candidate'),dict(from_task='producer',output='delivery/producer.md',path='schedule.md',purpose='Exact summary',authority='Unaccepted candidate')]
        reviewer['criteria']=producer['criteria']+[fixture.sourcing.SEMANTIC_CRITERION]
        rt.create(dict(id='catalog',brief='Synthetic continued records',backend=BACKEND,tasks=[producer,reviewer]));rt.tick('catalog')
        old=rt.task('catalog','producer')['latest'];f=factory.sessions[old]['frozen'];store=records.Store(f,rt.db)
        store.submit(dict(request_key='one',revision=0,records=[self.item()]),None)
        factory.sessions[old]['status']=dict(status='finished',exit_code=1,reason='Provider request budget exhausted',external_outcome='no_pending_response',pending_requests=[])
        rt.tick('catalog');new=rt.task('catalog','producer')['latest'];f=factory.sessions[new]['frozen'];store=records.Store(f,rt.db)
        self.assertEqual(store.progress()['next_items'][0]['item'],'item-b')
        store.submit(dict(request_key='two',revision=store.progress()['revision'],records=[self.item('item-b',0)]),None);store.export(Files(f),complete=True)
        def finish(f,decision):
            checks={'c'+str(n):dict(passed=True,evidence='Synthetic checkpoint review') for n in range(1,len(f['criteria'])+1)}
            (Path(f['workspace'])/'.relay/result.json').write_text(c.encoded(dict(summary='Synthetic records complete',decision=decision,instruction='',findings=[],checks=checks)))
            factory.sessions[f['assignment_id']]['status']=dict(status='finished',exit_code=0,reason=None)
            rt.tick('catalog')
        finish(f,'delivered');old_review=rt.task('catalog','reviewer')['latest'];f=factory.sessions[old_review]['frozen'];store=records.Store(f,rt.db);unit=store.progress()['next_items'][0]
        store.submit(dict(request_key='uncertain',revision=0,records=[dict(item=unit['item'],item_sha256=unit['item_sha256'],verdict='uncertain',reason='Variant unresolved')]),None)
        factory.sessions[old_review]['status']=dict(status='finished',exit_code=1,reason='time_limit',external_outcome='no_pending_response',pending_requests=[])
        rt.tick('catalog');new_review=rt.task('catalog','reviewer')['latest'];self.assertNotEqual(new_review,old_review)
        f=factory.sessions[new_review]['frozen'];store=records.Store(f,rt.db);unit=store.progress()['next_items'][0]
        store.submit(dict(request_key='second',revision=store.progress()['revision'],records=[dict(item=unit['item'],item_sha256=unit['item_sha256'],verdict='supported',reason='Honest unresolved coverage')]),None);store.export(Files(f),complete=True)
        with self.assertRaisesRegex(ValueError,'cannot be accepted'):store.verify_result(dict(decision='accept'))
        catalog=rt.output('catalog','producer','delivery/producer.json');data=json.loads(Path(catalog['blob']).read_text())
        self.assertEqual(data['items'][0]['options'][0]['fields']['price']['amount'],'349.99')
        self.assertEqual(rt.db.execute('select count(*) from operation_record_resumes').fetchone()[0],2)

    def test_full_native_chain_delivers_and_corrects_without_user_choices(self):
        self.native_chain()

    def test_native_specification_review_corrects_before_deck_without_invalid_input_growth(self):
        self.native_chain(specification=True)

    def native_chain(self,specification=False):
        from tests.workflow_decoys import NativeBoundary,SourcePages,edge
        from tests.test_web_sources import CONFIG as source_config
        from orchestrator.corrections import compile as compile_corrections
        def native(ident,capability,inputs,dependencies,output,media,params):
            spec=execution.REGISTRY[capability]
            return dict(id=ident,role='procedure',objective='Synthetic native chain',instruction='Execute declared operation',execution=dict(capability=capability,version=1,parameters=params),inputs=inputs,outputs=[dict(path=output,purpose='Native typed result',media_type=media)],dependencies=dependencies,criteria=spec['criteria'].copy(),limits=dict(seconds=spec['seconds'],tool_calls=1,output_bytes=spec['output_bytes']),max_attempts=2)
        def reviewer(op):
            return dict(id=op['id']+'-review',role='reviewer',objective='Review exact native artifact',instruction='Synthetic independent review',review_of=op['id'],inputs=[edge(op['id'],op['outputs'][0]['path'],op['outputs'][0]['media_type'])],outputs=[dict(path=op['id']+'-review.md',purpose='Review report')],dependencies=[op['id']],criteria=op['criteria'].copy(),tools=['files'],limits=dict(seconds=600,tool_calls=24,provider_requests=24,response_tokens=4096,output_bytes=200000),max_attempts=2)
        def extra(rt,image,observer):
            source=native('sources','web.sources',[],['inventory','inventory-review'],'delivery/source-pack.json','application/json',dict(queries=['synthetic walnut desk'],domains=[],model=source_config['models']['text']))
            source['max_attempts']=1
            source_review=reviewer(source)
            catalog=self.assignment('catalog')
            for k in ('assignment_id','workspace','operation_store','backend','report_contract'):catalog.pop(k,None)
            catalog.update(id='catalog',role='producer',objective='Source inventory',max_attempts=2,dependencies=['inventory','inventory-review','sources','sources-review'],inputs=[edge('inventory','delivery/inventory.json',binding='inventory.json'),edge('sources','delivery/source-pack.json',binding='sources.json')])
            catalog['operation_contract']['parameters']['max_records']=3
            for o in catalog['outputs']:o['media_type']='application/json' if o['path'].endswith('.json') else 'text/markdown'
            review=reviewer(catalog);review['id']='catalog-review';review['criteria']=catalog['criteria']+[fixture.sourcing.SEMANTIC_CRITERION]
            review['inputs']=[edge('catalog','delivery/catalog.json',binding='catalog.json'),edge('catalog','delivery/catalog.md','text/markdown',binding='schedule.md'),edge('inventory','delivery/inventory.json',binding='inventory.json'),edge('sources','delivery/source-pack.json',binding='sources.json')];review['dependencies']=['catalog','inventory','sources']
            review['outputs']=[dict(path='delivery/catalog-review.json',purpose='Code verdicts',media_type='application/json'),dict(path='delivery/catalog-review.md',purpose='Code review',media_type='text/markdown')]
            review['operation_contract']=dict(id='sourcing.review',version=1,parameters=dict(inventory_path='inventory.json',source_pack_path='sources.json',max_records=3,summary_path='delivery/catalog-review.md',catalog_path='catalog.json'),output='delivery/catalog-review.json')
            assembly=native('slides','sourcing.slides',[edge('inventory','delivery/inventory.json',binding='inventory.json'),edge('catalog','delivery/catalog.json',binding='catalog.json'),edge('sources','delivery/source-pack.json',binding='sources.json'),image],['inventory','catalog','catalog-review','sources'],'slides.json','application/json',dict(inventory_path='inventory.json',catalog_path='catalog.json',source_pack_path='sources.json',image_path='room.png',photos_path=None,title='Synthetic room',layout='standard'))
            deck=native('deck','pptx.create',[edge('slides','slides.json'),image],['slides','slides-review'],'deck.pptx','application/vnd.openxmlformats-officedocument.presentationml.presentation',{})
            final_review=reviewer(deck)
            from orchestrator import worker_capabilities
            from tests.test_worker_capabilities import CODEX
            final_review['inputs'].append(image);final_review.pop('tools')
            final_review['worker']={'requires':['files.text','files.binary','code.execute','images.view']}
            worker_capabilities.resolve(final_review,[worker_capabilities.entry(CODEX)],CODEX)
            tasks=[source,source_review,catalog,review,assembly,reviewer(assembly),deck,final_review]
            compile_corrections(tasks,sourcing_layout=True)
            self.assertIn('review_correction',deck)
            return tasks
        raw_text=self.text
        def fetch(_self,url,offset,limit,allowed_hosts=None):
            return dict(url=url,requested_url=url,title='Synthetic desk',text=raw_text,retrieved_at=124.,sha256=hashlib.sha256(raw_text.encode()).hexdigest(),content_type='text/html',total_characters=len(raw_text),next_offset=None)
        with patch('task_relay.gemini.read_config',return_value=source_config),patch.object(SourcePages,'fetch',fetch):
            rt,factory=self.runtime_inventory(extra=extra);native_factory=NativeBoundary();rt.factory=ExecutionFactory(factory,native_factory,factory)
            rt.factory.available=lambda _:None  # Scripted agent seam; no configured provider check.
            def finish(tid,values=None,decision='accept'):
                aid=rt.task('room',tid)['latest'];self.assertIsNotNone(aid,tid);f=factory.sessions[aid]['frozen']
                if values is not None:
                    store=records.Store(f,rt.db)
                    for r in values:store.submit(dict(request_key='item-'+str(store.progress()['revision']),revision=store.progress()['revision'],records=[r]),None)
                    store.export(Files(f),complete=True)
                    report=dict(summary='Synthetic independent data seam',decision=decision,instruction='',findings=[],checks={'c'+str(n):dict(passed=True,evidence='Synthetic checkpoint') for n in range(1,len(f['criteria'])+1)})
                    (Path(f['workspace'])/'.relay/result.json').write_text(c.encoded(report));factory.sessions[aid]['status']=dict(status='finished',exit_code=0,reason=None,usage=[dict(input_tokens=20,output_tokens=5)])
                else:
                    factory.finish(aid,decision=decision)
                    if decision=='revise':
                        result_path=Path(f['workspace'])/'.relay/result.json'
                        result=json.loads(result_path.read_text());result['findings']=[dict(category='requirements',code='layout_readability',message='Synthetic clipped content',evidence='Exact controlled layout')]
                        result_path.write_text(c.encoded(result))
                rt.tick('room');rt.tick('room')
            finish('inventory',[self.observation(),dict(kind='end',count=1,notes='Synthetic coverage')],'delivered');finish('inventory-review')
            for _ in range(4):rt.tick('room')
            self.assertIsNotNone(rt.output('room','sources','delivery/source-pack.json'))
            finish('sources-review')
            aid=rt.task('room','catalog')['latest'];ctx=records.Store(factory.sessions[aid]['frozen'],rt.db).sourcing_context()
            observed=next(iter(ctx['observations']));option=self.option();option['observation']=observed
            finish('catalog',[dict(item='VIS-001',options=[option],gap='Two alternatives unavailable in synthetic search')],'delivered')
            aid=rt.task('room','catalog-review')['latest'];store=records.Store(factory.sessions[aid]['frozen'],rt.db);unit=store.progress()['next_items'][0]
            finish('catalog-review',[dict(item=unit['item'],item_sha256=unit['item_sha256'],verdict='supported',reason='Synthetic correspondence and gap review')])
            for _ in range(3):rt.tick('room')
            if specification:
                first=rt.output('room','slides','slides.json')
                finish('slides-review',decision='revise')
                self.assertEqual(rt.task('room','slides')['attempts'],2)
                self.assertEqual(len(rt.spec(rt.task('room','slides'))['inputs']),4)
                finish('slides-review');finish('deck-review')
            else:
                finish('slides-review')
                for _ in range(3):rt.tick('room')
                first=rt.output('room','deck','deck.pptx');self.assertIsNotNone(first)
                finish('deck-review',decision='revise')
                for _ in range(3):rt.tick('room')
                finish('slides-review')
                for _ in range(3):rt.tick('room')
                finish('deck-review')
            self.assertEqual(rt.status('room')['status'],'completed')
            final=rt.output('room','deck','deck.pptx');self.assertNotEqual(first['attempt'],final['attempt'])
            self.assertEqual(rt.task('room','catalog')['attempts'],1)
            self.assertEqual(rt.task('room','slides')['attempts'],2)
            self.assertEqual(rt.db.execute('select count(*) from production_decisions').fetchone()[0],0)
            from pptx import Presentation
            deck=Presentation(io.BytesIO(Path(final['blob']).read_bytes()))
            self.assertIn('$349.99','\n'.join(s.text for page in deck.slides for s in page.shapes if s.has_text_frame))
            self.assertEqual(len(native_factory.calls),4 if specification else 5)

    def test_unknown_stop_zero_progress_and_exhausted_attempts_do_not_resume(self):
        for reason,external,pending,limit,commit in [('time_limit','unknown',['api-1'],2,True),('time_limit','no_pending_response',[],2,False),('time_limit','no_pending_response',[],1,True),('cancelled','no_pending_response',[],2,True)]:
            with self.subTest(external=external,limit=limit,commit=commit,reason=reason):
                # Independent runtime per case, original receipt and attempts retained.
                rt,factory=self.runtime_inventory(limit);old=rt.task('room','inventory')['latest'];f=factory.sessions[old]['frozen']
                if commit:records.Store(f,rt.db).submit(dict(request_key='one',revision=0,records=[self.observation()]),None)
                factory.sessions[old]['status']=dict(status='finished',exit_code=1,reason=reason,external_outcome=external,pending_requests=pending)
                rt.tick('room');self.assertEqual(rt.task('room','inventory')['latest'],old);self.assertEqual(len(factory.calls),1)
                rt.db.close();self.root=self.root/('case-'+str(limit)+'-'+str(commit)+'-'+reason);self.root.mkdir()


from tests import test_production_planning as planning_fixture

class PlanningTests(unittest.TestCase):
    setUp=planning_fixture.Tests.setUp
    tearDown=planning_fixture.Tests.tearDown
    request=planning_fixture.Tests.request
    action=planning_fixture.Tests.action
    queue=planning_fixture.Tests.queue
    row=planning_fixture.Tests.row

    def test_actual_inventory_intent_compiles_pixel_worker_and_independent_review(self):
        from task_relay import production_planning as planning
        with patch.object(executors,'configured',return_value=(CONFIG,BACKEND)),patch.object(executors,'available'):
            row=dict(self.queue(action=self.action(executor='gemini-agent'),text='Inspect supplied room image and inventory visible furniture.'))
        self.assertEqual(json.loads(row['options'])['sourcing_records_version'],2)
        source_file=self.rt.root/'room.png';buffer=io.BytesIO();Image.new('RGB',(12,8),'blue').save(buffer,format='PNG');source_file.write_bytes(buffer.getvalue())
        aid=self.rt.register(source_file,'Synthetic room');artifact=self.rt.artifact(aid)
        source=dict(artifact=aid,path='room.png',purpose='Room pixels',authority='Image reference',sha256=artifact['sha256'],bytes=artifact['bytes'],media_type='image/png',visual_reference=True)
        payload=json.loads(row['context']);payload['sources'].append(source);row['context']=c.encoded(payload)
        inp={k:source[k] for k in ('artifact','path','purpose','authority','media_type','visual_reference')}
        limits=dict(seconds=600,tool_calls=24,provider_requests=24,response_tokens=4096,output_bytes=200000)
        p=dict(id='inventory',role='observer',objective='Observe pixels',instruction='Use bounded visual fragments',inputs=[inp],outputs=[dict(path='inventory.json',purpose='Inventory',media_type='application/json'),dict(path='inventory.md',purpose='Schedule',media_type='text/markdown')],dependencies=[],criteria=['Visible furniture coverage'],limits=limits.copy(),max_attempts=1,worker={'requires':['files.text','images.view'],'executor':'gemini-agent'},sourcing=dict(version=2,kind='inventory',image={'source':aid},max_records=3))
        review=dict(id='review',role='reviewer',objective='Review exact pixels and inventory',instruction='Inspect image; metadata is insufficient',inputs=[inp]+[dict(from_task='inventory',output=o['path'],path='upstream/'+o['path'],purpose=o['purpose'],authority='Unaccepted candidate',media_type=o['media_type']) for o in p['outputs']],outputs=[dict(path='review.md',purpose='Independent review')],dependencies=['inventory'],review_of='inventory',criteria=p['criteria'].copy(),limits=limits.copy(),max_attempts=1,worker={'requires':['files.text','images.view'],'executor':'gemini-agent'})
        value=dict(decision='ready',message='Typed visual inventory',plan=dict(brief='Inspect room',tasks=[p,review]),input_basis=dict(mode='new',artifacts=[]))
        _,compiled=planning.validate_result(c.encoded(value),row)
        self.assertEqual(compiled['tasks'][0]['operation_contract']['id'],'sourcing.inventory')
        self.assertIn(inventory.CRITERION,compiled['tasks'][1]['criteria'])
        self.assertTrue(next(i for i in compiled['tasks'][0]['inputs'] if i.get('artifact')==aid)['visual_reference'])
        self.assertEqual(self.factory.calls,[])
        payload['sourcing_records_version']=1;row['context']=c.encoded(payload)
        with self.assertRaises(ValueError):planning.validate_result(c.encoded(value),row)

    def test_new_pptx_scope_exposes_native_assembler_without_expanding_legacy_scope(self):
        with patch.object(executors,'available'):
            row=self.queue(action=self.action(step_capabilities=['pptx.create']))
        payload=json.loads(row['context'])
        self.assertIn('sourcing.slides',{o['id'] for o in payload['graph_operations']})
        self.assertEqual(payload['sourcing_records_version'],2)
