import copy
import io
import unittest
from orchestrator import pptx_document as ppt, slide_templates as templates
from tests.test_gemini import PNG


class Tests(unittest.TestCase):
    def fixture(self):
        return dict(version=2,title='Sample',brand=dict(body_font='Arial',heading_font='Georgia',accent='123456',
            logo=dict(path='logo.png',x=11.5,y=6.95,w=1.2,h=.35)),slides=[
            dict(layout='title',content=dict(title='A real title',subtitle='A subtitle')),
            dict(layout='image_text',content=dict(title='Plant',image='photo.png',caption='Source credit',body='Description'))])

    def test_brand_layouts_preserve_native_content_and_exact_images(self):
        from pptx import Presentation
        value=self.fixture();original=copy.deepcopy(value)
        raw,evidence=ppt.create(value,{'logo.png':PNG,'photo.png':PNG})
        deck=Presentation(io.BytesIO(raw))
        self.assertEqual(evidence['slide_count'],2)
        self.assertEqual(deck.slides[0].shapes[0].text,'A real title')
        self.assertEqual(deck.slides[0].shapes[0].text_frame.paragraphs[0].font.name,'Georgia')
        self.assertEqual(str(deck.slides[0].shapes[0].text_frame.paragraphs[0].font.color.rgb),'123456')
        self.assertEqual(deck.slides[1].shapes[-1].image.blob,PNG)
        self.assertEqual(value,original)

    def test_frozen_worker_validator_supports_custom_layout_and_brand(self):
        value=self.fixture()
        value['templates']={'our_cover':[dict(slot='headline',type='text',x=1,y=1,w=10,h=2,role='heading')]}
        value['slides']=[dict(layout='our_cover',content={'headline':'Customer template'})]
        namespace={'__name__':'frozen_validator'};exec(ppt.validator_source(),namespace)
        result=namespace['validate'](value,['logo.png'])
        self.assertEqual(result['slides'][0]['elements'][0]['text'],'Customer template')

    def test_all_builtin_layouts_compile_without_silently_dropping_content(self):
        slots=dict(title='Title',subtitle='Subtitle',body='Body',left='Left',right='Right',image='photo.png',caption='Credit',
                   rows=[['A','B'],['C','D']],chart=dict(chart='bar',categories=['A'],series=[dict(name='B',values=[1])]))
        document=dict(version=2,title='All layouts',slides=[dict(layout=name,content={s['slot']:slots[s['slot']] for s in spec}) for name,spec in templates.CATALOG.items()])
        _,evidence=ppt.create(document,{'photo.png':PNG})
        self.assertEqual(evidence['slide_count'],7)
        self.assertEqual(evidence['native_objects']['chart'],1)
        self.assertEqual(evidence['native_objects']['table'],1)

    def test_missing_extra_slots_logo_collision_and_undeclared_assets_rejected(self):
        for mutate,pattern in [
            (lambda d:d['slides'][0]['content'].pop('subtitle'),'slots'),
            (lambda d:d['slides'][0]['content'].update(unused='Do not lose this'),'slots'),
            (lambda d:d['brand']['logo'].update(x=1,y=2.2),'overlaps'),
            (lambda d:d['brand']['logo'].update(path='unregistered.png'),'exact declared')]:
            with self.subTest(pattern=pattern):
                value=self.fixture();mutate(value)
                with self.assertRaisesRegex(ValueError,pattern):ppt.validate(value,['logo.png','photo.png'])
