"""Code-owned composition and native intent boundaries; no model/network calls."""
import copy
import hashlib
import io
import json
from pathlib import Path
import tempfile
import unittest
from types import SimpleNamespace as NS
from unittest.mock import patch

from orchestrator import document_compose as compose, rhino_document as rhino


def native_document():
    return dict(version=1,units='Millimeters',spatial=dict(units='Millimeters',basis='estimated',note='Synthetic dimensions; not fabrication measurements'),objects=[
        dict(type='box',name='Body',min=[0,0,0],max=[600,400,500]),
        dict(type='polyline',name='Front',points=[[800,0,0],[1400,0,0],[1400,500,0],[800,500,0],[800,0,0]]),
        dict(type='text',name='Note',text='Estimated dimensions — not for fabrication',position=[800,600,0],height=20),
        dict(type='dimension',name='Width',start=[800,0,0],end=[1400,0,0],offset=[1100,-100,0],height=20)],
        views=[dict(name='Drawing',projection='Top')],preview=dict(resolution=[512,384],named_view='Drawing'))


def pixel(color):
    from PIL import Image
    buf=io.BytesIO();Image.new('RGB',(20,15),color).save(buf,format='PNG');return buf.getvalue()


class Tests(unittest.TestCase):
    def fixture(self):
        raw=pixel('blue');digest=hashlib.sha256(raw).hexdigest()
        doc=dict(version=1,title='Saved evidence',slides=[dict(notes='Source https://example.org/item; dimensions estimated',elements=[
            dict(type='text',text='Retained exact claim',x=.5,y=.5,w=10,h=1),
            dict(type='image',path='old/preview.png',x=.5,y=2,w=6,h=4)])],asset_sha256={'old/preview.png':digest})
        blobs={'source.json':json.dumps(doc).encode(),'new/image.png':raw}
        frozen=dict(inputs=[dict(path=p,sha256=hashlib.sha256(b).hexdigest(),media_type='application/json' if p.endswith('.json') else 'image/png') for p,b in blobs.items()],
            execution=dict(parameters=dict(document_paths=['source.json'],image_paths=['new/image.png'],image_titles=['Native drawing'],title='Combined')))
        return doc,blobs,frozen

    def test_compose_rebases_exact_pixels_and_preserves_claims_notes_and_order(self):
        doc,blobs,frozen=self.fixture();raw,receipt=compose.build(frozen,blobs.__getitem__);actual=json.loads(raw)
        self.assertEqual(actual['slides'][0]['notes'],doc['slides'][0]['notes'])
        self.assertEqual(actual['slides'][0]['elements'][0],doc['slides'][0]['elements'][0])
        self.assertEqual(actual['slides'][0]['elements'][1]['path'],'new/image.png')
        self.assertEqual(actual['slides'][1]['elements'][1]['fit'],'contain')
        self.assertEqual(receipt['model_calls'],0)
        self.assertEqual(receipt['sections'][0]['sha256'],frozen['inputs'][0]['sha256'])
        self.assertEqual(compose.build(frozen,blobs.__getitem__),(raw,receipt))
        # A corrected preview is a new version; old section pixels cannot be
        # silently exchanged even if the selected filename remains identical.
        blobs['new/image.png']=pixel('red');frozen['inputs'][1]['sha256']=hashlib.sha256(blobs['new/image.png']).hexdigest()
        with self.assertRaisesRegex(ValueError,'Exact saved section image'):compose.build(frozen,blobs.__getitem__)

    def test_stale_source_and_unbound_image_reject_before_output(self):
        doc,blobs,frozen=self.fixture();blobs['source.json']+=b' '
        with self.assertRaisesRegex(ValueError,'section version changed'):compose.build(frozen,blobs.__getitem__)
        doc.pop('asset_sha256');blobs['source.json']=json.dumps(doc).encode();frozen['inputs'][0]['sha256']=hashlib.sha256(blobs['source.json']).hexdigest()
        with self.assertRaisesRegex(ValueError,'lacks saved pixel bindings'):compose.build(frozen,blobs.__getitem__)

    def test_render_reopens_exact_composition_and_rejects_pixel_substitution(self):
        from orchestrator import pptx_document
        from pptx import Presentation
        doc,blobs,frozen=self.fixture();raw,_=compose.build(frozen,blobs.__getitem__);doc=json.loads(raw)
        pptx,evidence=pptx_document.create(doc,{'new/image.png':blobs['new/image.png']})
        deck=Presentation(io.BytesIO(pptx));self.assertEqual(len(deck.slides),2)
        self.assertIn('Retained exact claim',deck.slides[0].shapes[0].text)
        with self.assertRaisesRegex(ValueError,'image version changed'):
            pptx_document.create(doc,{'new/image.png':pixel('red')})

    def test_native_invalid_or_unsupported_data_never_reaches_host(self):
        value=native_document();self.assertEqual(rhino.checks(value)['expected_dimensions']['Body'],[600,400,500])
        for mutate in (lambda x:x['objects'][0].update(type='eval'),
                       lambda x:x['objects'][0].update(script='danger'),
                       lambda x:x['objects'][0]['min'].__setitem__(0,float('nan')),
                       lambda x:x['preview'].update(named_view='Missing'),
                       lambda x:x['objects'][3].update(offset=[1100,0,0])):
            bad=copy.deepcopy(value);mutate(bad)
            with self.assertRaises(ValueError):rhino.validate(bad)

    def test_missing_native_capability_is_detected_before_any_geometry_mutation(self):
        from unittest.mock import Mock
        add=Mock();rs=NS(AddBox=add,AddPolyline=add,AddText=add,ObjectName=add)
        with patch.dict('sys.modules',{'Rhino':NS(RhinoApp=NS(Version=NS(Major=8))),'rhinoscriptsyntax':rs}):
            with self.assertRaisesRegex(ValueError,'before execution'):rhino.build(None,native_document())
        add.assert_not_called()

    def test_native_text_projection_rejects_content_loss_even_when_plain_text_matches(self):
        original=' Estimated — mm\n\nFront / side; keep  spacing and \\n literally.\n'
        expected=dict(type='text',name='Disclosure',text=original,height=16,position=[0,0,0])
        value=dict(objects=[expected])
        projected=original.replace('\n','\r\n').replace('Front /',' Front /')
        probe=NS(RichText='exact native encoded content',PlainText=projected,Dispose=lambda:None)
        saved=NS(RichText=probe.RichText,PlainText=projected,TextHeight=16)
        doc=NS(Objects=[NS(Attributes=NS(Name='Disclosure'),Geometry=saved)])
        with patch.dict('sys.modules',{'rhinoscriptsyntax':NS()}),patch.object(rhino,'native_text',return_value=probe):
            self.assertEqual(rhino.verify(doc,value),[])
            for changed in (projected.strip(),projected.replace('\r\n\r\n','\r\n'),
                            projected.replace('  ',' '),projected.replace('mm','cm'),
                            projected.replace('\\n','\n')):
                saved.PlainText=changed
                self.assertEqual(rhino.verify(doc,value),['Text differs: Disclosure'])
            saved.PlainText=projected;saved.RichText='different encoded content'
            self.assertEqual(rhino.verify(doc,value),['Text differs: Disclosure'])
            saved.RichText=probe.RichText;saved.TextHeight=15
            self.assertEqual(rhino.verify(doc,value),['Text differs: Disclosure'])
        self.assertEqual(expected['text'],original)

    def test_missing_literal_text_encoder_blocks_before_creating_any_objects(self):
        from unittest.mock import Mock
        add=Mock()
        rs=NS(**{name:add for name in ('AddBox','AddPolyline','ObjectName','SetDocumentUserText',
            'GetDocumentUserText','AddAlignedDimension','AddDimStyle','DimStyleTextHeight',
            'DimStyleArrowSize','DimensionValue')})
        entity=NS(PlainTextToRtf=lambda x:x)
        native=NS(RhinoApp=NS(Version=NS(Major=8)),Geometry=NS(TextEntity=entity))
        with patch.dict('sys.modules',{'Rhino':native,'rhinoscriptsyntax':rs}):
            with self.assertRaisesRegex(ValueError,'before execution: TextEntity.CreateWithRichText'):
                rhino.build(None,native_document())
        add.assert_not_called()

    def test_new_operations_have_shared_provider_independent_interfaces(self):
        from orchestrator import execution, prepared_schemas, parameter_contracts
        from orchestrator.operation_lifecycle import traits
        doc=native_document();fields={k:v for k,v in doc.items() if k!='version'}
        parameter_contracts.validate(fields,prepared_schemas.REGISTRY['rhino.document']['schema'])
        self.assertTrue(traits('rhino.document')['external_effects'])
        self.assertEqual(traits('rhino.document')['resources'],['host-rhino'])
        self.assertFalse(traits('document.compose')['external_effects'])
        self.assertNotIn('document.compose',execution.MODEL_PROVIDERS)
        self.assertNotIn('rhino.document',execution.MODEL_PROVIDERS)

    def test_native_composition_inherits_exact_assets_and_old_grants_stay_narrow(self):
        from orchestrator import handoff_closure as closure
        assets=[dict(artifact='room',path='room.png',media_type='image/png',purpose='Reference',authority='Exact original'),
                dict(artifact='photos',path='photos.zip',media_type='application/zip',purpose='Product photos',authority='Exact saved bundle')]
        producer=dict(id='section',inputs=assets+[dict(artifact='catalog',path='catalog.json',media_type='application/json',purpose='Facts',authority='Exact evidence')])
        tasks=[dict(id='compose',execution=dict(capability='document.compose'),inputs=[dict(from_task='section',output='section.json',path='section.json',media_type='application/json',purpose='Section',authority='Saved projection')]),
               dict(id='render',execution=dict(capability='pptx.create'),inputs=[dict(from_task='compose',output='result.json',path='result.json',media_type='application/json',purpose='Composition',authority='Saved projection')])]
        grant=closure.capture(['document.compose','pptx.create'])
        closure.bind(tasks,lambda _:producer,lambda _:None,preserve_paths=lambda _:True,
            inherit=lambda task,item:item['media_type'] in grant.get(task['execution']['capability'],[]))
        for task in tasks:
            self.assertEqual([i for i in task['inputs'] if i.get('artifact')],[dict(a,purpose=next(i['purpose'] for i in task['inputs'] if i.get('artifact')==a['artifact'])) for a in assets])
            self.assertNotIn('catalog',[i.get('artifact') for i in task['inputs']])
        old={'pptx.create':grant['pptx.create'][:]};before=copy.deepcopy(old)
        closure.validate(old,['document.compose','pptx.create'])
        self.assertEqual(old,before)
        self.assertNotIn('document.compose',old)
        with self.assertRaises(ValueError):closure.validate({'document.compose':['application/x-blender']},['document.compose'])


if __name__=='__main__':unittest.main()

from tests import test_worker_tools as wf
from tests import test_owner_compilation as of

class LifecycleTests(unittest.TestCase):
    create=of.Tests.create
    source=of.Tests.source
    session=of.Tests.session
    invoke=of.Tests.invoke
    setUp=wf.Tests.setUp
    report=wf.Tests.report

    def test_composition_runs_through_worker_graph_and_reuses_committed_result_after_restart(self):
        from orchestrator import executors,worker_tools
        from orchestrator.agent_loop import execute
        from orchestrator.workers import atomic
        from tests.test_operation_session import Script,CONFIG
        self.grant['operations'].append('document.compose')
        doc=dict(version=1,title='Evidence',slides=[dict(elements=[dict(type='text',text='Source fact retained',x=1,y=1,w=10,h=1)])])
        ref=self.source(json.dumps(doc),'section.json','application/json')
        self.create(inputs=[ref]);self.rt.tick('job')
        self.session().delegate(dict(request_id='compose',tasks=[dict(id='make',objective='Compose exact saved evidence and render',inputs=[dict(artifact=ref['artifact'])],
            work=dict(operation='document.compose',parameters=dict(document_paths=[dict(artifact=ref['artifact'])],image_paths=[],image_titles=[],title='Combined'),
                      steps=[dict(id='deck',operation='pptx.create',from_steps=['build']),dict(id='preview',operation='pptx.render',from_steps=['deck'])]))]))
        self.rt.tick('job');state=self.factory.sessions[self.rt.task('job','make')['latest']]
        f=state['frozen'];control=Path(state['session']['control']);control.mkdir(parents=True,exist_ok=True)
        atomic(control/'launch.json',dict(credential_fingerprint=executors.fingerprint(CONFIG,f['backend'])))
        script=Script([('work_run',dict(step='build',request_key='compose')),
            ('work_run',dict(step='build',request_key='compose')), # exact receipt reuse
            ('work_run',dict(step='deck',request_key='deck')),
            ('work_run',dict(step='preview',request_key='preview')),
            ('image_view',dict(path='outputs/work-preview/slide-previews.zip',members=['slide-001.png'])),self.report(f)])
        save=worker_tools.Session.save
        def stop(session,contents,history,requests,tools,**kw):
            save(session,contents,history,requests,tools,**kw)
            if requests==1 and not kw.get('pending'):raise InterruptedError('committed compose restart')
        with patch.object(worker_tools.Session,'save',stop):
            with self.assertRaisesRegex(InterruptedError,'compose restart'):execute(f,control,script,lambda:(CONFIG,f['backend']))
        try:execute(f,control,script,lambda:(CONFIG,f['backend']))
        except BaseException:
            for path in sorted(control.glob('tool-*.json')):print(path.name,path.read_text()[:2400])
            raise
        worker_tools.verify_delivery(f,control)
        self.assertEqual(len(list((control/'work-operations').glob('*/outcome.json'))),3)
        self.assertTrue((Path(f['workspace'])/'outputs/work-preview/preview.pdf').is_file())
        first=json.loads((Path(f['workspace'])/'outputs/work-build/result.json').read_text())
        self.assertEqual(first['slides'][0]['elements'][0]['text'],'Source fact retained')
        for path in control.glob('tool-*.json'):
            result=json.loads(path.read_text());self.assertNotIn('error',result.get('result',{}),str(result))

    def test_typed_native_batches_fit_frozen_capacity_and_survive_reopen(self):
        from orchestrator import operation_records,prepared_contract,parameter_contracts
        self.grant['operations'].append('rhino.document')
        f,control,_=wf.Tests.admit(self,cap='rhino.document')
        value=native_document();value['objects']=[dict(type='box',name='Box '+str(i),min=[i*2,0,0],max=[i*2+1,1,1]) for i in range(32)]
        records=[dict(kind='field',field='units',value=value['units']),dict(kind='field',field='spatial',value=value['spatial'])]
        records += [dict(kind='append',field='objects',index=i,value=obj) for i,obj in enumerate(value['objects'])]
        records += [dict(kind='append',field='views',index=0,value=value['views'][0]),dict(kind='field',field='preview',value=value['preview']),dict(kind='end')]
        store=operation_records.Store(f,self.rt.db);revision=0
        for start in range(0,len(records),8):
            store.submit(dict(request_key='batch'+str(start),revision=revision,records=records[start:start+8]),{})
            revision+=1
        resumed=operation_records.Store(f,self.rt.db)
        self.assertTrue(resumed.progress()['sealed'])
        self.assertEqual(resumed.progress()['revision'],5)
        self.assertEqual(f['limits']['provider_requests'],24)
        self.assertEqual(resumed.progress()['committed'],len(records))
