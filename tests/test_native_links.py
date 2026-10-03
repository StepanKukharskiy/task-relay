"""Small, controlled native-link and conservative withdrawal checks."""

import json
import io
import sys
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from orchestrator import pptx_document, pptx_edit
from orchestrator.runtime import Runtime
from task_relay.bridge import State
from task_relay import (impact_handoff, native_links, production_control as pc,
                        reviewed_links, workflow_files)
from tests.test_presentations import fixture


class NativeLinksTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name).resolve()
        self.state = State(self.root / 'data/state.sqlite')
        self.state.channel = 'fixture'
        self.rt = Runtime(pc.root(self.state), connection=self.state.db)
        self.job = 'job-' + 'c' * 24
        self.state.db.execute('''INSERT INTO relay_pipelines
            (id,request_id,request,title,spec,channel,provider,model,status,created)
            VALUES (?,?,?,?,?,?,?,?,?,?)''',
            (self.job, 1, 'Withdraw Oleander; no replacement selected.', 'Native-link fixture',
             '{}', 'fixture', 'none', 'none', 'active', 1))
        self.state.db.commit()
        source = self.root / 'research.json'
        source.write_text(json.dumps({'answer': 'Oleander\tNerium oleander\tToxic\nOther\tOther plant'}))
        spec = fixture()
        spec['slides'][1]['elements'][1]['rows'] = [
            ['Species', 'Botanical'], ['Oleander', 'Nerium oleander'],
            ['Other', 'Other plant']]
        deck = self.root / 'baseline.pptx'
        deck.write_bytes(pptx_document.create(spec)[0])
        self.source = self.rt.register(source, 'Fixture research')
        self.deck = self.rt.register(deck, 'Fixture deck')
        shape = next(item for item in pptx_edit.inspect(deck.read_bytes())['slides'][1]['shapes']
                     if 'table_cells' in item)
        self.locator = {'type': 'pptx_table_cell', 'slide': 2,
                        'shape_id': shape['shape_id'], 'row': 1,
                        'column': 0, 'text': 'Oleander'}
        self.source_locator = {'type': 'json_line', 'field': 'answer',
                               'prefix': 'Oleander\tNerium oleander\t'}
        self.other_locator = {'type': 'pptx_table_cell', 'slide': 2,
                              'shape_id': shape['shape_id'], 'row': 2,
                              'column': 1, 'text': 'Other plant'}

    def tearDown(self):
        self.state.db.close()
        self.temp.cleanup()

    def test_exact_link_projects_to_job_and_plans_without_mutation(self):
        link = native_links.record_link(self.state.db, job=self.job,
            entity_key='Nerium oleander', predicate='plant_selection',
            source_artifact=self.source, source_locator=self.source_locator,
            output_artifact=self.deck, output_locator=self.locator,
            evidence='Fixture row and native cell compared.', reviewer='fixture-reviewer',
            review_state='reviewed')
        native_links.record_coverage(self.state.db, job=self.job,
            output_artifact=self.deck, output_locator=self.locator,
            state='complete', note='Exact reviewed fixture link.')
        native_links.record_link(self.state.db, job=self.job,
            entity_key='Other plant', predicate='plant_selection',
            source_artifact=self.source,
            source_locator={'type': 'json_line', 'field': 'answer',
                            'prefix': 'Other\tOther plant'},
            output_artifact=self.deck, output_locator=self.other_locator,
            evidence='Independent fixture row.', reviewer='fixture-reviewer',
            review_state='reviewed')
        native_links.record_coverage(self.state.db, job=self.job,
            output_artifact=self.deck, output_locator=self.other_locator,
            state='complete', note='Independent reviewed fixture link.')
        climate = {'type': 'pptx_slide', 'slide': 1}
        native_links.record_coverage(self.state.db, job=self.job,
            output_artifact=self.deck, output_locator=climate,
            state='incomplete', note='No independent source link.')
        with self.assertRaisesRegex(ValueError, 'Whole-slide completeness'):
            native_links.record_coverage(self.state.db, job=self.job,
                output_artifact=self.deck, output_locator=climate,
                state='complete', note='A slide cannot be certified from one link.')
        changes = self.state.db.total_changes
        plan = native_links.plan_withdrawal(self.state.db, job=self.job,
            entity_key='Nerium oleander', baseline_artifact=self.deck)
        self.assertEqual(self.state.db.total_changes, changes)
        self.assertEqual([len(plan[k]) for k in ('affected', 'unaffected', 'unknown')], [1, 1, 1])
        self.assertEqual(plan['affected'][0]['link_ids'], [link])
        model = reviewed_links.records(self.state.db, self.job)
        self.assertIn(self.locator, [item['output']['locator'] for item in model['links']])
        self.assertEqual(sorted(item['state'] for item in model['coverage']),
                         ['complete', 'complete', 'incomplete'])
        self.state.db.execute('''INSERT INTO relay_pipeline_steps
            (pipeline,position,id,status,sources) VALUES (?,?,?,?,?)''',
            (self.job, 0, 'baseline', 'completed',
             json.dumps([{'artifact': self.source}, {'artifact': self.deck}])))
        self.state.db.commit()
        view = workflow_files.sync(self.state, self.job)
        import sqlite3
        with sqlite3.connect(view / '.relay/job.sqlite') as db:
            self.assertEqual(db.execute("SELECT count(*) FROM reviewed_links WHERE kind='native_subject'").fetchone()[0], 2)
            self.assertEqual(db.execute("SELECT count(*) FROM reviewed_coverage WHERE kind='native_subject'").fetchone()[0], 3)

    def test_missing_source_and_unreviewed_coverage_remain_unknown(self):
        with self.assertRaisesRegex(ValueError, 'source line is missing'):
            native_links.record_link(self.state.db, job=self.job,
                entity_key='Different plant', predicate='plant_selection',
                source_artifact=self.source, source_locator=self.source_locator,
                output_artifact=self.deck, output_locator=self.locator,
                evidence='Wrong subject must not bind.', reviewer='fixture-reviewer')
        duplicate = self.root / 'duplicate.json'
        duplicate.write_text('{"answer":"Oleander\\tNerium oleander",'
                             '"answer":"Oleander\\tNerium oleander"}')
        duplicate_id = self.rt.register(duplicate, 'Ambiguous source')
        with self.assertRaisesRegex(ValueError, 'Duplicate source JSON key'):
            native_links.record_link(self.state.db, job=self.job,
                entity_key='Nerium oleander', predicate='plant_selection',
                source_artifact=duplicate_id, source_locator=self.source_locator,
                output_artifact=self.deck, output_locator=self.locator,
                evidence='Duplicate keys must not bind.', reviewer='fixture-reviewer')
        native_links.record_link(self.state.db, job=self.job,
            entity_key='Nerium oleander', predicate='plant_selection',
            source_artifact=self.source, source_locator=self.source_locator,
            output_artifact=self.deck, output_locator=self.locator,
            evidence='Fixture candidate link.', reviewer='fixture-reviewer')
        with self.assertRaisesRegex(ValueError, 'requires one reviewed link'):
            native_links.record_coverage(self.state.db, job=self.job,
                output_artifact=self.deck, output_locator=self.locator,
                state='complete', note='Cannot certify a pending link.')
        native_links.record_coverage(self.state.db, job=self.job,
            output_artifact=self.deck, output_locator=self.locator,
            state='incomplete', note='Review pending.')
        pending = native_links.plan_withdrawal(self.state.db, job=self.job,
            entity_key='Nerium oleander', baseline_artifact=self.deck)
        self.assertEqual([len(pending[k]) for k in ('affected', 'unaffected', 'unknown')], [0, 0, 1])
        self.assertIn('Incomplete', pending['unknown'][0]['reason'])
        blob = Path(self.rt.artifact(self.source)['blob'])
        blob.chmod(0o600)
        blob.write_text('{}')
        plan = native_links.plan_withdrawal(self.state.db, job=self.job,
            entity_key='Nerium oleander', baseline_artifact=self.deck)
        self.assertEqual([len(plan[k]) for k in ('affected', 'unaffected', 'unknown')], [0, 0, 1])
        self.assertIn('missing, changed or too large', plan['unknown'][0]['reason'])

    def test_reviewed_pdf_claim_can_preserve_one_exact_run(self):
        from reportlab.pdfgen import canvas

        stream = io.BytesIO()
        pdf = canvas.Canvas(stream)
        pdf.drawString(72, 720, 'Bullhead City annual mean 74.2 F')
        pdf.showPage()
        pdf.drawString(72, 720, 'Unrelated station annual mean 75.0 F')
        pdf.save()
        source = self.root / 'normals.pdf'
        source.write_bytes(stream.getvalue())
        source_id = self.rt.register(source, 'Fixture station normals')
        spec = fixture()
        spec['slides'][1]['elements'][1]['rows'] = [
            ['Place', 'Annual mean'], ['Bullhead City', '74.2 F']]
        climate_deck = self.root / 'climate.pptx'
        climate_deck.write_bytes(pptx_document.create(spec)[0])
        climate_id = self.rt.register(climate_deck, 'Fixture climate deck')
        table = next(item for item in pptx_edit.inspect(climate_deck.read_bytes())['slides'][1]['shapes']
                     if 'table_cells' in item)
        output_locator = {'type': 'pptx_table_cell', 'slide': 2,
                          'shape_id': table['shape_id'], 'row': 1,
                          'column': 1, 'text': '74.2 F'}
        locator = {'type': 'pdf_text', 'page': 1,
                   'needle': 'Bullhead City annual mean 74.2 F'}
        with self.assertRaisesRegex(ValueError, 'missing or ambiguous'):
            native_links.record_link(self.state.db, job=self.job,
                entity_key='Bullhead City', predicate='annual_mean',
                source_artifact=source_id, source_locator={**locator, 'page': 2},
                output_artifact=climate_id, output_locator=output_locator,
                evidence='Wrong page.', reviewer='fixture-reviewer', review_state='reviewed')
        native_links.record_link(self.state.db, job=self.job,
            entity_key='Bullhead City', predicate='annual_mean',
            source_artifact=source_id, source_locator=locator,
            output_artifact=climate_id, output_locator=output_locator,
            evidence='Exact fixture page and run.', reviewer='fixture-reviewer',
            review_state='reviewed')
        native_links.record_coverage(self.state.db, job=self.job,
            output_artifact=climate_id, output_locator=output_locator,
            state='complete', note='One reviewed output run only.')
        plan = native_links.plan_withdrawal(self.state.db, job=self.job,
            entity_key='Nerium oleander', baseline_artifact=climate_id)
        self.assertEqual([len(plan[k]) for k in ('affected', 'unaffected', 'unknown')],
                         [0, 1, 0])

    def test_two_native_text_runs_in_one_shape_export_separately(self):
        from pptx import Presentation

        slides = Presentation()
        slide = slides.slides.add_slide(slides.slide_layouts[6])
        box = slide.shapes.add_textbox(0, 0, 4000000, 500000)
        box.text_frame.paragraphs[0].add_run().text = '74.2 F'
        box.text_frame.paragraphs[0].add_run().text = '4503 CDD'
        path = self.root / 'two-runs.pptx'
        slides.save(path)
        deck = self.rt.register(path, 'Two native runs')
        source = self.root / 'two-values.json'
        source.write_text(json.dumps({'answer': 'Bullhead City 74.2 F\n'
                                               'Bullhead City 4503 CDD'}))
        evidence = self.rt.register(source, 'Two source values')
        for value in ('74.2 F', '4503 CDD'):
            locator = {'type': 'pptx_text_run', 'slide': 1,
                       'shape_id': box.shape_id, 'text': value}
            native_links.record_link(self.state.db, job=self.job,
                entity_key='Bullhead City', predicate='climate_value',
                source_artifact=evidence,
                source_locator={'type': 'json_line', 'field': 'answer',
                                'prefix': 'Bullhead City ' + value},
                output_artifact=deck, output_locator=locator,
                evidence='Two distinct source lines.', reviewer='fixture-reviewer',
                review_state='reviewed')
            native_links.record_coverage(self.state.db, job=self.job,
                output_artifact=deck, output_locator=locator,
                state='complete', note='This run only.')
        self.state.db.execute('''INSERT INTO relay_pipeline_steps
            (pipeline,position,id,status,sources) VALUES (?,?,?,?,?)''',
            (self.job, 0, 'two-runs', 'completed',
             json.dumps([{'artifact': deck}, {'artifact': evidence}])))
        self.state.db.commit()
        view = workflow_files.sync(self.state, self.job)
        import sqlite3
        with sqlite3.connect(view / '.relay/job.sqlite') as projection:
            rows = projection.execute('''SELECT output_location FROM reviewed_coverage
                WHERE kind='native_subject' ORDER BY output_location''').fetchall()
        self.assertEqual(len(rows), 2)
        self.assertNotEqual(rows[0][0], rows[1][0])

    def test_exact_change_request_handoff_recovers_export_and_detects_staleness(self):
        native_links.record_link(self.state.db, job=self.job,
            entity_key='Nerium oleander', predicate='plant_selection',
            source_artifact=self.source, source_locator=self.source_locator,
            output_artifact=self.deck, output_locator=self.locator,
            evidence='Exact fixture row.', reviewer='fixture-reviewer',
            review_state='reviewed')
        native_links.record_coverage(self.state.db, job=self.job,
            output_artifact=self.deck, output_locator=self.locator,
            state='complete', note='Reviewed exact run.')
        request = 'Withdraw Oleander. Keep unrelated reviewed values.\n'
        ident = impact_handoff.record_native_withdrawal(self.state.db,
            job=self.job, request_key='request-1', exact_request=request,
            actor='fixture-user', entity_key='Nerium oleander',
            baseline_artifact=self.deck)
        self.assertEqual(impact_handoff.inspect(self.state.db, ident)['status'], 'current')
        with patch('task_relay.native_links.plan_withdrawal',
                   side_effect=ValueError('PPTX reader unavailable')):
            self.assertEqual(impact_handoff.inspect(self.state.db, ident)['status'],
                             'unverifiable')
        self.assertEqual(impact_handoff.record_native_withdrawal(self.state.db,
            job=self.job, request_key='request-1', exact_request=request,
            actor='fixture-user', entity_key='Nerium oleander',
            baseline_artifact=self.deck), ident)
        with self.assertRaisesRegex(ValueError, 'different frozen plan'):
            impact_handoff.record_native_withdrawal(self.state.db,
                job=self.job, request_key='request-1', exact_request='Different request',
                actor='fixture-user', entity_key='Nerium oleander',
                baseline_artifact=self.deck)
        self.state.db.execute('''INSERT INTO relay_pipeline_steps
            (pipeline,position,id,status,sources) VALUES (?,?,?,?,?)''',
            (self.job, 0, 'baseline', 'completed',
             json.dumps([{'artifact': self.source}, {'artifact': self.deck}])))
        self.state.db.commit()
        with patch('task_relay.workflow_files.sync', side_effect=OSError('export interrupted')):
            with self.assertRaisesRegex(OSError, 'export interrupted'):
                workflow_files.sync(self.state, self.job)
        self.assertEqual(impact_handoff.record_native_withdrawal(self.state.db,
            job=self.job, request_key='request-1', exact_request=request,
            actor='fixture-user', entity_key='Nerium oleander',
            baseline_artifact=self.deck), ident)
        self.assertEqual(len(impact_handoff.records(self.state.db, self.job)), 1)
        view = workflow_files.sync(self.state, self.job)
        import sqlite3
        with sqlite3.connect(view / '.relay/job.sqlite') as projection:
            row = projection.execute('''SELECT exact_request,plan_digest FROM impact_handoffs
                WHERE id=?''', (ident,)).fetchone()
            process = projection.execute('''SELECT count(*) FROM process_records
                WHERE source_table='relay_impact_handoffs' ''').fetchone()[0]
        self.assertEqual(row[0], request)
        self.assertEqual(row[1], impact_handoff.inspect(self.state.db, ident)['plan_digest'])
        self.assertEqual(process, 1)
        with sqlite3.connect(view / '.relay/job.sqlite') as projection:
            snapshot = projection.execute("SELECT value FROM meta WHERE key='snapshot'").fetchone()[0]
        state_view = json.loads((view / '.relay/snapshots' / snapshot / 'job-state.json').read_text())
        self.assertEqual(state_view['impact_handoffs'][0]['exact_request'], request)
        self.assertEqual(state_view['impact_handoffs'][0]['plan']['affected'][0]['location'],
                         self.locator)
        native_links.record_coverage(self.state.db, job=self.job,
            output_artifact=self.deck, output_locator={'type': 'pptx_slide', 'slide': 1},
            state='incomplete', note='New unknown scope.')
        self.assertEqual(impact_handoff.inspect(self.state.db, ident)['status'], 'stale')
        with self.assertRaisesRegex(ValueError, 'different frozen plan'):
            impact_handoff.record_native_withdrawal(self.state.db,
                job=self.job, request_key='request-1', exact_request=request,
                actor='fixture-user', entity_key='Nerium oleander',
                baseline_artifact=self.deck)

    def test_missing_source_reader_makes_handoff_unverifiable(self):
        with patch.dict(sys.modules, {'pypdf': None}):
            with self.assertRaisesRegex(native_links.VerifierUnavailable,
                                        'requires pypdf'):
                native_links._source(b'', {'type': 'pdf_text', 'page': 1,
                    'needle': 'Nerium oleander'}, 'Nerium oleander')
        native_links.record_link(self.state.db, job=self.job,
            entity_key='Nerium oleander', predicate='plant_selection',
            source_artifact=self.source, source_locator=self.source_locator,
            output_artifact=self.deck, output_locator=self.locator,
            evidence='Fixture source.', reviewer='fixture-reviewer', review_state='reviewed')
        native_links.record_coverage(self.state.db, job=self.job,
            output_artifact=self.deck, output_locator=self.locator,
            state='complete', note='Reviewed fixture location.')
        ident = impact_handoff.record_native_withdrawal(self.state.db,
            job=self.job, request_key='missing-reader', exact_request='Withdraw Oleander.',
            actor='fixture-user', entity_key='Nerium oleander',
            baseline_artifact=self.deck)
        self.assertEqual(impact_handoff.inspect(self.state.db, ident)['status'], 'current')
        with patch('task_relay.native_links._source',
                   side_effect=native_links.VerifierUnavailable('Reader missing.')):
            self.assertEqual(impact_handoff.inspect(self.state.db, ident)['status'],
                             'unverifiable')

    def test_notes_run_is_explicitly_linked_before_withdrawal(self):
        from pptx import Presentation
        baseline = self.root / 'with-notes.pptx'
        deck = Presentation(self.root / 'baseline.pptx')
        deck.slides[1].notes_slide.notes_text_frame.text = (
            'Image: Oleander\nImage: Other plant')
        deck.save(baseline)
        artifact = self.rt.register(baseline, 'Fixture with notes')
        notes = pptx_edit.inspect(baseline.read_bytes())['slides'][1]['notes_text_runs']
        locator = {'type': 'pptx_notes_run', 'slide': 2,
                   'index': notes.index('Image: Oleander'), 'text': 'Image: Oleander'}
        native_links.record_link(self.state.db, job=self.job,
            entity_key='Nerium oleander', predicate='photo_attribution',
            source_artifact=self.source, source_locator=self.source_locator,
            output_artifact=artifact, output_locator=locator,
            evidence='Exact notes run and fixture subject.', reviewer='fixture-reviewer',
            review_state='reviewed')
        native_links.record_coverage(self.state.db, job=self.job,
            output_artifact=artifact, output_locator=locator,
            state='complete', note='Exact notes source reviewed.')
        plan = native_links.plan_withdrawal(self.state.db, job=self.job,
            entity_key='Nerium oleander', baseline_artifact=artifact)
        self.assertEqual([item['location'] for item in plan['affected']], [locator])
        self.assertEqual(pptx_edit.inspect(baseline.read_bytes())['slides'][1]['notes_text_runs'], notes)


if __name__ == '__main__':
    unittest.main()
