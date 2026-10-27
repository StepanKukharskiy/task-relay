"""Shared boundaries: real files/Runtime with scripted decisions, no paid models."""
import copy
import io
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch
from orchestrator import artifact_structure as structure,working_context as context,handoff_closure
from orchestrator import contracts as c,job_owner
from tests import test_job_owner_loop as fixture


class ContextTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.view=context.WorkingView(Path(self.tmp.name))

    def test_growth_bounded_lossless_restart_and_tool_pairs(self):
        initial=dict(role='user',parts=[dict(text='Exact request and frozen grant')])
        messages=[initial]
        for n in range(80):
            messages += [dict(role='user',parts=[dict(text=context.STATE_PREFIX+str(n))]),
                dict(role='model',parts=[dict(functionCall=dict(name='job_delegate',args={'value':'x'*4000}),thoughtSignature='signature'+str(n))]),
                dict(role='user',parts=[dict(functionResponse=dict(name='job_delegate',response={'receipt':n,'data':'y'*4000}))])]
        before=copy.deepcopy(messages);view=self.view.project(messages,'gemini',20000)
        self.assertLessEqual(context.size(view),20000);self.assertEqual(messages,before)
        self.assertEqual(sum(context.state(m) for m in view),1)
        self.assertEqual(view[-2:],messages[-2:]);self.assertEqual(view[0],initial)
        for i,m in enumerate(view):
            if m.get('role')=='model':self.assertIn('functionResponse',view[i+1]['parts'][0])
        files=list(self.view.folder.glob('*.json'));self.assertEqual(len(files),1)
        reopened=context.WorkingView(Path(self.tmp.name));ident=files[0].stem;text='';offset=0
        while offset is not None:
            part=reopened.read(dict(context=ident,offset=offset));text+=part['text'];offset=part['next_offset']
        archived=json.loads(text)
        self.assertEqual(archived[0]['parts'][0]['thoughtSignature'],'signature0')
        files[0].write_text('[]')
        with self.assertRaisesRegex(ValueError,'changed'):reopened.read(dict(context=ident,offset=0))

    def test_native_multicall_reasoning_pairs_and_hard_capacity(self):
        for provider in ('openai','openrouter'):
            initial=dict(role='user',content='Exact request')
            if provider=='openai':
                group=[dict(type='reasoning',encrypted_content='opaque'),dict(type='function_call',call_id='a',name='read',arguments='{}'),dict(type='function_call',call_id='b',name='read',arguments='{}'),dict(type='function_call_output',call_id='a',output='x'*4000),dict(type='function_call_output',call_id='b',output='y'*4000)]
            else:
                group=[dict(role='assistant',tool_calls=[dict(id='a'),dict(id='b')]),dict(role='tool',tool_call_id='a',content='x'*4000),dict(role='tool',tool_call_id='b',content='y'*4000)]
            result=self.view.project([initial]+group*8,provider,15000)
            self.assertEqual(result[-len(group):],group)
            with self.assertRaisesRegex(ValueError,'capacity'):self.view.project([initial]+group,provider,1000)


class StructureTests(unittest.TestCase):
    def deck(self,text='Useful content'):
        from pptx import Presentation
        p=Presentation()
        if text is not None:p.slides.add_slide(p.slide_layouts[6]).shapes.add_textbox(0,0,1000000,1000000).text=text
        b=io.BytesIO();p.save(b);return b.getvalue()

    def test_metadata_repack_does_not_count_as_progress_but_text_edit_does(self):
        output=dict(path='deck.pptx',media_type=structure.PPTX);raw=self.deck();b=io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(raw)) as old,zipfile.ZipFile(b,'w') as new:
            for n in reversed(old.namelist()):
                i=zipfile.ZipInfo(n,(2001,1,1,0,0,0));new.writestr(i,old.read(n))
        self.assertNotEqual(raw,b.getvalue())
        a=structure.inspect(output,raw);repacked=structure.inspect(output,b.getvalue())
        self.assertEqual(a['fingerprint'],repacked['fingerprint']);self.assertNotEqual(a['sha256'],repacked['sha256'])
        self.assertNotEqual(a['fingerprint'],structure.inspect(output,self.deck('Corrected evidence'))['fingerprint'])
        state,errors=structure.progress([output],lambda _:self.deck(None))
        self.assertEqual(state,{});self.assertIn('no slides',errors[0])

    def test_broken_internal_asset_relationship_is_rejected(self):
        raw=self.deck();b=io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(raw)) as old,zipfile.ZipFile(b,'w') as new:
            for n in old.namelist():
                if n!='ppt/slides/slide1.xml':new.writestr(n,old.read(n))
        with self.assertRaisesRegex(ValueError,'relationship'):structure.inspect(dict(path='deck.pptx'),b.getvalue())

    def test_structural_receipt_reused_only_for_exact_bytes_and_validator(self):
        with tempfile.TemporaryDirectory() as directory:
            output=dict(path='deck.pptx');raw=self.deck()
            a=structure.inspect_cached(output,raw,directory)
            with patch.object(structure,'inspect',side_effect=AssertionError('Repeated structural parse')):
                self.assertEqual(a,structure.inspect_cached(output,raw,directory))
            self.assertNotEqual(a['sha256'],structure.inspect_cached(output,self.deck('Changed'),directory)['sha256'])
            receipt=next(p for p in Path(directory).glob('*.json') if json.loads(p.read_text())['receipt']['sha256']==a['sha256']);saved=json.loads(receipt.read_text());saved['identity']['sha256']='wrong';receipt.write_text(json.dumps(saved))
            with self.assertRaisesRegex(ValueError,'identity changed'):structure.inspect_cached(output,raw,directory)

    def test_json_pdf_and_new_registered_format_share_boundary(self):
        from pypdf import PdfWriter
        output=dict(path='data.json');a=structure.inspect(output,b'{"a":1,"b":2}')
        self.assertEqual(a['fingerprint'],structure.inspect(output,b'{"b": 2, "a": 1}')['fingerprint'])
        p=PdfWriter();b=io.BytesIO();p.write(b)
        with self.assertRaisesRegex(ValueError,'no pages'):structure.inspect(dict(path='preview.pdf'),b.getvalue())
        with patch.dict(structure.REGISTRY,{'application/x-fixture':dict(extensions=['.fixture'],inspect=lambda raw:dict(units=1,fingerprint=raw.decode()))}):
            self.assertEqual(structure.inspect(dict(path='model.fixture'),b'qualified')['fingerprint'],'qualified')


class HandoffTests(unittest.TestCase):
    def test_shared_graph_carries_exact_sources_and_collision_safe_paths(self):
        source=dict(artifact='original',path='evidence.txt',purpose='Original observations',authority='Source only',media_type='text/plain')
        tasks=[dict(id='extract',inputs=[source],outputs=[dict(path='data.json')],dependencies=[]),
               dict(id='build',inputs=[dict(from_task='extract',output='data.json',path='data.json',purpose='Data',authority='Unaccepted',media_type='application/json')],outputs=[dict(path='report.txt')],dependencies=[]),
               dict(id='review',inputs=[dict(from_task='build',output='report.txt',path='evidence.txt',purpose='Candidate',authority='Unaccepted',media_type='text/plain')],outputs=[dict(path='review.txt')],dependencies=[])]
        handoff_closure.bind(tasks,lambda _:None,lambda _:None)
        self.assertIn('extract',tasks[1]['dependencies']);self.assertIn('build',tasks[2]['dependencies'])
        inherited=next(i for i in tasks[2]['inputs'] if i.get('artifact')=='original')
        self.assertTrue(inherited['path'].startswith('evidence/'));self.assertEqual(inherited['authority'],'Source only')
        before=copy.deepcopy(tasks);handoff_closure.bind(tasks,lambda _:None,lambda _:None);self.assertEqual(tasks,before)
        tasks[0]['inputs'].append(dict(from_task='review',output='review.txt',path='cycle.txt'))
        with self.assertRaisesRegex(ValueError,'Cyclic'):handoff_closure.bind(tasks,lambda _:None,lambda _:None)


    def test_native_assets_keep_required_path_and_conflicts_fail_before_dispatch(self):
        source=dict(artifact='image',path='assets/photo.png',purpose='Selected pixels',authority='Source',media_type='image/png')
        tasks=[dict(id='content',inputs=[source],outputs=[dict(path='spec.json')]),dict(id='native',execution={},inputs=[dict(from_task='content',output='spec.json',path='spec.json',purpose='Data'),dict(source,path='renamed.png')],outputs=[dict(path='deck.pptx')])]
        handoff_closure.bind(tasks,lambda _:None,lambda _:None,preserve_paths=lambda t:t['id']=='native')
        self.assertIn('assets/photo.png',[i['path'] for i in tasks[1]['inputs']])
        tasks[1]['inputs'][-1]['artifact']='different-image'
        with self.assertRaisesRegex(ValueError,'path conflicts'):handoff_closure.bind(tasks,lambda _:None,lambda _:None,preserve_paths=lambda t:t['id']=='native')


class RuntimeTests(unittest.TestCase):
    setUp=fixture.Tests.setUp
    create=fixture.Tests.create
    invoke=fixture.Tests.invoke
    batch=fixture.Tests.batch

    def test_default_catalog_is_compact_and_large_sources_stay_retrievable(self):
        self.grant['working_context_version']=1 # Keep coverage of the older excerpt transport.
        self.grant['operations']=['pptx.create','pptx.render','images.crop','collection.slides','records.merge']
        source=self.root/'source.txt';source.write_text('evidence '*30000);ident=self.rt.register(source,'Large exact source')
        ref=dict(artifact=ident,path='source.txt',purpose='Source',authority='Evidence',media_type='text/plain')
        self.create(inputs=[ref]);self.rt.tick('job')
        _,script=self.invoke([('job_catalog',{}),('job_read',dict(artifact=ident,offset=16000,limit=100)),('job_delegate',dict(request_id='finish-fixture',batch_json=c.encoded(self.batch('next'))))])
        first=json.loads(script.requests[0]['contents'][0]['parts'][0]['text'])
        self.assertEqual(len(first['source_pack']['files'][0]['text']),16000)
        self.assertEqual(first['source_pack']['files'][0]['next_offset'],16000)
        reply=script.requests[1]['contents'][-1]['parts'][0]['functionResponse']['response']
        self.assertEqual(reply['operations'],{});self.assertEqual(len(reply['available_operations']),5)
        self.assertLess(len(c.encoded(reply)),10000)
        read=script.requests[2]['contents'][-1]['parts'][0]['functionResponse']['response']
        self.assertEqual(read['text'],source.read_text()[16000:16100])

    def test_owner_auto_binds_upstream_evidence_without_another_model_turn(self):
        source=self.root/'source.txt';source.write_text('Exact evidence')
        ident=self.rt.register(source,'Source');ref=dict(artifact=ident,path='source.txt',purpose='Source',authority='Observed evidence',media_type='text/plain')
        self.create(inputs=[ref]);self.rt.tick('job');batch=self.batch(inputs=[ref]);batch[1]['dependencies']=[]
        _,script=self.invoke([('job_delegate',dict(request_id='closure',batch_json=c.encoded(batch)))])
        self.assertEqual(len(script.requests),1)
        spec=self.rt.spec(self.rt.task('job','v1_review'))
        self.assertIn(ref['artifact'],[i.get('artifact') for i in spec['inputs']]);self.assertIn('v1',spec['dependencies'])

    def test_native_renderer_gets_assets_while_reviewer_gets_full_evidence(self):
        from orchestrator import worker_capabilities
        from tests.test_document_authoring import Tests as LayoutFixture
        from tests.test_browser_screenshots import png
        self.grant['operations']=['pptx.create'];self.grant['deliverables']={'deck':structure.PPTX}
        self.grant['workers'].append(worker_capabilities.entry(dict(type='gemini-code',model='fixture',runtime='a'*64)))
        sources=[]
        for name,raw,media in [('research.json',b'{"observation":"Exact source"}','application/json'),('image.png',png(),'image/png')]:
            p=self.root/name;p.write_bytes(raw);ident=self.rt.register(p,'Exact evidence')
            sources.append(dict(artifact=ident,path=name,purpose='Evidence',authority='Observed source',media_type=media,**({'visual_reference':True} if media=='image/png' else {})))
        self.create(inputs=sources);self.rt.tick('job')
        content=LayoutFixture.content(self,inputs=sources);render=LayoutFixture.renderer(self)
        review=dict(id='review',objective='Review final file',instruction='Judge the full request against exact original evidence.',review_of='render',
            inputs=[dict(from_task='render',output='deck.pptx',path='deck.pptx',purpose='Candidate',authority='Unaccepted',media_type=structure.PPTX)],
            outputs=[dict(path='review.md',purpose='Findings',media_type='text/markdown')],dependencies=[],worker=dict(requires=['code.execute','files.binary']))
        _,script=self.invoke([('job_delegate',dict(request_id='assets',batch_json=c.encoded([content,render,review]))),('job_delegate',dict(request_id='finish-fixture',batch_json=c.encoded(self.batch('next'))))])
        self.assertEqual(len(script.requests),1,c.encoded(script.requests[-1]['contents'][-1]))
        native=self.rt.spec(self.rt.task('job','render'));review=self.rt.spec(self.rt.task('job','review'))
        self.assertEqual([i['path'] for i in native['inputs']],['slides.json','image.png'])
        self.assertIn('research.json',[i['path'] for i in review['inputs']])
        self.assertIn('image.png',[i['path'] for i in review['inputs']])
        self.assertTrue(review['final_quality']['checks'])
        self.assertFalse(next(iter(review['final_quality']['checks'].values()))['target_gaps'])

    def test_empty_custom_document_rejected_before_paid_reviewer_dispatch(self):
        from orchestrator import worker_capabilities,document_authoring
        from pptx import Presentation
        self.grant['workers'].append(worker_capabilities.entry(dict(type='gemini-code',model='fixture',runtime='a'*64)))
        self.grant['authoring_policy']=document_authoring.capture([],{structure.PPTX:'Explicit synthetic custom fallback'})
        self.create();self.rt.tick('job');batch=self.batch();author,review=batch
        author.update(worker=dict(requires=['code.execute','files.binary']),authoring_fallback=dict(reason='Fixture'))
        author['outputs']=[dict(path='deck.pptx',purpose='Deck',media_type=structure.PPTX)]
        review['inputs']=[dict(from_task='v1',output='deck.pptx',path='deck.pptx',purpose='Candidate',authority='Unaccepted',media_type=structure.PPTX)]
        review['worker']=dict(requires=['code.execute','files.binary'])
        self.invoke([('job_delegate',dict(request_id='empty',batch_json=c.encoded(batch)))])
        self.rt.tick('job');aid=self.rt.task('job','v1')['latest'];s=self.factory.sessions[aid]
        self.factory.finish(aid);Presentation().save(str(Path(s['frozen']['workspace'])/'deck.pptx'))
        self.rt.tick('job')
        self.assertEqual(self.rt.task('job','v1')['status'],'blocked')
        self.assertIsNone(self.rt.task('job','v1_review')['latest'])
        self.assertNotIn('v1_review',[v['frozen']['id'] for v in self.factory.sessions.values()]) # owner may wake to repair
