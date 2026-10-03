"""Small controlled PPTX revision fixtures; no native application or provider."""
import hashlib
import io
import json
from pathlib import Path
import tempfile
import time
import unittest
import zipfile

from orchestrator import contracts, execution, pptx_document, pptx_edit
from orchestrator.adapters import ExecutionFactory, RegisteredFactory
from task_relay import presentations
from tests.test_presentations import fixture
from tests import test_mixed_execution as mixed
from tests import test_production_planning as planning_fixtures
from task_relay import production_planning as planning
from tests.test_orchestrator import plan


def manifest(raw, slide=1, shape_id=2, old='Quarterly results', new='Annual results'):
    return {'version':1,'source_sha256':hashlib.sha256(raw).hexdigest(),
            'edits':[{'kind':'replace_text','slide':slide,'shape_id':shape_id,'old':old,'new':new}]}


class EditTests(unittest.TestCase):
    def setUp(self):
        self.raw,_=pptx_document.create(fixture())

    def test_edits_exact_shape_without_rebuilding_other_package_parts(self):
        from pptx import Presentation
        listing=pptx_edit.inspect(self.raw)
        self.assertEqual(listing['slides'][0]['shapes'][0]['text_runs'],['Quarterly results'])
        spec=manifest(self.raw,shape_id=listing['slides'][0]['shapes'][0]['shape_id'])
        result,evidence=pptx_edit.edit(self.raw,spec)
        self.assertEqual(evidence['changed_slides'],[1])
        self.assertEqual(Presentation(io.BytesIO(result)).slides[0].shapes[0].text,'Annual results')
        self.assertEqual(Presentation(io.BytesIO(self.raw)).slides[0].shapes[0].text,'Quarterly results')
        with zipfile.ZipFile(io.BytesIO(self.raw)) as before,zipfile.ZipFile(io.BytesIO(result)) as after:
            self.assertEqual(set(before.namelist()),set(after.namelist()))
            for name in before.namelist():
                if name!='ppt/slides/slide1.xml':
                    self.assertEqual(before.read(name),after.read(name),name)

    def test_stale_or_ambiguous_edit_fails_without_candidate(self):
        spec=manifest(self.raw)
        spec['source_sha256']='0'*64
        with self.assertRaisesRegex(ValueError,'hash changed'):pptx_edit.edit(self.raw,spec)
        spec=manifest(self.raw,old='Unknown')
        with self.assertRaisesRegex(ValueError,'missing or ambiguous'):pptx_edit.edit(self.raw,spec)
        spec=manifest(self.raw)
        spec['edits'].append(dict(spec['edits'][0]))
        with self.assertRaisesRegex(ValueError,'missing or ambiguous'):pptx_edit.edit(self.raw,spec)

    def test_inspect_retains_native_table_cell_locations_for_repeated_text(self):
        source = fixture()
        source['slides'][1]['elements'][1]['rows'] = [
            ['Species', 'Botanical name'],
            ['Oleander', 'Nerium oleander'],
            ['Other', 'Oleander'],
        ]
        raw, _ = pptx_document.create(source)
        shapes = pptx_edit.inspect(raw)['slides'][1]['shapes']
        table = next(shape for shape in shapes if 'table_cells' in shape)
        cells = {(cell['row'], cell['column']): cell['text_runs'] for cell in table['table_cells']}
        self.assertEqual(cells[(1, 0)], ['Oleander'])
        self.assertEqual(cells[(1, 1)], ['Nerium oleander'])
        self.assertEqual(cells[(2, 1)], ['Oleander'])
        self.assertEqual(table['text_runs'].count('Oleander'), 2)

    def test_edit_targets_one_table_cell_when_text_repeats(self):
        source = fixture()
        source['slides'][1]['elements'][1]['rows'] = [
            ['Species', 'Botanical name'], ['Oleander', 'Nerium oleander'],
            ['Other', 'Oleander'],
        ]
        raw, _ = pptx_document.create(source)
        listing = pptx_edit.inspect(raw)
        table = next(shape for shape in listing['slides'][1]['shapes'] if 'table_cells' in shape)
        spec = manifest(raw, slide=2, shape_id=table['shape_id'], old='Oleander', new='Withdrawn')
        with self.assertRaisesRegex(ValueError, 'ambiguous'):
            pptx_edit.edit(raw, spec)
        spec['edits'][0].update(row=1, column=0)
        result, evidence = pptx_edit.edit(raw, spec)
        revised_table = next(shape for shape in pptx_edit.inspect(result)['slides'][1]['shapes']
                             if shape['shape_id'] == table['shape_id'])
        cells = {(cell['row'], cell['column']): cell['text_runs']
                 for cell in revised_table['table_cells']}
        self.assertEqual(cells[(1, 0)], ['Withdrawn'])
        self.assertEqual(cells[(2, 1)], ['Oleander'])
        self.assertEqual(evidence['changed_slides'], [2])
        with zipfile.ZipFile(io.BytesIO(raw)) as before, zipfile.ZipFile(io.BytesIO(result)) as after:
            for name in before.namelist():
                if name != 'ppt/slides/slide2.xml':
                    self.assertEqual(before.read(name), after.read(name), name)
        wrong = manifest(raw, slide=2, shape_id=table['shape_id'], old='Oleander', new='Withdrawn')
        wrong['edits'][0].update(row=9, column=0)
        with self.assertRaisesRegex(ValueError, 'row is missing'):
            pptx_edit.edit(raw, wrong)

    def test_notes_runs_are_addressable_and_removable_without_slide_change(self):
        from pptx import Presentation
        presentation = Presentation(io.BytesIO(self.raw))
        notes = presentation.slides[0].notes_slide.notes_text_frame
        notes.text = 'Image: Oleander\nAuthor: Example\nImage: Hop bush'
        source = io.BytesIO()
        presentation.save(source)
        raw = source.getvalue()
        listing = pptx_edit.inspect(raw)
        runs = listing['slides'][0]['notes_text_runs']
        first = runs.index('Image: Oleander')
        author = runs.index('Author: Example')
        spec = {'version': 1, 'source_sha256': hashlib.sha256(raw).hexdigest(),
                'edits': [
                    {'kind': 'remove_notes_run', 'slide': 1, 'index': first,
                     'old': 'Image: Oleander'},
                    {'kind': 'remove_notes_run', 'slide': 1, 'index': author,
                     'old': 'Author: Example'}]}
        changed, receipt = pptx_edit.edit(raw, spec)
        after = pptx_edit.inspect(changed)['slides'][0]['notes_text_runs']
        self.assertNotIn('Image: Oleander', after)
        self.assertNotIn('Author: Example', after)
        self.assertIn('Image: Hop bush', after)
        self.assertEqual(receipt['changed_slides'], [1])
        with zipfile.ZipFile(io.BytesIO(raw)) as old, zipfile.ZipFile(io.BytesIO(changed)) as new:
            self.assertEqual(old.read('ppt/slides/slide1.xml'), new.read('ppt/slides/slide1.xml'))
            changed_parts = {name for name in old.namelist() if old.read(name) != new.read(name)}
        self.assertEqual(changed_parts, {'ppt/notesSlides/notesSlide1.xml'})
        spec['edits'][1]['old'] = 'Wrong author'
        with self.assertRaisesRegex(ValueError, 'notes run changed'):
            pptx_edit.edit(raw, spec)

    def test_removes_exact_table_row_and_unique_picture_without_rebuilding_deck(self):
        from PIL import Image
        from pptx import Presentation
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        table = slide.shapes.add_table(3, 2, 0, 0, 6000000, 1500000).table
        for row, values in enumerate((('Species','Botanical'),
                                      ('Oleander','Nerium oleander'),
                                      ('Hop bush','Dodonaea viscosa'))):
            for column, value in enumerate(values):
                table.cell(row,column).text = value
        with tempfile.TemporaryDirectory() as folder:
            photo = Path(folder) / 'oleander.png'
            Image.new('RGB', (4, 4), 'pink').save(photo)
            slide.shapes.add_picture(str(photo), 0, 1700000)
            buffer = io.BytesIO()
            presentation.save(buffer)
        raw = buffer.getvalue()
        listing = pptx_edit.inspect(raw)['slides'][0]['shapes']
        table_shape = next(shape for shape in listing if shape.get('table_cells'))
        picture = next(shape for shape in listing if shape.get('image_sha256'))
        spec = {'version': 1, 'source_sha256': hashlib.sha256(raw).hexdigest(),
                'edits': [
                    {'kind': 'remove_table_row', 'slide': 1,
                     'shape_id': table_shape['shape_id'], 'row': 1,
                     'old_runs': ['Oleander','Nerium oleander']},
                    {'kind': 'remove_shape', 'slide': 1,
                     'shape_id': picture['shape_id'],
                     'old_xml_sha256': picture['xml_sha256']}]}
        changed, receipt = pptx_edit.edit(raw, spec)
        after = pptx_edit.inspect(changed)['slides'][0]['shapes']
        self.assertEqual(len(next(shape for shape in after if shape.get('table_cells'))['table_cells']), 4)
        self.assertFalse(any(shape.get('image_sha256') for shape in after))
        self.assertEqual(len(Presentation(io.BytesIO(changed)).slides), 1)
        self.assertEqual(len(receipt['removed_parts']), 1)
        with zipfile.ZipFile(io.BytesIO(raw)) as old, zipfile.ZipFile(io.BytesIO(changed)) as new:
            for name in old.namelist():
                if name not in {'ppt/slides/slide1.xml', 'ppt/slides/_rels/slide1.xml.rels',
                                *receipt['removed_parts']}:
                    self.assertEqual(old.read(name), new.read(name), name)
        spec['edits'][0]['old_runs'][0] = 'Wrong plant'
        with self.assertRaisesRegex(ValueError, 'row content changed'):
            pptx_edit.edit(raw, spec)

    def test_replaces_one_shared_picture_without_changing_the_other(self):
        from PIL import Image
        from pptx import Presentation
        from tests.test_gemini import PNG
        source=fixture();source['slides']=[{'elements':[
            {'type':'image','path':'old.png','x':1,'y':1,'w':3,'h':3},
            {'type':'image','path':'old.png','x':5,'y':1,'w':3,'h':3}]}]
        raw,_=pptx_document.create(source,{'old.png':PNG})
        picture=io.BytesIO();Image.new('RGB',(1,1),'blue').save(picture,format='PNG')
        replacement=picture.getvalue();listing=pptx_edit.inspect(raw)
        first,second=listing['slides'][0]['shapes']
        self.assertEqual(first['image_sha256'],second['image_sha256'])
        spec={'version':1,'source_sha256':hashlib.sha256(raw).hexdigest(),'edits':[
            {'kind':'replace_image','slide':1,'shape_id':first['shape_id'],
             'old_sha256':first['image_sha256'],'path':'replacement.png',
             'new_sha256':hashlib.sha256(replacement).hexdigest()}]}
        result,evidence=pptx_edit.edit(raw,spec,images={'replacement.png':replacement})
        deck=Presentation(io.BytesIO(result))
        self.assertEqual(deck.slides[0].shapes[0].image.blob,replacement)
        self.assertEqual(deck.slides[0].shapes[1].image.blob,PNG)
        self.assertEqual(Presentation(io.BytesIO(raw)).slides[0].shapes[0].image.blob,PNG)
        self.assertEqual(len(evidence['added_parts']),1)
        with tempfile.TemporaryDirectory() as directory:
            folder=Path(directory)
            (folder/'baseline.pptx').write_bytes(raw)
            (folder/'edits.json').write_text(json.dumps(spec))
            (folder/'replacement.png').write_bytes(replacement)
            self.assertEqual(presentations.main(['edit',str(folder/'baseline.pptx'),str(folder/'edits.json'),
                '--image','replacement.png','--output-dir',str(folder/'candidate-v2')]),0)
            self.assertEqual(Presentation(str(folder/'candidate-v2'/'presentation.pptx')).slides[0].shapes[0].image.blob,replacement)
        spec['edits'][0]['old_sha256']='0'*64
        with self.assertRaisesRegex(ValueError,'Original embedded picture hash changed'):
            pptx_edit.edit(raw,spec,images={'replacement.png':replacement})

    def test_contain_portrait_replacement_preserves_pixels_and_removes_orphan(self):
        from PIL import Image
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as directory:
            old_file = Path(directory) / 'old.png'
            Image.new('RGB', (800, 400), 'green').save(old_file)
            new_file = Path(directory) / 'new.png'
            Image.new('RGB', (400, 800), 'red').save(new_file)
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_picture(str(old_file), 1_000_000, 2_000_000,
                                     width=4_000_000, height=2_000_000)
            source = io.BytesIO()
            deck.save(source)
            raw, picture = source.getvalue(), new_file.read_bytes()
        selected = next(s for s in pptx_edit.inspect(raw)['slides'][0]['shapes']
                        if s.get('image_sha256'))
        spec = {'version': 1, 'source_sha256': hashlib.sha256(raw).hexdigest(),
                'edits': [{'kind': 'replace_image', 'slide': 1,
                           'shape_id': selected['shape_id'],
                           'old_sha256': selected['image_sha256'],
                           'new_sha256': hashlib.sha256(picture).hexdigest(),
                           'path': 'new.png', 'fit': 'contain'}]}
        changed, receipt = pptx_edit.edit(raw, spec, images={'new.png': picture})
        placed = Presentation(io.BytesIO(changed)).slides[0].shapes[0]
        self.assertEqual(placed.image.blob, picture)
        self.assertEqual((placed.left, placed.top, placed.width, placed.height),
                         (2_500_000, 2_000_000, 1_000_000, 2_000_000))
        self.assertEqual(len(receipt['removed_parts']), 1)
        with zipfile.ZipFile(io.BytesIO(changed)) as package:
            self.assertNotIn(selected['image_sha256'],
                             [hashlib.sha256(package.read(name)).hexdigest()
                              for name in package.namelist() if name.startswith('ppt/media/')])

    def test_cli_creates_distinct_version_and_refuses_overwrite(self):
        with tempfile.TemporaryDirectory() as root:
            source=Path(root)/'baseline.pptx';source.write_bytes(self.raw)
            edits=Path(root)/'edits.json';edits.write_text(json.dumps(manifest(self.raw)))
            target=Path(root)/'candidate-v2'
            self.assertEqual(presentations.main(['edit',str(source),str(edits),'--output-dir',str(target)]),0)
            first=(target/'presentation.pptx').read_bytes()
            self.assertEqual(presentations.main(['edit',str(source),str(edits),'--output-dir',str(target)]),1)
            self.assertEqual((target/'presentation.pptx').read_bytes(),first)
            self.assertEqual(source.read_bytes(),self.raw)
            self.assertEqual(json.loads((target/'receipt.json').read_text())['validation']['baseline_sha256'],hashlib.sha256(self.raw).hexdigest())


class OperationTests(unittest.TestCase):
    setUp=mixed.Tests.setUp
    tearDown=mixed.Tests.tearDown

    def test_real_worker_records_exact_baseline_and_new_artifact(self):
        raw,_=pptx_document.create(fixture())
        source=self.root/'baseline.pptx';source.write_bytes(raw)
        edits=self.root/'edits.json';edits.write_text(json.dumps(manifest(raw)))
        baseline=self.rt.register(source,'Selected baseline',path='baseline.pptx')
        edit_id=self.rt.register(edits,'Reviewed edit manifest',path='edits.json')
        step={'id':'edit','role':'procedure','objective':'Revise exact PPTX','instruction':'Apply guarded edits.',
              'execution':{'capability':'pptx.edit','version':1,'parameters':{}},'inputs':[
                  {'artifact':baseline,'path':'baseline.pptx','purpose':'Baseline','authority':'Selected','media_type':pptx_edit.MIME},
                  {'artifact':edit_id,'path':'edits.json','purpose':'Edits','authority':'Reviewed','media_type':'application/json'}],
              'outputs':[{'path':'candidate-v2.pptx','purpose':'Revised candidate','media_type':pptx_edit.MIME}],
              'criteria':execution.REGISTRY['pptx.edit']['criteria'].copy(),'user_gate':'Select revised deck'}
        factory=RegisteredFactory();self.rt.factory=ExecutionFactory(self.agent,factory)
        self.rt.create(plan([step]));self.rt.tick('demo')
        deadline=time.monotonic()+15
        while time.monotonic()<deadline:
            state=self.rt.tick('demo')
            if state['status'] in ('awaiting_user','blocked'):break
            time.sleep(.05)
        for child in factory.children:child.wait(timeout=10)
        self.assertEqual(state['status'],'awaiting_user',state)
        receipt=json.loads(state['attempts'][0]['receipt'])
        self.assertEqual(receipt['operation']['validation']['baseline_sha256'],hashlib.sha256(raw).hexdigest())
        self.assertEqual(source.read_bytes(),raw)
        self.assertIsNotNone(self.rt.output('demo','edit','candidate-v2.pptx'))


class PlanningTests(unittest.TestCase):
    setUp=planning_fixtures.Tests.setUp
    tearDown=planning_fixtures.Tests.tearDown
    request=planning_fixtures.Tests.request
    action=planning_fixtures.Tests.action
    queue=planning_fixtures.Tests.queue
    row=planning_fixtures.Tests.row
    response=planning_fixtures.Tests.response

    def test_edit_plan_binds_selected_baseline_and_reviewed_manifest(self):
        raw,_=pptx_document.create(fixture())
        path=self.rt.root/'baseline.pptx';path.write_bytes(raw)
        aid=self.rt.register(path,'Exact baseline',path='baseline.pptx')
        row=dict(self.queue(action=self.action(step_capabilities=['pptx.edit']),text='Revise the existing deck'))
        payload=json.loads(row['context'])
        source=planning.source_entry(self.rt,aid,'baseline.pptx','Exact baseline','Selected prior version')
        payload.setdefault('available_sources',[]).append(source)
        row['context']=json.dumps(payload)
        value=self.response();value['input_basis']={'mode':'modify_existing','artifacts':[aid]}
        author,review=value['plan']['tasks']
        author.pop('user_gate',None)
        author['inputs']=[{k:source[k] for k in ('artifact','path','purpose','authority')}]
        author['outputs']=[{'path':'edits.json','purpose':'Exact edit manifest','media_type':'application/json'}]
        review['inputs']=[{'from_task':'produce','output':'edits.json','path':'edits.json','purpose':'Review edits','authority':'Candidate','media_type':'application/json'}]
        operation=contracts.assignment({'id':'edit','role':'procedure','objective':'Revise PPTX','instruction':'Apply exact edits.',
            'execution':{'capability':'pptx.edit','version':1,'parameters':{}},'inputs':[
                {'artifact':aid,'path':'baseline.pptx','purpose':'Baseline','authority':'Selected','media_type':pptx_edit.MIME},
                {'from_task':'produce','output':'edits.json','path':'edits.json','purpose':'Edits','authority':'Reviewed','media_type':'application/json'}],
            'outputs':[{'path':'candidate-v2.pptx','purpose':'Revised deck','media_type':pptx_edit.MIME}],
            'dependencies':['produce','review'],'criteria':execution.REGISTRY['pptx.edit']['criteria'].copy(),
            'user_gate':'Select revised deck'})
        final=dict(review,id='edit-review',review_of='edit',dependencies=['edit'],
                   criteria=operation['criteria'].copy(),inputs=[{'from_task':'edit','output':'candidate-v2.pptx',
                   'path':'candidate-v2.pptx','purpose':'Actual edited deck','authority':'Candidate','media_type':pptx_edit.MIME}])
        value['plan']['tasks'] += [operation,final]
        _,compiled=planning.validate_result(json.dumps(value),row)
        bound=next(t for t in compiled['tasks'] if t['id']=='edit')
        self.assertEqual({i['media_type'] for i in bound['inputs']},{'application/json',pptx_edit.MIME})
        self.assertIn('operation-support/pptx.edit/edit-schema.txt',row['context'])
        bad=json.loads(json.dumps(value));bad['plan']['tasks'][2]['inputs'][0]['artifact']='missing'
        with self.assertRaises(ValueError):planning.validate_result(json.dumps(bad),row)
        bad=json.loads(json.dumps(value));bad['plan']['tasks'][0]['inputs']=[]
        with self.assertRaisesRegex(ValueError,'baseline'):
            planning.validate_result(json.dumps(bad),row)


if __name__=='__main__':unittest.main()
