"""Small observed-source fixtures; no live browsers, providers or network."""
from contextlib import nullcontext
import copy
import hashlib
import io
import json
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch
from urllib.parse import urlencode
from orchestrator import browser_images as images, image_sources, contracts as c, executors, pptx_document
from orchestrator.browser_contract import validate_captures
from orchestrator.browser_worker import run
from orchestrator.gemini_worker import Files
from orchestrator.workers import atomic
from task_relay.general_browser import Session
from tests.test_general_browser import Driver,policy,browser_graph
from tests.test_gemini_executor import CONFIG,report
from tests.test_image_sources import photo,subject


def graph():
    p=browser_graph();t=p['tasks'][0]
    t['browser']=policy(image_sources=[dict(path='sources.json',subjects=[subject()])])
    t['outputs']=[dict(path='sources.json',purpose='Observed original image sources',media_type='application/json')]
    p['tasks'][1]['inputs']=[dict(from_task=t['id'],output='sources.json',path='sources.json',purpose='Review source manifest',authority='untrusted source',media_type='application/json')]
    return p


class ImageDriver(Driver):
    def open(self,tab,url):
        super().open(tab,url)
        self.pages[tab]['images']=[dict(src='https://images.example.org/oak.jpg',href='',alt='Example oak')]


class Client:
    def __init__(self,frozen,provider):self.frozen=frozen;self.provider=provider;self.calls=[]
    def request(self,endpoint,payload,**kwargs):
        self.calls.append(copy.deepcopy(payload));n=len(self.calls)
        if n==1:name,args='browser_open',{'url':'https://example.test/photo'}
        elif n==2:
            if self.provider=='gemini':previous=payload['contents'][-1]['parts'][0]['functionResponse']['response']
            else:previous=json.loads(payload['input'][-1]['output'] if self.provider=='openai' else payload['messages'][-1]['content'])
            name,args='browser_image_source',dict(tab=previous['tab'],observation=previous['observation'],ref='0',path='sources.json',subject='oak')
        else:name,args='finish',report(self.frozen)
        if self.provider=='gemini':return {'candidates':[{'finishReason':'STOP','content':{'role':'model','parts':[{'functionCall':{'name':name,'args':args}}]}}]}
        call=dict(name=name,arguments=json.dumps(args))
        if self.provider=='openai':return {'id':str(n),'status':'completed','output':[dict(type='function_call',call_id=str(n),**call)]}
        return {'choices':[{'finish_reason':'tool_calls','message':{'role':'assistant','content':None,'tool_calls':[dict(id=str(n),type='function',function=call)]}}]}


class Tests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name).resolve()
        self.ws=self.root/'ws';self.ws.mkdir();(self.ws/'.relay').mkdir()
        self.f=c.assignment(graph()['tasks'][0]);self.f.update(assignment_id='image-job',workspace=str(self.ws),backend=graph()['backend'])
        self.files=Files(self.f);self.driver=ImageDriver();self.db=sqlite3.connect(':memory:')
        self.session=Session(self.db,'image-job',self.f['browser'],self.driver,self.files)
        self.page=self.session.call('open','browser_open',{'url':'https://example.test/photo'})
    def tearDown(self):self.db.close();self.temp.cleanup()
    def args(self,**changes):return dict(tab=self.page['tab'],observation=self.page['observation'],ref='0',path='sources.json',subject='oak',**changes)
    def exported(self):
        self.session.call('export','browser_image_source',self.args())
        return json.loads((self.ws/'sources.json').read_text())

    def test_observed_export_cannot_invent_urls_or_write_reserved_manifest(self):
        with self.assertRaisesRegex(ValueError,'browser_image_source'):self.files.call('file_write',dict(path='sources.json',text='{}'))
        for change in (dict(subject='unapproved'),dict(ref='9'),dict(path='other.json')):
            with self.assertRaises(ValueError):self.session.call('bad','browser_image_source',{**self.args(),**change})
        result=self.session.call('export','browser_image_source',self.args())
        self.assertEqual(self.session.call('export','browser_image_source',self.args()),result)
        value=json.loads((self.ws/'sources.json').read_text())
        self.assertEqual(value['candidates'][0]['download_url'],'https://images.example.org/oak.jpg')
        self.assertEqual(value['candidates'][0]['observed_url'],'https://example.test/photo')
        self.assertEqual(len(value['candidates']),1);validate_captures(self.f,self.ws)
        with self.assertRaisesRegex(ValueError,'Duplicate'):self.session.call('duplicate','browser_image_source',self.args())
        self.assertFalse(self.session.journal.pending('fixture'))
        value['subjects'][0]['query']='changed';(self.ws/'sources.json').write_text(json.dumps(value))
        with self.assertRaisesRegex(ValueError,'assignment'):validate_captures(self.f,self.ws)

    def test_stale_page_and_thumbnail_are_rejected_before_export(self):
        self.driver.pages[self.page['tab']]['images'][0]['src']='https://images.example.org/changed.jpg'
        with self.assertRaisesRegex(ValueError,'changed'):self.session.call('stale','browser_image_source',self.args())
        for url in ('data:image/png;base64,abc','https://encrypted-tbn0.gstatic.com/a','https://tse1.mm.bing.net/a','http://images.example.org/a','https://127.0.0.1/a'):
            with self.subTest(url=url),self.assertRaises(ValueError):images.observed_image(dict(src=url,href=''), 'https://example.test/photo')
        self.assertFalse((self.ws/'sources.json').exists())

    def test_google_explicit_original_and_publisher_links_are_decoded(self):
        url='https://www.google.com/imgres?'+urlencode(dict(imgurl='https://images.example.org/oak.jpg',imgrefurl='https://example.test/photo'))
        value=images.observed_image(dict(src='data:thumbnail',href=url,alt='Oak'),'https://www.google.com/search')
        self.assertEqual(value['download_url'],'https://images.example.org/oak.jpg')
        self.assertEqual(value['source_url'],'https://example.test/photo')
        with self.assertRaises(ValueError):images.observed_image(dict(src='https://images.example.org/a.jpg',href='https://www.google.com/search'), 'https://www.google.com/search')

    def test_source_bundle_handoff_embeds_exact_photo_with_unknown_rights_explicit(self):
        value=self.exported();value['subjects'].append(dict(id='missing',label='Missing subject',query='Other subject'))
        raw,m=images.fetch_bundle(value,fetch=lambda url,*args:(photo(),'image/jpeg',url))
        content,credits,_=image_sources.unpack(raw)
        self.assertEqual(content['images/oak.jpg'],photo());self.assertFalse(m['complete'])
        self.assertIn('Unknown',credits['images/oak.jpg']['license'])
        spec=dict(version=1,title='Fixture',slides=[dict(elements=[dict(type='image',x=1,y=1,w=4,h=3,path='photos.zip/images/oak.jpg')])])
        result,_=pptx_document.create(spec,bundles={'photos.zip':raw})
        from pptx import Presentation
        deck=Presentation(io.BytesIO(result))
        self.assertEqual(deck.slides[0].shapes[0].image.blob,photo())
        self.assertIn('https://example.test/photo',deck.slides[0].notes_slide.notes_text_frame.text)
        self.assertIn('Unknown',deck.slides[0].notes_slide.notes_text_frame.text)

    def test_download_failure_is_a_gap_and_duplicate_candidates_are_rejected(self):
        value=self.exported();calls=[]
        def failed(url,*args):calls.append(url);raise OSError('offline')
        raw,m=images.fetch_bundle(value,fetch=failed)
        self.assertEqual(len(calls),1);self.assertEqual(m['coverage']['found'],0)
        self.assertIn('offline',m['subjects'][0]['reason']);self.assertEqual(image_sources.unpack(raw)[0],{})
        value['candidates'].append(value['candidates'][0])
        with self.assertRaisesRegex(ValueError,'Duplicate'):images.fetch_bundle(value,fetch=failed)
        self.assertEqual(len(calls),1)

    def test_all_provider_loops_export_identical_observed_source_contracts(self):
        for provider in ('gemini','openai','qwen'):
            with self.subTest(provider=provider):
                ws=self.root/provider;ws.mkdir();(ws/'.relay').mkdir();control=ws/'control';control.mkdir()
                f=copy.deepcopy(self.f);f['workspace']=str(ws);f['backend']={'type':provider+'-browser','model':'fixture-model'}
                config=CONFIG if provider=='gemini' else {'api_key':'fixture-only','model':'fixture-model'}
                atomic(control/'launch.json',{'credential_fingerprint':executors.fingerprint(config,f['backend'])})
                client=Client(f,provider)
                with sqlite3.connect(':memory:') as db:
                    result=run(f,control,db,self.root,client=client,config_reader=lambda:(config,f['backend']),driver_context=nullcontext(ImageDriver()))
                self.assertEqual(result['decision'],'delivered');self.assertEqual(len(client.calls),3)
                validate_captures(f,ws)
                self.assertEqual(json.loads((ws/'sources.json').read_text())['candidates'][0]['subject'],'oak')

    def test_registered_fetch_executes_manifest_and_preserves_zip_receipt(self):
        from orchestrator.runtime import Runtime
        from orchestrator.adapters import ExecutionFactory
        from tests.test_mixed_execution import LocalRegistered,Client as TextClient,operation
        from tests.test_orchestrator import FakeFactory,plan
        value=self.exported();source=self.ws/'sources.json'
        def fetch(url,*args):return photo(),'image/jpeg',url
        rt=Runtime(self.root/'runtime',ExecutionFactory(FakeFactory(),LocalRegistered(TextClient())))
        try:
            aid=rt.register(source,'Observed images',path='sources.json')
            inp=dict(artifact=aid,path='sources.json',purpose='Exact observed references',authority='Untrusted browser evidence',media_type='application/json')
            task=operation('images','images.fetch',[inp]);task['execution']['parameters']={}
            task['outputs']=[dict(path='photos.zip',purpose='Exact photos and attribution',media_type='application/zip')]
            with patch.object(image_sources,'download',side_effect=fetch):
                rt.create(plan([task]));rt.tick('demo');rt.tick('demo');rt.tick('demo')
            artifact=rt.output('demo','images','photos.zip')
            self.assertEqual(image_sources.unpack(Path(artifact['blob']).read_bytes())[0]['images/oak.jpg'],photo())
            receipt=json.loads(rt.status('demo')['attempts'][0]['receipt'])
            self.assertEqual(receipt['operation']['validation']['provider'],'browser-sources')
        finally:rt.close()

    def test_public_downloader_rejects_private_dns_and_cross_host_redirects(self):
        import socket,time
        from unittest.mock import MagicMock
        with patch('task_relay.orchestrator_web.socket.getaddrinfo',return_value=[(socket.AF_INET,socket.SOCK_STREAM,6,'',('127.0.0.1',443))]),patch('task_relay.orchestrator_web.PublicHTTPS') as connection:
            with self.assertRaisesRegex(ValueError,'private/reserved'):
                image_sources.download('https://images.example.org/a.jpg',100,time.monotonic()+5,{'remaining':8},{'images.example.org'})
            connection.assert_not_called()
        conn=MagicMock();conn.getresponse.return_value.status=302
        conn.getresponse.return_value.getheader.return_value='https://unobserved.example.org/redirect'
        with patch('task_relay.orchestrator_web.public_addresses',return_value=['8.8.8.8']),patch('task_relay.orchestrator_web.PublicHTTPS',return_value=conn):
            with self.assertRaisesRegex(ValueError,'permitted source hosts'):
                image_sources.download('https://images.example.org/a.jpg',100,time.monotonic()+5,{'remaining':8},{'images.example.org'})
        self.assertEqual(conn.request.call_count,1)
