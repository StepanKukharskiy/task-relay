"""CLI wire -> shared owner/compiler/worker/review flow, no model/network calls."""
import copy
import base64
import hashlib
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from orchestrator import contracts as c, executors, job_owner, worker_capabilities
from orchestrator.codex_transport import Client, POLICY
from orchestrator.gemini_worker import execute
from orchestrator.runtime import Runtime
from orchestrator.workers import atomic
from tests.test_orchestrator import FakeFactory

BACKEND=dict(type='codex-agent',model='fixture',reasoning='high',budget_semantics=POLICY)
CONFIG=dict(executable='fixture',executable_sha256='a'*64)


class Tests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name);self.factory=FakeFactory();self.rt=Runtime(self.root/'runtime',self.factory)
        self.addCleanup(self.rt.close)
        self.grant=dict(workers=[worker_capabilities.entry(BACKEND)],operations=[],deliverables={'report':'text/plain'},
            max_tasks=10,task_limits=dict(seconds=300,tool_calls=24,provider_requests=24,response_tokens=4096,output_bytes=200000),task_attempts=2,owner_turns=6)

    def create(self):
        job_owner.create(self.rt,'job','Complete and independently review the report.',BACKEND,[],self.grant)
        self.rt.tick('job')

    def invoke(self,task,steps):
        self.rt.tick('job');state=self.factory.sessions[self.rt.task('job',task)['latest']]
        frozen=state['frozen'];control=Path(state['session']['control']);control.mkdir(parents=True,exist_ok=True)
        selected=frozen['backend']
        atomic(control/'launch.json',dict(credential_fingerprint=executors.fingerprint(CONFIG,selected)))
        path=control/'script.json';atomic(path,steps)
        command=[sys.executable,'-B',str(Path(__file__).with_name('codex_protocol_fixture.py')),str(path)]
        with Client(CONFIG,selected,control,20,command=command) as client:
            try:result=execute(frozen,control,client,lambda:(CONFIG,selected))
            except BaseException:
                for p in sorted(control.glob('tool-*.json')):print(p.name,p.read_text()[:4000])
                raise
        state['status']=dict(status='finished',exit_code=0,external_outcome='no_pending_response',pending_requests=[])
        return result,control,frozen

    def report(self,task,decision='delivered'):
        f=self.factory.sessions[self.rt.task('job',task)['latest']]['frozen']
        result=dict(summary='Exact fixture inspected',decision=decision,instruction='',findings=[],
            checks={f'c{i}':dict(passed=True,evidence='Inspected exact candidate and source') for i in range(1,len(f['criteria'])+1)})
        from orchestrator.final_quality import DIMENSIONS
        for number,binding in f.get('final_quality',{}).get('checks',{}).items():
            refs=[dict(path=p,location='Exact fixture page 1') for p in dict.fromkeys([binding['candidate'],*binding['inspection_inputs']])]
            result['checks']['c'+number]['evidence']={dim:dict(verdict='satisfied',reason='Scripted synthetic judgment',inspected=refs) for dim in DIMENSIONS}
        return result

    def worker(self):
        return dict(id='write',objective='Write a substantive report',outputs=[dict(name='result',media_type='text/plain',purpose='Requested report')])

    def test_owner_rejects_repairs_delegates_and_completes_after_independent_review(self):
        from orchestrator.owner_compilation import label
        self.create();bad=self.worker();bad['inputs']=[dict(artifact='nonexistent')]
        _,control,_=self.invoke('owner',[
            ('job_delegate',dict(request_id='bad',tasks=[bad])),
            ('job_history',dict(context='invented',offset=0)),
            ('job_delegate',dict(request_id='correct',tasks=[self.worker()]))])
        receipts=[json.loads(p.read_text()) for p in sorted(control.glob('tool-*.json'))]
        self.assertIn('correction',receipts[0]['result'])
        self.assertEqual([item['name'] for item in receipts],['job_delegate','job_delegate'])
        self.assertEqual(json.loads((control/'api-02.validation.json').read_text())['code'],'tool_authority_denied')
        sent=[json.loads(line) for line in (control/'codex-sent.jsonl').read_text().splitlines()]
        denial=[item for item in sent if item.get('id')==102 and 'result' in item]
        self.assertEqual(len(denial),1)
        self.assertEqual(json.loads(denial[0]['result']['contentItems'][0]['text'])['effect'],'none')
        self.assertEqual(len(self.rt.status('job')['tasks']),3)
        self.rt.tick('job')
        self.invoke('write',[('file_write',dict(path='outputs/result.txt',text='A substantive fixture report.')),('finish',self.report('write'))])
        self.rt.tick('job');review=label('write','review')
        self.invoke(review,[('finish',self.report(review,'accept'))]);self.rt.tick('job')
        result=self.rt.output('job','write','outputs/result.txt')
        self.invoke('owner',[('job_complete',dict(deliverables_json=json.dumps(dict(report=result['id'])),summary='Independently reviewed result'))])
        self.rt.tick('job');self.assertEqual(self.rt.status('job')['status'],'completed')
        self.assertEqual(job_owner.original(self.rt,'job')['backend'],BACKEND)
        self.assertNotIn('fixture-secret', (control/'codex-events.jsonl').read_text())
        self.assertIsNone(json.loads((control/'codex-usage.json').read_text())['provider_requests'])

    def test_disconnected_cli_does_not_execute_or_replay_a_missing_reply(self):
        from task_relay.gemini import ProviderError
        self.create()
        with self.assertRaises(ProviderError):self.invoke('owner',[dict(disconnect=True)])
        state=self.factory.sessions[self.rt.task('job','owner')['latest']];control=Path(state['session']['control'])
        self.assertEqual(len(list(control.glob('api-*.request.json'))),1)
        self.assertFalse(list(control.glob('tool-*.json')))
        self.assertEqual(json.loads((control/'api-01.outcome.json').read_text())['outcome'],'uncertain')
        self.assertEqual(len(self.rt.status('job')['tasks']),1)

    def test_native_permission_request_is_denied_without_delegate(self):
        from task_relay.gemini import ProviderError
        self.create()
        with self.assertRaisesRegex(ProviderError,'unsupported tool or permission'):
            self.invoke('owner',[dict(events=[dict(id=500,method='item/permissions/requestApproval',params={})])])
        self.assertEqual(len(self.rt.status('job')['tasks']),1)

    def test_transcript_notifications_do_not_allow_native_tool_execution(self):
        from task_relay.gemini import ProviderError
        self.create()
        event=dict(method='item/started',params=dict(threadId='thread',turnId='turn',item=dict(type='commandExecution',id='ungranted',command='forbidden')))
        with self.assertRaisesRegex(ProviderError,'Unexpected native CLI operation: commandExecution'):
            self.invoke('owner',[dict(events=[event])])
        self.assertEqual(len(self.rt.status('job')['tasks']),1)

    def test_registered_build_inspect_correct_keeps_one_cli_thread_and_exact_versions(self):
        from orchestrator.document_authoring import PPTX
        from pptx import Presentation
        self.grant.update(operations=['pptx.create','pptx.render'],deliverables={'deck':PPTX,'pdf':'application/pdf'})
        self.grant['task_limits']['output_bytes']=20000000
        self.grant['workers'].append(worker_capabilities.entry({**BACKEND,'type':'codex-code','runtime':'a'*64}))
        self.create()
        self.invoke('owner',[('job_delegate',dict(request_id='deck',tasks=[dict(id='deck',objective='Build inspect and correct one slide',work=dict(mode='create',operation='pptx.create',parameters={},content_parameters=dict(max_records=1)))]))])
        self.rt.tick('job')
        page=dict(kind='page',slot=1,layout='text',content=dict(title='Wrong title',body='Synthetic evidence.'),notes='Fixture')
        corrected=copy.deepcopy(page);corrected['content']['title']='Corrected title'
        finish=('finish',self.report('deck'))
        _,control,frozen=self.invoke('deck',[
            ('operation_submit',dict(request_key='page',revision=0,records=[page])),
            ('operation_submit',dict(request_key='end',revision=1,records=[dict(kind='end',count=1)])),
            ('work_run',dict(step='build',request_key='v1')),('work_run',dict(step='inspect',request_key='inspect1')),
            ('image_view',dict(path='outputs/work-inspect/slide-previews.zip',members=['slide-001.png'])),
            ('operation_amend',dict(request_key='repair',revision=2,expected_record_sha256=c.digest(page),record=corrected)),finish,
            ('work_run',dict(step='build',request_key='v2')),('work_run',dict(step='inspect',request_key='inspect2')),finish,
            ('image_view',dict(path='outputs/work-inspect/slide-previews.zip',members=['slide-001.png'])),finish])
        self.assertIn('stale',json.loads((control/'tool-07-00.json').read_text())['result']['error'])
        self.assertIn('error',json.loads((control/'tool-10-00.json').read_text())['result'])
        presentation=Presentation(str(Path(frozen['workspace'])/'outputs/work-build/result.pptx'))
        self.assertIn('Corrected title','\n'.join(s.text for s in presentation.slides[0].shapes if s.has_text_frame))
        self.assertEqual(len(list((control/'work-operations').glob('*/outcome.json'))),4)
        usage=json.loads((control/'codex-usage.json').read_text())
        self.assertEqual(usage['decision_exchanges'],12)
        self.assertEqual(usage['thread'],'thread')
        sent=[json.loads(line) for line in (control/'codex-sent.jsonl').read_text().splitlines()]
        self.assertEqual(sum(e.get('method')=='thread/start' for e in sent),1)
        self.assertEqual(sum(e.get('method')=='turn/start' for e in sent),1)
        self.assertTrue(any(p.get('type')=='inputImage' for e in sent for p in e.get('result',{}).get('contentItems',[])))
        self.assertFalse(any('inputSchema' in p.get('text','') for e in sent for p in e.get('result',{}).get('contentItems',[])))
        self.rt.tick('job')
        from orchestrator.owner_compilation import label
        review=label('deck','review')
        f=self.factory.sessions[self.rt.task('job',review)['latest']]['frozen']
        images=next(i['path'] for i in f['inputs'] if i['path'].endswith('slide-previews.zip'))
        with patch('task_relay.code_runtime.available',return_value={'tools':{}}):
            self.invoke(review,[('image_view',dict(path=images,members=['slide-001.png'])),('finish',self.report(review,'accept'))])
        self.rt.tick('job')
        outputs={k:self.rt.output('job','deck',p)['id'] for k,p in [('deck','outputs/work-build/result.pptx'),('pdf','outputs/work-inspect/preview.pdf')]}
        self.invoke('owner',[('job_complete',dict(deliverables_json=json.dumps(outputs),summary='Reviewed exact revised files'))])
        self.rt.tick('job');self.assertEqual(self.rt.status('job')['status'],'completed')

    def test_cli_policy_is_explicit_and_old_cli_retains_original_shell_profile(self):
        with self.assertRaisesRegex(ValueError,'accounting'):executors.validate({k:v for k,v in BACKEND.items() if k!='budget_semantics'})
        self.assertEqual(executors.validate(BACKEND),['files'])
        self.assertEqual(executors.validate(dict(type='codex-cli',model='fixture',reasoning='high')),['files','shell'])

    def test_api_key_login_stops_before_any_model_turn(self):
        control=self.root/'wrong-auth';control.mkdir()
        path=control/'script.json';atomic(path,[])
        command=[sys.executable,'-B',str(Path(__file__).with_name('codex_protocol_fixture.py')),str(path),'apiKey']
        with self.assertRaisesRegex(ValueError,'API-key billing is not authorized'):
            with Client(CONFIG,BACKEND,control,10,command=command):pass
        sent=(control/'codex-sent.jsonl').read_text()
        self.assertNotIn('turn/start',sent)
        self.assertNotIn('fixture-private-account',(control/'codex-events.jsonl').read_text())

    def test_queued_tool_call_cannot_observe_pixels_delivered_after_its_generation(self):
        control=self.root/'pixels';control.mkdir()
        def call(ident):return dict(id=ident,method='item/tool/call',params=dict(threadId='thread',turnId='turn',callId=str(ident),namespace=None,tool='finish',arguments={}))
        # A parallel call and a foreign thread's usage cannot prove a new
        # generation saw the image result. A later local generation can.
        foreign=dict(method='thread/tokenUsage/updated',params=dict(threadId='other',turnId='turn',tokenUsage=dict(total=dict(totalTokens=999))))
        path=control/'script.json';atomic(path,[('image_view',{}),dict(events=[foreign,call(102)]),('finish',{})])
        command=[sys.executable,'-B',str(Path(__file__).with_name('codex_protocol_fixture.py')),str(path)]
        definitions=[dict(name=n,description=n,parameters=dict(type='object',properties={})) for n in ['image_view','finish']]
        payload=dict(model='fixture',tools=[dict(type='function',function=t) for t in definitions],messages=[dict(role='system',content='Fixture'),dict(role='user',content='Inspect')])
        raw=b'pixel-transport-fixture';url='data:image/png;base64,'+base64.b64encode(raw).decode()
        with Client(CONFIG,BACKEND,control,10,command=command) as client:
            client.set_tools(definitions)
            first=client.request('chat/completions',payload)
            payload['messages'].extend([dict(role='tool',tool_call_id=first['id'],content='Image result'),dict(role='user',content=[dict(type='image_url',image_url=dict(url=url))])])
            queued=client.request('chat/completions',payload)
            self.assertEqual(queued['relay_transport']['observed_pixel_sha256'],[])
            payload['messages'].append(dict(role='tool',tool_call_id=queued['id'],content='Inspection incomplete'))
            observed=client.request('chat/completions',payload)
            self.assertEqual(observed['relay_transport']['observed_pixel_sha256'],[hashlib.sha256(raw).hexdigest()])

    def test_other_provider_operations_are_rejected_before_owner_dispatch(self):
        self.grant['operations']=['web.sources']
        with self.assertRaisesRegex(ValueError,'other model providers'):self.create()
        self.assertEqual(self.factory.calls,[])

    def test_codex_owner_uses_explicit_relay_search_provider_and_completes(self):
        from orchestrator import step_runner
        from orchestrator.owner_intents import catalog
        from orchestrator.owner_compilation import label
        from task_relay import orchestrator_web
        self.grant['operations']=['web.sources']
        search_backend=dict(type='gemini-agent',model='fixture-search')
        self.grant['operation_backends']={'web.sources':search_backend}
        configured=patch('task_relay.gemini.read_config',return_value=dict(api_key='fixture',models=dict(text='fixture-search')))
        configured.start();self.addCleanup(configured.stop)
        source=self.root.resolve()/'quoted-context.txt';source.write_text('"\\'*30000)
        ref=dict(artifact=self.rt.register(source,'Quoted context fixture'),path='quoted-context.txt',purpose='Exact input',authority='Synthetic evidence',media_type='text/plain')
        job_owner.create(self.rt,'job','Complete and independently review the report.',BACKEND,[ref],self.grant)
        self.rt.tick('job')
        captured=copy.deepcopy(job_owner.original(self.rt,'job')['job_owner'])
        self.assertNotIn('model',catalog(captured,['web.sources'])['operations']['web.sources']['parameters']['properties'])
        reads=[('job_read',dict(artifact=ref['artifact'],offset=n*12000,limit=12000)) for n in range(4)]
        _,owner_control,_=self.invoke('owner',reads+[('job_delegate',dict(request_id='research',tasks=[dict(id='search',objective='Retrieve evidence',
            operation='web.sources',parameters=dict(queries=['fixture one','fixture two'],domains=['example.org'],model='invented-by-owner'))]))])
        self.assertIn('Saved tool decisions and observations',(owner_control/'owner-session.json').read_text())
        self.rt.tick('job');state=self.factory.sessions[self.rt.task('job','search')['latest']]
        frozen=state['frozen'];control=Path(state['session']['control']);control.mkdir(parents=True,exist_ok=True)
        self.assertEqual(frozen['backend'],search_backend)
        self.assertEqual(frozen['execution']['parameters']['model'],'fixture-search')
        requests=[]
        def search(endpoint,payload,timeout=None):
            requests.append(copy.deepcopy(payload))
            query=payload['contents'][0]['parts'][0]['text'].split('Query: ',1)[1]
            return dict(candidates=[dict(content=dict(parts=[dict(text='https://example.org/not-grounded')]),
                groundingMetadata=dict(webSearchQueries=[query],groundingChunks=[dict(web=dict(uri='https://example.org/item',title='Observed item'))]))],
                usageMetadata=dict(promptTokenCount=10,candidatesTokenCount=1,totalTokenCount=11))
        html=b'<html><title>Observed item</title><p>Literal evidence</p></html>'
        with patch('task_relay.gemini.Client') as provider,patch.object(orchestrator_web,'download',return_value=('https://example.org/item',html,'text/html','utf-8')) as download:
            provider.return_value.request.side_effect=search
            step_runner.execute(frozen,control,config_reader=lambda:dict(api_key='fixture',models=dict(text='fixture-search')))
        self.assertEqual(len(requests),2);self.assertEqual(download.call_count,1)
        self.assertTrue(all('URL discovery only' in r['contents'][0]['parts'][0]['text'] for r in requests))
        saved=[json.loads(p.read_text()) for p in sorted(control.glob('query-*/search-*.json'))]
        self.assertEqual([r['request'] for r in saved],requests)
        pack=json.loads((Path(frozen['workspace'])/'delivery/source-pack.json').read_text())
        self.assertNotIn('not-grounded',json.dumps(pack))
        self.assertEqual(pack['entries'][1]['pages'][0]['reused_from'],'entries[0].pages[0]')
        state['status']=dict(status='finished',exit_code=0,external_outcome='response_received',pending_requests=[])
        self.rt.tick('job')
        evidence=self.rt.output('job','search','delivery/source-pack.json')
        # Resume the real controller and owner loop with the source artifact.
        self.rt.close();self.rt=Runtime(self.root/'runtime',self.factory)
        self.addCleanup(self.rt.close)
        task=self.worker();task['inputs']=[dict(artifact=evidence['id'])]
        self.invoke('owner',[('job_delegate',dict(request_id='write',tasks=[task]))])
        self.rt.tick('job')
        self.invoke('write',[('file_write',dict(path='outputs/result.txt',text='Evidence-based fixture report.')),('finish',self.report('write'))])
        self.rt.tick('job');review=label('write','review')
        self.invoke(review,[('finish',self.report(review,'accept'))]);self.rt.tick('job')
        artifact=self.rt.output('job','write','outputs/result.txt')
        self.invoke('owner',[('job_complete',dict(deliverables_json=json.dumps(dict(report=artifact['id'])),summary='Reviewed report'))])
        self.rt.tick('job')
        self.assertEqual(self.rt.status('job')['status'],'completed')
        for taskid in ('owner','write',review):
            f=self.factory.sessions[self.rt.task('job',taskid)['latest']]['frozen']
            self.assertEqual(f['backend'],BACKEND)
        self.assertEqual(job_owner.original(self.rt,'job')['job_owner'],captured)

    def test_tool_provider_binding_requires_explicit_compatible_authority(self):
        from orchestrator.model_binding import validate_grant
        for configured in ({'web.sources':BACKEND},{'ungranted':dict(type='gemini-agent',model='fixture')}):
            with self.subTest(configured=configured),self.assertRaises(ValueError):
                validate_grant(dict(operations=['web.sources'],operation_backends=configured))
        # The task model cannot add a tool grant in its delegation proposal.
        self.create()
        self.invoke('owner',[('job_delegate',dict(request_id='bad',tasks=[dict(self.worker(),operation_backends={'web.sources':dict(type='gemini-agent',model='fixture')})])),
                             ('job_delegate',dict(request_id='correct',tasks=[self.worker()]))])
        self.assertNotIn('operation_backends',job_owner.original(self.rt,'job')['job_owner'])

    def test_ui_distinguishes_cli_exchange_limit_from_api_token_limit(self):
        from task_relay.production_activity import snapshot
        value=snapshot(self.root,None,dict(objective='Fixture',limits=self.grant['task_limits']),BACKEND)
        self.assertEqual(value['request_unit'],'CLI decision exchanges')
        self.assertIsNone(value['response_limit'])

    def test_codex_search_uses_recorded_sources_and_common_page_reader_without_gemini(self):
        selector=patch('orchestrator.codex_sources.configuration',side_effect=lambda b:(CONFIG,b))
        selector.start();self.addCleanup(selector.stop)
        from orchestrator import codex_sources,step_runner,execution
        from task_relay import orchestrator_web
        self.grant['operations']=['web.sources.codex']
        self.create()
        self.invoke('owner',[('job_delegate',dict(request_id='search',tasks=[dict(id='search',objective='Research the exact fixture',
            operation='web.sources.codex',parameters=dict(queries=['fixture part'],domains=['example.org'],model='fixture',reasoning='high'))]))])
        self.rt.tick('job');state=self.factory.sessions[self.rt.task('job','search')['latest']]
        f=state['frozen'];control=Path(state['session']['control']);control.mkdir(parents=True,exist_ok=True)
        events=[dict(method='item/completed',params=dict(threadId='thread',turnId='turn',item=dict(id='search',type='webSearch',query='fixture part',results=[dict(url='https://example.org/item',title='Fixture part')]))),
                dict(method='turn/completed',params=dict(threadId='thread',turn=dict(id='turn',status='completed')))]
        path=control/'search-script.json';atomic(path,[dict(events=events)])
        command=[sys.executable,'-B',str(Path(__file__).with_name('codex_protocol_fixture.py')),str(path)]
        def transport(config,backend,folder,seconds):return Client(CONFIG,backend,folder,seconds,command=command)
        html=b'<html><title>Fixture part</title><p>Exact observed fixture evidence.</p></html>'
        # FakeFactory does not run native availability hooks; freeze the same
        # executable identity that the real registered factory would bind.
        f['codex_executor']=CONFIG
        with patch.object(codex_sources,'configuration',side_effect=lambda b:(CONFIG,b)),patch.object(codex_sources,'Client',side_effect=transport), \
             patch.object(orchestrator_web,'download',return_value=('https://example.org/item',html,'text/html','utf-8')), \
             patch('task_relay.gemini.Client',side_effect=AssertionError('Hidden Gemini call')):
            step_runner.execute(f,control,config_reader=lambda:(_ for _ in ()).throw(AssertionError('Hidden Gemini config')))
        raw=Path(f['workspace'],'delivery/source-pack.json').read_text();pack=json.loads(raw)
        self.assertEqual(pack['entries'][0]['grounding'],[dict(url='https://example.org/item',title='Fixture part')])
        self.assertIn('Exact observed fixture evidence.',pack['entries'][0]['pages'][0]['text'])
        self.assertEqual(execution.adapter('web.sources.prepared.codex').handoffs,execution.adapter('web.sources.prepared').handoffs)

    def test_search_summary_urls_do_not_become_recorded_source_evidence(self):
        control=self.root/'search';control.mkdir()
        events=[dict(method='item/completed',params=dict(threadId='thread',turnId='turn',item=dict(type='agentMessage',text='https://example.org/invented'))),
            dict(method='turn/completed',params=dict(threadId='thread',turn=dict(id='turn',status='completed')))]
        path=control/'script.json';atomic(path,[dict(events=events)])
        command=[sys.executable,'-B',str(Path(__file__).with_name('codex_protocol_fixture.py')),str(path)]
        with Client(CONFIG,BACKEND,control,10,command=command) as client:
            with self.assertRaisesRegex(ValueError,'no recorded source URLs'):client.search('fixture')

    def test_completed_empty_search_preserves_missing_outcome_without_retry(self):
        from orchestrator import codex_sources,web_sources
        from orchestrator.codex_transport import SearchEvidenceError
        config=dict(models=dict(text='fixture'),backend=BACKEND,cli_executor=CONFIG)
        with patch.object(codex_sources,'configuration',return_value=(CONFIG,BACKEND)),patch.object(codex_sources,'Client') as factory:
            factory.return_value.__enter__.return_value.search.side_effect=SearchEvidenceError('No recorded source URLs')
            raw,_=web_sources.collect(['fixture'],[],'fixture',self.root/'missing',config,session_factory=codex_sources.Session)
            self.assertEqual(factory.return_value.__enter__.return_value.search.call_count,1)
        pack=json.loads(raw)
        self.assertEqual(pack['entries'][0]['status'],'unresolved')
        self.assertEqual(pack['entries'][0]['pages'],[])
        receipt=next((self.root/'missing').rglob('search-0.json'))
        self.assertEqual(json.loads(receipt.read_text())['status'],'responded')

    @staticmethod
    def rejection_events(status=400):
        error=dict(message=json.dumps(dict(type='error',status=status,error=dict(type='invalid_request_error',
            message="The 'gpt-5' model is not supported when using Codex with a ChatGPT account."))),codexErrorInfo='other',additionalDetails=None,misalignment=None)
        return [dict(method='error',params=dict(threadId='thread',turnId='turn',willRetry=False,error=error)),
            dict(method='turn/completed',params=dict(threadId='thread',turn=dict(id='turn',status='failed',error=error,items=[])))]

    def test_model_binding_omits_model_choices_and_preserves_raw_intent(self):
        from orchestrator.owner_intents import catalog
        from orchestrator.model_binding import bind,verify
        self.grant['operations']=['web.sources.codex','web.sources.prepared.codex']
        self.create()
        for cap in self.grant['operations']:
            self.assertNotIn('model',catalog(self.grant,[cap])['operations'][cap]['parameters']['properties'])
            self.assertNotIn('reasoning',catalog(self.grant,[cap])['operations'][cap]['parameters']['properties'])
            self.assertEqual(bind(cap,dict(model='gpt-5',reasoning='low'),BACKEND),dict(model='fixture',reasoning='high'))
            with self.assertRaisesRegex(ValueError,'frozen backend'):verify(cap,dict(model='gpt-5',reasoning='high'),BACKEND)
            with self.assertRaisesRegex(ValueError,'authorized provider'):bind(cap,{},dict(type='gemini-agent',model='other'))
        params=dict(queries=['fixture'],domains=[],model='gpt-5',reasoning='low')
        proposal=dict(request_id='bound',tasks=[dict(id='search',objective='Search fixture',operation='web.sources.codex',parameters=params)])
        _,control,_=self.invoke('owner',[('job_delegate',proposal)])
        spec=self.rt.spec(self.rt.task('job','search'))
        self.assertEqual(spec['execution']['parameters'],dict(queries=['fixture'],domains=[],model='fixture',reasoning='high'))
        self.assertEqual(json.loads(next(control.glob('tool-*.json')).read_text())['arguments'],proposal)

    def test_missing_model_choices_are_filled_before_successful_delegation(self):
        self.grant['operations']=['web.sources.codex'];self.create()
        self.invoke('owner',[('job_delegate',dict(request_id='no-model',tasks=[dict(id='search',objective='Research fixture',
            operation='web.sources.codex',parameters=dict(queries=['fixture'],domains=[]))]))])
        spec=self.rt.spec(self.rt.task('job','search'))
        self.assertEqual(spec['execution']['parameters']['model'],BACKEND['model'])
        self.assertEqual(spec['execution']['parameters']['reasoning'],BACKEND['reasoning'])

    def rejection_client(self,control,events):
        control.mkdir(parents=True,exist_ok=True);path=control/'script.json';atomic(path,[dict(events=events)])
        return Client(CONFIG,BACKEND,control,10,command=[sys.executable,'-B',str(Path(__file__).with_name('codex_protocol_fixture.py')),str(path)])

    @staticmethod
    def overloaded(turn='turn', retries=0):
        error=dict(message='Selected model is at capacity.',codexErrorInfo='serverOverloaded')
        progress=[dict(method='error',params=dict(threadId='thread',turnId=turn,
            error=dict(message='Reconnecting',codexErrorInfo=dict(responseStreamDisconnected=dict(httpStatusCode=503))),willRetry=True)) for _ in range(retries)]
        return progress+[dict(method='error',params=dict(threadId='thread',turnId=turn,error=error,willRetry=False)),
                dict(method='turn/completed',params=dict(threadId='thread',turn=dict(id=turn,status='failed',error=error,items=[],itemsView='notLoaded')))]

    def test_overload_continues_same_thread_after_committed_tool_without_replay(self):
        from orchestrator.codex_rejection import generation_unavailable
        self.create()
        self.invoke('owner',[('job_delegate',dict(request_id='write',tasks=[self.worker()]))])
        self.rt.tick('job')
        _,control,_=self.invoke('write',[
            ('file_write',dict(path='outputs/result.txt',text='Committed report before overload.')),
            dict(events=self.overloaded(retries=4)),('finish',self.report('write'))])
        sent=[json.loads(line) for line in (control/'codex-sent.jsonl').read_text().splitlines()]
        self.assertEqual(len([x for x in sent if x.get('method')=='thread/start']),1)
        turns=[x for x in sent if x.get('method')=='turn/start'];self.assertEqual(len(turns),2)
        self.assertTrue(all(x['params']['threadId']=='thread' and x['params']['model']=='fixture' for x in turns))
        receipts=[json.loads(p.read_text()) for p in control.glob('tool-*.json')]
        self.assertEqual(sum(r.get('name')=='file_write' for r in receipts),1)
        recovery=json.loads((control/'api-02.recovery.json').read_text())
        self.assertTrue(recovery['continued']);self.assertIsNone(recovery['http_status'])
        self.assertEqual(recovery['provider_error']['acknowledged_calls'],['call-1'])
        self.assertEqual(recovery['provider_error']['internal_retry_notifications'],4)
        self.assertIsNone(recovery['provider_error']['usage'])
        self.assertEqual(len(list(control.glob('api-*.request.json'))),3)
        # Incomplete tool delivery, server-side search or a foreign native item
        # cannot establish safe continuation, even with an overload terminal.
        raw={n:(control/n).read_text() for n in ('codex-sent.jsonl','codex-events.jsonl')}
        for case in ('missing_response','native','search','mismatched_terminal','missing_terminal','ambiguous_failure'):
            target=self.root/case;target.mkdir()
            s=[json.loads(line) for line in raw['codex-sent.jsonl'].splitlines()]
            e=[json.loads(line) for line in raw['codex-events.jsonl'].splitlines()]
            if case=='missing_response':s=[x for x in s if x.get('id')!=101]
            if case=='search':next(x for x in s if x.get('method')=='thread/start')['params']['config']['web_search']='live'
            if case=='native':e.append(dict(method='item/completed',params=dict(threadId='thread',turnId='turn',item=dict(type='webSearch'))))
            if case=='mismatched_terminal':next(x for x in e if x.get('method')=='turn/completed')['params']['turn']['error']={}
            if case=='missing_terminal':e=[x for x in e if x.get('method')!='turn/completed']
            if case=='ambiguous_failure':next(x for x in e if x.get('method')=='error')['params']['willRetry']=False
            for name,values in [('codex-sent.jsonl',s),('codex-events.jsonl',e)]:
                (target/name).write_text(''.join(json.dumps(x)+'\n' for x in values))
            self.assertIsNone(generation_unavailable(target,'thread','turn'),case)

    def test_search_overload_before_any_tool_continues_once_on_same_thread(self):
        control=self.root/'search-continuation';control.mkdir()
        success=[dict(method='item/completed',params=dict(threadId='thread',turnId='turn-2',item=dict(type='webSearch',query='fixture',results=[dict(url='https://example.org/item',title='Synthetic source')]))),
                 dict(method='turn/completed',params=dict(threadId='thread',turn=dict(id='turn-2',status='completed')))]
        path=control/'script.json';atomic(path,[dict(events=self.overloaded(retries=2)),dict(events=success)])
        command=[sys.executable,'-B',str(Path(__file__).with_name('codex_protocol_fixture.py')),str(path)]
        with Client(CONFIG,BACKEND,control,10,command=command) as client:result=client.search('fixture')
        self.assertEqual(result['sources'][0]['url'],'https://example.org/item')
        sent=[json.loads(line) for line in (control/'codex-sent.jsonl').read_text().splitlines()]
        self.assertEqual(len([s for s in sent if s.get('method')=='thread/start']),1)
        self.assertEqual(len([s for s in sent if s.get('method')=='turn/start']),2)
        saved=json.loads((control/'codex-search-recovery.json').read_text())
        self.assertEqual(saved['proof']['scope'],'before_search');self.assertIsNone(saved['provider_requests'])
        self.assertIsNone(saved['proof']['usage'])

    def test_completed_search_survives_later_generation_failure_without_replay(self):
        from task_relay.gemini import ProviderError
        for pending in (False,True):
            control=self.root/('completed-read-'+str(pending));control.mkdir()
            events=[dict(method='item/completed',params=dict(threadId='thread',turnId='turn',item=dict(
                id='read-1',type='webSearch',query='fixture',results=[dict(url='https://example.org/item',title='Recorded source')])))]
            if pending:events.append(dict(method='item/started',params=dict(threadId='thread',turnId='turn',item=dict(id='read-2',type='webSearch'))))
            events+=self.overloaded()
            path=control/'script.json';atomic(path,[dict(events=events)])
            command=[sys.executable,'-B',str(Path(__file__).with_name('codex_protocol_fixture.py')),str(path)]
            with Client(CONFIG,BACKEND,control,10,command=command) as client:
                if pending:
                    with self.assertRaises(ProviderError) as error:client.search('fixture')
                    self.assertTrue(error.exception.uncertain)
                else:
                    result=client.search('fixture')
                    self.assertEqual(result['sources'][0]['url'],'https://example.org/item')
                    self.assertEqual(result['observation_outcome']['generation_status'],'failed')
                    self.assertFalse(result['observation_outcome']['replayed'])
                    self.assertFalse(result['observation_outcome']['acceptance'])
                    self.assertEqual(result['summary'],'')
            sent=[json.loads(line) for line in (control/'codex-sent.jsonl').read_text().splitlines()]
            self.assertEqual(len([s for s in sent if s.get('method')=='turn/start']),1)
            self.assertEqual((control/'codex-search-observation.json').exists(),not pending)

    def test_search_overload_after_any_tool_stays_uncertain_and_repeat_is_bounded(self):
        from task_relay.gemini import ProviderError
        for kind in ('webSearch','dynamicToolCall','commandExecution','repeat'):
            with self.subTest(kind=kind):
                control=self.root/kind;control.mkdir();path=control/'script.json'
                events=self.overloaded()
                if kind!='repeat':events.insert(0,dict(method='item/started',params=dict(threadId='thread',turnId='turn',item=dict(type=kind,id='started'))))
                steps=[dict(events=events)]+([dict(events=self.overloaded('turn-2'))] if kind=='repeat' else [])
                atomic(path,steps);command=[sys.executable,'-B',str(Path(__file__).with_name('codex_protocol_fixture.py')),str(path)]
                with Client(CONFIG,BACKEND,control,10,command=command) as client:
                    with self.assertRaises((ProviderError,ValueError)) as error:client.search('fixture')
                if isinstance(error.exception,ProviderError):self.assertEqual(error.exception.uncertain,kind!='repeat')
                else:self.assertIn('ungranted',str(error.exception))
                self.assertEqual((control/'codex-search-recovery.json').exists(),kind=='repeat')
                sent=[json.loads(line) for line in (control/'codex-sent.jsonl').read_text().splitlines()]
                self.assertEqual(len([s for s in sent if s.get('method')=='turn/start']),2 if kind=='repeat' else 1)

    def test_repeated_overload_stops_after_one_bounded_continuation(self):
        from task_relay.gemini import ProviderError
        self.create()
        with self.assertRaises(ProviderError):
            self.invoke('owner',[dict(events=self.overloaded()),dict(events=self.overloaded('turn-2'))])
        control=Path(self.factory.sessions[self.rt.task('job','owner')['latest']]['session']['control'])
        self.assertEqual(len(list(control.glob('api-*.request.json'))),2)
        self.assertTrue(json.loads((control/'api-01.recovery.json').read_text())['continued'])
        self.assertFalse(json.loads((control/'api-02.recovery.json').read_text())['continued'])
        self.assertFalse(list(control.glob('tool-*.json')))

    def test_worker_terminal_rejection_is_confirmed_without_replay_or_zero_usage(self):
        from task_relay.gemini import ProviderError
        self.create()
        with self.assertRaises(ProviderError) as error:self.invoke('owner',[dict(events=self.rejection_events())])
        self.assertFalse(error.exception.uncertain);self.assertEqual(error.exception.status,400)
        state=self.factory.sessions[self.rt.task('job','owner')['latest']];control=Path(state['session']['control'])
        outcome=json.loads(next(control.glob('api-*.outcome.json')).read_text())
        self.assertEqual(outcome['outcome'],'rejected')
        self.assertFalse(list(control.glob('tool-*.json')))
        self.assertIsNone(json.loads((control/'codex-usage.json').read_text())['usage'])

    def test_registered_search_rejection_reaches_controller_and_preserves_journals(self):
        selector=patch('orchestrator.codex_sources.configuration',side_effect=lambda b:(CONFIG,b))
        selector.start();self.addCleanup(selector.stop)
        from orchestrator import codex_sources,step_runner,execution
        from task_relay.gemini import ProviderError
        self.grant['operations']=['web.sources.codex'];self.create()
        self.invoke('owner',[('job_delegate',dict(request_id='search',tasks=[dict(id='search',objective='Search fixture',
            operation='web.sources.codex',parameters=dict(queries=['fixture'],domains=[]))]))])
        self.rt.tick('job');state=self.factory.sessions[self.rt.task('job','search')['latest']]
        f=state['frozen'];control=Path(state['session']['control']);control.mkdir(parents=True,exist_ok=True);f['codex_executor']=CONFIG
        with patch.object(codex_sources,'configuration',return_value=(CONFIG,BACKEND)),patch.object(codex_sources,'Client',side_effect=lambda a,b,folder,s:self.rejection_client(folder,self.rejection_events())):
            with self.assertRaises(ProviderError) as error:step_runner.execute(f,control)
        self.assertFalse(error.exception.uncertain)
        record=next(control.glob('query-*/search-*.json'));self.assertEqual(json.loads(record.read_text())['status'],'rejected')
        details=json.loads((control/'operation.json').read_text());self.assertEqual(details['outcome'],'rejected')
        # Legacy saved receipts are reconciled read-only by the same inspector.
        legacy=json.loads(record.read_text());legacy['status']='uncertain';atomic(record,legacy)
        details['outcome']='uncertain';atomic(control/'operation.json',details)
        before={p:hashlib.sha256(p.read_bytes()).hexdigest() for p in control.rglob('*') if p.is_file()}
        result=dict(status='finished',exit_code=1,operation=details,execution=f['execution'])
        execution.adapter('web.sources.codex').inspect(control,result,details)
        self.assertIn('not supported',result['reason'])
        self.assertEqual(result['external_outcome'],'rejected');self.assertEqual(details['outcome'],'rejected')
        self.assertTrue(result['rejection_reconciliation']);self.assertEqual(before,{p:hashlib.sha256(p.read_bytes()).hexdigest() for p in before})
        state['status']=result;self.rt.tick('job',dispatch=False)
        self.assertEqual(self.rt.task('job','search')['status'],'blocked')
        self.assertNotEqual(self.rt.status('job')['status'],'uncertain')
        self.assertEqual(len(list(control.glob('query-*/search-*.json'))),1)
        prior=copy.deepcopy(self.rt.spec(self.rt.task('job','search')))
        self.invoke('owner',[('job_delegate',dict(request_id='repair',tasks=[dict(id='replacement',objective='Repair rejected search',
            operation='web.sources.codex',parameters=dict(queries=['fixture'],domains=[],model='gpt-5',reasoning='low'))]))])
        self.assertEqual(self.rt.spec(self.rt.task('job','replacement'))['execution']['parameters']['model'],BACKEND['model'])
        self.assertEqual(self.rt.spec(self.rt.task('job','search')),prior)
        self.assertEqual(self.rt.task('job','owner')['attempts'],2)
        self.assertEqual(job_owner.original(self.rt,'job')['job_owner']['owner_turns'],6)

    def test_terminal_error_after_tool_or_without_matching_proof_stays_uncertain(self):
        from task_relay.gemini import ProviderError
        from orchestrator.codex_rejection import proof
        cases=[]
        cases.append(self.rejection_events(500))
        e=self.rejection_events();e[0]['params']['willRetry']=True;cases.append(e)
        e=self.rejection_events();e[0]['params']['turnId']='foreign';cases.append(e)
        e=self.rejection_events();e.insert(0,dict(method='item/completed',params=dict(threadId='thread',turnId='turn',item=dict(type='webSearch',results=[]))));cases.append(e)
        for n,events in enumerate(cases):
            with self.subTest(case=n),self.rejection_client(self.root/str(n),events) as client:
                with self.assertRaises(ProviderError) as error:client.search('fixture')
                self.assertTrue(error.exception.uncertain)
                self.assertIsNone(proof(self.root/str(n)))



if __name__=='__main__':unittest.main()
