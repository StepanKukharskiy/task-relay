"""Actual owner/fragment/native-render boundaries; scripted models, synthetic data."""
import copy
import io
import json
import unittest
from pathlib import Path
from unittest.mock import patch

from orchestrator import contracts as c,document_authoring as authoring,layout_contract as layout
from orchestrator import job_owner,operation_records,worker_capabilities,executors
from orchestrator.gemini_worker import execute
from orchestrator.workers import atomic
from tests import test_job_owner_loop as owner_fixture
from tests.test_operation_session import Script,CONFIG


class Tests(unittest.TestCase):
    setUp=owner_fixture.Tests.setUp
    create=owner_fixture.Tests.create
    invoke=owner_fixture.Tests.invoke
    batch=owner_fixture.Tests.batch

    def prepare(self):
        self.grant['operations']=['pptx.create','pptx.render']
        self.grant['deliverables']={'deck':authoring.PPTX}
        self.create();self.rt.tick('job')

    def content(self,inputs=None):
        return dict(id='content',objective='Select content and layouts',instruction='Use the saved source; one layout group per call.',
            inputs=inputs or [],outputs=[dict(path='slides.json',purpose='Prepared layout content',media_type='application/json')],
            dependencies=[],worker=dict(requires=['files.text']),
            records=dict(family=layout.IDENTITY.copy(),inputs={},parameters=dict(renderer='pptx.create',title='Evidence',style='clean_minimal_v1',max_records=3)))

    def renderer(self,inputs=None):
        return dict(id='render',objective='Render prepared content',instruction='Use registered geometry and file creation',
            execution=dict(capability='pptx.create',version=1,parameters={}),
            inputs=inputs or [dict(from_task='content',output='slides.json',path='slides.json',purpose='Layout data',authority='Code-owned projection',media_type='application/json')],
            outputs=[dict(path='deck.pptx',purpose='Editable deck',media_type=authoring.PPTX)],dependencies=['content'])

    def session(self):
        aid=self.rt.task('job','owner')['latest'];s=self.factory.sessions[aid]
        control=Path(s['session']['control']);control.mkdir(parents=True,exist_ok=True)
        session=job_owner.Session(s['frozen'],control,job_owner.files(s['frozen'],control));self.addCleanup(session.rt.close)
        return session

    def test_custom_code_is_rejected_atomically_and_cannot_self_authorize(self):
        self.prepare();session=self.session();bad=self.batch()[0]
        bad['outputs']=[dict(path='deck.pptx',purpose='Deck',media_type=authoring.PPTX)]
        for change in ({},{'authoring_fallback':{'reason':'I prefer custom Python'}},
                       {'outputs':[dict(path='deck.pptx',purpose='Deck disguised as text',media_type='text/plain')]}):
            with self.subTest(change=change),self.assertRaisesRegex(ValueError,'registered|Registered'):
                session.delegate(dict(request_id='bad',batch_json=c.encoded([self.content(),{**bad,**change}])))
            self.assertEqual([t['id'] for t in self.rt.status('job')['tasks']],['owner'])
        self.assertEqual(len(self.factory.calls),1) # only the already-running owner

    def test_explicit_fallback_is_format_scoped_and_preserved(self):
        self.grant['operations']=['pptx.create','pptx.render']
        self.grant['workers'].append(worker_capabilities.entry(dict(type='gemini-code',model='fixture',runtime='a'*64)))
        self.grant['authoring_policy']=authoring.capture(self.grant['operations'],{authoring.PPTX:'User authorized unsupported animation authoring'})
        self.create();self.rt.tick('job');session=self.session();task=self.batch()[0]
        task['worker']=dict(requires=['code.execute','files.binary'])
        task.update(outputs=[dict(path='deck.pptx',purpose='Explicit custom file',media_type=authoring.PPTX)],authoring_fallback={'reason':'Template renderer has no animation support'})
        session.delegate(dict(request_id='fallback',batch_json=c.encoded([task])))
        self.assertEqual(self.rt.spec(self.rt.task('job','v1'))['authoring_fallback'],task['authoring_fallback'])
        task['outputs']=[dict(path='preview.pdf',purpose='Not authorized',media_type='application/pdf')]
        with self.assertRaises(ValueError):authoring.check(task,self.grant['authoring_policy'],self.grant['operations'])

    def test_same_policy_accepts_new_registered_format_without_controller_branch(self):
        with patch.dict(authoring.REGISTRY,{'application/x-fixture':dict(extensions=['.fixture'],operations=['fixture.render'],preparation=None,renderer='fixture.render')}):
            policy=authoring.capture(['fixture.render'])
            task=dict(outputs=[dict(path='result.fixture',media_type='application/x-fixture')])
            with self.assertRaisesRegex(ValueError,'registered document'):authoring.check(task,policy,['fixture.render'])
            task['execution']={'capability':'fixture.render'};authoring.check(task,policy,['fixture.render'])

    def test_legacy_plan_is_not_rewritten_and_arbitrary_json_cannot_bypass_layouts(self):
        self.prepare();plan=job_owner.original(self.rt,'job');plan['job_owner'].pop('authoring_policy')
        before=c.encoded(plan);job_owner.validate_plan(plan);self.assertEqual(c.encoded(plan),before)
        session=self.session();content=self.content();content.pop('records')
        session.delegate(dict(request_id='auto-json',batch_json=c.encoded([content,self.renderer()])))
        bound=self.rt.spec(self.rt.task('job','content'))
        self.assertEqual(bound['operation_contract']['id'],'document.layout')
        self.assertEqual(bound['tools'],['files'])

    def test_owner_to_fragment_store_to_real_native_pptx_and_revision(self):
        from orchestrator import step_runner,pptx_document
        from pptx import Presentation
        self.prepare();self.invoke([('job_delegate',dict(request_id='layout',batch_json=c.encoded([self.content(),self.renderer()])))])
        self.rt.tick('job');aid=self.rt.task('job','content')['latest'];session=self.factory.sessions[aid];frozen=session['frozen']
        control=Path(session['session']['control']);control.mkdir(parents=True,exist_ok=True)
        atomic(control/'launch.json',dict(credential_fingerprint=executors.fingerprint(CONFIG,owner_fixture.BACKEND)))
        page=dict(kind='page',slot=1,layout='text',content=dict(title='Synthetic source',body='Exact initial observation'),notes='Synthetic fixture')
        script=Script([('operation_submit',dict(request_key='page',revision=0,records=[page])),
            ('operation_submit',dict(request_key='seal',revision=1,records=[dict(kind='end',count=1)])),
            ('finish',dict(report_json=c.encoded(dict(summary='Prepared content',decision='delivered',instruction='',findings=[],
                checks={f'c{i}':dict(passed=True,evidence='Saved exact content') for i in range(1,len(frozen['criteria'])+1)}))))])
        result=execute(frozen,control,script,lambda:(CONFIG,owner_fixture.BACKEND));self.assertEqual(result['decision'],'delivered')
        declarations=script.requests[0]['tools'][0]['functionDeclarations']
        self.assertNotIn('python_run',{d['name'] for d in declarations})
        # Reopen the transactional record store; committed data survives the tool object.
        store=operation_records.Store(frozen);self.assertTrue(store.progress()['sealed'])
        before=(Path(frozen['workspace'])/'slides.json').read_bytes()
        # An identical retry is idempotent after reopening; a conflicting seal
        # cannot partially change the saved projection.
        receipt=store.submit(dict(request_key='page',revision=0,records=[page]),{})
        self.assertEqual(receipt['committed'],1)
        with self.assertRaisesRegex(ValueError,'Conflicting request'):
            store.submit(dict(request_key='seal',revision=1,records=[dict(kind='end',count=2)]),{})
        self.assertEqual((Path(frozen['workspace'])/'slides.json').read_bytes(),before)
        self.assertEqual(json.loads(before)['version'],2)
        session['status']=dict(status='finished',exit_code=0,external_outcome='no_pending_response',pending_requests=[])
        self.rt.tick('job');native=self.factory.sessions[self.rt.task('job','render')['latest']]
        ncontrol=Path(native['session']['control']);ncontrol.mkdir(parents=True,exist_ok=True)
        details=step_runner.execute(native['frozen'],ncontrol)
        self.assertEqual(details['outcome'],'completed')
        raw=(Path(native['frozen']['workspace'])/'deck.pptx').read_bytes();deck=Presentation(io.BytesIO(raw))
        self.assertEqual(len(deck.slides),1)
        self.assertIn('Exact initial observation','\n'.join(s.text for s in deck.slides[0].shapes if s.has_text_frame))
        # Revision is data-only; original content/file remain byte-identical.
        updated=json.loads(before);updated['slides'][0]['content']['body']='Exact revised observation'
        changed,_=pptx_document.create(updated)
        self.assertIn('Exact revised observation','\n'.join(s.text for s in Presentation(io.BytesIO(changed)).slides[0].shapes if s.has_text_frame))
        self.assertEqual((Path(frozen['workspace'])/'slides.json').read_bytes(),before)
        self.assertEqual((Path(native['frozen']['workspace'])/'deck.pptx').read_bytes(),raw)

    def test_layout_validation_pagination_and_assets_are_code_owned(self):
        from orchestrator import pptx_document
        p=dict(renderer='pptx.create',title='Gallery',style='clean_minimal_v1',max_records=3)
        rows=[dict(image='item.png',label=f'Item {n}',caption='Observed image',credit='Synthetic fixture') for n in range(7)]
        record=dict(kind='page',slot=1,layout='image_grid',content=dict(title='Gallery',items=rows,per_page=4,fit='contain'),notes='')
        layout.record(record,p,{'images':{'item.png':b''}})
        with self.assertRaises(ValueError):layout.record(record,p,{'images':{}})
        with self.assertRaises(ValueError):layout.record({**record,'x':2},p)
        with self.assertRaises(ValueError):layout.document([record,dict(kind='end',count=2)],p)
        doc=layout.document([record,dict(kind='end',count=1)],p)
        expanded=pptx_document.expand(doc);self.assertEqual(len(expanded['slides']),2)
        labels=[e['text'] for s in expanded['slides'] for e in s['elements'] if e.get('name','').endswith('-label')]
        self.assertEqual(labels,[f'Item {n}' for n in range(7)])
        with self.assertRaises(ValueError):layout.validate_document({**doc,'templates':{}})

    def test_asset_bundle_handoff_uses_file_worker_without_python(self):
        from tests.test_image_sources import bundle,photo
        from orchestrator import step_runner
        from pptx import Presentation
        raw,manifest=bundle();source=self.root/'photos.zip';source.write_bytes(raw)
        ident=self.rt.register(source,'Synthetic validated photos')
        ref=dict(artifact=ident,path='photos.zip',purpose='Images',authority='Observed assets',media_type='application/zip')
        self.grant['operations']=['pptx.create'];self.grant['deliverables']={'deck':authoring.PPTX}
        self.create(inputs=[ref]);self.rt.tick('job')
        self.invoke([('job_delegate',dict(request_id='bundle',batch_json=c.encoded([self.content([ref]),self.renderer()])))])
        self.rt.tick('job');session=self.factory.sessions[self.rt.task('job','content')['latest']];frozen=session['frozen']
        self.assertEqual(frozen['backend']['type'],'gemini-agent')
        control=Path(session['session']['control']);control.mkdir(parents=True,exist_ok=True)
        atomic(control/'launch.json',dict(credential_fingerprint=executors.fingerprint(CONFIG,owner_fixture.BACKEND)))
        path=operation_records.Store(frozen).progress()['image_paths'][0]
        page=dict(kind='page',slot=1,layout='image_grid',content=dict(title='Assets',items=[dict(image=path,label='Example',caption='Synthetic',credit='Fixture')],per_page=4,fit='contain'),notes='')
        script=Script([('file_read',dict(path='photos.zip',offset=0,limit=100)),
            ('operation_submit',dict(request_key='page',revision=0,records=[page])),
            ('operation_submit',dict(request_key='seal',revision=1,records=[dict(kind='end',count=1)])),
            ('finish',dict(report_json=c.encoded(dict(summary='Prepared assets',decision='delivered',instruction='',findings=[],
                checks={f'c{i}':dict(passed=True,evidence='Exact fixture') for i in range(1,len(frozen['criteria'])+1)}))))])
        result=execute(frozen,control,script,lambda:(CONFIG,owner_fixture.BACKEND));self.assertEqual(result['decision'],'delivered')
        self.assertNotIn('python_run',{d['name'] for d in script.requests[0]['tools'][0]['functionDeclarations']})
        initial=json.loads(script.requests[0]['contents'][0]['parts'][0]['text'])
        self.assertEqual(initial['source_pack']['files'],[]) # Binary assets stay local.
        self.assertEqual(frozen['artifact_roles']['inputs']['photos.zip']['inspection'],dict(tool='image_view',path='photos.zip'))
        session['status']=dict(status='finished',exit_code=0,external_outcome='no_pending_response',pending_requests=[])
        self.rt.tick('job');native=self.factory.sessions[self.rt.task('job','render')['latest']]
        self.assertIn('photos.zip',[i['path'] for i in native['frozen']['inputs']])
        nc=Path(native['session']['control']);nc.mkdir(parents=True,exist_ok=True)
        self.assertEqual(step_runner.execute(native['frozen'],nc)['outcome'],'completed')
        deck=Presentation(str(Path(native['frozen']['workspace'])/'deck.pptx'))
        self.assertEqual([s.image.blob for p in deck.slides for s in p.shapes if s.shape_type==13],[photo()])
        # The representation is registered and finite, not general binary authority.
        from orchestrator.gemini_worker import Files
        changed=copy.deepcopy(frozen);changed['inputs'][0]['media_type']='application/octet-stream'
        with self.assertRaises(UnicodeError):Files(changed)
        source=Path(frozen['workspace'])/'photos.zip';source.chmod(0o600);source.write_bytes(raw+b'changed')
        with self.assertRaisesRegex(ValueError,'Frozen input changed'):Files(frozen)

    def test_cards_render_exact_links_images_and_disclosed_gaps(self):
        from orchestrator import pptx_document
        from tests.test_image_sources import photo
        from pptx import Presentation
        for density in (1,3):
            with self.subTest(density=density):
                rows=[dict(label=f'Observed option {n}',body='Observed dimensions: 60 × 30 inches. Price: USD 900. Delivery unknown.',
                           image='item.jpg' if n!=1 else None,link=f'https://example.com/item/{n}',credit='Synthetic source') for n in range(4)]
                doc=dict(version=2,title='Fixture',style='clean_minimal_v1',slides=[dict(layout='content_cards',content=dict(title='Options',items=rows,per_page=density))])
                raw,_=pptx_document.create(doc,{'item.jpg':photo()});deck=Presentation(io.BytesIO(raw))
                self.assertEqual(len(deck.slides),(4+density-1)//density)
                shapes=[s for page in deck.slides for s in page.shapes]
                self.assertEqual([s.image.blob for s in shapes if s.shape_type==13],[photo()]*3)
                self.assertEqual([s.click_action.hyperlink.address for s in shapes if s.name.endswith('-label')],[r['link'] for r in rows])
                self.assertIn('Image unavailable',[s.text for s in shapes if s.has_text_frame])
                validator={};exec(pptx_document.validator_source(),validator);validator['validate'](doc,['item.jpg'])
                changed=copy.deepcopy(doc);changed['slides'][0]['content']['items'][0]['link']='javascript:alert(1)'
                with self.assertRaises(ValueError):pptx_document.create(changed,{'item.jpg':photo()})

    def test_record_author_native_evidence_admission_dispatch_and_sealed_delivery(self):
        source=self.root/'native.fixture';source.write_bytes(b'\x00\xff controlled native evidence')
        aid=self.rt.register(source,'Native evidence for context')
        ref=dict(artifact=aid,path='native.fixture',media_type='application/x-fixture',purpose='Preserved evidence',authority='Metadata only; no content observation')
        self.grant['operations']=['pptx.create'];self.grant['deliverables']={'deck':authoring.PPTX}
        self.create(inputs=[ref]);self.rt.tick('job')
        self.invoke([('job_delegate',dict(request_id='context',batch_json=c.encoded([self.content([ref])])))])
        self.rt.tick('job');state=self.factory.sessions[self.rt.task('job','content')['latest']]
        frozen=state['frozen'];control=Path(state['session']['control']);control.mkdir(parents=True,exist_ok=True)
        atomic(control/'launch.json',dict(credential_fingerprint=executors.fingerprint(CONFIG,owner_fixture.BACKEND)))
        page=dict(kind='page',slot=1,layout='text',content=dict(title='Evidence retained',body='Native bytes remain preserved; no semantic observation is claimed.'),notes='Controlled fixture')
        script=Script([('file_read',dict(path='native.fixture',offset=0,limit=1000)),
            ('artifact_inspect',dict(path='native.fixture')),
            ('operation_submit',dict(request_key='page',revision=0,records=[page])),
            ('operation_submit',dict(request_key='seal',revision=1,records=[dict(kind='end',count=1)])),
            ('finish',dict(report_json=c.encoded(dict(summary='Prepared section',decision='delivered',instruction='',findings=[],checks={f'c{i}':dict(passed=True,evidence='Sealed exact section; native context retained without semantic claims') for i in range(1,len(frozen['criteria'])+1)}))))])
        result=execute(frozen,control,script,lambda:(CONFIG,owner_fixture.BACKEND))
        self.assertEqual(result['decision'],'delivered')
        self.assertTrue(operation_records.Store(frozen).progress()['sealed'])
        names={d['name'] for d in script.requests[0]['tools'][0]['functionDeclarations']}
        self.assertIn('artifact_inspect',names);self.assertNotIn('python_run',names)
        self.assertEqual(json.loads((control/'tool-02-00.json').read_text())['result']['semantic_acceptance'],False)

    def test_changed_source_and_missing_images_fail_before_fragment_commit(self):
        import hashlib
        from tests.test_image_sources import photo
        raw=photo();frozen=dict(inputs=[dict(path='photo.jpg',media_type='image/jpeg',sha256=hashlib.sha256(raw).hexdigest())])
        self.assertEqual(set(layout.context(frozen,lambda _:raw)['images']),{'photo.jpg'})
        with self.assertRaisesRegex(ValueError,'source version changed'):layout.context(frozen,lambda _:raw+b'x')


if __name__=='__main__':unittest.main()
