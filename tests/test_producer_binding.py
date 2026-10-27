"""Positive composed flows: real compiler/store/scheduler/renderers, scripted models.

No live provider, browser, host application, media render, or service mutation.
"""
import copy
import hashlib
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from orchestrator import contracts as c, execution, executors, prepared_contract, prepared_schemas, producer_binding
from orchestrator import operation_records, parameter_contracts
from orchestrator.adapters import ExecutionFactory
from orchestrator.runtime import Runtime
from orchestrator.gemini_worker import execute
from orchestrator.workers import atomic
from tests.test_orchestrator import FakeFactory
from tests.test_mixed_execution import LocalRegistered, Client
from tests.test_operation_session import Script, CONFIG
from tests.test_job_owner_loop import BACKEND
from tests.test_operation_adapters import assignment as native


def ref(task,output,path=None,media='application/json'):
    return dict(from_task=task,output=output,path=path or output,media_type=media,purpose='Exact upstream data',authority='Synthetic evidence')


def writer(ident,output='data.json',inputs=(),**extra):
    return dict(id=ident,role='specialist',objective='Synthetic document',instruction='Use exact synthetic inputs and registered fragments.',
        inputs=list(inputs),outputs=[dict(path=output,purpose='Prepared data',media_type='application/json')],
        dependencies=list(dict.fromkeys(i['from_task'] for i in inputs if 'from_task' in i)),
        tools=['files'],limits=dict(seconds=300,tool_calls=24,provider_requests=24,response_tokens=8192,output_bytes=200000),
        max_attempts=1,criteria=['Preserve exact content and disclosed gaps'],**extra)


def operation(ident,capability,inputs,parameters=None):
    value=native(capability,parameters or {},inputs,ident,list(dict.fromkeys(i['from_task'] for i in inputs if 'from_task' in i)))
    value['max_attempts']=1
    if execution.REGISTRY[capability].get('output_type') and not execution.REGISTRY[capability].get('outputs'):
        value['outputs']=[dict(path='deck.pptx' if capability.startswith('pptx.') else 'result.json',purpose='Native result',media_type=execution.REGISTRY[capability]['output_type'])]
    return value


class Tests(unittest.TestCase):
    def test_provider_variant_retains_registered_evidence_identity(self):
        source=operation('source','web.sources.codex',[])
        consumer=operation('consumer','web.sources.rendered',[ref('source','delivery/source-pack.json')])
        producer_binding.compile(dict(backend=BACKEND,tasks=[source,consumer]))
        self.assertEqual(source['execution']['capability'],'web.sources.codex')
        self.assertEqual(execution.adapter('web.sources.codex').backend_binding['provider'],'codex')
        bad=copy.deepcopy(source);bad['execution']['capability']='text.bundle'
        with self.assertRaisesRegex(ValueError,'Evidence port requires'):
            producer_binding.compile(dict(backend=BACKEND,tasks=[bad,copy.deepcopy(consumer)]))

    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name).resolve();self.agent=FakeFactory();self.ops=LocalRegistered(Client())
        self.factory=ExecutionFactory(self.agent,self.ops,gemini=self.agent)
        self.factory.available=lambda backend:None # synthetic worker connection; never launch a provider process
        self.rt=Runtime(self.root/'runtime',self.factory)
        self.addCleanup(lambda:self.rt.close())
        config=patch('task_relay.gemini.read_config',return_value={**CONFIG,'models':{'text':BACKEND['model']}})
        config.start();self.addCleanup(config.stop)

    def start(self,tasks,ident='job'):
        self.rt.create(dict(id=ident,brief='Exact synthetic request',backend=BACKEND,concurrency=3,tasks=tasks,producer_binding_version=1))
        self.rt.tick(ident)

    def staged(self,ident,run='job'):
        aid=self.rt.task(run,ident)['latest'];self.assertIsNotNone(aid,[tuple(r) for r in self.rt.db.execute('SELECT kind,data FROM production_events WHERE run=? ORDER BY rowid DESC LIMIT 1',(run,))])
        session=self.agent.sessions[aid];f=session['frozen'];control=Path(session['session']['control']);control.mkdir(parents=True,exist_ok=True)
        atomic(control/'launch.json',dict(credential_fingerprint=executors.fingerprint(CONFIG,BACKEND)))
        return session,f,control

    def fragments(self,ident,records,run='job'):
        session,f,control=self.staged(ident,run)
        steps=[('operation_submit',dict(request_key='part'+str(i),revision=i,records=[record])) for i,record in enumerate(records)]
        decision='accept' if f.get('review_of') else 'delivered'
        steps.append(('finish',dict(report_json=c.encoded(dict(summary='Synthetic content',decision=decision,instruction='',findings=[],
            checks={f'c{i}':dict(passed=True,evidence='Exact synthetic fragment') for i in range(1,len(f['criteria'])+1)})))))
        script=Script(steps)
        result=execute(f,control,script,lambda:(CONFIG,BACKEND));self.assertEqual(result['decision'],decision,result)
        tools={d['name'] for d in script.requests[0]['tools'][0]['functionDeclarations']}
        self.assertNotIn('python_run',tools)
        session['status']=dict(status='finished',exit_code=0,external_outcome='no_pending_response',pending_requests=[])
        self.rt.tick(run);self.rt.tick(run)
        return f,script

    def artifact(self,task,path,run='job'):
        value=self.rt.output(run,task,path);return value,Path(value['blob']).read_bytes()

    def saved(self,path,raw,media='application/json'):
        file=self.root/path;file.write_bytes(raw);aid=self.rt.register(file,'Synthetic source')
        return dict(artifact=aid,path=path,purpose='Exact baseline',authority='Synthetic evidence',media_type=media)

    def test_automatic_layout_to_real_pptx_then_bound_revision_preserves_baseline(self):
        from pptx import Presentation
        pptx='application/vnd.openxmlformats-officedocument.presentationml.presentation'
        self.start([writer('content'),operation('render','pptx.create',[ref('content','data.json')])])
        page=dict(kind='page',slot=1,layout='text',content=dict(title='Evidence',body='Initial observation'),notes='')
        f,_=self.fragments('content',[page,dict(kind='end',count=1)])
        self.assertEqual(f['operation_contract']['id'],'document.layout')
        self.assertEqual(self.rt.status('job')['status'],'completed',self.rt.status('job'))
        baseline,raw=self.artifact('render','deck.pptx');original_spec=self.rt.spec(self.rt.task('job','content'))
        deck=Presentation(io.BytesIO(raw));shape=next(s for s in deck.slides[0].shapes if s.has_text_frame and s.text=='Initial observation')
        source=dict(artifact=baseline['id'],path='baseline.pptx',media_type=pptx,purpose='Exact delivered baseline',authority='Immutable prior result')
        self.start([writer('change',inputs=[source]),operation('edit','pptx.edit',[source,ref('change','data.json')])],'revision')
        # Reconstruct Runtime before writing the revision; no assignment replay.
        self.rt.close();self.rt=Runtime(self.root/'runtime',self.factory)
        edit=dict(kind='replace_text',slide=1,shape_id=shape.shape_id,old='Initial observation',new='Revised observation')
        frozen,script=self.fragments('change',[dict(kind='append',field='edits',index=0,value=edit),dict(kind='end')],'revision')
        manifest=json.loads(self.artifact('change','data.json','revision')[1])
        self.assertEqual(manifest,dict(version=1,source_sha256=baseline['sha256'],edits=[edit]))
        advertised=c.encoded([tool for group in script.requests[0]['tools'] for tool in group.get('functionDeclarations',[]) if tool['name'].startswith('operation_')])
        self.assertNotIn('source_sha256',advertised) # caller cannot author baseline identity
        self.assertEqual(self.rt.status('revision')['status'],'completed',self.rt.status('revision'))
        output=next(a for a in self.rt.status('revision')['artifacts'] if a['task']=='edit' and a['path'].endswith('.pptx'))
        revised=Presentation(self.rt.artifact(output['id'])['blob'])
        self.assertIn('Revised observation',[s.text for s in revised.slides[0].shapes if s.has_text_frame])
        self.assertEqual(Path(baseline['blob']).read_bytes(),raw)
        self.assertEqual(self.rt.spec(self.rt.task('job','content')),original_spec)
        self.assertEqual(len(self.ops.calls),2)

    def test_query_fragments_flow_to_actual_source_collector_without_envelope_writing(self):
        from task_relay import orchestrator_web,gemini
        from tests.test_web_sources import FakeSession,CONFIG as SEARCH_CONFIG
        query=dict(id='part-1',queries=['ABC-9182 OEM material'],required_options=1)
        self.start([writer('query'),operation('search','web.sources.prepared',[ref('query','data.json')],dict(domains=[],model='fixture-text'))])
        # LocalRegistered's config callback is used only at this transport boundary.
        from orchestrator.step_runner import execute as native_execute
        def run(frozen,control,*args):return native_execute(frozen,control,config_reader=lambda:SEARCH_CONFIG)
        with patch.object(orchestrator_web,'Session',FakeSession),patch.object(orchestrator_web,'public_url',side_effect=lambda u:u),patch.object(gemini,'read_config',return_value=SEARCH_CONFIG),patch('tests.test_mixed_execution.execute',side_effect=run):
            f,_=self.fragments('query',[dict(kind='append',field='targets',index=0,value=query),dict(kind='end')])
        self.assertEqual(f['operation_contract']['parameters']['target'],'web.sources.prepared')
        self.assertEqual(self.rt.status('job')['status'],'completed',self.rt.status('job'))
        pack=json.loads(self.artifact('search','delivery/source-pack.json')[1])
        self.assertEqual(pack['query_specification']['data'],dict(version=1,targets=[query]))
        self.assertEqual(pack['entries'][0]['target_ids'],['part-1'])
        self.assertEqual(len(self.ops.calls),1)

    def test_no_history_rebinding_conflicting_consumers_or_partial_admission(self):
        self.start([writer('old')])
        before=self.rt.spec(self.rt.task('job','old'))
        with self.assertRaisesRegex(ValueError,'Saved producer'):
            self.rt.add_futures('job',[operation('late','pptx.create',[ref('old','data.json')])])
        self.assertEqual(self.rt.spec(self.rt.task('job','old')),before)
        self.assertEqual([t['id'] for t in self.rt.status('job')['tasks']],['old'])
        tasks=[writer('ambiguous'),operation('a','pptx.create',[ref('ambiguous','data.json')]),
            operation('b','web.sources.prepared',[ref('ambiguous','data.json')],dict(domains=[],model='fixture-text'))]
        with self.assertRaisesRegex(ValueError,'incompatible'):
            self.rt.create(dict(id='conflict',brief='Fixture',backend=BACKEND,tasks=tasks,producer_binding_version=1))
        self.assertIsNone(self.rt.db.execute('SELECT id FROM production_runs WHERE id=?',('conflict',)).fetchone())

    def test_generic_collection_binding_review_projection_and_real_deck(self):
        from pptx import Presentation
        from task_relay import orchestrator_web,gemini
        from tests.test_web_sources import FakeSession,CONFIG as SEARCH_CONFIG
        from orchestrator.step_runner import execute as native_execute
        def run_native(frozen,control,*args):return native_execute(frozen,control,config_reader=lambda:SEARCH_CONFIG)
        for field,label in [('botanical_name','Acer rubrum'),('designation','ABC-9182'),('name','Sofa option')]:
            with self.subTest(field=field):
                run='flow-'+field
                subjects=self.saved(field+'-subjects.json',c.encoded(dict(items=[dict(id='S0',label=label)])).encode())
                page='Observed '+label
                sources=ref('search','delivery/source-pack.json','sources.json')
                query=writer('query')
                search=operation('search','web.sources.prepared',[ref('query','data.json')],dict(domains=[],model='fixture-text'))
                class Session(FakeSession):
                    def fetch(self,url,offset,limit,allowed_hosts=None):
                        return {**super().fetch(url,offset,limit,allowed_hosts), 'text':page,
                                'sha256':hashlib.sha256(page.encode()).hexdigest(),'total_characters':len(page)}
                source=writer('collect',inputs=[subjects,sources],content_parameters=dict(field_definitions={field:'text'},identity_field=field,
                    options_per_item=2,count_policy='target',max_records=1))
                review=writer('review','review.md',[subjects,sources,ref('collect','data.json','candidate.json')],review_of='collect')
                review['outputs'][0]['media_type']='text/markdown'
                slides=operation('slides','collection.slides',[subjects,sources,ref('collect','data.json')],dict(
                    inventory_path=subjects['path'],source_pack_path=sources['path'],catalog_path='data.json',
                    image_path=None,photos_path=None,title='Synthetic comparison',layout='readable'))
                slides['dependencies'].append('review')
                output=slides['outputs'][0]['path']
                render=operation('render','pptx.create',[ref('slides',output)])
                with patch.object(orchestrator_web,'Session',Session),patch.object(orchestrator_web,'public_url',side_effect=lambda u:u),patch.object(gemini,'read_config',return_value=SEARCH_CONFIG),patch('tests.test_mixed_execution.execute',side_effect=run_native):
                    self.start([query,search,source,review,slides,render],run)
                    self.fragments('query',[dict(kind='append',field='targets',index=0,value=dict(id='S0',queries=[label],required_options=2)),dict(kind='end')],run)
                fragment=dict(item='S0',options=[dict(observation='page-0-0',fields={field:dict(status='observed',start=9,end=len(page))},comparison='Synthetic match')],gap='Only one observed option; second remains unavailable')
                f,_=self.fragments('collect',[fragment],run)
                self.assertEqual(f['operation_contract']['id'],'collection.options')
                catalog=json.loads(self.artifact('collect','data.json',run)[1]);item=catalog['items'][0]
                self.assertEqual(item['shortfall'],1)
                self.assertEqual(self.rt.task(run,'slides')['attempts'],0)
                self.fragments('review',[dict(item='S0',item_sha256=c.digest(item),verdict='supported',reason='Exact captured literal; gap explicit')],run)
                for _ in range(3):self.rt.tick(run)
                self.assertEqual(self.rt.status(run)['status'],'completed',self.rt.status(run))
                deck=Presentation(io.BytesIO(self.artifact('render','deck.pptx',run)[1]))
                text='\n'.join(s.text for page in deck.slides for s in page.shapes if s.has_text_frame)
                self.assertIn(label,text)
                self.assertIn('Only one observed option',text)
                self.assertEqual(self.rt.spec(self.rt.task(run,'collect'))['limits'],source['limits'])

    def test_store_reopen_idempotence_and_failed_seal_preserve_committed_content(self):
        self.start([writer('query'),operation('search','web.sources.prepared',[ref('query','data.json')],dict(domains=[],model='fixture-text'))])
        _,f,_=self.staged('query');store=operation_records.Store(f)
        first=dict(kind='append',field='targets',index=0,value=dict(id='p',queries=['query'],required_options=1))
        args=dict(request_key='first',revision=0,records=[first]);receipt=store.submit(args,{})
        reopened=operation_records.Store(f);self.assertEqual(reopened.submit(args,{}),receipt)
        with self.assertRaisesRegex(ValueError,'consecutive'):
            reopened.submit(dict(request_key='skip',revision=1,records=[{**first,'index':2}]),{})
        self.assertEqual(reopened.progress()['committed'],1)
        reopened.submit(dict(request_key='duplicate-target',revision=1,records=[{**first,'index':1}]),{})
        with self.assertRaisesRegex(ValueError,'distinct'):
            reopened.submit(dict(request_key='bad-seal',revision=2,records=[dict(kind='end')]),{})
        self.assertEqual(reopened.progress()['committed'],2)
        self.assertEqual(self.rt.task('job','search')['attempts'],0)
        self.assertEqual(self.ops.calls,[])

    def test_every_operation_has_explicit_handoff_policy_and_preparation_wire_schema(self):
        self.assertEqual(set(execution.REGISTRY),{cap for cap in execution.REGISTRY if producer_binding.requirements(cap) is not None})
        for cap in execution.REGISTRY:
            self.assertEqual(execution.adapter(cap).manifest()['handoffs']['version'],1)
        for cap,spec in prepared_schemas.REGISTRY.items():
            with self.subTest(cap=cap):
                p=dict(target=cap,max_records=100,hash_bindings={k:'baseline' for k in spec['hash_inputs']})
                parameter_contracts.check(prepared_contract.schema(p))
                parameter_contracts.check(prepared_schemas.document_schema(cap))
        from orchestrator.connector_protocol import capture,verify
        saved=capture('web.sources.prepared');changed=copy.deepcopy(prepared_schemas.REGISTRY['web.sources.prepared'])
        changed['schema']['properties']['targets']['maxItems']=10
        with patch.dict(prepared_schemas.REGISTRY,{'web.sources.prepared':changed}):
            with self.assertRaisesRegex(ValueError,'changed since capture'):verify('web.sources.prepared',saved)

    def test_all_preparation_adapters_compile_commit_reopen_export_and_validate(self):
        from orchestrator.gemini_worker import Files
        from tests.test_blender_operations import scene_data
        from tests.test_blender_edit import checks as blender_checks
        from tests.test_rhino_operations import checks as rhino_checks
        from tests.test_automatic_composition import native_document
        from tests.test_rhino3dm_script import checks as library_checks
        from tests.test_sketchup_operations import checks as sketchup_checks
        from tests.test_grasshopper import checks as grasshopper_checks
        from tests.test_rhino3dm import fixture as geometry
        from scripts.qualify_hyperframes import fixture as frontend
        baseline=b'Synthetic exact native file bytes';digest=hashlib.sha256(baseline).hexdigest()
        documents={
            'web.sources.prepared':dict(version=1,targets=[dict(id='part',queries=['exact part'],required_options=1)]),
            'blender.scene':scene_data(),'blender.mesh_scene':{**scene_data(),'version':2},
            'pptx.edit':dict(version=1,source_sha256=digest,edits=[dict(kind='replace_text',slide=1,shape_id=2,old='Old',new='New')]),
            'media.compose':dict(version=1,width=320,height=320,fps=30,background='#101820',foreground='#ffffff',accent='#82cfff',audio=None,
                scenes=[dict(duration=2,title='Fixture',body='',image=None,entrance='fade',narration='')]),
            'hyperframes.preview':frontend(),'rhino3dm.create':geometry(),'blender.run_python':blender_checks(),
            'rhino.run_python':rhino_checks(),'rhino.document':native_document(),'rhino3dm.run_python':library_checks(),'rhino.grasshopper':grasshopper_checks(),
            'sketchup.run_ruby':sketchup_checks(),'rhino.render':dict(version=1,named_view='Camera',resolution=[64,64],engine='rhino_render'),
            'blender.animate':dict(version=1,scene_sha256=digest,mode='preview',frame_start=1,frame_end=3,fps=3,camera='Camera',resolution=[64,64],samples=1,tracks=[]),
            'blender.import_asset':dict(version=1,scene='scene.blend',files=[dict(path='scene.blend',sha256=digest,kind='scene',provenance=dict(source='Fixture',license=None))],imports=[],preview=dict(camera='Camera',resolution=[64,64],samples=1))}
        self.assertEqual(set(documents),set(prepared_schemas.REGISTRY))
        for cap,expected in documents.items():
            with self.subTest(cap=cap):
                spec=prepared_schemas.REGISTRY[cap];inputs=[];workspace=self.root/cap;workspace.mkdir()
                for field,media in spec['hash_inputs'].items():
                    path='baseline.pptx' if cap=='pptx.edit' else 'scene.blend'
                    (workspace/path).write_bytes(baseline)
                    inputs.append(dict(artifact='baseline',path=path,media_type=media,purpose='Exact fixture',authority='Synthetic',sha256=digest))
                if cap=='blender.import_asset':
                    (workspace/'scene.blend').write_bytes(baseline)
                    inputs.append(dict(artifact='baseline',path='scene.blend',media_type='application/x-blender',purpose='Exact fixture',authority='Synthetic',sha256=digest))
                author=writer('prepare',inputs=inputs)
                consumer=dict(id='consume',execution=dict(capability=cap,parameters={}),inputs=[ref('prepare','data.json'),*inputs])
                producer_binding.compile(dict(tasks=[author,consumer],backend=BACKEND))
                f={**author,'backend':BACKEND,'assignment_id':cap,'workspace':str(workspace)}
                # Store uses this explicit synthetic database; no installed state.
                self.rt.db.execute('INSERT INTO production_attempts VALUES (?,?,?,?,?,?,?,?,?,?)',(cap,'matrix',cap,'assignment','running',None,c.encoded(f),'{}',None,None))
                store=operation_records.Store(f,self.rt.db);fragments=[]
                for field,rule in spec['schema']['properties'].items():
                    if field not in expected:continue
                    value=copy.deepcopy(expected[field])
                    if rule.get('type')=='array':
                        for index,item in enumerate(value):
                            for binding in spec.get('item_hashes',[]):
                                if binding['field']==field and binding['hash'] in item:item.pop(binding['hash'])
                            fragments.append(dict(kind='append',field=field,index=index,value=item))
                    elif rule.get('type')=='object' and isinstance(rule.get('additionalProperties'),dict):
                        fragments.extend(dict(kind='entry',field=field,key=k,value=v) for k,v in value.items())
                    else:fragments.append(dict(kind='field',field=field,value=value))
                fragments.append(dict(kind='end'))
                for index,record in enumerate(fragments):store.submit(dict(request_key=str(index),revision=index,records=[record]),{})
                reopened=operation_records.Store(f,self.rt.db);files=Files(f)
                reopened.export(files,complete=True);reopened.verify_export(files.read_bytes)
                actual=json.loads(files.read_bytes('data.json'));self.assertEqual(actual,expected)
                prepared_schemas.validate(cap,actual,{i['path']:i['media_type'] for i in inputs})
                self.assertTrue(reopened.progress()['sealed'])
        self.assertEqual(self.ops.calls,[]) # host/video execution expressly outside this matrix

    def test_new_registered_consumer_uses_same_writer_without_controller_branch(self):
        from dataclasses import replace
        cap='fixture.search';spec=copy.deepcopy(execution.REGISTRY['web.sources.prepared'])
        adapter=replace(execution.adapter('web.sources.prepared'),handoffs=dict(version=1,ports=[producer_binding.port('json','preparation',target='web.sources.prepared')]))
        with patch.dict(execution.REGISTRY,{cap:spec}),patch.dict(execution.ADAPTERS,{cap:adapter}):
            tasks=[writer('query'),operation('new',cap,[ref('query','data.json')],dict(domains=[],model='fixture-text'))]
            graph=dict(id='extension',brief='Fixture',backend=BACKEND,tasks=tasks,producer_binding_version=1)
            bound=c.plan(graph)
            self.assertEqual(bound['tasks'][0]['operation_contract']['parameters']['target'],'web.sources.prepared')
            self.assertNotIn('operation_contract',graph['tasks'][0])

    def test_cli_transport_uses_workspace_only_then_shared_commit_and_real_render(self):
        import subprocess,sys
        from orchestrator import record_tool,report_builder
        source=writer('content');source['tools']=['files','shell']
        backend=dict(type='codex-cli',model='fixture',reasoning='high')
        self.rt.create(dict(id='cli',brief='Fixture',backend=backend,producer_binding_version=1,
            tasks=[source,operation('render','pptx.create',[ref('content','data.json')])]))
        self.rt.tick('cli');session=self.agent.sessions[self.rt.task('cli','content')['latest']];f=session['frozen']
        workspace=Path(f['workspace']);instructions=record_tool.prepare(workspace,f)
        self.assertIn('record_tool.py',instructions)
        records=[dict(kind='page',slot=1,layout='text',content=dict(title='CLI fixture',body='Exact fragment'),notes=''),dict(kind='end',count=1)]
        for index,record in enumerate(records):
            result=subprocess.run([sys.executable,str(workspace/'.relay/record_tool.py'),'submit'],cwd=workspace,
                input=c.encoded(dict(request_key=str(index),revision=index,records=[record])),text=True,capture_output=True,timeout=20)
            self.assertEqual(result.returncode,0,result.stderr)
        self.assertEqual(self.rt.db.execute('SELECT count(*) FROM operation_records WHERE attempt=?',(f['assignment_id'],)).fetchone()[0],0)
        store=operation_records.Store(f,self.rt.db)
        with self.assertRaisesRegex(ValueError,'confirmed successful'):
            store.ingest_local(records)
        row=self.rt.db.execute('SELECT receipt FROM production_attempts WHERE id=?',(f['assignment_id'],)).fetchone()
        self.rt.db.execute('UPDATE production_attempts SET receipt=? WHERE id=?',(c.encoded(dict(status='finished',exit_code=0)),f['assignment_id']))
        with self.assertRaisesRegex(ValueError,'seal differs'):
            store.ingest_local([records[0],dict(kind='end',count=2)])
        self.assertEqual(self.rt.db.execute('SELECT count(*) FROM operation_records WHERE attempt=?',(f['assignment_id'],)).fetchone()[0],0)
        self.rt.db.execute('UPDATE production_attempts SET receipt=? WHERE id=?',(row[0],f['assignment_id']))
        report=dict(summary='Synthetic CLI preparation',decision='delivered',instruction='',findings=[],
                    checks={'c1':dict(passed=True,evidence='Exact fragment projection')})
        atomic(workspace/'.relay/result.json',report_builder.build(report,f) if f.get('report_contract') else dict(assignment_id=f['assignment_id'],
            summary=report['summary'],decision='delivered',instruction='',checks=[dict(criterion=1,passed=True,evidence='Exact fragment projection')]))
        session['status']=dict(status='finished',exit_code=0,external_outcome='no_pending_response',pending_requests=[])
        self.rt.tick('cli');self.rt.tick('cli')
        self.assertEqual(self.rt.status('cli')['status'],'completed',self.rt.status('cli'))
        self.assertEqual(self.rt.db.execute('SELECT count(*) FROM operation_records WHERE attempt=?',(f['assignment_id'],)).fetchone()[0],2)
        self.assertTrue(self.artifact('render','deck.pptx','cli')[1].startswith(b'PK'))

    def test_consumer_asset_paths_are_bound_before_writer_and_render(self):
        from tests.test_image_sources import photo
        from pptx import Presentation
        source=self.saved('photo.jpg',photo(),'image/jpeg')
        consumer={**source,'path':'assets/selected.jpg'}
        self.start([writer('content',inputs=[source]),operation('render','pptx.create',[ref('content','data.json'),consumer])])
        f,_=self.fragments('content',[dict(kind='page',slot=1,layout='content_cards',content=dict(title='Bound photo',
            items=[dict(label='Exact image',body='Observed fixture',image=consumer['path'],link=None,credit='Fixture')],per_page=1),notes=''),dict(kind='end',count=1)])
        self.assertEqual(f['inputs'][0]['path'],consumer['path'])
        deck=Presentation(io.BytesIO(self.artifact('render','deck.pptx')[1]))
        self.assertEqual([s.image.blob for p in deck.slides for s in p.shapes if s.shape_type==13],[photo()])
        self.assertEqual(self.rt.status('job')['status'],'completed')


    def test_unsealed_preparation_automatically_continues_exact_checkpoint_after_restart(self):
        source=writer('prepare');source['max_attempts']=2
        self.start([source,operation('collect','web.sources.prepared',[ref('prepare','data.json')],dict(model='fixture',domains=[]))])
        session,f,control=self.staged('prepare');store=operation_records.Store(f)
        record=dict(kind='append',field='targets',index=0,value=dict(id='item',queries=['query'],required_options=1))
        script=Script([('operation_submit',dict(request_key='first',revision=0,records=[record]))]+[('file_read',dict(path='data.json',offset=0,limit=1000))]*22+[
            ('finish',dict(report_json=c.encoded(dict(summary='Confirmed stop with one saved target',decision='blocked',instruction='Continue exact preparation seal within original grant',findings=[],checks={'c1':dict(passed=False,evidence='Target is saved but required seal remains')}))))])
        execute(f,control,script,lambda:(CONFIG,BACKEND))
        original=copy.deepcopy(f);original_records=store.progress();self.assertEqual(original_records['remaining'],1)
        # A confirmed exhausted invocation can continue only its existing grant.
        session['status']=dict(status='finished',exit_code=0,reason=None,external_outcome='no_pending_response',pending_requests=[],tool_calls=len(script.requests),api_requests=len(script.requests))
        self.rt.tick('job',dispatch=False)
        self.assertEqual(self.rt.task('job','prepare')['status'],'queued')
        self.rt.close();self.rt=Runtime(self.root/'runtime',self.factory)
        self.rt.tick('job');new,newf,newcontrol=self.staged('prepare')
        self.assertNotEqual(newf['assignment_id'],f['assignment_id'])
        for key in ('operation_contract','limits','backend','inputs'):self.assertEqual(newf[key],original[key])
        inherited=operation_records.Store(newf);self.assertEqual(inherited.progress()['committed'],1)
        steps=[('operation_submit',dict(request_key='seal',revision=inherited.progress()['revision'],records=[dict(kind='end')])),
            ('finish',dict(report_json=c.encoded(dict(summary='Sealed preserved target',decision='delivered',instruction='',findings=[],checks={'c1':dict(passed=True,evidence='Exact original target retained and sealed')}))))]
        execute(newf,newcontrol,Script(steps),lambda:(CONFIG,BACKEND))
        self.assertTrue(inherited.progress()['sealed']);self.assertEqual(store.progress(),original_records)
        self.assertEqual(self.rt.task('job','prepare')['attempts'],2)
        self.assertEqual(self.rt.task('job','collect')['attempts'],0)

    def test_preparation_batches_final_fragments_and_seal_before_last_report(self):
        self.start([writer('prepare'),operation('collect','web.sources.prepared',[ref('prepare','data.json')],dict(model='fixture',domains=[]))])
        _,f,control=self.staged('prepare')
        def target(index):return dict(kind='append',field='targets',index=index,value=dict(id='item-'+str(index),queries=['query-'+str(index)],required_options=1))
        steps=[('operation_submit',dict(request_key='first',revision=0,records=[target(0)]))]
        # Prior inspection has consumed the allowance; the final content call
        # must still commit the last fragment AND explicit seal atomically.
        steps += [('file_read',dict(path='data.json',offset=0,limit=1000))]*21
        steps += [('operation_submit',dict(request_key='seal',revision=1,records=[target(1),dict(kind='end')])),
            ('finish',dict(report_json=c.encoded(dict(summary='Sealed exact preparation',decision='delivered',instruction='',findings=[],checks={'c1':dict(passed=True,evidence='Both targets are retained in the sealed preparation')}))))]
        script=Script(steps);result=execute(f,control,script,lambda:(CONFIG,BACKEND))
        self.assertEqual(result['decision'],'delivered');self.assertEqual(len(script.requests),24)
        store=operation_records.Store(f);self.assertTrue(store.progress()['sealed'])
        self.assertFalse(store.progress()['completion']['pending'])
        receipt=store.submit(dict(request_key='seal',revision=1,records=[target(1),dict(kind='end')]),{})
        self.assertEqual(receipt['committed'],3)
        self.assertEqual(self.rt.task('job','collect')['attempts'],0)

    def test_failed_batched_seal_commits_no_partial_fragments_and_corrects_after_reopen(self):
        self.start([writer('prepare'),operation('collect','web.sources.prepared',[ref('prepare','data.json')],dict(model='fixture',domains=[]))])
        _,f,_=self.staged('prepare');store=operation_records.Store(f)
        first=dict(kind='append',field='targets',index=0,value=dict(id='item',queries=['query'],required_options=1))
        store.submit(dict(request_key='first',revision=0,records=[first]),{})
        second={**first,'index':1}
        with self.assertRaisesRegex(ValueError,'distinct'):
            store.submit(dict(request_key='bad',revision=1,records=[second,dict(kind='end')]),{})
        reopened=operation_records.Store(f)
        self.assertEqual(reopened.progress()['committed'],1)
        second['value']={**second['value'],'id':'other'}
        args=dict(request_key='correct',revision=1,records=[second,dict(kind='end')])
        receipt=reopened.submit(args,{})
        self.assertEqual(operation_records.Store(f).submit(args,{}),receipt)
        self.assertEqual(reopened.progress()['committed'],3)
        self.assertEqual(self.rt.task('job','collect')['attempts'],0)

    def test_invalid_combined_preparation_can_be_amended_restarted_and_sealed(self):
        from orchestrator import agent_loop,web_sources
        self.start([writer('prepare'),operation('collect','web.sources.prepared',[ref('prepare','data.json')],dict(model='fixture',domains=[]))])
        _,f,control=self.staged('prepare');store=operation_records.Store(f)
        def target(index,count):return dict(kind='append',field='targets',index=index,value=dict(id='item-'+str(index),queries=['query-'+str(index)+'-'+str(n) for n in range(count)],required_options=1))
        first,second=target(0,11),target(1,11)
        for n,record in enumerate((first,second)):store.submit(dict(request_key=str(n),revision=n,records=[record]),{})
        with self.assertRaisesRegex(ValueError,'1–20'):
            store.submit(dict(request_key='invalid-seal',revision=2,records=[dict(kind='end')]),{})
        self.assertFalse(store.progress()['sealed']);self.assertEqual(store.progress()['revision'],2)
        declared={t['name'] for t in agent_loop.definitions(f)}
        self.assertIn('operation_amend',declared)
        proof=store.progress()['fragments'][0]['record_sha256']
        change=dict(request_key='reduce',revision=2,expected_record_sha256=proof,record=target(0,9))
        receipt=store.amend(change,{})
        self.assertEqual(store.amend(change,{}),receipt)
        with self.assertRaisesRegex(ValueError,'Stale'):
            store.amend(dict(change,request_key='stale'),{})
        restarted=operation_records.Store(f)
        self.assertEqual(restarted.progress()['revision'],3)
        self.assertEqual(restarted.progress()['fragments'][0]['record_sha256'],c.digest(target(0,9)))
        restarted.submit(dict(request_key='seal',revision=3,records=[dict(kind='end')]),{})
        with restarted.connection() as db:
            _,records=restarted._state(db)
            versions=[json.loads(r[0]) for r in db.execute('SELECT value FROM operation_record_versions WHERE attempt=?',(f['assignment_id'],))]
        self.assertEqual(versions,[first]);self.assertTrue(restarted.progress()['sealed'])
        self.assertEqual(len(web_sources.validate_prepared(restarted.document(records))),20)
        self.assertEqual(f['limits']['provider_requests'],24)
        self.assertEqual(self.rt.task('job','collect')['attempts'],0)


if __name__=='__main__':unittest.main()

