import copy
import io
import unittest
from orchestrator import slide_templates, pptx_document, table_layout, product_evidence, sourcing_contract
from orchestrator import operation_records, contracts
from tests import test_worker_tools as fixture

class Tests(unittest.TestCase):
    def test_html_table_fields_keep_visible_separation(self):
        from task_relay.orchestrator_web import PageText
        parser=PageText('https://example.org/item')
        parser.feed('<table><tr><th>Diameter</th><td>21 inches</td></tr><tr><th>Overall</th><td>1000 mm</td></tr></table>')
        text=''.join(parser.parts)
        self.assertNotIn('Diameter21',text);self.assertNotIn('Overall1000',text)
        self.assertEqual(text.split(),['Diameter','21','inches','Overall','1000','mm'])

    def test_overflow_rejects_before_commit_without_fragment_pages(self):
        case=fixture.Tests();case.setUp()
        try:
            frozen,control,_=case.admit();store=operation_records.Store(frozen,case.rt.db)
            page=dict(kind='page',slot=1,layout='text',content=dict(title='Evidence',body='Exact finding. '*250),notes='')
            with self.assertRaisesRegex(ValueError,'Text slot body needs'):
                store.submit(dict(request_key='overflow',revision=0,records=[page]),{})
            self.assertEqual(store.progress()['revision'],0)
            page['content']['body']='Supported finding with explicit unknowns.'
            store.submit(dict(request_key='fit',revision=0,records=[page,dict(kind='end',count=1)]),{})
            self.assertTrue(store.progress()['sealed'])
        finally:case.tearDown()

    def test_fixed_layout_frozen_native_text_and_clickable_sources(self):
        from pptx import Presentation
        body='Target | Source product.\nPrice: unknown. Dimensions: unknown.\nVisible differences remain explicit.\nhttps://example.org/item'
        doc=dict(version=2,title='Evidence',slides=[dict(layout='text',content=dict(title='Source',body=body))])
        namespace={'__name__':'frozen'};exec(pptx_document.validator_source(),namespace)
        self.assertEqual(namespace['validate'](doc),pptx_document.validate(doc))
        raw=pptx_document.create(doc,{})
        if isinstance(raw,tuple):raw=raw[0]
        deck=Presentation(io.BytesIO(raw));self.assertEqual(len(deck.slides),1)
        shape=deck.slides[0].shapes[1];self.assertEqual(shape.text,body)
        self.assertEqual([r.hyperlink.address for p in shape.text_frame.paragraphs for r in p.runs if r.hyperlink.address],['https://example.org/item'])
        shape.text_frame.paragraphs[-1].runs[0].hyperlink.address='https://example.org/wrong'
        changed=io.BytesIO();deck.save(changed)
        with self.assertRaisesRegex(ValueError,'literal source hyperlink changed'):
            pptx_document.inspect(changed.getvalue(),slide_templates.expand(doc),{})
        doc['slides']=[dict(layout='image_grid',content=dict(title='Drawing',per_page=1,fit='contain',items=[dict(image='drawing.png',label='Three views',caption='Estimated dimensions',credit='Synthetic')]))]
        page=slide_templates.expand(doc)['slides'][0]
        picture=next(e for e in page['elements'] if e['type']=='image')
        self.assertGreater(picture['w'],12);self.assertGreater(picture['h'],5.8)
        self.assertEqual(len(slide_templates.expand(doc)['slides']),1)

    def test_batched_records_reject_atomically_then_seal_and_restart(self):
        case=fixture.Tests();case.setUp()
        try:
            frozen,control,_=case.admit();store=operation_records.Store(frozen,case.rt.db)
            page=dict(kind='page',slot=1,layout='text',content=dict(title='Correct candidate',body='Observed evidence.'),notes='')
            bad=dict(kind='end',count=2)
            with self.assertRaises(ValueError):store.submit(dict(request_key='bad',revision=0,records=[page,bad]),{})
            self.assertEqual(store.progress()['revision'],0)
            args=dict(request_key='good',revision=0,records=[page,dict(kind='end',count=1)])
            result=store.submit(args,{})
            resumed=operation_records.Store(frozen,case.rt.db)
            self.assertTrue(resumed.progress()['sealed']);self.assertEqual(resumed.progress()['revision'],1)
            resumed.submit(args,{})
            self.assertEqual(resumed.progress()['revision'],1)
        finally:case.tearDown()

    def test_source_metadata_is_observed_separately_without_changing_text_offsets(self):
        raw=dict(url='https://example.org/item',text='Item\nPrice $42\n',images=[dict(url='https://example.org/photo.jpg',basis='observed-page-metadata')],products=[dict(name='Item')])
        pack=dict(entries=[dict(pages=[raw])])
        journal=product_evidence.source_pack_journal(pack)
        ctx=dict(observations={a['id']:a['result'] for a in journal['actions']})
        result=sourcing_contract.source_read(ctx,dict(observation='page-0-0',offset=0,limit=100))
        self.assertEqual(result['text'],raw['text']);self.assertEqual(result['image_candidates'],raw['images'])
        self.assertEqual(result['price_spans'][0],dict(start=11,end=14,text='$42'))
        self.assertIn('not downloaded or visually verified',result['metadata_authority'])
        self.assertEqual(sourcing_contract.source_index(ctx)[0]['image_candidates'],1)
        ctx['observations']['page-0-0'].pop('images')
        self.assertEqual(raw['images'][0]['url'],'https://example.org/photo.jpg')

if __name__=='__main__':unittest.main()
