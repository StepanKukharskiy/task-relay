"""Controlled recovered-data → specification → real PPTX → correction → delivery."""
import copy
import io
import json
import unittest
from pathlib import Path
from unittest.mock import patch
from orchestrator.storage import transaction
from orchestrator.adapters import ExecutionFactory
from task_relay import pipelines as pipe,production_planning as planning,production_control as pc,delivery_completion as delivery
from task_relay.workflow_usage import report as usage
from tests import test_delivery_completion as fixtures
from tests import test_presentations as presentations
from tests.test_mixed_execution import LocalRegistered,Client
from tests.test_photo_recovery import bundle
from orchestrator.image_sources import merge_bundles

class Tests(unittest.TestCase):
    setUp=fixtures.Tests.setUp
    tearDown=fixtures.Tests.tearDown
    request=fixtures.Tests.request
    stage=fixtures.Tests.stage
    step=fixtures.Tests.step
    answer=fixtures.Tests.answer
    response=fixtures.Tests.response
    create=fixtures.Tests.create
    action=fixtures.Tests.action
    complete_stage=fixtures.Tests.complete_stage

    def test_recovered_versions_reach_real_deck_correction_and_delivery_without_choices(self):
        action=self.action();inventory=action['stages'][0]
        inventory['handoff']['outputs']['outline']['media_type']='application/json'
        catalog=self.stage('catalog',route='production');catalog['handoff']['outputs']['catalog']['media_type']='application/json'
        catalog['handoff']['inputs']=[dict(stage='outline',deliverable='outline',media_type='application/json',consumer='context')]
        photos=self.stage('photos',route='production');photos['handoff']['outputs']['photos']['media_type']='application/zip'
        final=action['stages'][-1];final['capabilities']=['pptx.create'];final['handoff']['outputs']['finish']['media_type']='application/vnd.openxmlformats-officedocument.presentationml.presentation'
        final['handoff']['inputs']=[dict(stage=s,deliverable=s,media_type=m,consumer='context') for s,m in [('outline','application/json'),('catalog','application/json'),('photos','application/zip')]]
        for stage in (inventory,catalog,photos):stage['research_mode']='none'
        action['stages']=[inventory,catalog,photos,final]
        action['completion']['coverage']=dict(inventory=dict(stage='outline',output='outline'),catalog=dict(stage='catalog',output='catalog'),options_per_item=3)
        p=self.create(action=action)
        self.complete_stage(p,'outline',body=dict(items=[dict(id='F01')]))
        old_run,_,_=self.complete_stage(p,'catalog',body=dict(items=[dict(id='F01',options=[dict(product_url='https://store.example/one')])]),completion_expected=False)
        recovered=dict(items=[dict(id='F01',options=[dict(product_url='https://store.example/'+s) for s in ('one','two','three')])])
        new_run,_,_=self.complete_stage(p,'catalog',body=recovered)
        self.assertNotEqual(old_run,new_run)
        raw,_=merge_bundles([bundle(['one','two'],['one']),bundle(['two'],['two'])])
        self.complete_stage(p,'photos',body=raw)
        expected={s:json.loads(self.step(p,s)['result'])['deliverables'][s] for s in ('outline','catalog','photos')}
        pipe.tick(self.state)
        self.answer(dict(kind='plan_production',template='custom',project=None,reference_pack_id=None,research_ids=[],artifact_ids=[],planning_only=False,step_capabilities=['pptx.create'],deliverables=final['deliverables']))
        step=self.step(p,'finish')
        response=presentations.PlanningTests.deck_response(self)
        context=json.loads(self.state.db.execute('SELECT context FROM production_plans WHERE id=?',(step['target'],)).fetchone()[0])
        aliases={key:next(s for s in context['sources'] if s['sha256']==value['sha256']) for key,value in expected.items()}
        response['deliverable_map']={'finish':dict(task='deck',output='candidate.pptx')}
        for t in response['plan']['tasks']:t.pop('user_gate',None)
        # Bind all exact final context versions to author, operation and reviews.
        for t in response['plan']['tasks']:
            if t.get('execution'):continue
            t.setdefault('inputs',[]).extend(dict(artifact=a['artifact'],path='evidence/'+s+('.zip' if s=='photos' else '.json'),purpose='Exact '+s+' source',authority='Exact reviewed upstream version') for s,a in aliases.items())
        deck=response['plan']['tasks'][2]
        deck['inputs'].append(dict(artifact=aliases['photos']['artifact'],path='photos.zip',purpose='Exact recovered source images',authority='Reviewed source ZIP',media_type='application/zip'))
        planning.Worker(self.state,lambda *_:(json.dumps(response),{})).tick()
        row=self.state.db.execute('SELECT * FROM production_plans WHERE id=?',(step['target'],)).fetchone()
        self.assertEqual(row['status'],'ready',row['error'])
        pipe.tick(self.state)
        self.assertEqual(self.state.db.execute('SELECT status FROM production_plans WHERE id=?',(row['id'],)).fetchone()[0],'started')
        native=LocalRegistered(Client());self.rt.factory=ExecutionFactory(self.factory,native)
        worker=pc.Worker(self.state,lambda _:self.rt);worker.tick();run=row['run'] or self.state.db.execute('SELECT run FROM production_plans WHERE id=?',(row['id'],)).fetchone()[0]
        def finish_author(text):
            aid=self.rt.task(run,'produce')['latest'];frozen=self.factory.sessions[aid]['frozen']
            self.assertTrue({a['sha256'] for a in expected.values()} <= {i['sha256'] for i in frozen['inputs']})
            self.factory.finish(aid)
            current_catalog=json.loads((self.factory.sessions[aid]['workspace']/'evidence/catalog.json').read_text())
            source_urls=[o['product_url'] for o in current_catalog['items'][0]['options']]
            text=text+'\n'+'\n'.join(source_urls)
            value=dict(version=1,title='Controlled recovery fixture',slides=[dict(elements=[dict(type='text',x=1,y=1,w=10,h=1,text=text,font_size=28),dict(type='image',x=1,y=2,w=4,h=3,path='photos.zip/images/two.jpg')])])
            (self.factory.sessions[aid]['workspace']/'slides.json').write_text(json.dumps(value));worker.tick()
        def review_spec():
            self.factory.finish(self.rt.task(run,'review')['latest'],decision='accept')
            for _ in range(6):
                worker.tick()
                if self.rt.task(run,'deck-review')['status'] in ('launching','running'):return
            self.fail(str(self.rt.status(run)))
        finish_author('First candidate');review_spec()
        first=self.rt.output(run,'deck','candidate.pptx');self.assertIsNotNone(first)
        op=self.state.db.execute('SELECT frozen FROM production_attempts WHERE id=?',(first['attempt'],)).fetchone()[0]
        self.assertIn(expected['photos']['sha256'],{i['sha256'] for i in json.loads(op)['inputs']})
        self.factory.finish(self.rt.task(run,'deck-review')['latest'],decision='revise')
        for _ in range(8):
            worker.tick()
            if self.rt.task(run,'produce')['attempts']==2:break
        self.assertEqual(self.rt.task(run,'produce')['attempts'],2)
        finish_author('Corrected candidate');review_spec()
        self.factory.finish(self.rt.task(run,'deck-review')['latest'],decision='accept')
        for _ in range(8):
            worker.tick();pipe.tick(self.state)
            if self.step(p,'finish')['status']=='completed':break
        pipe.tick(self.state)
        refreshed=self.state.db.execute('SELECT * FROM relay_pipelines WHERE id=?',(p['id'],)).fetchone()
        self.assertEqual(refreshed['status'],'completed',self.step(p,'finish')['error'])
        delivery.check_final(self.state,refreshed)
        final_source=json.loads(self.step(p,'finish')['result'])['deliverables']['finish']
        from pptx import Presentation
        pptx=Presentation(io.BytesIO(Path(final_source['path']).read_bytes()))
        self.assertEqual(pptx.slides[0].shapes[0].text,'Corrected candidate\nhttps://store.example/one\nhttps://store.example/two\nhttps://store.example/three')
        self.assertEqual(pptx.slides[0].shapes[1].image.blob,__import__('tests.test_image_sources',fromlist=['photo']).photo())
        self.assertEqual(len(native.calls),4)
        self.assertEqual(self.state.db.execute('SELECT count(*) FROM production_decisions').fetchone()[0],0)
        measured=usage(self.state,p['id'],now=999999999999)
        self.assertIn(old_run,measured['runs']);self.assertIn(new_run,measured['runs']);self.assertIn(run,measured['runs'])
        self.assertIsNotNone(measured['finished_at'])
        self.assertEqual(measured['elapsed_seconds'],usage(self.state,p['id'],now=1)['elapsed_seconds'])
        self.assertGreater(measured['totals']['input_tokens']['reported'],0)
        self.assertGreater(measured['totals']['output_tokens']['unknown_records'],0)
