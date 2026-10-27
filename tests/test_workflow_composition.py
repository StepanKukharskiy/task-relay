"""Full synthetic room workflow through admission, dispatch and final delivery."""
import io
import json
from pathlib import Path
import unittest
from unittest.mock import patch
from orchestrator.adapters import ExecutionFactory
from task_relay import pipelines as pipe, production_control as pc, production_planning as planning
from task_relay import delivery_completion as delivery, gemini
from task_relay.workflow_usage import report as usage
from tests import test_delivery_completion as fixtures, test_presentations as presentations
from tests.workflow_decoys import (PRODUCTS,SourcePages,NativeBoundary,reviewed,
                                  source_graph,photo_graph,observed_manifest,bind_context,bound_json)
from tests.test_image_sources import photo
from tests.test_web_sources import CONFIG

class Tests(unittest.TestCase):
    setUp=fixtures.Tests.setUp
    tearDown=fixtures.Tests.tearDown
    request=fixtures.Tests.request
    stage=fixtures.Tests.stage
    step=fixtures.Tests.step
    answer=fixtures.Tests.answer
    response=fixtures.Tests.response
    create=fixtures.Tests.create

    def assemble(self,count_policy='required'):
        stages=[self.stage(ident,route='production') for ident in ('inventory','catalog','photos','deck')]
        for stage,media,caps in zip(stages,['application/json','application/json','application/zip',
                'application/vnd.openxmlformats-officedocument.presentationml.presentation'],
                [[],['web.sources.prepared'],['images.fetch'],['pptx.create']]):
            ident=stage['id'];stage['handoff']['outputs'][ident]['media_type']=media
            stage['capabilities']=caps
        stages[0]['research_mode']='none'
        for index,stage in enumerate(stages):
            stage['handoff']['inputs']=[dict(stage=prior['id'],deliverable=prior['id'],
                media_type=prior['handoff']['outputs'][prior['id']]['media_type'],consumer='context')
                for prior in stages[:index]]
        return dict(kind='plan_pipeline',contract_version=1,title='Synthetic complete sourcing',planning_only=False,
                    stages=stages,completion=dict(version=1,stage='deck',output='deck',coverage=dict(
                    inventory=dict(stage='inventory',output='inventory'),catalog=dict(stage='catalog',output='catalog'),options_per_item=3,count_policy=count_policy)))

    def compile_stage(self,p,ident,value):
        pipe.tick(self.state)
        stage=self.step(p,ident);spec=next(s for s in json.loads(p['spec'])['stages'] if s['id']==ident)
        self.answer(dict(kind='plan_production',template='custom',project=None,reference_pack_id=None,
                         research_ids=[],reference_ids=[],artifact_ids=[],planning_only=False,step_capabilities=spec['capabilities'],
                         deliverables=spec['deliverables'],**({'research_mode':'none'} if ident=='inventory' else {})))
        stage=self.step(p,ident)
        row=self.state.db.execute('SELECT * FROM production_plans WHERE id=?',(stage['target'],)).fetchone()
        self.assertIsNotNone(row, [tuple(r) for r in self.state.db.execute('SELECT * FROM orchestrator_chat_errors')])
        value=bind_context(value,row)
        for task in value['plan']['tasks']:task.pop('user_gate',None)
        planning.Worker(self.state,lambda *_:(json.dumps(value),{'input_tokens':17,'output_tokens':9})).tick()
        row=self.state.db.execute('SELECT * FROM production_plans WHERE id=?',(stage['target'],)).fetchone()
        self.assertEqual(row['status'],'ready',row['error'])
        if ident=='catalog' and json.loads(row['context']).get('sourcing_count_policy')=='target':
            compiled=json.loads(row['plan'])
            reviewer=next(t for t in compiled['tasks'] if t['id']=='sources-review')
            self.assertTrue(any('not a blocking requirement' in criterion for criterion in reviewer['criteria']))
        pipe.tick(self.state)
        row=self.state.db.execute('SELECT * FROM production_plans WHERE id=?',(stage['target'],)).fetchone()
        self.assertEqual(row['status'],'started',self.step(p,ident)['error'])
        return row['run']

    def drive(self,run,fragments,revise=None,expect_uncertain=False,shortage=False):
        handled=set();revised=False
        worker=pc.Worker(self.state,lambda _:self.rt)
        for _ in range(60):
            worker.tick()
            for aid,session in list(self.factory.sessions.items()):
                frozen=session['frozen']
                attempt=self.state.db.execute('SELECT run,task FROM production_attempts WHERE id=?',(aid,)).fetchone()
                if attempt['run']!=run or aid in handled:continue
                handled.add(aid);tid=attempt['task']
                decision='accept' if frozen.get('review_of') else 'delivered'
                if tid==revise and not revised:decision='revise';revised=True
                self.factory.finish(aid,decision=decision)
                if shortage and tid=='sources-review':
                    path=session['workspace']/'.relay/result.json'
                    report=json.loads(path.read_text());report.update(decision='blocked',findings=[
                        dict(category='requirements',code='missing_options',message='Two required options lack product evidence',evidence='Exact source pack has one product page'),
                        dict(category='advisory',code='stock_caveat',message='Stock may change',evidence='Synthetic observation timestamp')])
                    path.write_text(json.dumps(report))
                if tid in fragments:
                    value=fragments[tid](session)
                    target=session['workspace']/frozen['outputs'][0]['path']
                    target.write_text(json.dumps(value))
            status=self.rt.status(run)
            if status['status']=='blocked' and shortage:return
            if status['status']=='uncertain' and expect_uncertain:return
            if status['status']=='completed':return
            self.assertNotIn(status['status'],('blocked','uncertain','awaiting_user'),str(self.native.errors)+' '+str(status))
        self.fail('Synthetic run did not complete: '+str(self.rt.status(run)))

    def run_room(self,correction=False,unknown=False,shortage=False,available=None):
        products=PRODUCTS if available is None else [dict(id='option-'+str(i),product_url='https://retailer.example/option-'+str(i),price_usd=200+i,image_url='https://images.example/option-'+str(i)+'.jpg') for i in range(available)]
        SourcePages.calls=[];SourcePages.error=None;SourcePages.products=PRODUCTS[:1] if shortage else products
        self.addCleanup(setattr,SourcePages,'products',PRODUCTS)
        self.addCleanup(setattr,SourcePages,'error',None)
        native=NativeBoundary();self.native=native;self.rt.factory=ExecutionFactory(self.factory,native)
        # Registered uploaded bytes travel through normal retained-input dispatch.
        raw=photo();upload=self.state.media_dir.parent/'production-guides/123/room.jpg'
        upload.parent.mkdir(parents=True);upload.write_bytes(raw)
        import hashlib
        self.state.db.execute('''INSERT INTO production_uploads
            (id,run,file_id,filename,caption,declared_size,status,path,sha256,bytes)
            VALUES (123,'@orchestrator','synthetic','room.jpg','',?,'ready',?,?,?)''',
            (len(raw),str(upload),hashlib.sha256(raw).hexdigest(),len(raw)))
        self.state.put('orchestrator-attachments:1',[123]);self.state.db.commit()
        with patch.object(gemini,'read_config',return_value=CONFIG):
            p=self.create(action=self.assemble('required' if available is None else 'target'));runs=[]
            value=self.response();value['plan']['tasks']=reviewed(value,'inventory','inventory.json')
            value['deliverable_map']=dict(inventory=dict(task='inventory',output='inventory.json'))
            run=self.compile_stage(p,'inventory',value);runs.append(run)
            def inventory(session):
                self.assertIn(hashlib.sha256(raw).hexdigest(),{i['sha256'] for i in session['frozen']['inputs']})
                return dict(version=1,items=[dict(id='F01',label='Synthetic chair')])
            self.drive(run,dict(inventory=inventory));pipe.tick(self.state)
            inventory_version=json.loads(self.step(p,'inventory')['result'])['deliverables']['inventory']
            value=source_graph(self.response());run=self.compile_stage(p,'catalog',value);runs.append(run)
            def queries(session):
                support=[i for i in session['frozen']['inputs'] if i['path'].endswith('/validate.py')]
                self.assertEqual(len(support),1)
                ns={'__name__':'frozen'};exec((session['workspace']/support[0]['path']).read_text(),ns)
                observed=bound_json(session,inventory_version['sha256'])
                searches=['synthetic chair US']+(['synthetic chair US extra'] if len(products)>3 else [])
                data=dict(version=1,targets=[dict(id=observed['items'][0]['id'],queries=searches,required_options=3)])
                ns['validate_prepared'](data)
                return data
            def catalog(session):
                pack=json.loads((session['workspace']/'delivery/source-pack.json').read_text())
                self.assertEqual(pack['entries'][0]['target_ids'],['F01'])
                pages=list({page['url']:page for entry in pack['entries'] for page in entry['pages']}.values());self.assertEqual(len(pages),len(products))
                options=[json.loads(page['text']) for page in pages]
                self.assertEqual(options,products)
                return dict(version=1,items=[dict(id='F01',options=options)])
            if unknown:SourcePages.error=gemini.ProviderError('Synthetic unknown source submission',uncertain=True)
            self.drive(run,dict(queries=queries,catalog=catalog),expect_uncertain=unknown,shortage=shortage);pipe.tick(self.state)
            if shortage:
                receipt=self.state.db.execute('SELECT * FROM completion_recoveries WHERE parent=?',(run,)).fetchone()
                self.assertIsNotNone(receipt,self.step(p,'catalog')['error'])
                baseline=json.loads(receipt['baseline'])
                self.assertEqual([f['category'] for f in baseline['gaps'][0]['findings']],['requirements','advisory'])
                self.assertEqual(self.rt.task(run,'sources')['attempts'],1)
                self.assertEqual(self.rt.task(run,'catalog')['attempts'],0)
                self.assertEqual(self.step(p,'photos')['status'],'pending')
                self.assertEqual(self.step(p,'deck')['status'],'pending')
                self.assertEqual(len(native.calls),1)
                self.assertNotEqual(self.step(p,'catalog')['target'],run)
                self.assertEqual(self.state.db.execute('SELECT count(*) FROM production_decisions').fetchone()[0],0)
                SourcePages.products=PRODUCTS
                run=self.compile_stage(p,'catalog',source_graph(self.response()));runs.append(run)
                self.drive(run,dict(queries=queries,catalog=catalog));pipe.tick(self.state)
            if unknown:
                for _ in range(3):
                    pc.Worker(self.state,lambda _:self.rt).tick();pipe.tick(self.state)
                self.assertEqual(self.rt.task(run,'sources')['attempts'],1)
                self.assertEqual(SourcePages.calls,[('search','synthetic chair US')])
                self.assertEqual(len(native.calls),1)
                self.assertEqual(self.step(p,'photos')['status'],'pending')
                self.assertEqual(self.step(p,'deck')['status'],'pending')
                self.assertEqual(self.rt.task(run,'catalog')['attempts'],0)
                with self.assertRaisesRegex(ValueError,'Missing upstream artifact'):
                    self.rt.output(run,'sources','delivery/source-pack.json')
                self.assertEqual(self.step(p,'inventory')['status'],'completed')
                return
            run=self.compile_stage(p,'photos',photo_graph(self.response()));runs.append(run)
            catalog_version=json.loads(self.step(p,'catalog')['result'])['deliverables']['catalog']
            self.drive(run,dict(manifest=lambda s:observed_manifest(bound_json(s,catalog_version['sha256'])['items'][0]['options'])));pipe.tick(self.state)
            expected={ident:json.loads(self.step(p,ident)['result'])['deliverables'][ident]
                      for ident in ('inventory','catalog','photos')}
            value=presentations.PlanningTests.deck_response(self)
            value['deliverable_map']=dict(deck=dict(task='deck',output='candidate.pptx'))
            value['plan']['tasks'][2]['inputs'].append(dict(artifact=expected['photos']['artifact'],path='photos.zip',
                purpose='Exact reviewed product photos',authority='Synthetic source bytes',media_type='application/zip'))
            run=self.compile_stage(p,'deck',value);runs.append(run)
            def slides(session):
                self.assertTrue({s['sha256'] for s in expected.values()} <= {i['sha256'] for i in session['frozen']['inputs']})
                text='Corrected' if self.rt.task(run,'produce')['attempts']>1 else 'Initial'
                current=bound_json(session,expected['catalog']['sha256'])
                text+='\n'+'\n'.join(p['product_url']+' $'+str(p['price_usd']) for p in current['items'][0]['options'])
                return dict(version=1,title='Synthetic room',slides=[dict(elements=[
                    dict(type='text',x=1,y=1,w=10,h=1,text=text,font_size=18),
                    dict(type='image',x=1,y=2,w=4,h=3,path='photos.zip/images/'+current['items'][0]['options'][0]['id']+'.jpg')])])
            self.drive(run,dict(produce=slides),revise='deck-review' if correction else None)
            pipe.tick(self.state);pipe.tick(self.state)
        p=self.state.db.execute('SELECT * FROM relay_pipelines WHERE id=?',(p['id'],)).fetchone()
        self.assertEqual(p['status'],'completed',self.step(p,'deck')['error']);delivery.check_final(self.state,p)
        final=json.loads(self.step(p,'deck')['result'])['deliverables']['deck']
        from pptx import Presentation
        deck=Presentation(io.BytesIO(Path(final['path']).read_bytes()))
        self.assertTrue(deck.slides[0].shapes[0].text.startswith('Corrected' if correction else 'Initial'))
        self.assertEqual(deck.slides[0].shapes[1].image.blob,photo())
        expected_calls=[('search','synthetic chair US')]+[('fetch',p['product_url']) for p in products[:6]]
        if len(products)>3:expected_calls+=[('search','synthetic chair US extra')]
        if shortage:expected_calls=[('search','synthetic chair US'),('fetch',PRODUCTS[0]['product_url'])]+expected_calls
        self.assertEqual(SourcePages.calls,expected_calls)
        self.assertEqual(native.downloads,[p['image_url'] for p in products])
        self.assertEqual(len(native.calls),4+2*int(correction)+int(shortage))
        self.assertEqual(self.state.db.execute('SELECT count(*) FROM production_decisions').fetchone()[0],0)
        measured=usage(self.state,p['id']);self.assertEqual(set(measured['runs']),set(runs))
        self.assertIsNotNone(measured['finished_at']);self.assertGreaterEqual(measured['elapsed_seconds'],0)
        self.assertGreater(measured['totals']['input_tokens']['reported'],0)
        self.assertGreater(measured['totals']['output_tokens']['unknown_records'],0)

    def test_all_stages_execute_through_production_entry_points(self):self.run_room()
    def test_final_review_correction_does_not_repeat_source_or_photo_calls(self):self.run_room(correction=True)

    def test_unknown_source_submission_is_not_replayed_or_delivered(self):self.run_room(unknown=True)

    def test_incompatible_connection_is_rejected_before_planning_or_workers(self):
        from orchestrator.handoff_contracts import compile_workflow,ContractError
        stages=self.assemble()['stages']
        stages[1]['handoff']['inputs'][0]['media_type']='text/plain'
        with self.assertRaises(ContractError):compile_workflow(stages,required=True)
        self.assertEqual(self.factory.calls,[])
        self.assertEqual(self.state.db.execute('SELECT count(*) FROM production_plans').fetchone()[0],0)

    def test_short_source_pack_with_advisory_queues_recovery_without_acceptance(self):self.run_room(shortage=True)

    def test_target_count_delivers_one_supported_option_without_waiting_for_three(self):self.run_room(available=1)
    def test_target_count_delivers_all_five_supported_options_without_truncation(self):self.run_room(available=5)
