"""Controlled O14 transfer test: reviewed spreadsheet facts to native PPTX text."""

from contextlib import closing
import io
import json
import os
from pathlib import Path
import sqlite3
import tempfile
import unittest
from types import SimpleNamespace
import uuid
from zipfile import ZipFile

from openpyxl import Workbook
from orchestrator import pptx_document, pptx_edit
from orchestrator.runtime import Runtime
from task_relay.bridge import State
from task_relay import (presentation_revisions as slides, production_control as pc,
                        revision_review, workflow_files)


def source(path, plant):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Facts'
    sheet.append(['Key', 'Value'])
    sheet.append(['plant', plant])
    sheet.append(['climate', 'Temperate climate'])
    book.save(path)
    book.close()


def deck():
    spec = {'version': 1, 'title': 'Plant review', 'slides': [
        {'elements': [{'type': 'text', 'x': 0.7, 'y': 0.6, 'w': 11, 'h': 1,
                       'font_size': 36, 'text': 'North Plant'}]},
        {'elements': [{'type': 'text', 'x': 0.7, 'y': 0.6, 'w': 11, 'h': 1,
                       'font_size': 36, 'text': 'Temperate climate'}]},
        {'elements': [{'type': 'text', 'x': 0.7, 'y': 0.6, 'w': 11, 'h': 1,
                       'font_size': 36, 'text': 'Unverified estimate'}]},
    ]}
    return pptx_document.create(spec)[0]


class PresentationRevisionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name).resolve()
        self.state = State(self.root / 'data/state.sqlite')
        self.rt = Runtime(pc.root(self.state), connection=self.state.db)
        self.job = 'job-' + 'b' * 24
        self.state.db.execute('''INSERT INTO relay_pipelines
            (id,request_id,request,title,spec,channel,provider,model,status,created)
            VALUES (?,?,?,?,?,?,?,?,?,?)''',
            (self.job, 1, 'Revise the changed plant slide.', 'O14 PPTX fixture',
             '{}', 'telegram', 'fixture', 'fixture', 'active', 1))
        self.state.db.commit()
        source(self.root / 'old.xlsx', 'North Plant')
        source(self.root / 'replacement.xlsx', 'South Plant')
        (self.root / 'baseline.pptx').write_bytes(deck())
        self.old = self.rt.register(self.root / 'old.xlsx', 'Reviewed source')
        self.new = self.rt.register(self.root / 'replacement.xlsx', 'Replacement source')
        self.baseline = self.rt.register(self.root / 'baseline.pptx', 'Baseline deck')
        listing = pptx_edit.inspect((self.root / 'baseline.pptx').read_bytes())
        self.shapes = [page['shapes'][0]['shape_id'] for page in listing['slides']]
        for slide, key, cell, value in ((1, 'plant', 'B2', 'North Plant'),
                                         (2, 'climate', 'B3', 'Temperate climate')):
            slides.record_link(self.state.db, job=self.job, entity_key=key,
                predicate='displayed_value', value=value, source_artifact=self.old,
                source_sheet='Facts', source_key_column='A', source_cell=cell,
                evidence='Fixture source row reviewed against the exact slide run.',
                output_artifact=self.baseline, slide=slide,
                shape_id=self.shapes[slide-1], reviewer='source-reviewer')
            slides.record_coverage(self.state.db, job=self.job,
                output_artifact=self.baseline, slide=slide,
                shape_id=self.shapes[slide-1], state='complete', note='Exact reviewed link.')
        slides.record_coverage(self.state.db, job=self.job,
            output_artifact=self.baseline, slide=3, shape_id=self.shapes[2],
            state='incomplete', note='No reviewed source for this statement.')

    def tearDown(self):
        self.state.db.close()
        self.temp.cleanup()

    def plan(self):
        return slides.plan_impact(self.state.db, job=self.job, old_source=self.old,
            replacement_source=self.new, baseline_artifact=self.baseline)

    def test_candidate_preserves_independent_slide_and_requires_review(self):
        prior = self.state.db.total_changes
        plan = self.plan()
        self.assertEqual(self.state.db.total_changes, prior)
        self.assertEqual([len(plan[s]) for s in ('affected', 'unaffected', 'unknown')], [1, 1, 1])
        self.assertEqual(plan['reviewed_impact']['kind'], 'presentation_text')
        candidate = slides.revise_pptx(self.rt, plan=plan, reviewer='candidate-author')
        revision, artifact, checked, checks = slides.verified_candidate(self.state.db, candidate)
        self.assertEqual(checked, plan)
        self.assertEqual(checks['changed_slides'], [1])
        self.assertEqual(checks['visual_review'], 'not_performed')
        old = (self.root / 'baseline.pptx').read_bytes()
        new = Path(artifact['blob']).read_bytes()
        with ZipFile(io.BytesIO(old)) as before, ZipFile(io.BytesIO(new)) as after:
            self.assertEqual(set(before.namelist()), set(after.namelist()))
            self.assertEqual([name for name in before.namelist()
                if before.read(name) != after.read(name)], ['ppt/slides/slide1.xml'])
        listing = pptx_edit.inspect(new)
        self.assertEqual(listing['slides'][0]['shapes'][0]['text_runs'], ['South Plant'])
        self.assertEqual(listing['slides'][1]['shapes'][0]['text_runs'], ['Temperate climate'])
        self.assertEqual(listing['slides'][2]['shapes'][0]['text_runs'], ['Unverified estimate'])
        with self.assertRaisesRegex(ValueError, 'independently accepted'):
            slides.select_candidate(self.state.db, candidate_artifact=candidate,
                selected_by='fixture-user', receipt='Synthetic user action')
        self.assertEqual(slides.records(self.state.db, self.job)['selections'], [])
        report = workflow_files.sync(self.state, self.job)
        with closing(sqlite3.connect(report / '.relay/job.sqlite')) as db:
            self.assertEqual(db.execute('SELECT count(*) FROM presentation_links').fetchone()[0], 2)
            self.assertEqual(db.execute('SELECT count(*) FROM presentation_revisions').fetchone()[0], 1)
            self.assertEqual(db.execute('SELECT count(*) FROM reviewed_links WHERE kind=?',
                ('presentation_text',)).fetchone()[0], 2)
        view = revision_review.detail(candidate, type('Paths', (), {
            'state': self.root / 'data/state.sqlite',
            'generated': self.root / 'data/generated'})())
        self.assertTrue(view['verified'])
        self.assertEqual(view['kind'], 'presentation_text')
        self.assertTrue(view['can_review'])
        with self.assertRaisesRegex(ValueError, 'independent'):
            slides.review_candidate(self.state.db, candidate_artifact=candidate,
                decision='accept', reviewer='candidate-author', note='Self review')
        slides.review_candidate(self.state.db, candidate_artifact=candidate,
            decision='accept', reviewer='independent-reviewer', note='Fixture validated')
        slides.select_candidate(self.state.db, candidate_artifact=candidate,
            selected_by='fixture-user', receipt='Synthetic exact candidate selection')
        selected = json.loads((workflow_files.sync(self.state, self.job) / 'manifest.json').read_text())
        self.assertTrue(next(a['selected'] for a in selected['artifacts'] if a['id'] == candidate))

    def test_ambiguous_replacement_is_unknown_and_does_not_register_candidate(self):
        source(self.root / 'ambiguous.xlsx', 'South Plant')
        from openpyxl import load_workbook
        book = load_workbook(self.root / 'ambiguous.xlsx')
        book.active.append(['plant', 'Another Plant'])
        book.save(self.root / 'ambiguous.xlsx')
        book.close()
        replacement = self.rt.register(self.root / 'ambiguous.xlsx', 'Ambiguous source')
        plan = slides.plan_impact(self.state.db, job=self.job, old_source=self.old,
            replacement_source=replacement, baseline_artifact=self.baseline)
        self.assertEqual([len(plan[s]) for s in ('affected', 'unaffected', 'unknown')], [0, 1, 2])
        with self.assertRaisesRegex(ValueError, 'No reviewed affected'):
            slides.revise_pptx(self.rt, plan=plan, reviewer='candidate-author')
        self.assertEqual(slides.records(self.state.db, self.job)['revisions'], [])

    def test_unlinked_source_cannot_claim_the_deck_unaffected(self):
        source(self.root / 'unlinked.xlsx', 'North Plant')
        unrelated = self.rt.register(self.root / 'unlinked.xlsx', 'Unlinked source version')
        with self.assertRaisesRegex(ValueError, 'no reviewed link'):
            slides.plan_impact(self.state.db, job=self.job, old_source=unrelated,
                replacement_source=self.new, baseline_artifact=self.baseline)

    def test_changed_candidate_bytes_block_review_and_selection(self):
        candidate = slides.revise_pptx(self.rt, plan=self.plan(), reviewer='candidate-author')
        row = self.state.db.execute('SELECT blob FROM production_artifacts WHERE id=?',
                                    (candidate,)).fetchone()
        os.chmod(row['blob'], 0o600)
        Path(row['blob']).write_bytes(b'broken candidate')
        with self.assertRaisesRegex(ValueError, 'missing, changed or unsupported'):
            slides.review_candidate(self.state.db, candidate_artifact=candidate,
                decision='accept', reviewer='independent-reviewer', note='Cannot verify')
        with self.assertRaisesRegex(ValueError, 'missing, changed or unsupported'):
            slides.select_candidate(self.state.db, candidate_artifact=candidate,
                selected_by='fixture-user', receipt='Synthetic exact candidate selection')
        records = slides.records(self.state.db, self.job)
        self.assertEqual(records['reviews'], [])
        self.assertEqual(records['selections'], [])

    def test_desktop_decision_receipts_bind_exact_pptx_candidate(self):
        plan = self.plan()
        candidate = slides.revise_pptx(self.rt, plan=plan, reviewer='candidate-author')
        workflow_files.sync(self.state, self.job)
        paths = SimpleNamespace(state=self.root / 'data/state.sqlite',
                                generated=self.root / 'data/generated')
        view = revision_review.detail(candidate, paths)
        with self.assertRaisesRegex(ValueError, 'changed'):
            revision_review.decide(candidate_artifact=candidate, verb='accept',
                actor='independent-reviewer', note='Checked exact deck',
                expected_sha256=view['candidate_sha256'], expected_plan_digest='0' * 64,
                request_id=str(uuid.uuid4()), paths=paths)
        self.assertEqual(slides.records(self.state.db, self.job)['reviews'], [])
        request_id = str(uuid.uuid4())
        args = dict(candidate_artifact=candidate, verb='accept',
            actor='independent-reviewer', note='Checked exact deck',
            expected_sha256=view['candidate_sha256'],
            expected_plan_digest=plan['digest'], request_id=request_id, paths=paths)
        first = revision_review.decide(**args)
        second = revision_review.decide(**args)
        self.assertEqual(first['result_id'], second['result_id'])
        self.assertEqual(len(slides.records(self.state.db, self.job)['reviews']), 1)
        selection = revision_review.decide(candidate_artifact=candidate, verb='select',
            actor='fixture-user', note='Choose this exact candidate',
            expected_sha256=view['candidate_sha256'],
            expected_plan_digest=plan['digest'], request_id=str(uuid.uuid4()), paths=paths)
        self.assertTrue(selection['detail']['selection'])
        self.assertEqual(len(slides.records(self.state.db, self.job)['selections']), 1)


if __name__ == '__main__':
    unittest.main()
