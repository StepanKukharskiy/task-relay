"""Synthetic PPTX creation and recovery tests; no account or native app needed."""
import copy
import io
import json
from pathlib import Path
import tempfile
import time
import unittest
from unittest.mock import patch

from orchestrator import contracts, execution, pptx_document
from orchestrator.adapters import ExecutionFactory, RegisteredFactory
from orchestrator.runtime import Runtime
from task_relay import presentations, production_planning as planning
from tests import test_mixed_execution as mixed, test_production_planning as planning_fixtures
from tests.test_orchestrator import plan, task


def fixture():
    return {'version': 1, 'title': 'Quarterly results', 'slides': [
        {'notes': 'Synthetic fixture. Пример.', 'elements': [
            {'type': 'text', 'x': 0.7, 'y': 0.6, 'w': 12, 'h': 0.8,
             'text': 'Quarterly results', 'font_size': 40, 'bold': True},
            {'type': 'text', 'x': 0.7, 'y': 2, 'w': 11, 'h': 2,
             'text': 'Editable text\nEnglish and русский', 'font_size': 28},
            {'type': 'shape', 'shape': 'rectangle', 'x': 0.7, 'y': 5, 'w': 4, 'h': 0.8, 'text': 'Draft for review'}]},
        {'elements': [
            {'type': 'text', 'x': 0.7, 'y': 0.6, 'w': 12, 'h': 0.8, 'text': 'Data table', 'font_size': 36},
            {'type': 'table', 'x': 0.7, 'y': 2, 'w': 11.9, 'h': 2.7,
             'rows': [['Quarter', 'Units'], ['Q1', '12'], ['Q2', '18']], 'font_size': 24}]},
        {'elements': [
            {'type': 'text', 'x': 0.7, 'y': 0.6, 'w': 12, 'h': 0.8, 'text': 'Units by quarter', 'font_size': 36},
            {'type': 'chart', 'chart': 'column', 'x': 0.7, 'y': 1.8, 'w': 11.9, 'h': 5,
             'categories': ['Q1', 'Q2'], 'series': [{'name': 'Units', 'values': [12, 18]}]}]}]}


def pptx_step(inputs, **kwargs):
    return dict(id='deck', role='procedure', objective='Create an editable deck', instruction='Create the exact specified slides.',
        execution={'capability': 'pptx.create', 'version': 1, 'parameters': {}}, inputs=inputs,
        outputs=[{'path': 'candidate.pptx', 'purpose': 'Editable candidate', 'media_type': pptx_document.MIME}],
        criteria=execution.REGISTRY['pptx.create']['criteria'].copy(), user_gate='Select the deck', **kwargs)


class DocumentTests(unittest.TestCase):
    def test_notes_master_is_discoverable_and_missing_registration_is_rejected(self):
        import zipfile
        from lxml import etree
        ns = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
              'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
        value = fixture()
        raw, _ = pptx_document.create(value)
        with zipfile.ZipFile(io.BytesIO(raw)) as package:
            root = etree.fromstring(package.read('ppt/presentation.xml'))
            listing = root.find('p:notesMasterIdLst', ns)
            self.assertIsNotNone(listing)
            rid = listing[0].get('{'+ns['r']+'}id')
            rels = etree.fromstring(package.read('ppt/_rels/presentation.xml.rels'))
            relation = next(r for r in rels if r.get('Id') == rid)
            self.assertTrue(relation.get('Type').endswith('/notesMaster'))
            self.assertIn('ppt/'+relation.get('Target'), package.namelist())
            self.assertLess(root.index(listing), root.index(root.find('p:sldIdLst', ns)))
            root.remove(listing)
            broken = io.BytesIO()
            with zipfile.ZipFile(broken, 'w', zipfile.ZIP_DEFLATED) as dest:
                for item in package.infolist():
                    dest.writestr(item, etree.tostring(root) if item.filename == 'ppt/presentation.xml' else package.read(item))
        with self.assertRaisesRegex(ValueError, 'notes master registration'):
            pptx_document.inspect(broken.getvalue(), value, {})

    def test_no_notes_does_not_create_notes_master_and_registration_is_idempotent(self):
        from pptx import Presentation
        from pptx.oxml.ns import qn
        value = fixture()
        value['slides'][0].pop('notes')
        raw, _ = pptx_document.create(value)
        deck = Presentation(io.BytesIO(raw))
        self.assertIsNone(deck.part._element.find(qn('p:notesMasterIdLst')))
        raw, _ = pptx_document.create(fixture())
        deck = Presentation(io.BytesIO(raw)); before = deck.part._element.xml
        pptx_document._register_notes_master(deck)
        self.assertEqual(deck.part._element.xml, before)

    def test_native_content_survives_reopen_and_remains_editable(self):
        from pptx import Presentation
        raw, evidence = pptx_document.create(fixture())
        self.assertEqual(evidence['native_objects'], {'text': 4, 'shape': 1, 'table': 1, 'chart': 1, 'image': 0})
        deck = Presentation(io.BytesIO(raw))
        deck.slides[0].shapes[0].text = 'Revised title'
        deck.slides[1].shapes[1].table.cell(1, 1).text = '15'
        modified = io.BytesIO(); deck.save(modified)
        reopened = Presentation(io.BytesIO(modified.getvalue()))
        self.assertEqual(reopened.slides[0].shapes[0].text, 'Revised title')
        self.assertEqual(reopened.slides[1].shapes[1].table.cell(1, 1).text, '15')
        self.assertEqual(reopened.slides[2].shapes[1].chart.series[0].values, (12.0, 18.0))
        self.assertEqual(evidence['visual_review'], 'not_performed')

    def test_selected_image_embeds_exact_bytes_without_distortion(self):
        from pptx import Presentation
        from tests.test_gemini import PNG
        value = fixture()
        value['slides'] = [{'elements': [{'type': 'image', 'path': 'assets/selected.png', 'x': 1, 'y': 1, 'w': 8, 'h': 4}]}]
        raw, _ = pptx_document.create(value, {'assets/selected.png': PNG})
        image = Presentation(io.BytesIO(raw)).slides[0].shapes[0]
        self.assertEqual(image.image.blob, PNG)
        self.assertAlmostEqual(image.width / image.height, image.image.size[0] / image.image.size[1])
        with self.assertRaisesRegex(ValueError, 'exact declared'):
            pptx_document.create(value)

    def test_invalid_specifications_fail_before_serialization(self):
        variants = []
        bad = fixture(); bad['slides'][0]['elements'][0]['x'] = 99; variants.append(bad)
        bad = fixture(); bad['slides'][2]['elements'][1]['series'][0]['values'][0] = float('nan'); variants.append(bad)
        bad = fixture(); bad['slides'][1]['elements'][1]['rows'][1].append('extra'); variants.append(bad)
        bad = fixture(); bad['script'] = 'ignored code'; variants.append(bad)
        bad = fixture(); bad['slides'][2]['elements'][1]['categories'][0] = '=1+1'; variants.append(bad)
        bad = fixture(); bad['slides'][0]['elements'][0]['text'] = 'bad\x00text'; variants.append(bad)
        bad = fixture(); bad['slides'][0]['elements'] = [{'type': 'image', 'path': '../secret.png', 'x': 0, 'y': 0, 'w': 1, 'h': 1}]; variants.append(bad)
        for value in variants:
            with self.subTest(value=value), self.assertRaises(ValueError):
                pptx_document.validate(value, ['../secret.png'])
        with self.assertRaisesRegex(ValueError, 'Duplicate'):
            pptx_document.load('{"version":1,"version":2}')
        with self.assertRaisesRegex(ValueError, 'byte limit'):
            pptx_document.create(fixture(), max_bytes=10)

    def test_cli_preserves_prior_candidate_and_rejects_undeclared_files(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); source = root / 'slides.json'; source.write_text(json.dumps(fixture()))
            dest = root / 'result'
            with patch('builtins.print'):
                self.assertEqual(presentations.main(['create', str(source), '--output-dir', str(dest)]), 0)
                first = (dest / 'presentation.pptx').read_bytes()
                self.assertEqual(presentations.main(['create', str(source), '--output-dir', str(dest)]), 1)
                self.assertEqual((dest / 'presentation.pptx').read_bytes(), first)
                self.assertEqual(presentations.main(['create', str(source), '--image', '../secret.png', '--output-dir', str(root / 'bad')]), 1)
                self.assertFalse((root / 'bad').exists())
            receipt = json.loads((dest / 'receipt.json').read_text())
            self.assertEqual(receipt['validation']['slide_count'], 3)


class GraphTests(unittest.TestCase):
    setUp = mixed.Tests.setUp
    tearDown = mixed.Tests.tearDown

    def source_input(self):
        source = self.root / 'slides.json'; source.write_text(json.dumps(fixture()))
        aid = self.rt.register(source, 'Exact slide specification', path='slides.json')
        return dict(artifact=aid, path='slides.json', purpose='Slide data', authority='Selected source', media_type='application/json')

    def test_creation_review_selection_and_restart_preserve_identity(self):
        producer = task('prepare')
        producer['outputs'] = [{'path': 'slides.json', 'purpose': 'Slide specification', 'media_type': 'application/json'}]
        spec_review = task('spec-review', dependencies=['prepare'], review_of='prepare',
            inputs=[dict(from_task='prepare', output='slides.json', path='slides.json', purpose='Review slide data', authority='Candidate', media_type='application/json')])
        deck = pptx_step([dict(from_task='prepare', output='slides.json', path='slides.json', purpose='Slide data', authority='Reviewed source', media_type='application/json')], dependencies=['prepare', 'spec-review'])
        review = task('review', dependencies=['deck'], review_of='deck',
            inputs=[dict(from_task='deck', output='candidate.pptx', path='candidate.pptx', purpose='Inspect actual deck', authority='Candidate', media_type=pptx_document.MIME)])
        review['criteria'] = deck['criteria'].copy()
        self.rt.create(plan([producer, spec_review, deck, review])); self.rt.tick('demo')
        aid = self.rt.task('demo', 'prepare')['latest']; self.agent.finish(aid)
        (self.agent.sessions[aid]['workspace'] / 'slides.json').write_text(json.dumps(fixture()))
        self.rt.tick('demo')
        self.assertIsNone(self.rt.task('demo', 'deck')['latest'])
        self.agent.finish(self.rt.task('demo', 'spec-review')['latest'], decision='accept')
        self.rt.tick('demo'); self.rt.tick('demo')
        output = self.rt.output('demo', 'deck', 'candidate.pptx')
        self.assertIsNotNone(output)
        rid = self.rt.task('demo', 'review')['latest']; self.agent.finish(rid, decision='accept'); self.rt.tick('demo')
        self.assertEqual(self.rt.status('demo')['status'], 'awaiting_user')
        self.rt.close(); self.rt = Runtime(self.root / 'runtime', self.factory)
        self.rt.tick('demo')
        self.assertEqual(len(self.ops.calls), 1)
        self.assertEqual(self.rt.output('demo', 'deck', 'candidate.pptx')['id'], output['id'])
        self.rt.select('demo', 'deck', output['id'], 'Select the deck', 'Use this exact deck')
        self.rt.tick('demo')
        self.assertEqual(self.rt.status('demo')['status'], 'completed')
        self.assertEqual(self.client.calls, [])

    def test_missing_dependency_blocks_before_claim(self):
        self.rt.create(plan([pptx_step([self.source_input()])]))
        with patch.object(pptx_document, 'available', side_effect=ValueError('PPTX dependency missing')):
            self.rt.tick('demo')
        self.assertEqual(self.rt.status('demo')['status'], 'blocked')
        self.assertEqual(self.rt.status('demo')['attempts'], [])

    def test_invalid_specification_is_not_replayed_after_restart(self):
        source = self.root / 'bad.json'; source.write_text('{"version":1}')
        aid = self.rt.register(source, 'Bad slide data', path='bad.json')
        item = self.source_input(); item.update(artifact=aid, path='bad.json')
        self.rt.create(plan([pptx_step([item])]))
        self.rt.tick('demo'); self.rt.tick('demo')
        self.rt.close(); self.rt = Runtime(self.root / 'runtime', self.factory); self.rt.tick('demo')
        self.assertEqual(self.rt.status('demo')['status'], 'blocked')
        self.assertEqual(len(self.ops.calls), 1)
        with self.assertRaisesRegex(ValueError, 'Missing upstream artifact'):
            self.rt.output('demo', 'deck', 'candidate.pptx')

    def test_changed_input_blocks_without_output_or_replay(self):
        item = self.source_input()
        original = self.ops.submit
        def tamper(session):
            launch = json.loads((Path(session['control']) / 'launch.json').read_text())
            source = Path(launch['workspace']) / 'slides.json'
            source.chmod(0o600)
            source.write_text('{}')
            return original(session)
        self.ops.submit = tamper
        self.rt.create(plan([pptx_step([item])]))
        self.rt.tick('demo'); self.rt.tick('demo')
        self.assertEqual(self.rt.status('demo')['status'], 'blocked')
        self.rt.tick('demo')
        self.assertEqual(len(self.ops.calls), 1)
        with self.assertRaisesRegex(ValueError, 'Missing upstream artifact'):
            self.rt.output('demo', 'deck', 'candidate.pptx')

    def test_implementation_drift_blocks_before_creation(self):
        from orchestrator.runtime import file_hash
        self.rt.create(plan([pptx_step([self.source_input()])]))
        with patch('orchestrator.step_runner.file_hash', side_effect=lambda p: '0' * 64 if Path(p).name == 'pptx_document.py' else file_hash(p)):
            self.rt.tick('demo'); self.rt.tick('demo')
        self.assertEqual(self.rt.status('demo')['status'], 'blocked')
        with self.assertRaisesRegex(ValueError, 'Missing upstream artifact'):
            self.rt.output('demo', 'deck', 'candidate.pptx')

    def test_real_supervised_child_creates_pptx_and_retains_receipt(self):
        factory = RegisteredFactory(); self.rt.factory = ExecutionFactory(self.agent, factory)
        self.rt.create(plan([pptx_step([self.source_input()])]))
        self.rt.tick('demo')
        deadline = time.monotonic() + 15
        while time.monotonic() < deadline:
            state = self.rt.tick('demo')
            if state['status'] in ('awaiting_user', 'blocked'): break
            time.sleep(.05)
        for child in factory.children: child.wait(timeout=10)
        self.assertEqual(state['status'], 'awaiting_user', state)
        receipt = json.loads(state['attempts'][0]['receipt'])
        self.assertEqual(receipt['operation']['validation']['slide_count'], 3)
        self.assertEqual(receipt['operation']['validation']['visual_review'], 'not_performed')


class PlanningTests(unittest.TestCase):
    setUp = planning_fixtures.Tests.setUp
    tearDown = planning_fixtures.Tests.tearDown
    request = planning_fixtures.Tests.request
    action = planning_fixtures.Tests.action
    queue = planning_fixtures.Tests.queue
    row = planning_fixtures.Tests.row
    response = planning_fixtures.Tests.response

    def deck_response(self):
        result = self.response()
        producer, reviewer = result['plan']['tasks']
        producer.pop('user_gate', None)
        producer['outputs'] = [{'path': 'slides.json', 'purpose': 'Bounded slide data', 'media_type': 'application/json'}]
        reviewer['inputs'] = [dict(from_task='produce', output='slides.json', path='slides.json', purpose='Review specification', authority='Candidate', media_type='application/json')]
        deck = contracts.assignment(pptx_step([dict(from_task='produce', output='slides.json', path='slides.json', purpose='Slide data', authority='Reviewed data', media_type='application/json')], dependencies=['produce', 'review']))
        deck_review = copy.deepcopy(reviewer)
        deck_review.update(id='deck-review', review_of='deck', dependencies=['deck'], criteria=deck['criteria'].copy(),
            inputs=[dict(from_task='deck', output='candidate.pptx', path='candidate.pptx', purpose='Review actual deck', authority='Candidate', media_type=pptx_document.MIME)])
        result['plan']['tasks'] += [deck, deck_review]
        return result

    def test_plan_freezes_schema_and_reviews_exact_deck_and_source(self):
        self.queue(action=self.action(step_capabilities=['pptx.create']), text='Create an editable presentation')
        planning.Worker(self.state, lambda *_: (json.dumps(self.deck_response()), {})).tick()
        row = self.row(); self.assertEqual(row['status'], 'ready', row['error'])
        value = json.loads(row['plan'])
        review = next(t for t in value['tasks'] if t['id'] == 'deck-review')
        self.assertTrue(any(i.get('from_task') == 'produce' and i['output'] == 'slides.json' for i in review['inputs']))
        self.assertIn('operation-support/pptx.create/contract.json', row['context'])
        self.assertEqual(self.state.db.execute('SELECT count(*) FROM production_runs').fetchone()[0], 0)

    def test_missing_spec_review_dependency_or_deck_gate_is_rejected(self):
        row = self.queue(action=self.action(step_capabilities=['pptx.create']))
        result = self.deck_response(); result['plan']['tasks'][2]['dependencies'].remove('review')
        with self.assertRaisesRegex(ValueError, 'depend on independent review'):
            planning.validate_result(json.dumps(result), row)
        result = self.deck_response(); result['plan']['tasks'][2].pop('user_gate')
        with self.assertRaisesRegex(ValueError, 'candidate selection gate'):
            planning.validate_result(json.dumps(result), row)

    def context_with_native_and_images(self):
        row=dict(self.queue(action=self.action(step_capabilities=['pptx.create'])))
        payload=json.loads(row['context']);ids={}
        for name,data in [('candidate.3dm',b'native fixture'),('preview.png',b'image routing fixture'),('visual.jpg',b'image routing fixture')]:
            path=self.rt.root/name;path.write_bytes(data)
            aid=self.rt.register(path,'Exact selected workflow source',path=name);ids[name]=aid
            source=planning.source_entry(self.rt,aid,'upstream/'+name,'Selected workflow source','Exact selected source')
            payload['sources'].append(source);payload['required_artifacts'].append(aid)
        row['context']=json.dumps(payload)
        return row,ids

    def test_native_context_preserved_but_only_compatible_inputs_bound_to_pptx(self):
        row,ids=self.context_with_native_and_images()
        result,plan=planning.validate_result(json.dumps(self.deck_response()),row)
        deck=next(t for t in plan['tasks'] if t['id']=='deck')
        bound={i.get('artifact') for i in deck['inputs']}
        self.assertNotIn(ids['candidate.3dm'],bound)
        self.assertIn(ids['preview.png'],bound);self.assertIn(ids['visual.jpg'],bound)
        for task in plan['tasks']:
            if task.get('execution'):continue
            self.assertTrue(set(ids.values()) <= {i.get('artifact') for i in task['inputs']})
        self.assertIn(ids['candidate.3dm'],json.loads(row['context'])['required_artifacts'])

    def test_explicit_incompatible_native_operation_input_is_rejected(self):
        row,ids=self.context_with_native_and_images();response=self.deck_response()
        response['plan']['tasks'][2]['inputs'].append(dict(artifact=ids['candidate.3dm'],path='candidate.3dm',purpose='Invalid native input',authority='Fixture'))
        with self.assertRaisesRegex(ValueError,'incompatible'):
            planning.validate_result(json.dumps(response),row)

    def test_implicit_routing_uses_capability_types_across_operations(self):
        source={'path':'model.3dm','media_type':'application/octet-stream'}
        self.assertTrue(planning.compatible_implicit_source('rhino.inspect',source))
        for cap in ['pptx.create','text.bundle','gemini.image','blender.inspect']:
            self.assertFalse(planning.compatible_implicit_source(cap,source))
        png={'path':'preview.png'}
        self.assertTrue(planning.compatible_implicit_source('pptx.create',png))
        self.assertTrue(planning.compatible_implicit_source('gemini.image',png))
        self.assertFalse(planning.compatible_implicit_source('text.bundle',png))

    def test_operation_keeps_declared_asset_paths_and_exact_identity(self):
        row,ids=self.context_with_native_and_images();response=self.deck_response()
        deck=response['plan']['tasks'][2]
        for filename,alias in [('preview.png','images/viewport.png'),('visual.jpg','images/rendering.jpg')]:
            deck['inputs'].append(dict(artifact=ids[filename],path=alias,purpose='Exact image at declared path',authority='Selected source'))
        _,plan=planning.validate_result(json.dumps(response),row)
        bound=next(t for t in plan['tasks'] if t['id']=='deck')['inputs']
        self.assertEqual({i['path'] for i in bound if i.get('artifact')==ids['preview.png']},{'images/viewport.png'})
        self.assertEqual({i['path'] for i in bound if i.get('artifact')==ids['visual.jpg']},{'images/rendering.jpg'})
        response['plan']['tasks'][2]['inputs'][-1]['path']='../rendering.jpg'
        with self.assertRaises(ValueError):planning.validate_result(json.dumps(response),row)

    def test_declared_path_collision_is_not_resolved_by_substitution(self):
        row,ids=self.context_with_native_and_images();response=self.deck_response()
        for filename in ['preview.png','visual.jpg']:
            response['plan']['tasks'][2]['inputs'].append(dict(artifact=ids[filename],path='images/same.png',purpose='Collision fixture',authority='Selected'))
        with self.assertRaises(ValueError):planning.validate_result(json.dumps(response),row)

    def test_local_input_recovery_retains_failed_attempt_and_requires_path_change(self):
        from orchestrator.storage import transaction
        from task_relay import production_stages
        row=dict(self.queue(action=self.action(step_capabilities=['text.bundle'])))
        path=self.rt.root/'source.txt';path.write_text('Exact original bytes')
        aid=self.rt.register(path,'Input',path='source.txt');payload=json.loads(row['context'])
        payload['sources'].append(planning.source_entry(self.rt,aid,'context/source.txt','Input','Selected source'))
        payload['required_artifacts'].append(aid);row['context']=contracts.encoded(payload)
        response=self.response();producer,review=response['plan']['tasks']
        producer.update(execution={'capability':'text.bundle','version':1,'parameters':{}},tools=[],
            inputs=[dict(artifact=aid,path='declared/source.txt',purpose='Declared input',authority='Selected')],
            criteria=execution.REGISTRY['text.bundle']['criteria'].copy())
        producer['outputs'][0]['media_type']='text/plain'
        producer['limits']={'seconds':30,'tool_calls':0,'output_bytes':2100000}
        review['criteria']=producer['criteria'].copy()
        result,plan=planning.validate_result(json.dumps(response),row)
        # Reproduce the old compiler dropping the explicit alias before dispatch.
        bound=next(i for i in plan['tasks'][0]['inputs'] if i.get('artifact')==aid)
        bound['path']='context/source.txt'
        self.rt.create(plan);run=plan['id']
        with transaction(self.state.db):
            self.state.db.execute('UPDATE production_plans SET status=?,run=?,result=?,plan=?,context=?,context_hash=? WHERE id=?',
                ('started',run,contracts.encoded(result),contracts.encoded(plan),row['context'],contracts.digest(payload),row['id']))
        self.rt.tick(run);attempt=self.rt.task(run,'produce')['latest']
        self.factory.sessions[attempt]['status']={'status':'finished','exit_code':1,'operation':{'outcome':'failed','reason':'Declared input path missing'}}
        self.rt.tick(run);before=dict(self.state.db.execute('SELECT * FROM production_attempts WHERE id=?',(attempt,)).fetchone())
        with transaction(self.state.db):
            ident=planning.prepare_local_input_repair(self.state,run,'Repair the failed declared input path.')
        repaired=self.state.db.execute('SELECT * FROM production_plans WHERE id=?',(ident,)).fetchone()
        self.assertEqual(repaired['status'],'ready');self.assertEqual(repaired['calls'],0)
        task=json.loads(repaired['plan'])['tasks'][0]
        self.assertEqual(next(i['path'] for i in task['inputs'] if i.get('artifact')==aid),'declared/source.txt')
        self.assertEqual(before,dict(self.state.db.execute('SELECT * FROM production_attempts WHERE id=?',(attempt,)).fetchone()))
        self.assertEqual(len(self.factory.calls),1)
        with transaction(self.state.db),self.assertRaises(ValueError):planning.prepare_local_input_repair(self.state,run,'Again')
        with transaction(self.state.db):
            self.state.db.execute("UPDATE production_attempts SET state='uncertain' WHERE id=?",(attempt,))
            with self.assertRaisesRegex(ValueError,'uncertain replay'):
                production_stages.failed_execution_snapshot(self.state,self.rt,run,'telegram','local_inputs')


if __name__ == '__main__':
    unittest.main()
