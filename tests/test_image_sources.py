"""Small synthetic image fixtures; no public downloads or paid models."""
import io
import json
import zipfile
import unittest
from unittest.mock import patch
from PIL import Image
from orchestrator import image_sources as sources, pptx_document, execution


def photo():
    out=io.BytesIO();Image.new('RGB',(200,150),'green').save(out,format='JPEG');return out.getvalue()


def subject(id='oak'):return dict(id=id,label='Example oak',query='Quercus example')


def page():
    return dict(title='File:Quercus example.jpg',index=1,imageinfo=[dict(mime='image/jpeg',
        url='https://upload.wikimedia.org/wikipedia/commons/a/ab/Quercus_example.jpg',
        descriptionurl='https://commons.wikimedia.org/wiki/File:Quercus_example.jpg',
        extmetadata={k:{'value':v} for k,v in dict(Artist='<a href="https://example.org">Example Author</a>',
        LicenseShortName='CC BY-SA 4.0',LicenseUrl='https://creativecommons.org/licenses/by-sa/4.0/',
        ImageDescription='A Quercus example tree.').items()})])


def bundle():
    def fetch(url,limit,deadline):
        if 'api.php' in url:return json.dumps({'query':{'pages':{'1':page()}}}).encode(),'application/json',url
        return photo(),'image/jpeg',url
    return sources.collect([subject()],fetch=fetch)


class ImageSourcesTests(unittest.TestCase):
    def test_collect_keeps_exact_bytes_credits_and_explicit_missing_subjects(self):
        calls=[];journal=[]
        def fetch(url,limit,deadline):
            calls.append(url)
            if 'api.php' in url:
                pages={'1':page()} if len(calls)==1 else {}
                return json.dumps({'query':{'pages':pages}}).encode(),'application/json',url
            return photo(),'image/jpeg',url
        raw,manifest=sources.collect([subject(),dict(id='missing',label='Other subject',query='Unknown subject')],fetch=fetch,record=journal.append)
        self.assertFalse(manifest['complete']);self.assertEqual(len(calls),4)
        images,credits,_=sources.unpack(raw)
        self.assertEqual(images['images/oak.jpg'],photo())
        self.assertEqual(credits['images/oak.jpg']['author'],'Example Author')
        self.assertEqual(manifest['subjects'][1]['status'],'missing')
        self.assertEqual(journal[0]['status'],'requested')

        # A user-permitted omitted photo leaves its subject text intact. The
        # retained missing receipt does not prevent creating the editable deck.
        document=dict(version=1,title='Partial photo coverage',slides=[dict(elements=[
            dict(type='text',x=1,y=1,w=5,h=1,text='Other subject: retained research'),
            dict(type='image',x=1,y=2,w=3,h=2,path='selected.zip/images/oak.jpg')])])
        result,evidence=pptx_document.create(document,bundles={'selected.zip':raw})
        from pptx import Presentation
        deck=Presentation(io.BytesIO(result))
        self.assertEqual(sum(s.shape_type==13 for s in deck.slides[0].shapes),1)
        self.assertTrue(any(getattr(s,'text','')=='Other subject: retained research' for s in deck.slides[0].shapes))

    def test_no_unmatched_wrong_species_or_unlicensed_substitution(self):
        for change in ('subject','license','author'):
            p=page()
            if change=='subject':p['title']='File:Other species.jpg';p['imageinfo'][0]['extmetadata']['ImageDescription']['value']='Other species'
            elif change=='license':p['imageinfo'][0]['extmetadata']['LicenseShortName']['value']='All rights reserved'
            else:p['imageinfo'][0]['extmetadata']['Artist']['value']=''
            self.assertIsNone(sources.candidate(p,subject()))

    def test_failed_download_records_gap_without_retry(self):
        calls=[]
        def fetch(url,*args):calls.append(url);raise OSError('offline')
        raw,m=sources.collect([subject()],fetch=fetch)
        self.assertEqual(len(calls),1);self.assertEqual(m['subjects'][0]['reason'],'offline')
        self.assertEqual(sources.unpack(raw)[0],{})

    def test_bundle_tampering_paths_and_extra_files_rejected(self):
        raw,_=bundle()
        for variant in ('bytes','traversal','extra','duplicate'):
            buf=io.BytesIO()
            with zipfile.ZipFile(io.BytesIO(raw)) as original,zipfile.ZipFile(buf,'w') as out:
                for name in original.namelist():
                    data=original.read(name)
                    if variant=='bytes' and name.endswith('.jpg'):data=photo()+b'changed'
                    if variant=='traversal' and name=='manifest.json':
                        m=json.loads(data);m['subjects'][0]['path']='../escape.jpg';data=json.dumps(m).encode()
                    out.writestr(name,data)
                if variant=='extra':out.writestr('../escape.txt','not allowed')
                if variant=='duplicate':
                    import warnings
                    with warnings.catch_warnings():
                        warnings.simplefilter('ignore');out.writestr('manifest.json','{}')
            with self.subTest(variant=variant),self.assertRaises(ValueError):sources.unpack(buf.getvalue())

    def test_source_bundle_creates_editable_deck_with_exact_image_and_credit_notes(self):
        raw,_=bundle();value=dict(version=1,title='Sourced photos',slides=[dict(elements=[
            dict(type='text',text='Example oak',x=1,y=.5,w=8,h=1),
            dict(type='image',path='selected.zip/images/oak.jpg',x=1,y=2,w=6,h=4)])])
        result,evidence=pptx_document.create(value,bundles={'selected.zip':raw})
        from pptx import Presentation
        deck=Presentation(io.BytesIO(result))
        self.assertEqual(deck.slides[0].shapes[1].image.blob,photo())
        self.assertIn('Example Author',deck.slides[0].notes_slide.notes_text_frame.text)
        self.assertIn('CC BY-SA 4.0',deck.slides[0].notes_slide.notes_text_frame.text)
        self.assertNotIn('notes',value['slides'][0])
        self.assertEqual(evidence['native_objects']['image'],1)

    def test_host_restriction_precedes_any_network_access(self):
        with patch('task_relay.orchestrator_web.public_addresses') as dns:
            with self.assertRaises(ValueError):sources.download('https://example.org/a.jpg',500,999999999)
            dns.assert_not_called()

    def test_subject_validation_blocks_duplicates_and_search_operators(self):
        for subjects in ([subject(),subject()], [dict(subject(),query='x OR filetype:svg')]):
            with self.assertRaises(ValueError):sources.validate_subjects(subjects)

    def test_current_commons_thumbnail_host_and_slashless_licence(self):
        p=page();i=p['imageinfo'][0]
        i['thumburl']='https://thumb.wikimedia.org/wikipedia/commons/thumb/a/ab/example.jpg/1280px-example.jpg'
        i['extmetadata']['LicenseUrl']['value']='https://creativecommons.org/licenses/by-sa/4.0'
        self.assertIsNotNone(sources.candidate(p,subject()))

    def test_registered_collection_preserves_bundle_receipt_and_is_not_replayed(self):
        import tempfile
        from pathlib import Path
        from orchestrator.runtime import Runtime
        from orchestrator.adapters import ExecutionFactory
        from tests.test_mixed_execution import LocalRegistered,Client,operation
        from tests.test_orchestrator import FakeFactory,plan
        def fetch(url,*args):
            if 'api.php' in url:return json.dumps({'query':{'pages':{'1':page()}}}).encode(),'application/json',url
            return photo(),'image/jpeg',url
        with tempfile.TemporaryDirectory() as directory,patch.object(sources,'download',side_effect=fetch) as network:
            root=Path(directory).resolve();f=root/'request.txt';f.write_text('Find an image of the specified subject.')
            rt=Runtime(root/'runtime',ExecutionFactory(FakeFactory(),LocalRegistered(Client())))
            try:
                aid=rt.register(f,'Exact user request',path='request.txt')
                task=operation('images','images.collect',[])  # Queries are frozen parameters; no history blobs.
                task['execution']['parameters']={'subjects':[subject()]}
                task['outputs']=[dict(path='images.zip',purpose='Sourced image candidates',media_type='application/zip')]
                rt.create(plan([task]));rt.tick('demo');rt.tick('demo');rt.tick('demo')
                artifact=rt.output('demo','images','images.zip')
                images,credits,_=sources.unpack(Path(artifact['blob']).read_bytes())
                self.assertEqual(images['images/oak.jpg'],photo())
                self.assertEqual(network.call_count,2)
                receipt=json.loads(rt.status('demo')['attempts'][0]['receipt'])
                self.assertTrue(receipt['operation']['validation']['complete'])
            finally:rt.close()

    def test_review_can_exclude_exact_rejected_titles_and_maps_are_not_photos(self):
        p=page();self.assertIsNone(sources.candidate(p,dict(subject(),exclude_titles=[p['title']])))
        p['title']='File:Quercus exampleDistMap227.png'
        self.assertIsNone(sources.candidate(p,subject()))

    def test_keyword_fallback_preserves_all_words_and_optional_identity(self):
        from urllib.parse import parse_qs,urlsplit
        calls=[]
        scene=dict(id='scene',label='Winter waterfront',query='Example City winter snow',identity='Example City')
        p=page();p['title']='File:Snow in winter, Example City.jpg'
        p['imageinfo'][0]['extmetadata']['ImageDescription']['value']='Waterfront'
        def fetch(url,*args):
            calls.append(url)
            if 'api.php' in url:
                pages={} if len(calls)==1 else {'1':p}
                return json.dumps({'query':{'pages':pages}}).encode(),'application/json',url
            return photo(),'image/jpeg',url
        raw,m=sources.collect([scene],fetch=fetch)
        self.assertEqual(m['coverage'],dict(found=1,missing=0,total=1))
        query=parse_qs(urlsplit(calls[1]).query)['gsrsearch'][0]
        self.assertEqual(query,'"example" "city" "winter" "snow" filetype:bitmap')
        self.assertEqual(m['subjects'][0]['query'],scene['query'])
        self.assertEqual(sources.unpack(raw)[0]['images/scene.jpg'],photo())
        self.assertIsNone(sources.candidate(p,dict(scene,query='Other City winter snow'),keywords=True))
        self.assertIsNone(sources.candidate(p,dict(scene,identity='City Example'),keywords=True))
        p['title']='File:Quercus other example.jpg'
        self.assertIsNone(sources.candidate(p,subject(),keywords=True))

    def test_bad_candidate_tries_a_distinct_file_without_retrying_same_url(self):
        p=page();q=page();q['title']='File:Quercus example 2.jpg';q['index']=2
        q['imageinfo'][0]['url']=q['imageinfo'][0]['url'].replace('.jpg','2.jpg')
        calls=[]
        def fetch(url,*args):
            calls.append(url)
            if 'api.php' in url:return json.dumps({'query':{'pages':{'1':p,'2':q}}}).encode(),'application/json',url
            return (b'invalid' if url==p['imageinfo'][0]['url'] else photo()),'image/jpeg',url
        raw,m=sources.collect([subject()],fetch=fetch)
        self.assertEqual(len(calls),3);self.assertEqual(len(set(calls)),3)
        self.assertEqual(m['subjects'][0]['title'],q['title'])
        self.assertEqual(len(m['subjects'][0]['download_failures']),1)
        self.assertEqual(sources.unpack(raw)[0]['images/oak.jpg'],photo())

    def test_missing_reasons_distinguish_no_results_rejected_metadata_and_bad_downloads(self):
        for mode in ('empty','wrong','download'):
            p=page()
            if mode=='wrong':p['imageinfo'][0]['extmetadata']['Artist']['value']=''
            calls=[]
            def fetch(url,*args):
                calls.append(url)
                if 'api.php' in url:return json.dumps({'query':{'pages':{} if mode=='empty' else {'1':p}}}).encode(),'application/json',url
                raise OSError('bad image response')
            _,m=sources.collect([subject()],fetch=fetch)
            reason=m['subjects'][0]['reason']
            self.assertIn({'empty':'no files','wrong':'none met','download':'bad image response'}[mode],reason)
            self.assertLessEqual(len(calls),3)
            self.assertEqual(len(calls),len(set(calls)))

    def test_redirects_consume_the_same_bounded_subject_budget(self):
        import time
        from unittest.mock import MagicMock
        connection=MagicMock();connection.getresponse.return_value.status=302
        connection.getresponse.return_value.getheader.return_value='https://commons.wikimedia.org/next'
        budget={'remaining':2}
        with patch('task_relay.orchestrator_web.public_addresses',return_value=['8.8.8.8']),patch(
                'task_relay.orchestrator_web.PublicHTTPS',return_value=connection):
            with self.assertRaisesRegex(ValueError,'request budget exhausted'):
                sources.download('https://commons.wikimedia.org/first',100,time.monotonic()+30,budget)
        self.assertEqual(connection.request.call_count,2)
        self.assertEqual(budget['remaining'],0)

    def test_empty_registered_collection_retains_diagnostic_zip_but_fails(self):
        import tempfile
        from pathlib import Path
        from orchestrator.runtime import Runtime
        from orchestrator.adapters import ExecutionFactory
        from tests.test_mixed_execution import LocalRegistered,Client,operation
        from tests.test_orchestrator import FakeFactory,plan
        def fetch(url,*args):return b'{"query":{"pages":{}}}','application/json',url
        with tempfile.TemporaryDirectory() as directory,patch.object(sources,'download',side_effect=fetch):
            root=Path(directory).resolve()
            rt=Runtime(root/'runtime',ExecutionFactory(FakeFactory(),LocalRegistered(Client())))
            try:
                task=operation('images','images.collect',[])
                task['execution']['parameters']={'subjects':[subject()]}
                task['outputs']=[dict(path='images.zip',purpose='Photo candidates',media_type='application/zip')]
                rt.create(plan([task]));rt.tick('demo');rt.tick('demo');rt.tick('demo')
                status=rt.status('demo')
                self.assertEqual(status['tasks'][0]['status'],'blocked')
                receipt=json.loads(status['attempts'][0]['receipt'])
                self.assertEqual(receipt['operation']['outcome'],'failed')
                self.assertIn('No usable photos',receipt['operation']['reason'])
                self.assertEqual(receipt['operation']['validation']['coverage']['found'],0)
                bundles=list(root.rglob('images.zip'))
                self.assertTrue(bundles)
                self.assertEqual(sources.unpack(bundles[0].read_bytes())[0],{})
            finally:rt.close()

    def test_fallback_ranks_focused_title_before_incidental_long_caption(self):
        scene=dict(id='scene',label='Scene',query='Example City snowy waterfront')
        long=page();short=page()
        long['title']='File:A person near a shop with a very long unrelated description at the waterfront of snowy Example City.jpg'
        short['title']='File:Snowy waterfront, Example City.jpg';short['index']=2
        short['imageinfo'][0]['url']=short['imageinfo'][0]['url'].replace('.jpg','2.jpg')
        searches=[]
        def fetch(url,*args):
            if 'api.php' in url:
                searches.append(url)
                return json.dumps({'query':{'pages':{} if len(searches)==1 else {'1':long,'2':short}}}).encode(),'application/json',url
            return photo(),'image/jpeg',url
        _,m=sources.collect([scene],fetch=fetch)
        self.assertEqual(m['subjects'][0]['title'],short['title'])
