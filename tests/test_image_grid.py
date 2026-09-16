import copy
import io
import unittest
from pptx import Presentation
from orchestrator import pptx_document as ppt, slide_templates as templates
from tests.test_gemini import PNG


class Tests(unittest.TestCase):
    def document(self,n=7,per_page=6):
        return dict(version=2,title='Photo catalogue',style='clean_minimal_v1',slides=[dict(layout='image_grid',
            content=dict(title='Subjects',per_page=per_page,items=[dict(image=f'photos/{i}.png',label=f'Subject {i}',caption=f'Name {i}',credit=f'Credit {i}') for i in range(n)]))])

    def test_pagination_preserves_order_all_images_editable_labels_and_credits(self):
        for capacity in (4,6):
            with self.subTest(capacity=capacity):
                value=self.document(35,capacity);before=copy.deepcopy(value)
                paths={f'photos/{i}.png':PNG for i in range(35)}
                raw,evidence=ppt.create(value,paths)
                self.assertEqual(evidence['slide_count'],(35+capacity-1)//capacity)
                self.assertEqual(evidence['native_objects']['image'],35)
                deck=Presentation(io.BytesIO(raw))
                names=[shape.text for slide in deck.slides for shape in slide.shapes if shape.name.endswith('-label')]
                self.assertEqual(names,[f'Subject {i}' for i in range(35)])
                for i in range(35):self.assertIn(f'Credit {i}',deck.slides[i//capacity].notes_slide.notes_text_frame.text)
                self.assertEqual(value,before)

    def test_frozen_validator_paginates_identically_and_checks_each_page_asset(self):
        doc=self.document();ns={'__name__':'frozen_validator'};exec(ppt.validator_source(),ns)
        paths=[f'photos/{i}.png' for i in range(7)]
        self.assertEqual(ns['validate'](doc,paths),ppt.validate(doc,paths))
        with self.assertRaisesRegex(ValueError,'exact declared'):ns['validate'](doc,paths[:-1])

    def test_invalid_density_overlong_text_style_and_total_slide_limit_do_not_truncate(self):
        for mutate in [lambda d:d['slides'][0]['content'].update(per_page=True),
                       lambda d:d['slides'][0]['content']['items'][0].update(label='x'*61),
                       lambda d:d['slides'][0]['content']['items'][0].update(caption='a\nb'),
                       lambda d:d.update(style='unknown'),
                       lambda d:d.update(slides=d['slides']*50)]:
            value=self.document();mutate(value)
            with self.assertRaises(ValueError):ppt.validate(value,[f'photos/{i}.png' for i in range(7)])

    def test_logo_collision_checked_on_grids_and_custom_brand_applied(self):
        value=self.document();value['brand']=dict(accent='446633',logo=dict(path='logo.png',x=1,y=1.5,w=1,h=.4))
        paths=[f'photos/{i}.png' for i in range(7)]+['logo.png']
        with self.assertRaisesRegex(ValueError,'overlaps'):ppt.validate(value,paths)
        value['brand']['logo'].update(x=11.5,y=6.95)
        doc=ppt.validate(value,paths)
        self.assertEqual(doc['slides'][0]['elements'][0]['color'],'446633')
        self.assertEqual(doc['slides'][1]['elements'][-1]['path'],'logo.png')

    def test_explicit_cover_crops_native_picture_without_changing_source_bytes(self):
        value=self.document(1);value['slides'][0]['content']['fit']='cover'
        raw,_=ppt.create(value,{'photos/0.png':PNG});deck=Presentation(io.BytesIO(raw))
        image=next(s for s in deck.slides[0].shapes if s.shape_type==13)
        self.assertEqual(image.image.blob,PNG)
        self.assertGreater(image.crop_top+image.crop_left,0)
        del value['slides'][0]['content']['fit']
        raw,_=ppt.create(value,{'photos/0.png':PNG});image=next(s for s in Presentation(io.BytesIO(raw)).slides[0].shapes if s.shape_type==13)
        self.assertEqual(image.crop_top+image.crop_left,0)

    def test_bundle_credits_follow_pages_and_disclose_explicit_crop(self):
        from tests.test_image_sources import bundle,photo
        raw,_=bundle();value=self.document(7)
        for item in value['slides'][0]['content']['items']:item['image']='photos.zip/images/oak.jpg'
        value['slides'][0]['content']['fit']='cover'
        output,_=ppt.create(value,bundles={'photos.zip':raw})
        deck=Presentation(io.BytesIO(output))
        self.assertEqual(len(deck.slides),2)
        for slide in deck.slides:
            self.assertIn('Example Author',slide.notes_slide.notes_text_frame.text)
            self.assertIn('center-cropped',slide.notes_slide.notes_text_frame.text)
            for shape in slide.shapes:
                if shape.shape_type==13:self.assertEqual(shape.image.blob,photo())
