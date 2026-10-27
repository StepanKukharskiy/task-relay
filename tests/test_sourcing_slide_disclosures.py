"""Native cards retain source context and map regions rather than invent facts."""
import json
import unittest

from orchestrator import sourcing_slides as slides
from tests import test_typed_room_chain as fixtures


class Tests(unittest.TestCase):
    def setUp(self):
        self.f=fixtures.Tests();self.f.setUp()
        self.addCleanup(self.f.doCleanups)

    def test_map_marker_is_centered_and_full_observed_region_remains_editable(self):
        f=self.f;assignment=f.assembler_fixture(1)
        spec=json.loads(slides.build(assignment,lambda p:(f.root/p).read_bytes())[0])
        first=spec['slides'][0];image=next(e for e in first['elements'] if e['type']=='image')
        marker=next(e for e in first['elements'] if e.get('name')=='VIS-001')
        self.assertAlmostEqual(marker['x']+marker['w']/2,image['x']+.4*image['w'])
        self.assertAlmostEqual(marker['y']+marker['h']/2,image['y']+.5*image['h'])
        borders=[e for e in first['elements'] if e.get('name')=='VIS-001 region']
        self.assertEqual(len(borders),4)
        self.assertAlmostEqual(borders[0]['x'],image['x']+.1*image['w'])
        self.assertAlmostEqual(borders[0]['w'],.6*image['w'])
        self.assertIn('grouped region',first['notes'])

    def test_former_price_context_comparison_and_variant_caveats_survive_native_pptx(self):
        f=self.f;f.text=f.text.replace('USD Price','Was\nUSD Price')
        for entry in f.pack['entries']:
            entry['pages'][0]['text']=f.text
            entry['pages'][0]['retrieved_at']=1791881786
        f.inputs[1]=f.input('sources.json',f.pack)
        assignment=f.assembler_fixture(1)
        # Recreate the code-owned catalog with a larger comparison, retaining
        # all source spans; no authored amount is inserted into the projection.
        from orchestrator import sourcing_contract
        catalog=json.loads((f.root/'catalog.json').read_text())
        comparison='Exact mismatch and unresolved variant. '*12
        catalog['records'][0]['options'][0]['comparison']=comparison
        cp=catalog['parameters'];ctx=sourcing_contract.context({**assignment,'operation_contract':dict(id='sourcing.options',version=1,parameters=cp)},lambda p:(f.root/p).read_bytes())
        catalog=sourcing_contract.options_document(catalog['records'],cp,ctx)
        replaced=f.input('catalog.json',catalog);replaced['media_type']='application/json';assignment['inputs'][2]=replaced
        assignment['execution']['parameters']['layout']='readable'
        raw,_=slides.build(assignment,lambda p:(f.root/p).read_bytes());spec=json.loads(raw)
        workload=slides.review_workload(assignment,lambda p:(f.root/p).read_bytes() if p.endswith('.json') else None)
        self.assertEqual(workload['slides'],len(spec['slides']))
        option=next(s for s in spec['slides'] if 'Full comparison:' in s['notes'])
        self.assertIn(comparison,option['notes'])
        crop=next(e for e in option['elements'] if e.get('source_region'))
        self.assertEqual(crop['source_region'],json.loads((f.root/'inventory.json').read_text())['items'][0]['bbox'])
        self.assertEqual(crop['path'],'room.png')
        self.assertIn('USD Price',option['notes'])
        self.assertTrue(any('current/variant/currency unverified' in e.get('text','') for e in option['elements']))
        self.assertIn('$349.99', '\n'.join(e.get('text','') for e in option['elements']))
        self.assertTrue(any('Source label: Was' in e.get('text','') for e in option['elements']))
        self.assertTrue(any(e.get('hyperlink') for e in option['elements']))
        import io
        from orchestrator.pptx_document import create
        from pptx import Presentation
        native,_=create(spec,{'room.png':f.png()},5000000)
        reopened=Presentation(io.BytesIO(native))
        self.assertTrue(any(comparison in s.notes_slide.notes_text_frame.text for s in reopened.slides))
        visible='\n'.join(shape.text for page in reopened.slides for shape in page.shapes if shape.has_text_frame)
        self.assertIn('2026-10-13',visible)
        self.assertIn(comparison,visible)
        from orchestrator.table_layout import lines
        import math
        for page in spec['slides']:
            for e in page['elements']:
                if e['type']=='text' and e['x']==4.75 and e['y']==2.55:
                    self.assertLessEqual(lines(e['text'],e['w'],e['font_size']),math.floor((e['h']-.12)*72/(e['font_size']*1.3)))

    def test_shared_text_flow_retains_long_literals_at_readable_size(self):
        from orchestrator.table_layout import flow_blocks,lines
        original=('Observed fact and its full qualification. '*100)+'END'
        pages=list(flow_blocks([original],7.75,4,14))
        self.assertGreater(len(pages),1)
        self.assertEqual(''.join(''.join(page) for page in pages),original)
        self.assertTrue(all(lines('\n\n'.join(page),7.75,14)<=15 for page in pages))


if __name__=='__main__':unittest.main()
