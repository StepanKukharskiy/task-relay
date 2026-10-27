"""Purchase links must survive delivery, including tampering detection."""
import copy
import io
import unittest
from pptx import Presentation
from orchestrator import pptx_document
from tests.test_presentations import fixture

class PurchaseLinks(unittest.TestCase):
    def test_editable_objects_keep_exact_purchase_links_and_detect_changes(self):
        spec=fixture()
        url='https://shop.example/products/chair?variant=green&size=large#details'
        for item in spec['slides'][0]['elements']:
            item['hyperlink']=url
        raw,_=pptx_document.create(spec)
        deck=Presentation(io.BytesIO(raw))
        self.assertEqual([s.click_action.hyperlink.address for s in deck.slides[0].shapes], [url]*3)
        deck.slides[0].shapes[0].click_action.hyperlink.address='https://shop.example/wrong'
        changed=io.BytesIO();deck.save(changed)
        with self.assertRaisesRegex(ValueError,'hyperlink changed'):
            pptx_document.inspect(changed.getvalue(),spec,{})

    def test_product_photo_keeps_image_bytes_and_purchase_link(self):
        from tests.test_gemini import PNG
        spec=fixture();spec['slides'][0]['elements']=[{
            'type':'image','path':'product.png','x':1,'y':1,'w':3,'h':3,
            'hyperlink':'https://shop.example/products/chair?variant=blue'}]
        raw,_=pptx_document.create(spec,{'product.png':PNG})
        photo=Presentation(io.BytesIO(raw)).slides[0].shapes[0]
        self.assertEqual(photo.image.blob,PNG)
        self.assertEqual(photo.click_action.hyperlink.address,spec['slides'][0]['elements'][0]['hyperlink'])

    def test_frozen_validator_rejects_non_purchase_schemes_and_credentials(self):
        ns={'__name__':'frozen'};exec(pptx_document.validator_source(),ns)
        for url in ('javascript:alert(1)','file:///tmp/x','http://shop.example/x','https://u:p@shop.example/x','https://shop.example/\nx'):
            spec=fixture();spec['slides'][0]['elements'][0]['hyperlink']=url
            with self.subTest(url=url),self.assertRaises(ValueError):
                ns['validate'](spec)
        spec=fixture();spec['slides'][0]['elements'][0]['hyperlink']='https://shop.example/chair'
        ns['validate'](spec)


class SourceRegions(unittest.TestCase):
    def test_region_is_native_preserves_original_pixels_and_detects_tampering(self):
        from PIL import Image
        b=io.BytesIO();Image.new('RGB',(400,200),'orange').save(b,format='PNG');pixels=b.getvalue()
        for fit in ('contain','cover'):
            spec=fixture();item=dict(type='image',path='room.png',x=1,y=1,w=2,h=2,source_region=[.25,.1,.75,.9],fit=fit)
            spec['slides'][0]['elements']=[item]
            raw,_=pptx_document.create(spec,{'room.png':pixels});deck=Presentation(io.BytesIO(raw));shape=deck.slides[0].shapes[0]
            self.assertEqual(shape.image.blob,pixels)
            self.assertAlmostEqual(shape.crop_left,.25 if fit=='contain' else .3)
            self.assertAlmostEqual(shape.crop_top,.1)
            if fit=='contain':self.assertAlmostEqual(shape.width/shape.height,1.25,places=5)
            shape.crop_left=.05;changed=io.BytesIO();deck.save(changed)
            with self.assertRaisesRegex(ValueError,'source region changed'):pptx_document.inspect(changed.getvalue(),spec,{'room.png':pixels})
        frozen={'__name__':'frozen'};exec(pptx_document.validator_source(),frozen)
        for region in ([0,0,0,1],[0,0,2,1],[True,0,1,1],[.5,.5,.2,.7],[0,0,1]):
            item['source_region']=region
            with self.subTest(region=region),self.assertRaises(ValueError):frozen['validate'](spec,['room.png'])
