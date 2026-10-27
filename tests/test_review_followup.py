"""Scripted provider judgments through real admission, report and delivery paths."""
import copy,json,unittest
from pathlib import Path
from orchestrator import contracts as c,job_owner,review_followup,report_builder,executors
from orchestrator.runtime import Runtime
from orchestrator.gemini_worker import execute
from orchestrator.workers import atomic
from tests import test_owner_intents as intent
from tests.test_operation_session import Script,CONFIG


class Tests(unittest.TestCase):
    create=intent.Tests.create
    invoke=intent.Tests.invoke
    session=intent.Tests.session
    source=intent.Tests.source

    def setUp(self):
        intent.Tests.setUp(self)
        self.grant['max_tasks']=16

    def delegate(self,ident):
        self.rt.tick('job')
        tasks=[intent.worker(ident),dict(id=ident+'_review',objective='Inspect every requirement and prior finding',review_of=ident)]
        self.invoke([('job_delegate',dict(request_id=ident,tasks=tasks))]);self.rt.tick('job')

    def finish(self,tid,text,*,decision='delivered',defect=None,resolved=True):
        self.rt.tick('job');sid=self.rt.task('job',tid)['latest'];s=self.factory.sessions[sid];f=s['frozen']
        control=Path(s['session']['control']);control.mkdir(parents=True,exist_ok=True)
        atomic(control/'launch.json',dict(credential_fingerprint=executors.fingerprint(CONFIG,f['backend'])))
        r=dict(summary='Inspected fixture',decision=decision,instruction='Repair exact missing requirement' if decision=='revise' else '',findings=[],
            checks={f'c{i}':dict(passed=True,evidence='Inspected exact candidate') for i in range(1,len(f['criteria'])+1)})
        if defect:
            r['findings']=[dict(category='requirements',code='missing_content',message=defect,evidence='Required content absent from candidate')]
        if f.get('review_followup'):
            p=f['review_followup'];path=p['candidates'][0]
            r['checks']['c'+str(p['criterion'])]=dict(passed=resolved,evidence={k:dict(verdict='resolved' if resolved else 'unresolved',reason='Actual content repaired' if resolved else 'Disclaimer still omits mandatory content',inspected=[dict(path=path,location='Required content section')]) for k in p['findings']})
        steps=[('file_read',dict(path=i['path'],offset=0,limit=1000)) for i in f['inputs']]
        steps += [('file_write',dict(path=f['outputs'][0]['path'],text=text)),('finish',dict(report_json=c.encoded(r)))]
        execute(f,control,Script(steps),lambda:(CONFIG,f['backend']))
        s['status']=dict(status='finished',exit_code=0,external_outcome='no_pending_response',pending_requests=[])
        self.rt.tick('job');return f,r

    def test_rejected_disclaimer_repair_restart_independent_review_and_delivery(self):
        for missing,repair in [('Required product photograph','Product photograph: exact authorized image reference'),('Required measured component dimensions','Measured component dimensions: 10 x 20 mm')]:
            with self.subTest(missing=missing):
                # Two field configurations use the same lifecycle, with no domain branch.
                if self.rt.db.execute('select count(*) from production_runs').fetchone()[0]:
                    self.rt.close();self.root=self.root/'second';self.rt=Runtime(self.root,self.factory)
                self.create(request=missing);self.rt.tick('job');self.delegate('draft')
                first,_=self.finish('draft','Incomplete candidate')
                rejected,_=self.finish('draft_review','Missing required content',decision='revise',defect=missing)
                original=self.rt.output('job','draft','outputs/result.txt');raw=Path(original['blob']).read_bytes()
                self.delegate('disclaimer')
                cf,_=self.finish('disclaimer','Unavailable. See retailer or manufacturer.')
                self.assertIn(missing,c.encoded(cf['correction_context']))
                rf,report=self.finish('disclaimer_review','Still unresolved',decision='revise',resolved=False)
                forged=copy.deepcopy(report);forged['decision']='accept';forged['instruction']='';forged['checks']['c'+str(rf['review_followup']['criterion'])]['passed']=True
                forged['findings']=[dict(category='advisory',code='disclosed',message='Disclosed absence',evidence='Disclaimer')]
                with self.assertRaisesRegex(ValueError,'Unresolved review'):
                    c.report(forged,rf)
                omitted=copy.deepcopy(report);omitted['checks']['c'+str(rf['review_followup']['criterion'])]['evidence']={}
                with self.assertRaises(ValueError):c.report(omitted,rf)
                self.rt.close();self.rt=Runtime(self.root,self.factory)
                self.delegate('repaired');self.finish('repaired',repair)
                reviewed,_=self.finish('repaired_review','Repair independently verified',decision='accept')
                self.assertEqual(set(v['id'] for v in rf['review_followup']['findings'].values()),set(v['id'] for v in reviewed['review_followup']['findings'].values()))
                candidate=self.rt.output('job','repaired','outputs/result.txt')
                self.invoke([('job_complete',dict(deliverables_json=c.encoded({'report':candidate['id']}),summary='All repairs reviewed'))]);self.rt.tick('job')
                self.assertEqual(self.rt.status('job')['status'],'completed')
                self.assertEqual(Path(original['blob']).read_bytes(),raw)
                self.assertEqual(self.rt.task('job','repaired')['attempts'],1)
                self.assertEqual(job_owner.original(self.rt,'job')['job_owner']['task_limits'],self.grant['task_limits'])

    def test_same_bytes_new_task_and_stale_or_changed_evidence_cannot_close(self):
        self.create();self.rt.tick('job');self.delegate('draft');self.finish('draft','Missing content')
        self.finish('draft_review','Rejected',decision='revise',defect='Required image omitted')
        self.delegate('renamed');self.finish('renamed','Missing content')
        self.rt.tick('job');rf=self.factory.sessions[self.rt.task('job','renamed_review')['latest']]['frozen'];p=rf['review_followup'];path=p['candidates'][0]
        report=dict(summary='Purported repair',decision='accept',instruction='',findings=[],checks={f'c{i}':dict(passed=True,evidence='Inspected') for i in range(1,len(rf['criteria'])+1)})
        report['checks']['c'+str(p['criterion'])]['evidence']={k:dict(verdict='resolved',reason='Renamed',inspected=[dict(path=path,location='Whole file')]) for k in p['findings']}
        with self.assertRaisesRegex(ValueError,'Unchanged candidate'):c.report(report,rf)
        candidate=Path(rf['workspace'])/path;candidate.chmod(0o600);candidate.write_text('Changed after freezing')
        with self.assertRaisesRegex(ValueError,'evidence changed'):c.report(report,rf)
        self.assertFalse(review_followup.covers(self.rt,'job',dict(checks=[]),{'review_of':'renamed'}))

    def test_legacy_plan_does_not_acquire_new_report_obligations(self):
        from tests.test_orchestrator import pair
        graph=pair();self.rt.create(graph);self.rt.tick('demo')
        self.factory.finish(self.rt.task('demo','produce')['latest']);self.rt.tick('demo')
        self.factory.finish(self.rt.task('demo','review')['latest'],decision='revise');self.rt.tick('demo')
        f=self.factory.sessions[self.rt.task('demo','produce')['latest']]['frozen']
        self.assertNotIn('correction_context',f)

    def test_standard_controller_correction_uses_same_findings_after_restart(self):
        from tests.test_orchestrator import pair
        graph=pair();graph['review_followup_version']=1;self.rt.create(graph);self.rt.tick('demo')
        self.factory.finish(self.rt.task('demo','produce')['latest']);self.rt.tick('demo')
        self.factory.finish(self.rt.task('demo','review')['latest'],decision='revise');self.rt.tick('demo')
        self.rt.close();self.rt=Runtime(self.root,self.factory)
        self.factory.finish(self.rt.task('demo','produce')['latest']);self.rt.tick('demo')
        aid=self.rt.task('demo','review')['latest'];f=self.factory.sessions[aid]['frozen'];self.assertTrue(f['review_followup'])
        self.factory.finish(aid,decision='accept');resultpath=Path(f['workspace'])/'.relay/result.json';report=json.loads(resultpath.read_text())
        p=f['review_followup'];report['checks'][-1]['evidence']=c.encoded({k:dict(verdict='resolved',reason='Corrected candidate independently inspected',inspected=[dict(path=p['candidates'][0],location='Whole candidate')]) for k in p['findings']})
        resultpath.write_text(c.encoded(report));self.assertEqual(self.rt.tick('demo')['status'],'completed')
        self.assertEqual(self.rt.task('demo','produce')['attempts'],2)

    def test_exhausted_correction_keeps_findings_and_capacity(self):
        from tests.test_orchestrator import pair
        graph=pair();graph['review_followup_version']=1;self.rt.create(graph);self.rt.tick('demo')
        for cycle in (1,2):
            self.factory.finish(self.rt.task('demo','produce')['latest']);self.rt.tick('demo')
            aid=self.rt.task('demo','review')['latest'];self.factory.finish(aid,decision='revise')
            f=self.factory.sessions[aid]['frozen'];path=Path(f['workspace'])/'.relay/result.json'
            if cycle==2:
                report=json.loads(path.read_text());p=f['review_followup']
                report['checks'][-1]['evidence']=c.encoded({k:dict(verdict='unresolved',reason='Required content is still absent',inspected=[dict(path=p['candidates'][0],location='Whole file')]) for k in p['findings']})
                path.write_text(c.encoded(report))
            self.rt.tick('demo')
        count=len(self.factory.calls);self.rt.tick('demo')
        self.assertEqual(len(self.factory.calls),count)
        self.assertEqual(self.rt.task('demo','produce')['attempts'],2)
        self.assertEqual(self.rt.task('demo','produce')['status'],'blocked')
        self.assertTrue(review_followup.findings(self.rt,'demo'))


    def test_new_revision_request_is_not_lost_after_prior_findings_resolve(self):
        from tests.test_orchestrator import pair
        graph=pair(max_attempts=3);graph['review_followup_version']=1;self.rt.create(graph);self.rt.tick('demo')
        self.factory.finish(self.rt.task('demo','produce')['latest']);self.rt.tick('demo')
        self.factory.finish(self.rt.task('demo','review')['latest'],decision='revise');self.rt.tick('demo')
        self.factory.finish(self.rt.task('demo','produce')['latest']);self.rt.tick('demo')
        aid=self.rt.task('demo','review')['latest'];self.factory.finish(aid,decision='revise')
        f=self.factory.sessions[aid]['frozen'];p=f['review_followup'];path=Path(f['workspace'])/'.relay/result.json';report=json.loads(path.read_text())
        for check in report['checks']:check['passed']=True
        report['checks'][-1]['evidence']=c.encoded({k:dict(verdict='resolved',reason='Earlier defect repaired',inspected=[dict(path=p['candidates'][0],location='Whole current candidate')]) for k in p['findings']})
        report['instruction']='Newly observed source attribution is missing'
        path.write_text(c.encoded(report));self.rt.tick('demo')
        new=self.factory.sessions[self.rt.task('demo','produce')['latest']]['frozen']
        self.assertIn(report['instruction'],c.encoded(new['correction_context']))

    def test_authorized_successor_carries_findings_restart_and_remaining_capacity(self):
        self.create();self.rt.tick('job');self.delegate('draft');self.finish('draft','Missing content')
        self.rt.pause('job')
        self.finish('draft_review','Rejected',decision='revise',defect='Required image omitted')
        old=job_owner.original(self.rt,'job');status=self.rt.status('job')
        grant=copy.deepcopy(old['job_owner'])
        grant['owner_turns']-=self.rt.task('job','owner')['attempts']
        grant['max_tasks']-=len(status['tasks'])-1
        def admit(name,allocation):
            job_owner.create(self.rt,name,old['brief'],old['backend'],[],allocation)
            review_followup.inherit(self.rt,'job',name)
        expanded=copy.deepcopy(grant);expanded['owner_turns']+=1
        with self.assertRaisesRegex(ValueError,'remaining capacity'):
            with self.rt.transaction():admit('expanded',expanded)
        self.assertIsNone(self.rt.db.execute("SELECT id FROM production_runs WHERE id='expanded'").fetchone())
        with self.rt.transaction():admit('next',grant)
        review_followup.inherit(self.rt,'job','next') # exact retry is idempotent
        self.rt.close();self.rt=Runtime(self.root,self.factory)
        self.assertEqual(review_followup.findings(self.rt,'job'),review_followup.findings(self.rt,'next'))
        self.assertFalse(review_followup.covers(self.rt,'next',dict(checks=[]),dict(review_of='new_candidate')))
        with self.assertRaisesRegex(ValueError,'already has a continuation'):
            with self.rt.transaction():admit('duplicate',grant)
        self.assertEqual(self.rt.status('job'),status)
        frozen=dict(id='review',instruction='Inspect repair',criteria=['Requested quality'],review_of='fixed',
            inputs=[dict(path='candidate.txt',from_task='fixed',media_type='text/plain')])
        review_followup.bind(self.rt,job_owner.original(self.rt,'next'),frozen)
        self.assertIn('Required image omitted',c.encoded(frozen['correction_context']))
        self.assertEqual(set(v['id'] for v in frozen['review_followup']['findings'].values()),{v['id'] for v in review_followup.findings(self.rt,'job')})


    def test_continuation_cannot_change_independent_tool_provider_binding(self):
        self.grant['operations']=['web.sources']
        self.grant['operation_backends']={'web.sources':dict(type='gemini-agent',model='original-search')}
        self.create();self.rt.pause('job')
        old=job_owner.original(self.rt,'job');grant=copy.deepcopy(old['job_owner'])
        grant['operation_backends']['web.sources']['model']='different-search'
        with self.assertRaisesRegex(ValueError,'original authority: operation_backends'):
            with self.rt.transaction():
                job_owner.create(self.rt,'changed',old['brief'],old['backend'],[],grant)
                review_followup.inherit(self.rt,'job','changed')
        self.assertIsNone(self.rt.db.execute("SELECT id FROM production_runs WHERE id='changed'").fetchone())
        self.assertEqual(job_owner.original(self.rt,'job'),old)

    def test_explicit_capacity_receipt_preserves_findings_and_original_grant(self):
        self.create();self.rt.tick('job');self.delegate('draft');self.finish('draft','Missing content')
        self.rt.pause('job');self.finish('draft_review','Rejected',decision='revise',defect='Required image omitted')
        old=job_owner.original(self.rt,'job');status=self.rt.status('job');grant=copy.deepcopy(old['job_owner'])
        grant['owner_turns']=old['job_owner']['owner_turns']-self.rt.task('job','owner')['attempts']+1
        grant['max_tasks']=old['job_owner']['max_tasks']-(len(status['tasks'])-1)+2
        with self.rt.transaction():
            job_owner.create(self.rt,'approved',old['brief'],old['backend'],[],grant)
            proof=dict(version=1,previous='job',successor='approved',prior_plan_sha256=c.digest(old),
                successor_plan_sha256=c.digest(job_owner.original(self.rt,'approved')),
                additional_capacity=dict(owner_turns=1,max_tasks=2),
                user_request='Authorize one additional owner invocation and two task slots?',
                user_reply='Yes, authorize exactly that additional capacity.')
            review_followup.inherit(self.rt,'job','approved',capacity_authorization=proof)
        review_followup.inherit(self.rt,'job','approved',capacity_authorization=proof)
        self.rt.close();self.rt=Runtime(self.root,self.factory)
        self.assertEqual(job_owner.original(self.rt,'job'),old);self.assertEqual(self.rt.status('job'),status)
        self.assertEqual(review_followup.findings(self.rt,'approved'),review_followup.findings(self.rt,'job'))
        receipts=self.rt.db.execute("SELECT data FROM production_events WHERE run='approved' AND kind='review_findings_inherited'").fetchall()
        self.assertEqual(len(receipts),1);self.assertEqual(json.loads(receipts[0][0])['capacity_authorization'],proof)
        with self.assertRaisesRegex(ValueError,'authorization'):
            review_followup.inherit(self.rt,'job','approved',capacity_authorization={**proof,'additional_capacity':dict(owner_turns=2,max_tasks=2)})

    def test_mismatched_capacity_receipts_roll_back_successor(self):
        self.create();self.rt.pause('job');old=job_owner.original(self.rt,'job')
        grant=copy.deepcopy(old['job_owner']);grant['owner_turns']+=1
        for altered in ('prior_plan_sha256','successor_plan_sha256','user_reply','additional_capacity'):
            with self.subTest(altered=altered),self.assertRaisesRegex(ValueError,'authorization'):
                with self.rt.transaction():
                    job_owner.create(self.rt,'invalid',old['brief'],old['backend'],[],grant)
                    proof=dict(version=1,previous='job',successor='invalid',prior_plan_sha256=c.digest(old),
                        successor_plan_sha256=c.digest(job_owner.original(self.rt,'invalid')),
                        additional_capacity=dict(owner_turns=1,max_tasks=0),user_request='Add one owner turn?',user_reply='Yes')
                    proof[altered]=dict(owner_turns=True,max_tasks=0) if altered=='additional_capacity' else ''
                    review_followup.inherit(self.rt,'job','invalid',capacity_authorization=proof)
            self.assertIsNone(self.rt.db.execute("SELECT id FROM production_runs WHERE id='invalid'").fetchone())


    def test_unrelated_branch_review_passes_but_delivery_cannot_drop_its_findings(self):
        self.grant['deliverables']={'notes':'text/markdown'}
        source=self.source('Shared original reference')
        self.create(inputs=[source]);self.rt.tick('job')
        self.delegate('draft');self.finish('draft','Wrong auxiliary evidence')
        self.finish('draft_review','Auxiliary defect',decision='revise',defect='Auxiliary capture date is wrong')
        original=self.rt.output('job','draft','outputs/result.txt');raw=Path(original['blob']).read_bytes()
        self.rt.tick('job')
        tasks=[intent.worker('notes',[dict(artifact=source['artifact'])],media='text/markdown'),
            dict(id='notes_review',objective='Review the independent notes',review_of='notes')]
        self.invoke([('job_delegate',dict(request_id='notes',tasks=tasks))]);self.rt.tick('job')
        author,_=self.finish('notes','Independent correct notes')
        self.assertNotIn('correction_context',author)
        reviewer,_=self.finish('notes_review','Notes independently accepted',decision='accept')
        self.assertNotIn('review_followup',reviewer)
        self.assertEqual(self.rt.task('job','notes')['status'],'completed')
        candidate=self.rt.output('job','notes','outputs/result.md')
        with self.assertRaisesRegex(ValueError,'collectively resolve'):
            self.session().complete(dict(deliverables_json=c.encoded({'notes':candidate['id']}),summary='Attempt to omit failed branch'))
        self.assertFalse(self.rt.db.execute("SELECT 1 FROM production_events WHERE kind='owner_delivery_verified'").fetchone())
        self.assertEqual(Path(original['blob']).read_bytes(),raw)
        # Explicit derivation carries the defect across a format change.
        self.rt.tick('job')
        tasks=[intent.worker('derived',[dict(artifact=original['id'])],media='text/markdown'),
            dict(id='derived_review',objective='Review repaired derived evidence',review_of='derived')]
        self.invoke([('job_delegate',dict(request_id='derived',tasks=tasks))]);self.rt.tick('job')
        self.finish('derived','Auxiliary capture date corrected with exact evidence')
        self.rt.close();self.rt=Runtime(self.root,self.factory)
        fixed,_=self.finish('derived_review','Exact repair verified',decision='accept')
        self.assertIn('Auxiliary capture date is wrong',c.encoded(fixed['review_followup']))
        result=self.rt.output('job','derived','outputs/result.md')
        self.invoke([('job_complete',dict(deliverables_json=c.encoded({'notes':result['id']}),summary='Repair reviewed'))]);self.rt.tick('job')
        self.assertEqual(self.rt.status('job')['status'],'completed')


class NativeCorrectionTests(unittest.TestCase):
    def test_actual_pptx_missing_photo_rejection_repair_pagination_and_delivery(self):
        import io,tempfile
        from PIL import Image
        from pptx import Presentation
        from orchestrator import step_runner,pptx_document
        from tests.test_orchestrator import FakeFactory
        from tests.test_corrections import graph
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp).resolve();factory=FakeFactory();rt=Runtime(root/'runtime',factory)
            try:
                photo=root/'photo.png';Image.new('RGB',(4,3),'#bb5533').save(photo)
                ref=dict(artifact=rt.register(photo,'Authorized synthetic product pixels'),path='product.png',purpose='Required product photograph',authority='Exact synthetic fixture',media_type='image/png')
                plan=graph();plan['review_followup_version']=1;plan['tasks'][2].pop('user_gate',None)
                plan['tasks'][0]['inputs']=[ref];plan['tasks'][2]['inputs'].append(ref)
                rt.create(plan);rt.tick('demo')
                rows=[['Item','Details','Source','Price','Size','Status']]+[[str(n),'Verified texture. '*12,'Example retailer','USD 1200','62 inches','Availability requires confirmation. '*5] for n in range(1,7)]
                doc=dict(version=2,title='Correction fixture',slides=[dict(layout='table',content=dict(title='Options',rows=rows,caption='Synthetic evidence'))])
                original=None;limits={t['id']:t['max_attempts'] for t in job_owner.original(rt,'demo')['tasks']}
                for cycle in (1,2):
                    aid=rt.task('demo','produce')['latest'];factory.finish(aid)
                    if cycle==2:
                        self.assertIn('Missing product photo',c.encoded(factory.sessions[aid]['frozen']['correction_context']))
                        doc['slides'].append(dict(layout='image_text',content=dict(title='Required product photograph',image='product.png',caption='Exact authorized product pixels',body='Photo accompanies sourced item 1')))
                    (factory.sessions[aid]['workspace']/'slides.json').write_text(c.encoded(doc));rt.tick('demo')
                    factory.finish(rt.task('demo','review')['latest'],decision='accept');rt.tick('demo')
                    build=rt.task('demo','deck')['latest'];session=factory.sessions[build];control=Path(session['session']['control']);control.mkdir(parents=True,exist_ok=True)
                    result=step_runner.execute(session['frozen'],control);self.assertEqual(result['outcome'],'completed')
                    session['status']=dict(status='finished',exit_code=0,external_outcome='no_pending_response',pending_requests=[],operation=result)
                    rt.tick('demo');artifact=rt.output('demo','deck','deck.pptx');raw=Path(artifact['blob']).read_bytes();deck=Presentation(io.BytesIO(raw))
                    images=[s for slide in deck.slides for s in slide.shapes if s.shape_type==13]
                    self.assertEqual(len(images),cycle-1)
                    self.assertGreater(len(deck.slides),1) # table content is paginated by code
                    for slide in deck.slides:
                        for shape in slide.shapes:
                            self.assertLessEqual(shape.top+shape.height,deck.slide_height)
                    rid=rt.task('demo','final')['latest'];factory.finish(rid,decision='revise' if cycle==1 else 'accept')
                    frozen=factory.sessions[rid]['frozen'];path=Path(frozen['workspace'])/'.relay/result.json';report=json.loads(path.read_text())
                    if cycle==1:
                        report['findings']=[dict(category='requirements',code='missing_photo',message='Missing product photo',evidence='Reopened PPTX contains zero pictures')]
                        original=(artifact,raw)
                    else:
                        self.assertEqual(images[0].image.blob,photo.read_bytes())
                        p=frozen['review_followup'];report['checks'][-1]['evidence']=c.encoded({k:dict(verdict='resolved',reason='Reopened exact PPTX contains authorized photo and all table cells inside slide bounds',inspected=[dict(path=p['candidates'][0],location='Final image slide and every table page')]) for k in p['findings']})
                    path.write_text(c.encoded(report));rt.tick('demo')
                    if cycle==1:
                        rt.close();rt=Runtime(root/'runtime',factory)
                self.assertEqual(rt.status('demo')['status'],'completed')
                self.assertEqual(Path(original[0]['blob']).read_bytes(),original[1])
                self.assertEqual({t['id']:rt.spec(t)['max_attempts'] for t in rt.status('demo')['tasks']},limits)
                self.assertEqual(rt.task('demo','deck')['attempts'],2)
            finally:rt.close()

