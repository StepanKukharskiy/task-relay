"""One interface harness, multiple data schemas; no providers or production state."""
import copy
import hashlib
import json
from pathlib import Path
import sqlite3
import tempfile
import unittest

from orchestrator import collection_core as engine, operation_contracts as c
from orchestrator import operation_records as records, record_recovery, sourcing_contract as content
from orchestrator import sourcing_slides, sourcing_merge, pptx_document
from orchestrator import execution
from orchestrator import product_photos,image_sources
from task_relay import record_planning,planning_contract
from orchestrator.gemini_worker import Files, definitions
from tests.test_sourcing_records import BACKEND


CONFIGURATIONS={
    'plants':dict(fields={'botanical_name':'text','climate':'text'},identity='botanical_name'),
    'furniture':dict(fields={'name':'text','price':'price','dimensions':'text'},identity='name'),
    'components':dict(fields={'designation':'text','load_rating':'text'},identity='designation'),
}


class Harness:
    def __init__(self,root,config,count=1,options=1):
        self.root=Path(root).resolve();self.db=sqlite3.connect(self.root/'state.sqlite');self.db.row_factory=sqlite3.Row
        self.db.execute('CREATE TABLE production_attempts (id TEXT PRIMARY KEY,run TEXT,task TEXT,assignment TEXT,state TEXT,parent TEXT,frozen TEXT,receipt TEXT,session TEXT,error TEXT)')
        records.initialize(self.db);self.config=config;self.count=count;self.options=options
        self.inventory=dict(version=1,items=[dict(id='S'+str(i),label='Subject '+str(i)) for i in range(count)])
        entries=[]
        for subject in self.inventory['items']:
            for n in range(options):
                text='\n'.join(k+': '+('USD $123.45' if rule=='price' else subject['id']+' '+k+' '+str(n)) for k,rule in config['fields'].items())
                url='https://public.example/'+subject['id']+'/'+str(n)
                entries.append(dict(query=subject['label'],target_ids=[subject['id']],grounding=[],failures=[],pages=[dict(url=url,text=text,sha256=hashlib.sha256(text.encode()).hexdigest())]))
        self.pack=dict(version=1,entries=entries)
        self.inputs=[self.input('subjects.json',self.inventory),self.input('sources.json',self.pack)]

    def input(self,path,value):
        raw=(c.encoded(value)+'\n').encode();p=self.root/path;p.parent.mkdir(parents=True,exist_ok=True);p.write_bytes(raw)
        return dict(path=path,artifact=c.digest(path),sha256=hashlib.sha256(raw).hexdigest(),purpose='Bound synthetic evidence',authority='Test data',media_type='application/json')

    def assignment(self,ident,bounds=None,catalog=None):
        params=dict(inventory_path='subjects.json',source_pack_path='sources.json',max_records=max(1,self.count),
            summary_path=ident+'.md',field_definitions=self.config['fields'],identity_field=self.config['identity'])
        inputs=copy.deepcopy(self.inputs)
        if bounds:params['record_range']=bounds
        if catalog is None:params.update(options_per_item=3,count_policy='target')
        else:
            params['catalog_path']=ident+'-candidate.json';inputs.append(self.input(params['catalog_path'],catalog))
        f=dict(assignment_id=ident,run='synthetic',id=ident,workspace=str(self.root),operation_store=str(self.root/'state.sqlite'),
            backend=BACKEND.copy(),inputs=inputs,outputs=[dict(path=ident+'.json',purpose='Code projection'),dict(path=ident+'.md',purpose='Code summary')],
            operation_contract=dict(id='collection.review' if catalog is not None else 'collection.options',version=1,parameters=params,output=ident+'.json'),
            tools=['files'],criteria=['Preserve evidence and explicit gaps'],limits=dict(seconds=600,tool_calls=24,provider_requests=24,response_tokens=4096,output_bytes=5000000))
        if catalog is not None:f['review_of']='producer'
        self.db.execute('INSERT INTO production_attempts VALUES (?,?,?,?,?,?,?,?,?,?)',(ident,'synthetic',ident,'assignment','running',None,c.encoded(f),'{}',None,None));self.db.commit()
        return f

    def fragment(self,subject):
        options=[]
        for i,entry in enumerate(self.pack['entries']):
            if subject['id'] not in entry['target_ids']:continue
            text=entry['pages'][0]['text'];fields={}
            for field in self.config['fields']:
                quote=next(line for line in text.splitlines() if line.startswith(field+':'))
                start=text.index(quote);fields[field]=dict(status='observed',start=start,end=start+len(quote))
            options.append(dict(observation='page-'+str(i)+'-0',fields=fields,comparison='Synthetic correspondence judgment'))
        return dict(item=subject['id'],options=options,gap='Fewer alternatives found' if len(options)<3 else '')

    def produce(self,ident,bounds=None):
        frozen=self.assignment(ident,bounds);store=records.Store(frozen,self.db)
        subjects=self.inventory['items'][bounds[0]:bounds[1]] if bounds else self.inventory['items']
        def adapter(subject):
            fragment=self.fragment(subject);revision=store.progress()['revision']
            store.submit(dict(request_key=subject['id'],revision=revision,records=[fragment]),{})
            return dict(id=subject['id'],status='found' if fragment['options'] else 'missing')
        engine.collect(subjects,[dict(id='saved-pages',version=1,collect=adapter)],batch_size=2)
        files=Files(frozen);store.export(files,complete=True);store.verify_export(files.read_bytes)
        return store,json.loads(files.read_bytes(ident+'.json'))

    def review(self,ident,catalog,bounds=None):
        f=self.assignment(ident,bounds,catalog);store=records.Store(f,self.db)
        for item in catalog['items']:
            store.submit(dict(request_key=item['id'],revision=store.progress()['revision'],records=[dict(
                item=item['id'],item_sha256=c.digest(item),verdict='supported',reason='Synthetic independent verdict')]),{})
        files=Files(f);store.export(files,complete=True);store.verify_export(files.read_bytes)
        return json.loads(files.read_bytes(ident+'.json'))


class Tests(unittest.TestCase):
    def test_interface_inventory_is_serializable_and_distinguishes_authorship(self):
        catalog=json.loads(json.dumps(c.interface_catalog()))
        generic=next(r for r in catalog if r['id']=='collection.options')
        self.assertTrue(generic['lifecycle']['resumable'])
        self.assertIn('commits',generic['code_owned'])
        self.assertIn('independent judgments',generic['model_owned'])
        self.assertEqual(generic['planning_interface']['merge_operation'],'records.merge')
        self.assertIn('builder',generic['fragment_schema'])

    def test_saved_baseline_recovery_is_schema_neutral_and_retains_source_versions(self):
        from task_relay import sourcing_recovery
        from tests.test_sourcing_recovery_composition import CompilerTests as recovery_fixture
        for name,config in CONFIGURATIONS.items():
            with self.subTest(schema=name),tempfile.TemporaryDirectory() as root:
                # The production recovery compiler uses exact ports/archive, not room labels.
                plan,scope,coverage,known,documents=recovery_fixture().fixture()
                contract=dict(id='collection.options',version=1)
                scope['pipeline_step']['handoff']['outputs']['catalog']['content_contract']=contract
                port=scope['input_ports']['upstream.items'];port.pop('content_contract');port['media_type']='application/json'
                documents['catalog']['contract']=contract
                documents['catalog']['parameters'].update(field_definitions=config['fields'],identity_field=config['identity'])
                receipts=sourcing_recovery.bind(plan,scope,coverage,known,lambda s:documents[s['artifact']])
                author=plan['tasks'][1]
                self.assertEqual(author['records']['family'],contract)
                self.assertEqual(author['records']['parameters']['field_definitions'],config['fields'])
                self.assertEqual(author['record_composition']['base_catalog_path'],'catalog.json')
                self.assertEqual(receipts[0]['provider_calls'],0)
                # The same record lifecycle performs the actual evidence/fragment composition.
                h=Harness(root,config);self.addCleanup(h.db.close);_,baseline=h.produce('baseline')
                fresh=copy.deepcopy(h.pack);fresh['entries'][0]['pages'][0]['url']='https://public.example/recovery'
                f=h.assignment('recover');f['inputs']+= [h.input('baseline.json',baseline),h.input('fresh.json',fresh)]
                params=f['operation_contract']['parameters'];params.update(source_pack_path='fresh.json',base_source_pack_path='sources.json',base_catalog_path='baseline.json',source_pack_output='combined.json')
                f['outputs'].append(dict(path='combined.json',purpose='Exact code-composed evidence'))
                h.db.execute('UPDATE production_attempts SET frozen=? WHERE id=?',(c.encoded(f),'recover'));h.db.commit()
                store=records.Store(f,h.db);fragment=h.fragment(h.inventory['items'][0]);fragment['options'][0]['observation']='page-1-0'
                store.submit(dict(request_key='recover',revision=0,records=[fragment]),{})
                files=Files(f);store.export(files,complete=True);store.verify_export(files.read_bytes)
                recovered=json.loads(files.read_bytes('recover.json'))
                self.assertEqual(recovered['items'][0]['options'][0],baseline['items'][0]['options'][0])
                self.assertEqual(len(recovered['items'][0]['options']),2)
                self.assertEqual(recovered['source_versions']['source_pack_sha256'],hashlib.sha256(files.read_bytes('combined.json')).hexdigest())
                changed=copy.deepcopy(f);changed['operation_contract']['parameters']['identity_field']='other'
                with self.assertRaises(ValueError):records.Store(changed,h.db).context()

    def test_image_sources_are_configuration_and_commons_fallback_preserves_binding(self):
        from tests.test_image_sources import photo,page
        for name,config in CONFIGURATIONS.items():
            with self.subTest(schema=name),tempfile.TemporaryDirectory() as root:
                h=Harness(root,config);self.addCleanup(h.db.close);_,catalog=h.produce('producer')
                binding=h.input('catalog.json',catalog)
                frozen=dict(assignment_id='generic-images',inputs=h.inputs+[binding],limits=dict(seconds=600,output_bytes=5000000),
                    execution=dict(capability='images.catalog',parameters=dict(inventory_path='subjects.json',catalog_path='catalog.json',source_pack_path='sources.json',
                        count_policy='target',source_adapters=['observed-pages','commons'],image_identity='subject')))
                query=catalog['items'][0]['options'][0]['fields'][config['identity']]['excerpt'][:240];calls=[]
                def fetch(url,limit,deadline):
                    calls.append(url)
                    if '/w/api.php?' in url:
                        p=page();p['title']='File:'+query+'.jpg'
                        p['imageinfo'][0]['extmetadata']['ImageDescription']={'value':query}
                        return json.dumps({'query':{'pages':{'1':p}}}).encode(),'application/json',url
                    return photo(),'image/jpeg',url
                raw,manifest=product_photos.build(frozen,lambda p:(Path(root)/p).read_bytes(),fetch=fetch)
                self.assertEqual(manifest['coverage']['found'],1);self.assertEqual(len(calls),2)
                subject=manifest['subjects'][0]
                self.assertEqual(subject['catalog_subject_binding']['catalog_sha256'],binding['sha256'])
                self.assertEqual(subject['id'],catalog['items'][0]['options'][0]['option_id'])
                images,_,_=image_sources.unpack(raw)
                self.assertEqual(images[subject['path']],photo())
                photo_input=h.input('photos.zip',{})
                # Bind actual binary bytes, never relabel a JSON input as a ZIP.
                (h.root/'photos.zip').write_bytes(raw);photo_input.update(media_type='application/zip',sha256=hashlib.sha256(raw).hexdigest())
                f=dict(inputs=h.inputs+[binding,photo_input],execution=dict(capability='collection.slides',parameters=dict(inventory_path='subjects.json',catalog_path='catalog.json',source_pack_path='sources.json',image_path=None,photos_path='photos.zip',title=name,layout='readable')))
                spec,proof=sourcing_slides.build(f,lambda p:(h.root/p).read_bytes())
                self.assertEqual(proof['missing_photos'],[])
                pptx,_=pptx_document.create(json.loads(spec),{},5000000,bundles={'photos.zip':raw})
                self.assertTrue(pptx.startswith(b'PK'))

    def proposal(self,config,count=13):
        inventory=dict(artifact='subjects',path='subjects.json',purpose='Subjects',authority='Test data',media_type='application/json')
        sources=dict(artifact='pages',path='sources.json',purpose='Exact saved pages',authority='Test data',media_type='application/json')
        def output(path):return dict(path=path,purpose='Requested projection',media_type='application/json' if path.endswith('.json') else 'text/markdown')
        def task(ident,inputs,outputs):
            return dict(id=ident,role='writer',objective='Produce requested content',instruction='Use bound synthetic data',
                inputs=copy.deepcopy(inputs),outputs=[output(p) for p in outputs],dependencies=[],tools=['files'],
                worker={'requires':['files.text']},criteria=['Retain exact evidence'],limits=dict(seconds=600,tool_calls=24,provider_requests=24,response_tokens=4096,output_bytes=5000000),max_attempts=3)
        producer=task('collect',[inventory,sources],['catalog.json','catalog.md'])
        producer['records']=dict(family=dict(id='collection.options',version=1),
            inputs=dict(inventory_path={'source':'subjects'},source_pack_path={'source':'pages'}),
            parameters=dict(field_definitions=config['fields'],identity_field=config['identity'],options_per_item=3,count_policy='target',max_records=count))
        review=task('review',[inventory,sources,dict(from_task='collect',output='catalog.json',path='candidate.json',purpose='Exact candidate',authority='Unaccepted',media_type='application/json')],['review.md'])
        review.update(review_of='collect',dependencies=['collect'])
        return dict(backend=BACKEND.copy(),tasks=[producer,review])

    def test_registered_builder_compiles_ports_reviews_and_batches_for_all_schemas(self):
        for name,config in CONFIGURATIONS.items():
            with self.subTest(schema=name):
                plan=self.proposal(config)
                record_planning.bind(plan)
                self.assertEqual(plan['tasks'][0]['operation_contract']['parameters']['inventory_path'],'subjects.json')
                self.assertEqual(plan['tasks'][1]['operation_contract']['parameters']['catalog_path'],'candidate.json')
                self.assertEqual(plan['tasks'][1]['operation_contract']['id'],'collection.review')
                self.assertEqual(len(plan['tasks'][1]['outputs']),2)
                plan=self.proposal(config);native=execution.REGISTRY['records.merge']
                scope=dict(connector_companions_version=1,graph_operations=[dict(id='records.merge',available=True,**native)])
                limits=dict(max_tasks=30,limits=dict(seconds=600,output_bytes=5000000))
                deliverables={'catalog':dict(task='collect',output='catalog.json')}
                receipts=record_planning.partition(plan,scope,limits,{'collect':11,'review':11},deliverables)
                self.assertEqual(receipts[0]['ranges'],[[0,4],[4,8],[8,12],[12,13]])
                record_planning.bind(plan)
                self.assertEqual(plan['tasks'][0]['execution']['capability'],'records.merge')
                self.assertEqual(deliverables['catalog']['output'],'delivery/catalog.json')
                self.assertEqual(sum(t.get('operation_contract',{}).get('id')=='collection.options' for t in plan['tasks']),4)
                self.assertTrue(any(i.get('output')=='delivery/catalog.json' for i in plan['tasks'][1]['inputs']))
                execution.validate(plan['tasks'][0])
        old=planning_contract.operation_capabilities({'web.sources','pptx.create'},dict(connector_companions_version=1))
        self.assertNotIn('records.merge',old);self.assertNotIn('collection.slides',old)
        new=planning_contract.operation_capabilities({'web.sources','pptx.create'},dict(connector_companions_version=1,record_interfaces_version=1))
        self.assertIn('records.merge',new);self.assertIn('collection.slides',new)

    def test_same_engine_projects_three_schemas_zero_one_five_options_and_pptx(self):
        for name,config in CONFIGURATIONS.items():
            for count in (0,1,5):
                with self.subTest(schema=name,options=count),tempfile.TemporaryDirectory() as root:
                    h=Harness(root,config,options=count)
                    self.addCleanup(h.db.close)
                    store,catalog=h.produce('producer');h.review('review',catalog)
                    self.assertEqual(len(catalog['items'][0]['options']),count)
                    self.assertEqual(catalog['items'][0]['shortfall'],max(0,3-count))
                    self.assertEqual(set(catalog['items'][0]['options'][0]['fields']) if count else set(),set(config['fields']) if count else set())
                    if count and 'price' in config['fields']:self.assertEqual(catalog['items'][0]['options'][0]['fields']['price']['amount'],'123.45')
                    schema=c.tool(store.contract)['parameters']['properties']['records']['items']
                    self.assertEqual(set(schema['properties']['options']['items']['properties']['fields']['properties']),set(config['fields']))
                    tool_names={t['name'] for t in definitions(frozen=store.frozen)}
                    self.assertIn('source_read',tool_names);self.assertNotIn('file_write',tool_names)
                    source=h.input('catalog.json',catalog)
                    f=dict(inputs=h.inputs+[source],execution=dict(capability='collection.slides',parameters=dict(
                        inventory_path='subjects.json',catalog_path='catalog.json',source_pack_path='sources.json',
                        image_path=None,photos_path=None,title=name,layout='readable')))
                    raw,proof=sourcing_slides.build(f,lambda p:(Path(root)/p).read_bytes())
                    spec=json.loads(raw);pptx,checked=pptx_document.create(spec,{},5000000)
                    self.assertTrue(pptx.startswith(b'PK'));self.assertEqual(proof['options'],count)
                    self.assertEqual(proof['inventory_sha256'],h.inputs[0]['sha256'])

    def test_capacity_partitions_and_exact_reviewed_join_use_same_interface(self):
        for name,config in CONFIGURATIONS.items():
            with self.subTest(schema=name),tempfile.TemporaryDirectory() as root:
                h=Harness(root,config,count=13);self.addCleanup(h.db.close)
                ranges=engine.RECORD_CAPACITY.ranges(13,11);catalogs=[];reviews=[];inputs=copy.deepcopy(h.inputs)
                self.assertEqual(ranges,[[0,4],[4,8],[8,12],[12,13]])
                for index,bounds in enumerate(ranges):
                    _,catalog=h.produce('p'+str(index),bounds);review=h.review('r'+str(index),catalog,bounds)
                    for paths,prefix,value in ((catalogs,'cat',catalog),(reviews,'review',review)):
                        path=prefix+str(index)+'.json';paths.append(path);inputs.append(h.input(path,value))
                params=dict(inventory_path='subjects.json',source_pack_path='sources.json',max_records=13,
                    options_per_item=3,count_policy='target',field_definitions=config['fields'],identity_field=config['identity'],catalog_paths=catalogs,review_paths=reviews)
                frozen=dict(inputs=inputs,execution=dict(capability='records.merge',version=1,parameters=params))
                raw,_,proof=sourcing_merge.build(frozen,lambda p:(Path(root)/p).read_bytes())
                wrong=copy.deepcopy(frozen);wrong['execution']['capability']='sourcing.merge'
                with self.assertRaisesRegex(ValueError,'content family'):
                    sourcing_merge.build(wrong,lambda p:(Path(root)/p).read_bytes())
                merged=json.loads(raw);self.assertEqual([i['id'] for i in merged['items']],[i['id'] for i in h.inventory['items']])
                self.assertEqual(len(proof['partitions']),4)
                # The registered native output carries the same indexed review
                # interface as its fragment family. Identical source copies do
                # not require another model read or lose their artifact IDs.
                from types import SimpleNamespace
                from orchestrator import evidence_context
                candidate={**h.input('merged.json',merged), 'from_task':'merge','output':'delivery/catalog.json'}
                alias={**h.input('same-source.json',h.pack),'from_task':'merge','output':'delivery/source-pack.json'}
                resolved=[{**i,'blob':h.root/i['path'],'bytes':(h.root/i['path']).stat().st_size} for i in h.inputs+[candidate,alias]]
                rt=SimpleNamespace(task=lambda *args:dict(id='merge',latest='native-attempt',assignment='original'),spec=lambda t:frozen)
                review=dict(review_of='merge',outputs=[dict(path='review.md')],limits=dict(tool_calls=6,provider_requests=6,output_bytes=10000))
                evidence_context.bind(rt,'run',review,BACKEND,resolved)
                self.assertEqual(review['evidence_context']['aliases'],{'same-source.json':'sources.json'})
                f={**review,'workspace':str(h.root),'backend':BACKEND,'assignment_id':'review',
                   'inputs':[{k:v for k,v in i.items() if k not in ('blob','bytes')} for i in resolved]}
                files=Files(f);first=files.source_pack(max_bytes=16)
                with self.assertRaisesRegex(ValueError,'complete exact candidate'):files.validate_text_delivery(dict(decision='accept'))
                batch=files.call('evidence_next',{})
                self.assertEqual(batch['remaining'],[])
                self.assertNotIn('same-source.json',[part['path'] for part in batch['parts']])
                files.validate_text_delivery(dict(decision='accept'))
                self.assertEqual(files.observed['same-source.json'],files.observed['sources.json'])
                (h.root/'same-source.json').write_text('changed alias')
                with self.assertRaisesRegex(ValueError,'Frozen input changed'):files.call('evidence_next',{})
                unknown=copy.deepcopy(frozen);unknown['execution']['version']=99
                rt.spec=lambda t:unknown
                evidence_context.bind(rt,'run',review,BACKEND,resolved)
                self.assertNotIn('evidence_context',review)
                changed=copy.deepcopy(merged);changed['records'][0]['item']='foreign'
                with self.assertRaises(ValueError):engine.join([changed['records']],['S0'],key=lambda r:r['item'],validate=lambda r:None)

    def test_fallback_checkpoint_restart_and_unknown_effect_do_not_repeat_sources(self):
        for name in CONFIGURATIONS:
            with self.subTest(schema=name):
                subjects=[dict(id='A',label=name)];calls=[];saved=[]
                def missing(s):calls.append('commons');return dict(id=s['id'],status='missing',reason='No match')
                def found(s):calls.append('web');return dict(id=s['id'],status='found',fields={'name':name},usage={'tokens':7})
                adapters=[dict(id='commons',version=1,collect=missing),dict(id='web',version=1,collect=found)]
                result,state=engine.collect(subjects,adapters,commit=saved.append,policy={'fields':['name']})
                self.assertEqual(calls,['commons','web']);self.assertEqual(result[0]['usage']['tokens'],7)
                again,_=engine.collect(subjects,adapters,checkpoint=state,policy={'fields':['name']})
                self.assertEqual(result,again);self.assertEqual(calls,['commons','web'])
                with self.assertRaises(ValueError):engine.collect(subjects,adapters,checkpoint=state,policy={'fields':['other']})
                with self.assertRaises(engine.PendingCollection):engine.collect(subjects,adapters,checkpoint=saved[0],policy={'fields':['name']})
                def unavailable(s):raise OSError('No confirmed response')
                pending=[]
                with self.assertRaises(OSError):engine.collect(subjects,[dict(id='web',version=1,collect=unavailable)],commit=pending.append)
                self.assertIsNotNone(pending[-1]['pending'])
                with self.assertRaises(engine.PendingCollection):engine.collect(subjects,[dict(id='web',version=1,collect=found)],checkpoint=pending[-1])

    def test_generic_amendments_and_confirmed_yield_preserve_records(self):
        for name,config in CONFIGURATIONS.items():
            with self.subTest(schema=name),tempfile.TemporaryDirectory() as root:
                h=Harness(root,config,count=2);self.addCleanup(h.db.close)
                frozen=h.assignment('partial');store=records.Store(frozen,h.db)
                original=h.fragment(h.inventory['items'][0]);store.submit(dict(request_key='one',revision=0,records=[original]),{})
                changed=copy.deepcopy(original);changed['options'][0]['comparison']='Corrected synthetic judgment'
                args=dict(request_key='amend',revision=1,expected_record_sha256=c.digest(original),record=changed)
                receipt=store.amend(args,{});self.assertEqual(store.amend(args,{}),receipt)
                with self.assertRaises(ValueError):store.amend(dict(args,request_key='stale'),{})
                stopped=dict(status='finished',external_outcome='no_pending_response',pending_requests=[],reason='time_limit')
                proof=record_recovery.checkpoint(frozen,stopped,h.db)
                self.assertEqual(proof['committed'],1);self.assertEqual(proof['remaining'],1)
                self.assertIsNone(record_recovery.checkpoint(frozen,{**stopped,'pending_requests':['unknown']},h.db))
                h.inputs[1]=h.input('sources.json',dict(version=1,entries=[]))
                with self.assertRaises(ValueError):records.Store(frozen,h.db).context()

    def test_registered_runtime_delivers_and_corrects_three_schemas_without_user_gates(self):
        from unittest.mock import patch
        from orchestrator.runtime import Runtime
        from orchestrator.adapters import ExecutionFactory
        from orchestrator import corrections,worker_capabilities
        from tests.test_orchestrator import FakeFactory
        from tests.test_worker_capabilities import CODEX
        from tests.workflow_decoys import NativeBoundary,SourcePages,edge,registered
        from tests.test_web_sources import CONFIG
        from task_relay import gemini
        from pptx import Presentation
        import io
        for name,config in CONFIGURATIONS.items():
            with self.subTest(schema=name),tempfile.TemporaryDirectory() as root:
                h=Harness(root,config);self.addCleanup(h.db.close)
                agents=FakeFactory();native=NativeBoundary();factory=ExecutionFactory(agents,native,agents)
                factory.available=lambda _:None
                rt=Runtime(h.root/'runtime',factory);self.addCleanup(rt.db.close)
                aid=rt.register(h.root/'subjects.json','Exact synthetic subjects')
                proposal=self.proposal(config,count=1)
                producer,judge=proposal['tasks'];producer['inputs'][0]['artifact']=aid;judge['inputs'][0]['artifact']=aid
                for task in (producer,judge):
                    task['inputs']=[i for i in task['inputs'] if i.get('artifact')!='pages']+[edge('sources','delivery/source-pack.json',binding='sources.json')]
                    task['dependencies']=list(dict.fromkeys(task['dependencies']+['sources','sources-review']))
                    task['max_attempts']=2
                producer['records']['inputs'].update(inventory_path={'source':aid},source_pack_path={'producer':'sources','output':'delivery/source-pack.json'})
                record_planning.bind(proposal)
                from orchestrator import executors
                for task in (producer,judge):
                    worker_capabilities.resolve(task,[worker_capabilities.entry(BACKEND)],BACKEND)
                    for key,value in executors.limits_for(BACKEND,task).items():task['limits'][key]=min(task['limits'].get(key,value),value)
                def reviewer(op,binary=False):
                    t=dict(id=op['id']+'-review',role='reviewer',objective='Review exact synthetic output',instruction='Scripted independent judgment',review_of=op['id'],
                        inputs=[edge(op['id'],o['path'],o.get('media_type','application/json')) for o in op['outputs']],
                        outputs=[dict(path=op['id']+'-review.md',purpose='Independent judgment')],dependencies=[op['id']],criteria=op['criteria'].copy(),
                        limits=dict(seconds=600,tool_calls=24,provider_requests=24,response_tokens=4096,output_bytes=200000),max_attempts=2)
                    if binary:
                        t['worker']={'requires':['files.text','files.binary','code.execute','images.view']}
                        worker_capabilities.resolve(t,[worker_capabilities.entry(CODEX)],CODEX)
                    else:t['tools']=['files']
                    return t
                source=registered('sources','web.sources',[],[],'delivery/source-pack.json','application/json')
                source['execution']['parameters']=dict(queries=['Exact synthetic '+name],domains=[],model=CONFIG['models']['text'])
                ports=[dict(artifact=aid,path='subjects.json',purpose='Exact subjects',authority='Test data',media_type='application/json'),
                    edge('collect','catalog.json'),edge('sources','delivery/source-pack.json',binding='sources.json')]
                images=registered('images','images.catalog',ports,['collect','review','sources'],'photos.zip','application/zip')
                images['execution']['parameters']=dict(inventory_path='subjects.json',catalog_path='catalog.json',source_pack_path='sources.json',count_policy='target',source_adapters=['observed-pages'],image_identity='page')
                assembly=registered('slides','collection.slides',ports+[edge('images','photos.zip','application/zip')],['collect','review','sources','images','images-review'],'slides.json','application/json')
                assembly['execution']['parameters']=dict(inventory_path='subjects.json',catalog_path='catalog.json',source_pack_path='sources.json',image_path=None,photos_path='photos.zip',title=name,layout='standard');assembly['max_attempts']=2
                deck=registered('deck','pptx.create',[edge('slides','slides.json'),edge('images','photos.zip','application/zip')],['slides','slides-review','images'],'deck.pptx','application/vnd.openxmlformats-officedocument.presentationml.presentation');deck['max_attempts']=2
                tasks=[source,reviewer(source),producer,judge,images,reviewer(images,True),assembly,reviewer(assembly),deck,reviewer(deck,True)]
                corrections.compile(tasks,record_layout=True)
                self.assertIn('review_correction',deck)
                page=h.pack['entries'][0]['pages'][0]
                def fetch(_self,url,offset,limit,allowed_hosts=None):
                    return dict(url=url,requested_url=url,title=name,text=page['text'],retrieved_at=124.,sha256=page['sha256'],content_type='text/html',total_characters=len(page['text']),next_offset=None,
                        images=[dict(url='https://images.example/source.jpg',basis='observed-product-jsonld')])
                with patch.object(SourcePages,'products',[dict(id=name,product_url=page['url'])]),patch.object(SourcePages,'fetch',fetch),patch.object(gemini,'read_config',return_value=CONFIG):
                    rt.create(dict(id='shared',brief='Schema-configured collection delivery',backend=BACKEND,tasks=tasks))
                    def advance():
                        for _ in range(3):rt.tick('shared')
                        self.assertEqual(native.errors,[])
                    def finish(tid,decision='accept',fragments=None):
                        current=rt.task('shared',tid);self.assertIsNotNone(current['latest'],str(rt.status('shared')))
                        token=current['latest'];f=agents.sessions[token]['frozen']
                        if fragments is None:
                            agents.finish(token,decision=decision)
                            if decision=='revise':
                                target=Path(f['workspace'])/'.relay/result.json'
                                result=json.loads(target.read_text())
                                result['findings']=[dict(category='requirements',code='layout_readability',message='Synthetic clipped text',evidence='Exact controlled layout')]
                                target.write_text(c.encoded(result))
                        else:
                            store=records.Store(f,rt.db)
                            for fragment in fragments:store.submit(dict(request_key=str(store.progress()['revision']),revision=store.progress()['revision'],records=[fragment]),{})
                            store.export(Files(f),complete=True)
                            report=dict(summary='Scripted independent fragments',decision=decision,instruction='',findings=[],checks={'c'+str(i):dict(passed=True,evidence='Synthetic seam') for i in range(1,len(f['criteria'])+1)})
                            (Path(f['workspace'])/'.relay/result.json').write_text(c.encoded(report))
                            agents.sessions[token]['status']=dict(status='finished',exit_code=0,reason=None,usage=[dict(input_tokens=20,output_tokens=5)])
                        advance()
                    advance();finish('sources-review')
                    f=agents.sessions[rt.task('shared','collect')['latest']]['frozen'];ctx=records.Store(f,rt.db).context()
                    fragment=h.fragment(h.inventory['items'][0]);fragment['options'][0]['observation']=next(iter(ctx['observations']))
                    finish('collect','delivered',[fragment])
                    f=agents.sessions[rt.task('shared','review')['latest']]['frozen'];unit=records.Store(f,rt.db).progress()['next_items'][0]
                    finish('review',fragments=[dict(item=unit['item'],item_sha256=unit['item_sha256'],verdict='supported',reason='Scripted semantic review')])
                    finish('images-review');finish('slides-review')
                    first=rt.output('shared','deck','deck.pptx');self.assertIsNotNone(first,str(native.errors))
                    finish('deck-review','revise');finish('slides-review');finish('deck-review')
                    self.assertEqual(rt.status('shared')['status'],'completed',str(native.errors))
                    final=rt.output('shared','deck','deck.pptx');self.assertNotEqual(first['attempt'],final['attempt'])
                    self.assertEqual(rt.task('shared','collect')['attempts'],1)
                    self.assertEqual(rt.db.execute('SELECT count(*) FROM production_decisions').fetchone()[0],0)
                    rendered=Presentation(io.BytesIO(Path(final['blob']).read_bytes()))
                    self.assertTrue(any(s.shape_type==13 for page in rendered.slides for s in page.shapes))
                    self.assertEqual(len(native.downloads),1)  # Layout repair does not repeat source downloads.
                    source_version=rt.output('shared','sources','delivery/source-pack.json')['sha256']
                    photo_version=rt.output('shared','images','photos.zip')['sha256']
                    latest=json.loads(rt.db.execute('SELECT frozen FROM production_attempts WHERE id=?',(final['attempt'],)).fetchone()[0])
                    self.assertIn(photo_version,{i['sha256'] for i in latest['inputs']})
                    slides=rt.output('shared','slides','slides.json')
                    frozen=json.loads(rt.db.execute('SELECT frozen FROM production_attempts WHERE id=?',(slides['attempt'],)).fetchone()[0])
                    self.assertIn(source_version,{i['sha256'] for i in frozen['inputs']})
                    self.assertTrue(all(row['created'] is not None for row in rt.db.execute('SELECT created FROM production_events WHERE run=?',('shared',))))
                    self.assertTrue(any('input_tokens' in a['receipt'] for a in rt.status('shared')['attempts']))

    def test_empty_subjects_and_large_photo_catalog_are_valid_data(self):
        from tests.test_image_sources import photo
        for count in (0,42):
            with self.subTest(subjects=count),tempfile.TemporaryDirectory() as root:
                h=Harness(root,CONFIGURATIONS['plants'],count=14 if count else 0,options=3 if count else 1);self.addCleanup(h.db.close)
                for e in h.pack['entries']:e['pages'][0]['images']=[dict(url='https://images.example/'+e['target_ids'][0]+'.jpg',basis='observed-product-jsonld')]
                h.inputs[1]=h.input('sources.json',h.pack)
                _,catalog=h.produce('producer');review=h.review('review',catalog)
                self.assertTrue(catalog['complete']);self.assertTrue(review['complete'])
                inputs=h.inputs+[h.input('catalog.json',catalog)]
                frozen=dict(assignment_id='images',inputs=inputs,limits=dict(seconds=600,output_bytes=5000000),execution=dict(capability='images.catalog',parameters=dict(
                    inventory_path='subjects.json',catalog_path='catalog.json',source_pack_path='sources.json',count_policy='target',source_adapters=['observed-pages'],image_identity='page')))
                calls=[]
                def fetch(url,*_):calls.append(url);return photo(),'image/jpeg',url
                raw,manifest=product_photos.build(frozen,lambda p:(h.root/p).read_bytes(),fetch=fetch)
                self.assertEqual(manifest['coverage']['found'],count);self.assertEqual(len(calls),count)
                image_sources.unpack(raw)

    def test_checkpoint_validation_precedes_dispatch_and_commit_failure_is_reconciled(self):
        subjects=[dict(id='A'),dict(id='B')];calls=[];snapshots=[]
        def fetch(s):calls.append(s['id']);return dict(id=s['id'],status='found')
        adapters=[dict(id='public',version=1,collect=fetch)]
        _,state=engine.collect(subjects,adapters);calls.clear()
        changed=copy.deepcopy(state);changed['units'].pop('A');changed['units']['B']['public']['id']='wrong'
        with self.assertRaises(ValueError):engine.collect(subjects,adapters,checkpoint=changed)
        self.assertEqual(calls,[])
        def commit(snapshot):
            snapshots.append(snapshot)
            if snapshot['pending'] is None:raise OSError('Unconfirmed commit acknowledgement')
        with self.assertRaises(OSError):engine.collect(subjects[:1],adapters,commit=commit)
        self.assertEqual(calls,['A'])
        # An exact durable completion permits restart, while intent alone needs reconciliation.
        resumed,_=engine.collect(subjects[:1],adapters,checkpoint=snapshots[-1])
        self.assertEqual(resumed[0]['id'],'A');self.assertEqual(calls,['A'])
        with self.assertRaises(engine.PendingCollection):engine.collect(subjects[:1],adapters,checkpoint=snapshots[0])


from tests import test_sourcing_connector_compile as compiler_fixture

class CompilerTests(compiler_fixture.Tests):
    def test_shared_compiler_source_fallback_coverage_and_batches(self):
        for name,config in CONFIGURATIONS.items():
            with self.subTest(schema=name):
                self.configuration=config;self.fresh_scope=True;self.record_count=24;self.coverage_round=True
                self.test_prepared_helpers_review_edges_outputs_and_api_limits_compose()
