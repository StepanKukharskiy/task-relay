"""Offline shared compiler graphs: no providers, native hosts, browsers or media renders."""
import copy,json,unittest
from pathlib import Path
from orchestrator import owner_intents,worker_tools,artifact_roles,job_owner,review_followup,contracts as c
from orchestrator.runtime import Runtime
from tests import test_owner_intents as intent, test_review_followup as reviews

class Routing(unittest.TestCase):
    create=intent.Tests.create
    source=intent.Tests.source
    session=intent.Tests.session
    def setUp(self):
        intent.Tests.setUp(self)
        self.grant.update(delegation_version=4,worker_tools_version=3,operations=['pptx.create','pptx.render','text.bundle','blender.scene','rhino.document'])
        self.grant['task_limits'].update(provider_requests=24,tool_calls=24,output_bytes=2000000)
    def doc(self):return json.dumps(dict(version=1,title='Exact saved document',slides=[dict(elements=[dict(type='text',text='Original source-bound claim',x=1,y=1,w=4,h=1)])]))
    def start(self,inputs=()):
        self.create(inputs=list(inputs));self.rt.tick('job');return self.session()
    def request(self,mode,**kw):return dict(request_id='route',tasks=[dict(id='make',objective='Use registered operation',work=dict(operation='pptx.create',mode=mode,**kw))])
    def test_saved_document_is_bound_without_authoring_and_other_json_stays_evidence(self):
        data=self.source(self.doc(),'saved.json','application/json');context=self.source('{"unrelated":true}','history.json','application/json')
        session=self.start([data,context]);args=self.request('execute',data={'artifact':data['artifact']})
        args['tasks'][0]['inputs']=[{'artifact':context['artifact']}]
        result=session.delegate(args);self.assertIn('make',result['queued'])
        spec=self.rt.spec(self.rt.task('job','make'));self.assertIsNone(spec['work_tools']['preparation']);self.assertNotIn('operation_contract',spec)
        self.assertEqual([i['artifact'] for i in artifact_roles.operation_inputs(spec,'build')],[data['artifact']])
        self.assertEqual({i['artifact'] for i in spec['inputs']},{data['artifact'],context['artifact']})
        self.rt.close();self.rt=Runtime(self.root,self.factory);self.rt.tick('job')
        f=self.factory.sessions[self.rt.task('job','make')['latest']]['frozen']
        self.assertEqual(f['work_tools']['intent'],spec['work_tools']['intent'])
        self.assertNotIn('operation_contract',f)
        self.assertEqual((Path(f['workspace'])/spec['work_tools']['intent']['data']).read_text(),self.doc())
    def test_saved_document_executes_real_build_without_authoring(self):
        from pptx import Presentation
        data=self.source(self.doc(),'saved.json','application/json');session=self.start([data])
        session.delegate(self.request('execute',data={'artifact':data['artifact']}));self.rt.tick('job')
        state=self.factory.sessions[self.rt.task('job','make')['latest']];f=state['frozen']
        control=Path(state['session']['control']);control.mkdir(parents=True,exist_ok=True)
        operations=worker_tools.Session(worker_tools.files(f,control),control)
        receipt=operations.run(dict(step='build',request_key='saved-build'))
        self.assertNotIn('error',receipt)
        deck=Presentation(str(Path(f['workspace'])/'outputs/work-build/result.pptx'))
        self.assertIn('Original source-bound claim',' '.join(shape.text for shape in deck.slides[0].shapes if shape.has_text_frame))
        self.assertFalse((Path(f['workspace'])/worker_tools.CONTENT).exists())
        resumed=worker_tools.Session(worker_tools.files(f,control),control)
        resumed.run(dict(step='build',request_key='saved-build'))
        self.assertEqual(len(resumed.receipts),1)
        self.assertEqual(Path(self.rt.artifact(data['artifact'])['blob']).read_text(),self.doc())
    def test_invalid_future_document_is_blocked_before_worker_dispatch(self):
        session=self.start();producer=intent.worker('prepare',media='application/json')
        args=self.request('execute',data={'task':'prepare','output':'result'})
        args['tasks'].insert(0,producer)
        session.delegate(args);self.rt.tick('job')
        aid=self.rt.task('job','prepare')['latest'];self.assertIsNotNone(aid)
        self.factory.finish(aid)
        producer=self.factory.sessions[aid]['frozen']
        (Path(producer['workspace'])/producer['outputs'][0]['path']).write_text('{}') # valid JSON, invalid operation data
        self.rt.tick('job')
        for row in self.rt.status('job')['tasks']:
            if row['latest'] and self.rt.spec(self.rt.task('job',row['id'])).get('review_of')=='prepare':
                self.factory.finish(row['latest'],decision='accept')
        self.rt.tick('job')
        task=self.rt.task('job','make')
        self.assertEqual(task['status'],'blocked');self.assertIsNone(task['latest']);self.assertEqual(task['attempts'],0)
        self.assertFalse(any(v['frozen']['id']=='make' for v in self.factory.sessions.values()))
        event=self.rt.db.execute("SELECT data FROM production_events WHERE task='make' AND kind='input_preflight_blocked'").fetchone()
        self.assertFalse(json.loads(event[0])['dispatched'])
    def test_ambiguous_missing_or_invalid_data_never_commits_a_worker(self):
        bad=self.source('{"not":"a document"}','bad.json','application/json');session=self.start([bad]);before=len(self.factory.calls)
        variants=[self.request('execute'),self.request('execute',data={'artifact':bad['artifact']}),self.request('execute',content_parameters={'max_records':1}),self.request('revise')]
        missing=self.request('create');missing['tasks'][0]['work'].pop('mode');variants.append(missing)
        for n,args in enumerate(variants):
            args['request_id']='invalid'+str(n)
            with self.subTest(n=n),self.assertRaises(ValueError):session.delegate(args)
            self.assertIsNone(self.rt.db.execute("select 1 from production_tasks where id='make'").fetchone())
        self.assertEqual(len(self.factory.calls),before)
    def test_create_and_revision_are_explicit_and_preserve_exact_baseline(self):
        data=self.source(self.doc(),'baseline.json','application/json');session=self.start([data])
        for mode in ('create','revise'):
            args=self.request(mode,**({'baseline':{'artifact':data['artifact']}} if mode=='revise' else {}))
            task=owner_intents.compile(session,args)[0]
            self.assertIsNotNone(task['work_tools']['preparation'])
            if mode=='revise':self.assertEqual(next(i['artifact'] for i in task['inputs'] if i['path']==task['work_tools']['intent']['baseline']),data['artifact'])
        self.assertEqual(Path(self.rt.artifact(data['artifact'])['blob']).read_text(),self.doc())
    def test_same_explicit_contract_covers_native_and_text_operations(self):
        for cap in ['blender.scene','rhino.document','text.bundle']:
            task=dict(id='work',instruction='',objective='Fixture',inputs=[],outputs=[])
            worker_tools.configure(task,dict(operation=cap,mode='execute'),dict(worker_tools_version=3,operations=self.grant['operations'],task_limits=self.grant['task_limits']))
            self.assertIsNone(task['work_tools']['preparation'])
            if cap!='text.bundle':
                with self.assertRaisesRegex(ValueError,'explicit data'):worker_tools.validate_inputs(task)
                authored=dict(id='work',instruction='',objective='Fixture',inputs=[],outputs=[])
                worker_tools.configure(authored,dict(operation=cap,mode='create'),dict(worker_tools_version=3,operations=self.grant['operations'],task_limits=self.grant['task_limits']))
                self.assertEqual(authored['work_tools']['preparation']['parameters']['target'],cap)
    def test_legacy_scope_is_not_reinterpreted(self):
        task=dict(id='old',instruction='',objective='Fixture',inputs=[],outputs=[])
        worker_tools.configure(task,dict(operation='pptx.create'),dict(worker_tools_version=2,operations=self.grant['operations'],task_limits=self.grant['task_limits']))
        self.assertIsNotNone(task['work_tools']['preparation']);self.assertNotIn('intent',task['work_tools'])

class CorrectionScopes(unittest.TestCase):
    create=reviews.Tests.create
    invoke=reviews.Tests.invoke
    session=reviews.Tests.session
    finish=reviews.Tests.finish
    delegate=reviews.Tests.delegate
    def setUp(self):
        reviews.Tests.setUp(self)
        self.grant.update(delegation_version=4,worker_tools_version=3,correction_scope_version=1,max_tasks=28)
    def test_helper_repairs_one_obligation_restart_final_review_keeps_all(self):
        self.create();self.delegate('draft');self.finish('draft','Missing image and labels')
        self.finish('draft_review','Image absent',decision='revise',defect='Required image missing')
        first=review_followup.findings(self.rt,'job')[0]['id']
        self.delegate('second');self.finish('second','Still incomplete')
        self.finish('second_review','Labels absent',decision='revise',defect='Required labels missing',resolved=False)
        issues={v['id'] for v in review_followup.findings(self.rt,'job')};self.assertGreaterEqual(len(issues),2)
        prior=self.rt.output('job','draft','outputs/result.txt')
        helper=intent.worker('helper',[dict(artifact=prior['id'])],media='text/plain');helper['repairs']=[first]
        self.invoke([('job_delegate',dict(request_id='helper',tasks=[helper,dict(id='helper_review',objective='Review exact helper repair only',review_of='helper')]))])
        self.rt.tick('job');f=self.factory.sessions[self.rt.task('job','helper')['latest']]['frozen']
        self.assertEqual([v['id'] for v in f['correction_context']],[first])
        self.assertNotIn('Repair each concrete defect',f['instruction'])
        self.finish('helper','{"image":"exact bound replacement"}')
        self.rt.close();self.rt=Runtime(self.root,self.factory)
        rf,_=self.finish('helper_review','Scoped repair independently inspected',decision='accept')
        self.assertEqual({v['id'] for v in rf['review_followup']['findings'].values()},{first})
        self.assertFalse(review_followup.final_review(self.rt,job_owner.original(self.rt,'job'),rf))
        self.assertEqual({v['id'] for v in review_followup.findings(self.rt,'job')},issues)
        candidate=self.rt.output('job','helper','outputs/result.txt')
        self.rt.tick('job')
        with self.assertRaises(ValueError):self.session().complete(dict(deliverables_json=c.encoded({'report':candidate['id']}),summary='Incomplete helper cannot close the job'))
        self.delegate('final');self.finish('final','Image and labels repaired')
        self.rt.tick('job');rf=self.factory.sessions[self.rt.task('job','final_review')['latest']]['frozen']
        self.assertTrue(issues<={v['id'] for v in rf['review_followup']['findings'].values()})
        # Omitting an obligation fails the unchanged exact-candidate report gate.
        policy=rf['review_followup'];report=dict(summary='Forged partial approval',decision='accept',instruction='',findings=[],checks={f'c{i}':dict(passed=True,evidence='Inspected') for i in range(1,len(rf['criteria'])+1)})
        report['checks']['c'+str(policy['criterion'])]['evidence']={}
        with self.assertRaises(ValueError):c.report(report,rf)
        self.finish('final_review','All exact repairs independently inspected',decision='accept')
        candidate=self.rt.output('job','final','outputs/result.txt')
        self.invoke([('job_complete',dict(deliverables_json=c.encoded({'report':candidate['id']}),summary='All obligations covered'))]);self.rt.tick('job')
        self.assertEqual(self.rt.status('job')['status'],'completed')
        self.assertEqual(Path(prior['blob']).read_text(),'Missing image and labels')
    def test_unknown_or_reviewer_selected_scope_is_rejected_atomically(self):
        self.create();self.rt.tick('job');session=self.session();before=len(self.factory.calls)
        for task in [dict(intent.worker('helper'),repairs=['invented']),dict(id='review',objective='Skip obligations',review_of='owner',repairs=['invented'])]:
            with self.assertRaises(ValueError):session.delegate(dict(request_id='invalid',tasks=[task]))
        self.assertEqual(len(self.factory.calls),before)
        self.assertEqual(self.rt.db.execute('select count(*) from production_tasks').fetchone()[0],1)
