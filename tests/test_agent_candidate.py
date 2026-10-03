"""External-agent candidate admission, exact package checks and decisions."""

import io
import json
from contextlib import redirect_stdout
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

from orchestrator import pptx_document, pptx_edit
from orchestrator.runtime import Runtime
from task_relay.bridge import State
from task_relay import (agent_candidate, impact_handoff, native_links,
                        production_control as pc, workflow_files)
from task_relay.job_delete import preview, delete
from task_relay.relay_paths import Paths
from tests.test_presentations import fixture


class AgentCandidateTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()
        self.state = State(self.root / 'data/state.sqlite')
        self.addCleanup(self.state.db.close)
        self.state.channel = 'fixture'
        self.rt = Runtime(pc.root(self.state), connection=self.state.db)
        self.job = 'job-' + 'd' * 24
        self.state.db.execute('''INSERT INTO relay_pipelines
            (id,request_id,request,title,spec,channel,provider,model,status,created)
            VALUES (?,?,?,?,?,?,?,?,?,?)''',
            (self.job, 1, 'Controlled withdrawal', 'Candidate fixture', '{}',
             'fixture', 'none', 'none', 'active', 1))
        self.state.db.commit()
        source = self.root / 'research.json'
        source.write_text(json.dumps({'answer': 'Oleander\tNerium oleander\tToxic'}))
        spec = fixture()
        spec['slides'][1]['elements'][1]['rows'] = [
            ['Plant', 'Botanical'], ['Oleander', 'Nerium oleander']]
        deck = self.root / 'baseline.pptx'
        deck.write_bytes(pptx_document.create(spec)[0])
        self.source = self.rt.register(source, 'Fixture research')
        self.deck = self.rt.register(deck, 'Fixture baseline PPTX')
        table = next(s for s in pptx_edit.inspect(deck.read_bytes())['slides'][1]['shapes']
                     if s.get('table_cells'))
        self.location = {'type': 'pptx_table_cell', 'slide': 2,
                         'shape_id': table['shape_id'], 'row': 1,
                         'column': 0, 'text': 'Oleander'}
        native_links.record_link(self.state.db, job=self.job,
            entity_key='Nerium oleander', predicate='plant_selection',
            source_artifact=self.source,
            source_locator={'type': 'json_line', 'field': 'answer',
                            'prefix': 'Oleander\tNerium oleander\t'},
            output_artifact=self.deck, output_locator=self.location,
            evidence='Fixture line and cell.', reviewer='source-reviewer',
            review_state='reviewed')
        native_links.record_coverage(self.state.db, job=self.job,
            output_artifact=self.deck, output_locator=self.location,
            state='complete', note='Exact source and output reviewed.')
        self.handoff = impact_handoff.record_native_withdrawal(self.state.db,
            job=self.job, request_key='change-1',
            exact_request='Replace the Oleander entry in this controlled fixture.',
            actor='fixture-requester', entity_key='Nerium oleander',
            baseline_artifact=self.deck)
        self.digest = impact_handoff.inspect(self.state.db, self.handoff)['plan_digest']
        self.manifest = {'version': 1, 'source_sha256': self.rt.artifact(self.deck)['sha256'],
            'edits': [{'kind': 'replace_text', 'slide': 2,
                       'shape_id': table['shape_id'], 'row': 1, 'column': 0,
                       'old': 'Oleander', 'new': 'Desert willow'}]}
        candidate, _ = pptx_edit.edit(deck.read_bytes(), self.manifest)
        self.candidate_path = self.root / 'agent-candidate.pptx'
        self.candidate_path.write_bytes(candidate)

    def submit(self, key='submission-1', path=None, manifest=None):
        return agent_candidate.submit(self.rt, handoff_id=self.handoff,
            submission_key=key, candidate_path=path or self.candidate_path,
            manifest=manifest or self.manifest, image_files={}, submitted_by='fixture-agent')

    def test_submit_rezip_verify_review_and_select_are_separate(self):
        # An agent may repackage the same OOXML members with different ZIP metadata.
        repacked = self.root / 'repacked.pptx'
        with ZipFile(self.candidate_path) as source, ZipFile(repacked, 'w', ZIP_DEFLATED) as target:
            for name in reversed(source.namelist()):
                target.writestr(name, source.read(name))
        aid = self.submit(path=repacked)
        self.assertEqual(self.submit(path=repacked), aid)
        row, candidate, handoff, checks = agent_candidate.verified(self.state.db, aid)
        self.assertEqual(checks['affected_locations_matched'], 1)
        self.assertEqual(checks['semantic_review'], 'pending')
        self.assertEqual(handoff['plan_digest'], self.digest)
        self.assertFalse(agent_candidate.records(self.state.db, self.job)['reviews'])
        feedback = agent_candidate.record_feedback(self.state.db,
            candidate_artifact=aid, feedback_key='user-layout-review-1',
            actor='user', exact_feedback='the layout looks good, i reviewd it',
            expected_sha256=candidate['sha256'], expected_plan_digest=self.digest)
        self.assertEqual(feedback, agent_candidate.record_feedback(self.state.db,
            candidate_artifact=aid, feedback_key='user-layout-review-1',
            actor='user', exact_feedback='the layout looks good, i reviewd it',
            expected_sha256=candidate['sha256'], expected_plan_digest=self.digest))
        with self.assertRaisesRegex(ValueError, 'different exact feedback'):
            agent_candidate.record_feedback(self.state.db,
                candidate_artifact=aid, feedback_key='user-layout-review-1',
                actor='user', exact_feedback='Accept every claim',
                expected_sha256=candidate['sha256'], expected_plan_digest=self.digest)
        self.assertEqual(agent_candidate._status(self.state.db, aid), 'verified_pending_review')
        self.assertFalse(agent_candidate.records(self.state.db, self.job)['reviews'])
        with self.assertRaisesRegex(ValueError, 'Independent reviewer'):
            agent_candidate.review(self.state.db, candidate_artifact=aid,
                decision='accept', reviewer='fixture-agent', note='Self review',
                expected_sha256=candidate['sha256'], expected_plan_digest=self.digest)
        with self.assertRaisesRegex(ValueError, 'independent acceptance'):
            agent_candidate.select(self.state.db, candidate_artifact=aid,
                selected_by='fixture-requester', receipt='Fixture selection',
                expected_sha256=candidate['sha256'], expected_plan_digest=self.digest)
        agent_candidate.review(self.state.db, candidate_artifact=aid,
            decision='accept', reviewer='fixture-reviewer', note='Reviewed the controlled edit.',
            expected_sha256=candidate['sha256'], expected_plan_digest=self.digest)
        with self.assertRaisesRegex(ValueError, 'version changed'):
            agent_candidate.select(self.state.db, candidate_artifact=aid,
                selected_by='fixture-requester', receipt='Fixture selection',
                expected_sha256='0' * 64, expected_plan_digest=self.digest)
        agent_candidate.select(self.state.db, candidate_artifact=aid,
            selected_by='fixture-requester', receipt='Exact fixture candidate selected.',
            expected_sha256=candidate['sha256'], expected_plan_digest=self.digest)
        self.state.db.execute('''INSERT INTO relay_pipeline_steps
            (pipeline,position,id,status,sources) VALUES (?,?,?,?,?)''',
            (self.job, 0, 'baseline', 'completed',
             json.dumps([{'artifact': self.source}, {'artifact': self.deck}])))
        self.state.db.commit()
        view = workflow_files.sync(self.state, self.job)
        with sqlite3.connect(view / '.relay/job.sqlite') as projection:
            self.assertEqual(projection.execute('SELECT count(*) FROM agent_candidates').fetchone()[0], 1)
            self.assertEqual(projection.execute('SELECT count(*) FROM agent_candidate_reviews').fetchone()[0], 1)
            self.assertEqual(projection.execute('SELECT count(*) FROM agent_candidate_feedback').fetchone()[0], 1)
            self.assertEqual(projection.execute("SELECT count(*) FROM process_records WHERE source_table='relay_agent_candidate_feedback' AND category='decisions'").fetchone()[0], 1)
            self.assertEqual(projection.execute('SELECT count(*) FROM agent_candidate_selections').fetchone()[0], 1)
            self.assertEqual(projection.execute("SELECT count(*) FROM process_records WHERE source_table='relay_agent_candidates'").fetchone()[0], 2)
            self.assertEqual(projection.execute("SELECT count(*) FROM process_records WHERE source_table='production_events' AND category='events'").fetchone()[0], 1)

    def test_reject_unplanned_edit_and_changed_submission_key(self):
        extra = {'version': 1, 'source_sha256': self.manifest['source_sha256'],
            'edits': self.manifest['edits'] + [
                {'kind': 'replace_text', 'slide': 1, 'shape_id': 2,
                 'old': 'Quarterly results', 'new': 'Unexpected change'}]}
        with self.assertRaisesRegex(ValueError, 'exactly the handoff affected'):
            self.submit(key='extra-manifest', manifest=extra)
        changed, _ = pptx_edit.edit(Path(self.rt.artifact(self.deck)['blob']).read_bytes(), extra)
        wrong = self.root / 'wrong.pptx'
        wrong.write_bytes(changed)
        with self.assertRaisesRegex(ValueError, 'differs from the declared edit'):
            self.submit(key='extra-file', path=wrong)
        aid = self.submit()
        with self.assertRaisesRegex(ValueError, 'different candidate content'):
            self.submit(path=wrong)
        self.assertEqual(len(agent_candidate.records(self.state.db, self.job)['candidates']), 1)
        native_links.record_coverage(self.state.db, job=self.job,
            output_artifact=self.deck, output_locator={'type': 'pptx_slide', 'slide': 1},
            state='incomplete', note='New scope.')
        with self.assertRaisesRegex(ValueError, 'not current'):
            agent_candidate.verified(self.state.db, aid)

    def test_picture_input_is_pinned_and_rechecked(self):
        from PIL import Image
        from pptx import Presentation
        import hashlib

        def png(color):
            buffer = io.BytesIO()
            Image.new('RGB', (4, 4), color).save(buffer, format='PNG')
            return buffer.getvalue()

        old, new = png('red'), png('blue')
        old_file = self.root / 'old.png'
        old_file.write_bytes(old)
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_picture(str(old_file), 0, 0)
        deck = self.root / 'photo-baseline.pptx'
        presentation.save(deck)
        baseline = self.rt.register(deck, 'Photo fixture deck')
        shape = next(s for s in pptx_edit.inspect(deck.read_bytes())['slides'][0]['shapes']
                     if s.get('image_sha256'))
        old_sha, new_sha = hashlib.sha256(old).hexdigest(), hashlib.sha256(new).hexdigest()
        manifest_source = self.root / 'photo-manifest.json'
        manifest_source.write_text(json.dumps({'photos': [{
            'id': 'oleander', 'sha256': old_sha, 'query': 'Nerium oleander'}]}))
        source = self.rt.register(manifest_source, 'Photo source manifest')
        locator = {'type': 'pptx_picture', 'slide': 1,
                   'shape_id': shape['shape_id'], 'sha256': old_sha}
        native_links.record_link(self.state.db, job=self.job,
            entity_key='Nerium oleander', predicate='plant_photo',
            source_artifact=source,
            source_locator={'type': 'photo_manifest_entry', 'id': 'oleander', 'sha256': old_sha},
            output_artifact=baseline, output_locator=locator,
            evidence='Exact photo hash.', reviewer='photo-reviewer', review_state='reviewed')
        native_links.record_coverage(self.state.db, job=self.job,
            output_artifact=baseline, output_locator=locator,
            state='complete', note='Exact photo bytes reviewed.')
        replacement = self.root / 'replacement.png'
        replacement.write_bytes(new)
        pinned = self.rt.register(replacement, 'Pinned replacement picture',
            task='agent_candidate', path='inputs/replacement.png')
        handoff = impact_handoff.record_native_replacement(self.state.db,
            job=self.job, request_key='photo-change', exact_request='Replace the fixture photo.',
            actor='fixture-requester', entity_key='Nerium oleander',
            baseline_artifact=baseline, replacement_artifact=pinned)
        frozen = impact_handoff.inspect(self.state.db, handoff)
        self.assertEqual(frozen['status'], 'current')
        self.assertEqual(frozen['plan']['replacement_sha256'], new_sha)
        self.state.db.execute("UPDATE relay_pipelines SET status='completed' WHERE id=?", (self.job,))
        self.state.db.commit()
        paths = Paths(self.root / 'app', self.root / 'data',
                      self.root / 'workspaces', self.root / 'generated')
        before_candidate = preview(self.job, paths)
        self.assertEqual(before_candidate['blockers'], [])
        self.assertEqual(before_candidate['counts']['production_artifacts'], 1)
        manifest = {'version': 1, 'source_sha256': self.rt.artifact(baseline)['sha256'],
            'edits': [{'kind': 'replace_image', 'slide': 1,
                       'shape_id': shape['shape_id'], 'old_sha256': old_sha,
                       'path': 'replacement.png', 'new_sha256': new_sha}]}
        raw, _ = pptx_edit.edit(deck.read_bytes(), manifest, images={'replacement.png': new})
        candidate = self.root / 'photo-candidate.pptx'
        candidate.write_bytes(raw)
        aid = agent_candidate.submit(self.rt, handoff_id=handoff,
            submission_key='photo-submission', candidate_path=candidate,
            manifest=manifest, image_files={'replacement.png': replacement},
            submitted_by='fixture-agent')
        self.assertEqual(agent_candidate.records(self.state.db, self.job)['inputs'][0]['artifact'], pinned)
        wrong_manifest = json.loads(json.dumps(manifest))
        wrong_manifest['edits'][0]['new_sha256'] = '1' * 64
        with self.assertRaisesRegex(ValueError, 'exact pinned picture'):
            agent_candidate.submit(self.rt, handoff_id=handoff,
                submission_key='wrong-picture', candidate_path=candidate,
                manifest=wrong_manifest, image_files={'replacement.png': replacement},
                submitted_by='fixture-agent')
        self.assertEqual(len(agent_candidate.records(self.state.db, self.job)['inputs']), 1)
        agent_candidate.verified(self.state.db, aid)
        self.state.db.execute("UPDATE relay_pipelines SET status='completed' WHERE id=?", (self.job,))
        self.state.db.commit()
        deletion = preview(self.job, paths)
        self.assertEqual(deletion['blockers'], [])
        self.assertEqual(deletion['counts']['production_artifacts'], 2)
        self.assertEqual(deletion['counts']['production_events'], 2)
        other = 'job-' + 'e' * 24
        self.state.db.execute('''INSERT INTO relay_pipelines
            (id,request_id,request,title,spec,channel,provider,model,status,created)
            VALUES (?,?,?,?,?,?,?,?,?,?)''',
            (other, 2, 'Other', 'Other', '{}', 'fixture', 'none', 'none', 'completed', 1))
        self.state.db.execute('INSERT INTO relay_impact_handoffs VALUES (?,?,?,?,?,?,?,?,?,?,?)',
            ('other-replacement', other, 'other-key', 'Other request',
             'native_subject_replacement', baseline,
             json.dumps({'entity_key': 'Nerium oleander',
                         'baseline_artifact': baseline,
                         'replacement_artifact': pinned}),
             'digest', '{}', 'other', 1))
        self.state.db.commit()
        self.assertIn('Another job has a reviewed link to this job’s artifact.',
                      preview(self.job, paths)['blockers'])
        self.state.db.execute('DELETE FROM relay_impact_handoffs WHERE job=?', (other,))
        self.state.db.execute('DELETE FROM relay_pipelines WHERE id=?', (other,))
        self.state.db.commit()
        input_id = agent_candidate.records(self.state.db, self.job)['inputs'][0]['artifact']
        blob = Path(self.rt.artifact(input_id)['blob'])
        blob.chmod(0o600)
        blob.write_bytes(old)
        with self.assertRaisesRegex(ValueError, 'missing, changed or too large'):
            agent_candidate.verified(self.state.db, aid)
        self.state.db.execute('DELETE FROM relay_agent_candidate_inputs WHERE candidate_artifact=?',
                              (aid,))
        self.state.db.commit()
        view = workflow_files.sync(self.state, self.job)
        with sqlite3.connect(view / '.relay/job.sqlite') as projection:
            kinds = {row[0] for row in projection.execute('SELECT kind FROM process_gaps')}
        self.assertIn('missing_declared_candidate_input', kinds)

    def test_cli_export_failure_retries_committed_submission(self):
        manifest_file = self.root / 'candidate-manifest.json'
        manifest_file.write_text(json.dumps(self.manifest))
        args = ['submit-native', '--database', str(self.root / 'data/state.sqlite'),
                '--handoff-id', self.handoff, '--submission-key', 'cli-retry',
                '--candidate', str(self.candidate_path), '--manifest', str(manifest_file),
                '--actor', 'fixture-agent']
        with patch('task_relay.workflow_files.sync', side_effect=OSError('export interrupted')):
            with self.assertRaisesRegex(OSError, 'export interrupted'):
                agent_candidate.main(args)
        self.assertEqual(len(agent_candidate.records(self.state.db, self.job)['candidates']), 1)
        self.state.db.execute('''INSERT INTO relay_pipeline_steps
            (pipeline,position,id,status,sources) VALUES (?,?,?,?,?)''',
            (self.job, 0, 'baseline', 'completed',
             json.dumps([{'artifact': self.source}, {'artifact': self.deck}])))
        self.state.db.commit()
        with redirect_stdout(io.StringIO()):
            agent_candidate.main(args)
        saved = agent_candidate.records(self.state.db, self.job)['candidates']
        self.assertEqual(len(saved), 1)
        aid = saved[0]['candidate_artifact']
        output = io.StringIO()
        with redirect_stdout(output):
            agent_candidate.main(['verify', '--database', args[2], '--artifact', aid])
        self.assertEqual(json.loads(output.getvalue())['status'], 'verified_pending_review')
        note = self.root / 'review.txt'
        note.write_text('The controlled text change was independently reviewed.')
        receipt = self.root / 'selection.txt'
        receipt.write_text('The fixture requester explicitly selected this exact candidate.')
        with redirect_stdout(io.StringIO()):
            agent_candidate.main(['review', '--database', args[2], '--artifact', aid,
                '--decision', 'accept', '--reviewer', 'fixture-reviewer',
                '--note-file', str(note), '--expected-sha256', saved[0]['candidate_sha256'],
                '--expected-plan-digest', self.digest])
            agent_candidate.main(['select', '--database', args[2], '--artifact', aid,
                '--selected-by', 'fixture-requester', '--receipt-file', str(receipt),
                '--expected-sha256', saved[0]['candidate_sha256'],
                '--expected-plan-digest', self.digest])
        self.assertEqual(len(agent_candidate.records(self.state.db, self.job)['selections']), 1)
        output = io.StringIO()
        with redirect_stdout(output):
            agent_candidate.main(['verify', '--database', args[2], '--artifact', aid])
        self.assertEqual(json.loads(output.getvalue())['status'], 'selected')

    def test_delete_owns_registered_candidate_and_registration_receipt(self):
        aid = self.submit()
        candidate = self.rt.artifact(aid)
        agent_candidate.record_feedback(self.state.db,
            candidate_artifact=aid, feedback_key='layout-comment',
            actor='user', exact_feedback='The layout looks good.',
            expected_sha256=candidate['sha256'], expected_plan_digest=self.digest)
        self.state.db.execute("UPDATE relay_pipelines SET status='completed' WHERE id=?", (self.job,))
        self.state.db.commit()
        paths = Paths(self.root / 'app', self.root / 'data',
                      self.root / 'workspaces', self.root / 'generated')
        check = preview(self.job, paths)
        self.assertEqual(check['blockers'], [])
        self.assertEqual(check['counts']['relay_agent_candidates'], 1)
        self.assertEqual(check['counts']['relay_agent_candidate_feedback'], 1)
        self.assertEqual(check['counts']['production_artifacts'], 1)
        self.assertEqual(check['counts']['production_events'], 1)
        other = 'job-' + 'e' * 24
        self.state.db.execute('''INSERT INTO relay_pipelines
            (id,request_id,request,title,spec,channel,provider,model,status,created)
            VALUES (?,?,?,?,?,?,?,?,?,?)''',
            (other, 2, 'Other', 'Other', '{}', 'fixture', 'none', 'none', 'completed', 1))
        self.state.db.execute('INSERT INTO relay_impact_handoffs VALUES (?,?,?,?,?,?,?,?,?,?,?)',
            ('other-handoff', other, 'other-key', 'Other request',
             'native_subject_withdrawal', aid, '{}', 'digest', '{}', 'other', 1))
        self.state.db.commit()
        self.assertIn('Another job has a reviewed link to this job’s artifact.',
                      preview(self.job, paths)['blockers'])
        self.state.db.execute("DELETE FROM relay_impact_handoffs WHERE job=?", (other,))
        self.state.db.execute("DELETE FROM relay_pipelines WHERE id=?", (other,))
        self.state.db.commit()
        check = preview(self.job, paths)
        self.assertEqual(delete(self.job, check['digest'], paths)['status'], 'complete')
        self.assertIsNone(self.state.db.execute(
            'SELECT id FROM production_artifacts WHERE id=?', (aid,)).fetchone())
        self.assertEqual(self.state.db.execute(
            'SELECT count(*) FROM relay_agent_candidate_feedback').fetchone()[0], 0)

    def test_structural_row_removal_requires_every_removed_run_to_be_linked(self):
        table = next(shape for shape in pptx_edit.inspect(
            Path(self.rt.artifact(self.deck)['blob']).read_bytes())['slides'][1]['shapes']
            if shape.get('table_cells'))
        manifest = {'version': 1, 'source_sha256': self.manifest['source_sha256'],
                    'edits': [{'kind': 'remove_table_row', 'slide': 2,
                               'shape_id': table['shape_id'], 'row': 1,
                               'old_runs': ['Oleander', 'Nerium oleander']}]}
        raw, _ = pptx_edit.edit(Path(self.rt.artifact(self.deck)['blob']).read_bytes(), manifest)
        candidate = self.root / 'remove-row.pptx'
        candidate.write_bytes(raw)
        with self.assertRaisesRegex(ValueError, 'exactly the handoff affected'):
            agent_candidate.submit(self.rt, handoff_id=self.handoff,
                submission_key='incomplete-row', candidate_path=candidate,
                manifest=manifest, image_files={}, submitted_by='fixture-agent')
        botanical = {**self.location, 'column': 1, 'text': 'Nerium oleander'}
        native_links.record_link(self.state.db, job=self.job,
            entity_key='Nerium oleander', predicate='plant_selection',
            source_artifact=self.source,
            source_locator={'type': 'json_line', 'field': 'answer',
                            'prefix': 'Oleander\tNerium oleander\t'},
            output_artifact=self.deck, output_locator=botanical,
            evidence='Exact reviewed botanical name.', reviewer='source-reviewer',
            review_state='reviewed')
        native_links.record_coverage(self.state.db, job=self.job,
            output_artifact=self.deck, output_locator=botanical,
            state='complete', note='Second removed run reviewed.')
        handoff = impact_handoff.record_native_withdrawal(self.state.db,
            job=self.job, request_key='remove-row-2',
            exact_request='Remove the reviewed Oleander row in this fixture.',
            actor='fixture-requester', entity_key='Nerium oleander',
            baseline_artifact=self.deck)
        aid = agent_candidate.submit(self.rt, handoff_id=handoff,
            submission_key='complete-row', candidate_path=candidate,
            manifest=manifest, image_files={}, submitted_by='fixture-agent')
        self.assertEqual(agent_candidate.verified(self.state.db, aid)[3]['affected_locations_matched'], 2)
