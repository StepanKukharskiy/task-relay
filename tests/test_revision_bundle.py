"""Small controlled cross-artifact revision and recovery checks."""

import hashlib
import io
import json
from pathlib import Path
import tempfile
import unittest

from PIL import Image
from pptx import Presentation

from orchestrator import pptx_edit
from orchestrator.runtime import Runtime
from task_relay.bridge import State
from task_relay import (agent_candidate, bundle_continuation, impact_handoff,
                        native_links, production_control, revision_bundle,
                        workflow_files)
from task_relay.job_delete import preview
from task_relay.relay_paths import Paths


def png(size, color):
    buffer = io.BytesIO()
    Image.new('RGB', size, color).save(buffer, format='PNG')
    return buffer.getvalue()


class BundleTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()
        self.state = State(self.root/'data/state.sqlite')
        self.addCleanup(self.state.db.close)
        self.state.channel = 'fixture'
        self.job = 'job-' + 'd'*24
        self.state.db.execute('''INSERT INTO relay_pipelines
            (id,request_id,request,title,spec,channel,provider,model,status,created)
            VALUES (?,?,?,?,?,?,?,?,?,?)''',
            (self.job, 1, 'Replace this picture.', 'Fixture', '{}',
             'fixture', 'none', 'none', 'completed', 1))
        self.state.db.commit()
        self.rt = Runtime(production_control.root(self.state), connection=self.state.db)
        old, new = png((8, 4), 'green'), png((2, 4), 'red')
        self.old_sha, self.new_sha = (hashlib.sha256(item).hexdigest()
                                      for item in (old, new))
        old_image = self.root/'old.png'
        old_image.write_bytes(old)
        new_image = self.root/'new.png'
        new_image.write_bytes(new)
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_picture(str(old_image), 0, 0, width=4_000_000, height=2_000_000)
        slide.shapes.add_textbox(0, 2_100_000, 4_000_000, 500_000).text = 'Selection pending review'
        baseline = self.root/'baseline.pptx'
        deck.save(baseline)
        self.deck = self.rt.register(baseline, 'Baseline native deck', path='baseline.pptx')
        photo_baseline = {'photos': [{'id': 'old-plant', 'query': 'Nerium oleander',
                                      'sha256': self.old_sha, 'staged_path': 'old.png'}]}
        photo_file = self.root/'selected-photo-manifest.json'
        photo_file.write_text(json.dumps(photo_baseline))
        self.photo = self.rt.register(photo_file, 'Baseline photo manifest',
                                      path='baseline/selected-photo-manifest.json')
        slide_file = self.root/'slides.json'
        slide_file.write_text(json.dumps({'slides': [{'elements': [{
            'type': 'image', 'path': 'old.png', 'x': 0, 'y': 0,
            'w': 4_000_000/914400, 'h': 2_000_000/914400},
            {'type': 'text', 'text': 'Selection pending review'}]}]}))
        self.slides = self.rt.register(slide_file, 'Baseline slide spec',
                                       path='baseline/slides.json')
        locator = {'type': 'pptx_picture', 'slide': 1, 'shape_id': 2,
                   'sha256': self.old_sha}
        native_links.record_link(self.state.db, job=self.job,
            entity_key='Nerium oleander', predicate='plant_photo',
            source_artifact=self.photo,
            source_locator={'type': 'photo_manifest_entry',
                            'id': 'old-plant', 'sha256': self.old_sha},
            output_artifact=self.deck, output_locator=locator,
            evidence='Exact fixture photo.', reviewer='fixture-source-reviewer',
            review_state='reviewed')
        native_links.record_coverage(self.state.db, job=self.job,
            output_artifact=self.deck, output_locator=locator, state='complete',
            note='Exact fixture picture location.')
        self.image = self.rt.register(new_image, 'User replacement image',
                                      task='agent_candidate', path='inputs/new.png')
        self.handoff = impact_handoff.record_native_replacement(self.state.db,
            job=self.job, request_key='picture-1', exact_request='Replace this picture.',
            actor='user', entity_key='Nerium oleander',
            baseline_artifact=self.deck, replacement_artifact=self.image)
        manifest = {'version': 1,
                    'source_sha256': hashlib.sha256(baseline.read_bytes()).hexdigest(),
                    'edits': [{'kind': 'replace_image', 'slide': 1, 'shape_id': 2,
                               'old_sha256': self.old_sha, 'new_sha256': self.new_sha,
                               'path': 'new.png', 'fit': 'contain'}]}
        candidate_raw, _ = pptx_edit.edit(baseline.read_bytes(), manifest,
                                          images={'new.png': new})
        candidate_path = self.root/'candidate.pptx'
        candidate_path.write_bytes(candidate_raw)
        self.pptx = agent_candidate.submit(self.rt, handoff_id=self.handoff,
            submission_key='picture-candidate', candidate_path=candidate_path,
            manifest=manifest, image_files={'new.png': new_image},
            submitted_by='fixture-agent')
        placed = next(item for item in Presentation(io.BytesIO(candidate_raw)).slides[0].shapes
                      if item.shape_id == 2)
        self.new_slide = {'slides': [{'elements': [{
            'type': 'image', 'path': 'inputs/new.png',
            'x': placed.left/914400, 'y': placed.top/914400,
            'w': placed.width/914400, 'h': placed.height/914400},
            {'type': 'text', 'text': 'Selection pending review'}]}]}
        self.new_photo = {'photos': [{'id': 'rose-unverified',
            'status': 'user_provided_unverified',
            'sha256': self.new_sha, 'bytes': len(new), 'dimensions': [2, 4],
            'media_type': 'image/png', 'staged_path': 'inputs/new.png',
            'author': None, 'license': None, 'source_url': None}]}
        self.companions = [
            {'role': 'slides', 'baseline_artifact': self.slides,
             'patches': [{'op': 'replace', 'path': '/slides/0/elements/0',
                          'old': json.loads(slide_file.read_text())['slides'][0]['elements'][0],
                          'new': self.new_slide['slides'][0]['elements'][0]}]},
            {'role': 'photo_manifest', 'baseline_artifact': self.photo,
             'patches': [{'op': 'replace', 'path': '/photos/0',
                          'old': photo_baseline['photos'][0],
                          'new': self.new_photo['photos'][0]}]},
        ]
        self.cross_checks = [
            {'kind': 'picture', 'role': 'slides',
             'pointer': '/slides/0/elements/0', 'slide': 1, 'shape_id': 2,
             'manifest_role': 'photo_manifest', 'manifest_pointer': '/photos/0',
             'unknown_fields': ['author', 'license', 'source_url'],
             'status': 'user_provided_unverified'},
            {'kind': 'forbid_terms', 'roles': ['slides', 'photo_manifest'],
             'terms': ['Oleander', 'Nerium']},
        ]
        self.plan = revision_bundle.record_plan(self.state.db,
            handoff_id=self.handoff, request_key='bundle-1',
            exact_request='Replace this picture.', actor='user',
            companions=self.companions, cross_checks=self.cross_checks)
        self.files = {'slides': self.root/'candidate-slides.json',
                      'photo_manifest': self.root/'candidate-photo.json'}
        self.files['slides'].write_text(json.dumps(self.new_slide))
        self.files['photo_manifest'].write_text(json.dumps(self.new_photo))

    def test_scoped_companions_are_pinned_and_projected_without_selection(self):
        ident = revision_bundle.submit(self.rt, plan_id=self.plan,
            pptx_candidate_artifact=self.pptx, submission_key='bundle-candidate',
            companion_files=self.files, submitted_by='fixture-agent')
        self.assertEqual(ident, revision_bundle.submit(self.rt, plan_id=self.plan,
            pptx_candidate_artifact=self.pptx, submission_key='bundle-candidate',
            companion_files=self.files, submitted_by='fixture-agent'))
        _, checks = revision_bundle.verified(self.state.db, ident)
        self.assertEqual(checks['couplings']['passed'], 2)
        self.assertEqual(checks['review'], 'pending')
        self.assertEqual(revision_bundle.status(self.state.db, ident),
                         'verified_pending_review')
        view = workflow_files.sync(self.state, self.job)
        import sqlite3
        with sqlite3.connect(view/'.relay/job.sqlite') as projection:
            self.assertEqual(projection.execute(
                'SELECT count(*) FROM revision_bundle_files').fetchone()[0], 2)
            self.assertEqual(projection.execute(
                "SELECT count(*) FROM process_records WHERE source_table='relay_revision_bundles'").fetchone()[0], 2)
        deletion = preview(self.job, Paths(self.root/'app', self.root/'data',
                                          self.root/'workspaces', self.root/'generated'))
        self.assertEqual(deletion['blockers'], [])
        self.assertEqual(deletion['counts']['relay_revision_bundle_files'], 2)

    def test_extra_change_or_wrong_picture_is_rejected_before_registration(self):
        changed = json.loads(self.files['slides'].read_text())
        changed['unrelated'] = 'changed'
        self.files['slides'].write_text(json.dumps(changed))
        with self.assertRaisesRegex(ValueError, 'undeclared or missing JSON change'):
            revision_bundle.submit(self.rt, plan_id=self.plan,
                pptx_candidate_artifact=self.pptx, submission_key='bad-extra',
                companion_files=self.files, submitted_by='fixture-agent')
        self.files['slides'].write_text(json.dumps(self.new_slide))
        changed = json.loads(self.files['photo_manifest'].read_text())
        changed['photos'][0]['sha256'] = '0'*64
        self.files['photo_manifest'].write_text(json.dumps(changed))
        with self.assertRaisesRegex(ValueError, 'undeclared or missing JSON change'):
            revision_bundle.submit(self.rt, plan_id=self.plan,
                pptx_candidate_artifact=self.pptx, submission_key='bad-photo',
                companion_files=self.files, submitted_by='fixture-agent')
        self.assertFalse(revision_bundle.records(self.state.db, self.job)['bundles'])

    def test_stale_baseline_blocks_frozen_bundle(self):
        source = native_links._artifact(self.state.db, self.slides)
        Path(source['blob']).chmod(0o600)
        Path(source['blob']).write_text('{}')
        self.assertEqual(revision_bundle.inspect_plan(self.state.db, self.plan)['status'], 'stale')
        with self.assertRaisesRegex(ValueError, 'stale'):
            revision_bundle.submit(self.rt, plan_id=self.plan,
                pptx_candidate_artifact=self.pptx, submission_key='stale',
                companion_files=self.files, submitted_by='fixture-agent')

    def test_delete_blocks_another_jobs_frozen_reference_to_companion(self):
        ident = revision_bundle.submit(self.rt, plan_id=self.plan,
            pptx_candidate_artifact=self.pptx, submission_key='bundle-candidate',
            companion_files=self.files, submitted_by='fixture-agent')
        companion = next(item for item in revision_bundle.records(self.state.db, self.job)['files']
                         if item['bundle_id'] == ident and item['role'] == 'slides')
        other_job = 'job-' + 'e'*24
        self.state.db.execute('''INSERT INTO relay_revision_bundle_plans
            VALUES (?,?,?,?,?,?,?,?,?)''',
            ('other-plan', other_job, 'other-request', 'other-handoff',
             'Reference a shared artifact.', 'digest', json.dumps({
                 'image_artifact': self.image, 'companions': [{
                     'baseline_artifact': companion['candidate_artifact']}]},
                 sort_keys=True), 'other-agent', 2))
        self.state.db.commit()
        deletion = preview(self.job, Paths(self.root/'app', self.root/'data',
                                          self.root/'workspaces', self.root/'generated'))
        self.assertTrue(any('Another job' in item for item in deletion['blockers']))

    def test_exact_review_then_atomic_set_selection(self):
        ident = revision_bundle.submit(self.rt, plan_id=self.plan,
            pptx_candidate_artifact=self.pptx, submission_key='bundle-candidate',
            companion_files=self.files, submitted_by='fixture-agent')
        bundle, checks = revision_bundle.verified(self.state.db, ident)
        with self.assertRaisesRegex(ValueError, 'exact set digest'):
            revision_bundle.review(self.state.db, bundle_id=ident, decision='accept',
                reviewer='user', exact_feedback='Yes, all good.', note='Accept provisional rose.',
                expected_set_digest='0'*64)
        revision_bundle.review(self.state.db, bundle_id=ident, decision='accept',
            reviewer='user', exact_feedback='Yes, all good.', note='Accept provisional rose.',
            expected_set_digest=bundle['set_digest'])
        self.assertEqual(revision_bundle.status(self.state.db, ident), 'accepted')
        with self.assertRaisesRegex(ValueError, 'native candidate needs independent acceptance'):
            revision_bundle.select(self.state.db, bundle_id=ident, selected_by='user',
                receipt='Select exact approved provisional set.',
                expected_set_digest=bundle['set_digest'])
        native = agent_candidate.verified(self.state.db, self.pptx)
        agent_candidate.review(self.state.db, candidate_artifact=self.pptx,
            decision='accept', reviewer='user', note='Approved exact native candidate.',
            expected_sha256=native[1]['sha256'],
            expected_plan_digest=native[2]['plan_digest'])
        revision_bundle.select(self.state.db, bundle_id=ident, selected_by='user',
            receipt='Select exact approved provisional set.',
            expected_set_digest=bundle['set_digest'])
        self.assertEqual(revision_bundle.status(self.state.db, ident), 'selected')
        with self.assertRaisesRegex(ValueError, 'already selected'):
            revision_bundle.select(self.state.db, bundle_id=ident, selected_by='user',
                receipt='Select exact approved provisional set.',
                expected_set_digest=bundle['set_digest'])
        self.assertEqual(checks['selection'], 'none')  # immutable submission checks
        records = revision_bundle.records(self.state.db, self.job)
        self.assertEqual(len(records['reviews']), 1)
        self.assertEqual(len(records['selections']), 1)
        view = workflow_files.sync(self.state, self.job)
        marker = json.loads((view/'.relay-workflow.json').read_text())
        exported = json.loads((view/'.relay/snapshots'/marker['snapshot']/'job-state.json').read_text())
        selected = {item['id'] for item in exported['artifacts'] if item['selected']}
        self.assertIn(self.pptx, selected)
        self.assertNotIn(self.deck, selected)
        self.assertIn(next(item['candidate_artifact'] for item in records['files']
                           if item['role'] == 'slides'), selected)
        self.assertNotIn(self.slides, selected)
        import sqlite3
        with sqlite3.connect(view/'.relay/job.sqlite') as projection:
            self.assertEqual(projection.execute(
                'SELECT count(*) FROM revision_bundle_selections').fetchone()[0], 1)
            self.assertEqual(projection.execute(
                "SELECT count(*) FROM process_records WHERE source_table='relay_revision_bundle_selections'").fetchone()[0], 1)
        deletion = preview(self.job, Paths(self.root/'app', self.root/'data',
                                          self.root/'workspaces', self.root/'generated'))
        self.assertEqual(deletion['blockers'], [])
        self.assertEqual(deletion['counts']['relay_revision_bundle_selections'], 1)
        competing = revision_bundle.submit(self.rt, plan_id=self.plan,
            pptx_candidate_artifact=self.pptx, submission_key='second-bundle-candidate',
            companion_files=self.files, submitted_by='fixture-agent')
        competing_row, _ = revision_bundle.verified(self.state.db, competing)
        revision_bundle.review(self.state.db, bundle_id=competing,
            decision='accept', reviewer='user', exact_feedback='Accept second set.',
            note='Competing exact set.', expected_set_digest=competing_row['set_digest'])
        with self.assertRaisesRegex(ValueError, 'Another revision set'):
            revision_bundle.select(self.state.db, bundle_id=competing, selected_by='user',
                receipt='Try selecting competing set.',
                expected_set_digest=competing_row['set_digest'])


class ContinuationTests(unittest.TestCase):
    setUp = BundleTests.setUp

    def selected_parent(self):
        parent = revision_bundle.submit(self.rt, plan_id=self.plan,
            pptx_candidate_artifact=self.pptx, submission_key='parent-bundle',
            companion_files=self.files, submitted_by='fixture-agent')
        bundle, _ = revision_bundle.verified(self.state.db, parent)
        native = agent_candidate.verified(self.state.db, self.pptx)
        agent_candidate.review(self.state.db, candidate_artifact=self.pptx,
            decision='accept', reviewer='user', note='Accept exact provisional picture.',
            expected_sha256=native[1]['sha256'],
            expected_plan_digest=native[2]['plan_digest'])
        revision_bundle.review(self.state.db, bundle_id=parent, decision='accept',
            reviewer='user', exact_feedback='Yes, all good.',
            note='Accept exact provisional set.',
            expected_set_digest=bundle['set_digest'])
        receipt = json.dumps({'exact_rights_response': 'Have permission to',
                              'image_sha256': self.new_sha}, sort_keys=True)
        revision_bundle.select(self.state.db, bundle_id=parent, selected_by='user',
            receipt=receipt, expected_set_digest=bundle['set_digest'])
        return parent

    def followon_plan(self, parent):
        native = native_links._artifact(self.state.db, self.pptx)
        manifest = {'version': 1, 'source_sha256': native['sha256'], 'edits': [{
            'kind': 'replace_text', 'slide': 1, 'shape_id': 3,
            'old': 'Selection pending review', 'new': 'User-selected replacement'}]}
        self.follow_slides = json.loads(self.files['slides'].read_text())
        self.follow_slides['slides'][0]['elements'][1]['text'] = 'User-selected replacement'
        self.follow_photo = json.loads(self.files['photo_manifest'].read_text())
        self.follow_photo['photos'][0]['permission_attestation'] = {
            'exact_response': 'Have permission to', 'image_sha256': self.new_sha}
        companions = [
            {'role': 'slides', 'patches': [{'op': 'replace',
                'path': '/slides/0/elements/1/text',
                'old': 'Selection pending review', 'new': 'User-selected replacement'}]},
            {'role': 'photo_manifest', 'patches': [{'op': 'add',
                'path': '/photos/0/permission_attestation',
                'new': self.follow_photo['photos'][0]['permission_attestation']}]},
        ]
        checks = [*self.cross_checks, {'kind': 'text_run', 'role': 'slides',
            'pointer': '/slides/0/elements/1/text', 'slide': 1, 'shape_id': 3}]
        plan = bundle_continuation.record_plan(self.state.db,
            parent_bundle=parent, request_key='followon-request',
            exact_request='Confirm the selected picture.',
            intent='Update selection status while reusing the picture.', actor='user',
            manifest=manifest, companions=companions, cross_checks=checks,
            decision_projections=[{'receipt_key': 'exact_rights_response',
                'role': 'photo_manifest',
                'pointer': '/photos/0/permission_attestation/exact_response'}])
        baseline = Path(native['blob']).read_bytes()
        candidate, _ = pptx_edit.edit(baseline, manifest, images={})
        self.follow_pptx = self.root/'followon.pptx'
        self.follow_pptx.write_bytes(candidate)
        self.follow_files = {'slides': self.root/'followon-slides.json',
                             'photo_manifest': self.root/'followon-photos.json'}
        self.follow_files['slides'].write_text(json.dumps(self.follow_slides))
        self.follow_files['photo_manifest'].write_text(json.dumps(self.follow_photo))
        return plan

    def test_second_revision_reuses_image_and_preserves_unrelated_members(self):
        parent = self.selected_parent()
        plan = self.followon_plan(parent)
        self.assertEqual(bundle_continuation.inspect_plan(self.state.db, plan)['status'],
                         'current')
        candidate = bundle_continuation.submit(self.rt, plan_id=plan,
            submission_key='followon-candidate', pptx_file=self.follow_pptx,
            companion_files=self.follow_files, submitted_by='fixture-agent')
        row, checks = bundle_continuation.verified(self.state.db, candidate)
        self.assertEqual(checks['native_members']['changed'], ['ppt/slides/slide1.xml'])
        self.assertEqual(checks['image_reused_sha256'], self.new_sha)
        self.assertEqual(checks['couplings']['passed'], 3)
        self.assertEqual(checks['review'], 'pending')
        view = workflow_files.sync(self.state, self.job)
        import sqlite3
        with sqlite3.connect(view/'.relay/job.sqlite') as projection:
            self.assertEqual(projection.execute(
                'SELECT count(*) FROM bundle_continuations').fetchone()[0], 1)
        deletion = preview(self.job, Paths(self.root/'app', self.root/'data',
                                          self.root/'workspaces', self.root/'generated'))
        self.assertEqual(deletion['blockers'], [])
        self.assertEqual(deletion['counts']['relay_bundle_continuations'], 1)

    def test_changed_decision_receipt_stales_plan_and_blocks_submission(self):
        parent = self.selected_parent()
        plan = self.followon_plan(parent)
        self.state.db.execute('''UPDATE relay_revision_bundle_selections SET receipt=?
            WHERE bundle_id=?''', (json.dumps({'exact_rights_response': 'Unknown'}), parent))
        self.state.db.commit()
        self.assertEqual(bundle_continuation.inspect_plan(self.state.db, plan)['status'],
                         'stale')
        with self.assertRaisesRegex(ValueError, 'stale'):
            bundle_continuation.submit(self.rt, plan_id=plan,
                submission_key='stale-followon', pptx_file=self.follow_pptx,
                companion_files=self.follow_files, submitted_by='fixture-agent')

    def test_continuation_selection_promotes_exact_set_and_preserves_history(self):
        parent = self.selected_parent()
        plan = self.followon_plan(parent)
        candidate = bundle_continuation.submit(self.rt, plan_id=plan,
            submission_key='followon-for-review', pptx_file=self.follow_pptx,
            companion_files=self.follow_files, submitted_by='fixture-agent')
        row, _ = bundle_continuation.verified(self.state.db, candidate)
        with self.assertRaisesRegex(ValueError, 'acceptance is required'):
            bundle_continuation.select(self.state.db, candidate_id=candidate,
                selected_by='user', receipt='Select exact set.',
                expected_set_digest=row['set_digest'])
        bundle_continuation.review(self.state.db, candidate_id=candidate,
            decision='accept', reviewer='user', note='Reviewed exact second set.',
            expected_set_digest=row['set_digest'])
        bundle_continuation.select(self.state.db, candidate_id=candidate,
            selected_by='user', receipt='Select exact second set.',
            expected_set_digest=row['set_digest'])
        self.assertEqual(bundle_continuation.status(self.state.db, candidate), 'selected')
        view = workflow_files.sync(self.state, self.job)
        marker = json.loads((view/'.relay-workflow.json').read_text())
        exported = json.loads((view/'.relay/snapshots'/marker['snapshot']/'job-state.json').read_text())
        selected = {item['id'] for item in exported['artifacts'] if item['selected']}
        self.assertIn(row['pptx_artifact'], selected)
        self.assertIn(row['slides_artifact'], selected)
        self.assertIn(row['photo_manifest_artifact'], selected)
        self.assertNotIn(self.pptx, selected)
        import sqlite3
        with sqlite3.connect(view/'.relay/job.sqlite') as projection:
            self.assertEqual(projection.execute(
                'SELECT count(*) FROM bundle_continuation_selections').fetchone()[0], 1)
        deletion = preview(self.job, Paths(self.root/'app', self.root/'data',
                                          self.root/'workspaces', self.root/'generated'))
        self.assertEqual(deletion['blockers'], [])
        self.assertEqual(deletion['counts']['relay_bundle_continuation_selections'], 1)

    def test_unchanged_photo_manifest_is_reused_by_artifact_identity(self):
        parent = self.selected_parent()
        full_plan = self.followon_plan(parent)
        spec = bundle_continuation.inspect_plan(self.state.db, full_plan)['plan']
        slide_spec = next(item for item in spec['companions'] if item['role'] == 'slides')
        reuse_plan = bundle_continuation.record_plan(self.state.db,
            parent_bundle=parent, request_key='reuse-photo-manifest',
            exact_request='Update the status text only.',
            intent='Update exact selection text and reuse photo metadata.', actor='user',
            manifest=spec['manifest'],
            companions=[{'role': 'slides', 'patches': slide_spec['patches']}],
            cross_checks=spec['cross_checks'], decision_projections=[])
        candidate = bundle_continuation.submit(self.rt, plan_id=reuse_plan,
            submission_key='reuse-photo-candidate', pptx_file=self.follow_pptx,
            companion_files={'slides': self.follow_files['slides']},
            submitted_by='fixture-agent')
        row, checks = bundle_continuation.verified(self.state.db, candidate)
        parent_photo = next(item['candidate_artifact'] for item in
            revision_bundle.records(self.state.db, self.job)['files']
            if item['bundle_id'] == parent and item['role'] == 'photo_manifest')
        self.assertEqual(row['photo_manifest_artifact'], parent_photo)
        self.assertTrue(checks['companions']['photo_manifest']['reused'])
        self.assertEqual(checks['image_reused_sha256'], self.new_sha)

    def test_undeclared_companion_change_is_rejected_before_registration(self):
        parent = self.selected_parent()
        plan = self.followon_plan(parent)
        changed = json.loads(self.follow_files['slides'].read_text())
        changed['unrelated'] = True
        self.follow_files['slides'].write_text(json.dumps(changed))
        with self.assertRaisesRegex(ValueError, 'exact patches'):
            bundle_continuation.submit(self.rt, plan_id=plan,
                submission_key='extra-followon', pptx_file=self.follow_pptx,
                companion_files=self.follow_files, submitted_by='fixture-agent')
        self.assertFalse(bundle_continuation.records(self.state.db, self.job)['candidates'])


if __name__ == '__main__':
    unittest.main()
