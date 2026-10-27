"""Owner intents through real admission/workers/native files; no live providers."""
import copy,hashlib,io,json,unittest
from pathlib import Path
from unittest.mock import patch
from orchestrator import contracts as c,job_owner,owner_intents as intents,executors,operation_records
from orchestrator.runtime import Runtime
from orchestrator.gemini_worker import execute
from orchestrator.workers import atomic
from tests import test_job_owner_loop as fixture
from tests.test_operation_session import Script,CONFIG


def output(name='result',media='text/plain'):
    return dict(name=name,media_type=media,purpose='Requested result')

def worker(ident,inputs=None,media='text/plain'):
    return dict(id=ident,objective='Complete substantive work',inputs=inputs or [],outputs=[output(media=media)])

def native(ident,operation,inputs=None,parameters=None):
    return dict(id=ident,objective='Execute registered operation',operation=operation,inputs=inputs or [],parameters=parameters or {})

class Tests(unittest.TestCase):
    create=fixture.Tests.create
    invoke=fixture.Tests.invoke

    def setUp(self):
        fixture.Tests.setUp(self)
        self.grant['delegation_version']=2 # preserved historical intent contract

    def source(self,text='Exact observed source',name='input.txt',media='text/plain'):
        p=self.root/name;p.parent.mkdir(parents=True,exist_ok=True);p.write_text(text)
        aid=self.rt.register(p,'Original evidence')
        return dict(artifact=aid,path='prior/'+name,purpose='Saved evidence',authority='Observed fixture; no acceptance implied',media_type=media)

    def session(self):
        aid=self.rt.task('job','owner')['latest'];s=self.factory.sessions[aid];control=Path(s['session']['control']);control.mkdir(parents=True,exist_ok=True)
        value=job_owner.Session(s['frozen'],control,job_owner.files(s['frozen'],control));self.addCleanup(value.rt.close);return value

    def run_worker(self,ident,steps,decision='delivered'):
        self.rt.tick('job');s=self.factory.sessions[self.rt.task('job',ident)['latest']];f=s['frozen'];control=Path(s['session']['control']);control.mkdir(parents=True,exist_ok=True)
        atomic(control/'launch.json',dict(credential_fingerprint=executors.fingerprint(CONFIG,f['backend'])))
        report=dict(summary='Exact synthetic result',decision=decision,instruction='',findings=[],checks={f'c{i}':dict(passed=True,evidence='Inspected exact fixture') for i in range(1,len(f['criteria'])+1)})
        script=Script(steps(f)+[('finish',dict(report_json=c.encoded(report)))])
        execute(f,control,script,lambda:(CONFIG,f['backend']))
        s['status']=dict(status='finished',exit_code=0,external_outcome='no_pending_response',pending_requests=[])
        self.rt.tick('job');return f,script

    def test_text_delivery_review_restart_and_revision_without_model_paths(self):
        ref=self.source();self.create(inputs=[ref]);self.rt.tick('job')
        tasks=[worker('write',[{'artifact':ref['artifact']}]),dict(id='review',objective='Independently inspect the evidence and report',review_of='write')]
        result,script=self.invoke([('job_delegate',dict(request_id='first',tasks=tasks))])
        self.assertEqual(result['outcome'],'yielded');self.assertEqual(len(script.requests),1)
        prompt=json.loads(script.requests[0]['contents'][0]['parts'][0]['text'])
        self.assertNotIn('sha256',c.encoded(prompt));self.assertNotIn('prior/input.txt',c.encoded(prompt));self.assertLess(len(prompt['instruction']),1500)
        wire=next(t for t in script.requests[0]['tools'][0]['functionDeclarations'] if t['name']=='job_delegate')
        self.assertIn('tasks',wire['parameters']['properties']);self.assertNotIn('batch_json',c.encoded(wire))
        f,_=self.run_worker('write',lambda f:[('file_read',dict(path=f['inputs'][0]['path'],offset=0,limit=100)),('file_write',dict(path=f['outputs'][0]['path'],text='Exact observed source: a substantive result.'))])
        self.assertEqual(f['inputs'][0]['sha256'],hashlib.sha256(b'Exact observed source').hexdigest())
        rf,_=self.run_worker('review',lambda f:[('file_read',dict(path=i['path'],offset=0,limit=1000)) for i in f['inputs']]+[('file_write',dict(path=f['outputs'][0]['path'],text='Independent inspection passed'))],decision='accept')
        self.assertEqual(len(rf['inputs']),2) # candidate and inherited source, selected by code
        candidate=self.rt.output('job','write','outputs/result.txt');old=Path(candidate['blob']).read_bytes()
        self.rt.close();self.rt=Runtime(self.root,self.factory)
        self.invoke([('job_complete',dict(deliverables_json=c.encoded({'report':candidate['id']}),summary='Reviewed result'))]);self.rt.tick('job')
        self.assertEqual(self.rt.status('job')['status'],'completed')
        job_owner.revise(self.rt,'job','revision','Change only the wording, retain original evidence.')
        self.assertEqual(job_owner.original(self.rt,'revision')['job_owner']['delegation_version'],2)
        self.assertEqual(Path(candidate['blob']).read_bytes(),old)
        self.assertEqual(self.rt.status('job')['status'],'completed')

    def test_staged_authoring_receives_registered_contract_without_host_authority(self):
        self.grant['operations']=['rhino.run_python','sketchup.run_ruby']
        self.create();self.rt.tick('job');session=self.session()
        for cap,media in [('rhino.run_python','text/x-python'),('sketchup.run_ruby','text/x-ruby')]:
            with self.subTest(cap=cap):
                task=worker('prepare',media=media);task['prepares']=cap
                batch=intents.compile(session,dict(request_id='prepare',tasks=[task]))
                spec=batch[0]
                self.assertIn('Registered authoring contract supplied by code',spec['instruction'])
                self.assertIn('expected_object_count' if cap=='rhino.run_python' else 'expected_entity_count',spec['instruction'])
                self.assertIn(cap,spec['instruction'])
                context=json.loads(spec['instruction'].split('conferred.\n',1)[1])
                shape=context['preparation_document_schema']
                self.assertFalse(shape['additionalProperties'])
                self.assertNotIn('scene_sha256',shape['properties'])
                self.assertNotIn('execution',spec)
                self.assertNotIn('work_tools',spec)
        task['prepares']='blender.run_python'
        with self.assertRaisesRegex(ValueError,'outside this job grant'):
            intents.compile(session,dict(request_id='denied',tasks=[task]))

    def test_automatic_schema_binding_real_pptx_and_revision_preserve_baseline(self):
        from orchestrator import step_runner
        from pptx import Presentation
        from orchestrator.document_authoring import PPTX
        self.grant['operations']=['pptx.create','pptx.edit'];self.grant['deliverables']={'deck':PPTX}
        self.create();self.rt.tick('job')
        content=worker('content',media='application/json')
        self.invoke([('job_delegate',dict(request_id='build',tasks=[content,native('build','pptx.create',[dict(task='content',output='result')])]))])
        self.rt.tick('job');f=self.factory.sessions[self.rt.task('job','content')['latest']]['frozen']
        self.assertEqual(f['operation_contract']['id'],'document.layout')
        page=dict(kind='page',slot=1,layout='text',content=dict(title='Evidence',body='Original observation'),notes='Synthetic')
        self.run_worker('content',lambda f:[('operation_submit',dict(request_key='page',revision=0,records=[page])),('operation_submit',dict(request_key='end',revision=1,records=[dict(kind='end',count=1)]))])
        s=self.factory.sessions[self.rt.task('job','build')['latest']];control=Path(s['session']['control']);control.mkdir(parents=True,exist_ok=True)
        detail=step_runner.execute(s['frozen'],control);self.assertEqual(detail['outcome'],'completed')
        s['status']=dict(status='finished',exit_code=0,external_outcome='no_pending_response',pending_requests=[])
        self.rt.tick('job');baseline=self.rt.output('job','build','outputs/result.pptx');raw=Path(baseline['blob']).read_bytes()
        deck=Presentation(io.BytesIO(raw));shape=next(s for s in deck.slides[0].shapes if s.has_text_frame and s.text=='Original observation')
        ref=dict(artifact=baseline['id']);change=worker('change',[ref],media='application/json')
        self.invoke([('job_delegate',dict(request_id='change',tasks=[change,native('edit','pptx.edit',[ref,dict(task='change',output='result')])]))])
        self.rt.close();self.rt=Runtime(self.root,self.factory)
        edit=dict(kind='replace_text',slide=1,shape_id=shape.shape_id,old='Original observation',new='Revised observation')
        frozen,_=self.run_worker('change',lambda f:[('operation_submit',dict(request_key='edit',revision=0,records=[dict(kind='append',field='edits',index=0,value=edit)])),('operation_submit',dict(request_key='end',revision=1,records=[dict(kind='end')]))])
        manifest=json.loads((Path(frozen['workspace'])/frozen['operation_contract']['output']).read_text())
        self.assertEqual(manifest['source_sha256'],hashlib.sha256(raw).hexdigest())
        s=self.factory.sessions[self.rt.task('job','edit')['latest']];control=Path(s['session']['control']);control.mkdir(parents=True,exist_ok=True)
        self.assertEqual(step_runner.execute(s['frozen'],control)['outcome'],'completed')
        result=Presentation(str(Path(s['frozen']['workspace'])/s['frozen']['outputs'][0]['path']))
        self.assertIn('Revised observation','\n'.join(s.text for s in result.slides[0].shapes if s.has_text_frame))
        self.assertEqual(Path(baseline['blob']).read_bytes(),raw)

    def test_duplicate_names_paths_atomic_retries_changed_bytes_and_scope(self):
        a=self.source('A','a/same.txt');b=self.source('B','b/same.txt');other=self.source('private','outside.txt')
        self.create(inputs=[a,b]);self.rt.tick('job');s=self.session()
        args=dict(request_id='once',tasks=[worker('write',[dict(artifact=a['artifact']),dict(artifact=b['artifact'])])])
        before=len(self.factory.calls);one=s.delegate(args);two=s.delegate(args)
        self.assertEqual(one,two);self.assertEqual(len(self.factory.calls),before)
        spec=self.rt.spec(self.rt.task('job','write'));self.assertEqual(len({i['path'] for i in spec['inputs']}),2)
        bad=dict(request_id='different',tasks=[worker('good'),worker('bad',[dict(artifact=other['artifact'])])])
        with self.assertRaisesRegex(ValueError,'outside'):s.delegate(bad)
        self.assertEqual([t['id'] for t in self.rt.status('job')['tasks']],['owner','write'])
        changed=copy.deepcopy(args);changed['tasks'][0]['objective']='Other work'
        with self.assertRaises(ValueError):s.delegate(changed)
        artifact=self.rt.artifact(a['artifact']);p=Path(artifact['blob']);p.chmod(0o600);p.write_text('Changed')
        with self.assertRaisesRegex(ValueError,'bytes changed'):s.delegate(dict(request_id='changed',tasks=[worker('changed',[dict(artifact=a['artifact'])])]))

    def test_registered_parameters_use_references_and_no_guessed_paths_or_hashes(self):
        from orchestrator import execution
        a=self.source('{}','data.json','application/json');self.grant['operations']=['hyperframes.render','web.sources.coverage']
        self.create(inputs=[a]);self.rt.tick('job');s=self.session();ref=dict(artifact=a['artifact'])
        args=dict(request_id='ref',tasks=[native('render','hyperframes.render',parameters=dict(project_sha256=ref,preview_sha256=ref))])
        batch=intents.compile(s,args)
        self.assertEqual(batch[0]['execution']['parameters']['project_sha256'],self.rt.artifact(a['artifact'])['sha256'])
        args['tasks']=[native('collect','web.sources.coverage',parameters=dict(catalog_path=ref,source_pack_path=ref,domains=[],model='fixture'))]
        bound=intents.compile(s,args)[0]
        self.assertEqual(bound['execution']['parameters']['catalog_path'],bound['inputs'][0]['path'])
        for cap in execution.REGISTRY:
            schema=intents.parameter_schema(cap)
            from orchestrator.parameter_contracts import check
            check(schema)
        args['tasks'][0]['parameters']['catalog_path']='delivery/invented.json'
        with self.assertRaises(ValueError):intents.compile(s,args)

    def test_new_registered_operation_uses_shared_intents_and_frozen_legacy_stays_legacy(self):
        from orchestrator import execution
        from orchestrator.operation_adapters import Adapter,object_schema
        from orchestrator.producer_binding import PORTS
        cap='fixture.copy';spec=copy.deepcopy(execution.REGISTRY['text.bundle'])
        with patch.dict(execution.REGISTRY,{cap:spec}),patch.dict(execution.ADAPTERS,{cap:Adapter(object_schema({}),handoffs=dict(version=1,default='declared-context-or-native-file',ports=[]))}),patch.dict(PORTS,{cap:[]}):
            self.grant['operations']=[cap];ref=self.source();self.create(inputs=[ref]);self.rt.tick('job')
            self.invoke([('job_delegate',dict(request_id='copy',tasks=[native('copy',cap,[dict(artifact=ref['artifact'])])]))])
            self.assertEqual(self.rt.spec(self.rt.task('job','copy'))['execution']['capability'],cap)
        self.grant['operations']=[];self.grant['delegation_version']=1;self.create('legacy')
        p=job_owner.original(self.rt,'legacy');before=c.encoded(p);job_owner.validate_plan(p);self.assertEqual(c.encoded(p),before)
        self.assertIn('batch_json',job_owner.definitions(p['job_owner'])[3]['parameters']['properties'])
        self.assertIn('tasks',job_owner.definitions(job_owner.original(self.rt,'job')['job_owner'])[3]['parameters']['properties'])

    def test_content_subjects_use_same_intent_for_different_semantic_fields(self):
        for field in ('botanical_name','part_number'):
            with self.subTest(field=field):
                # Each new job keeps its own exact grant; no shared fixture state.
                run='subjects_'+field;ref=self.source('One observed subject',field+'.txt')
                self.create(run,inputs=[ref]);self.rt.tick(run)
                producer=worker('subjects',[dict(artifact=ref['artifact'])],media='application/json')
                producer.pop('outputs')
                producer['content']=dict(family='collection.subjects',sources=dict(source_path=dict(artifact=ref['artifact'])),parameters=dict(max_records=1,fields=[field]))
                review=dict(id='review',objective='Review exact subject coverage',review_of='subjects')
                result,script=self.invoke([('job_delegate',dict(request_id='subjects',tasks=[review,producer]))],run)
                self.assertEqual(result['outcome'],'yielded');self.assertEqual(len(script.requests),1)
                self.rt.tick(run);session=self.factory.sessions[self.rt.task(run,'subjects')['latest']];frozen=session['frozen']
                control=Path(session['session']['control']);control.mkdir(parents=True,exist_ok=True)
                atomic(control/'launch.json',dict(credential_fingerprint=executors.fingerprint(CONFIG,frozen['backend'])))
                record=dict(kind='item',slot=1,label='Observed subject',attributes={field:None},quote='One observed subject',bbox=None,notes='Unknown attribute')
                checks={f'c{i}':dict(passed=True,evidence='Exact source') for i in range(1,len(frozen['criteria'])+1)}
                execute(frozen,control,Script([('operation_submit',dict(request_key='item',revision=0,records=[record])),('operation_submit',dict(request_key='end',revision=1,records=[dict(kind='end',count=1,notes='All covered')])),('finish',dict(report_json=c.encoded(dict(summary='Subject evidence',decision='delivered',instruction='',findings=[],checks=checks))))]),lambda:(CONFIG,frozen['backend']))
                raw=(Path(frozen['workspace'])/frozen['operation_contract']['output']).read_bytes()
                from orchestrator.operation_contracts import validate_document
                validate_document(raw,dict(id='collection.subjects',version=1))
                self.assertEqual(json.loads(raw)['items'][0]['attributes'],{field:None})
                self.assertEqual(len(intents.ports(self.rt.spec(self.rt.task(run,'subjects')))),2)

    def test_visual_review_binds_all_outputs_and_requires_real_inspection(self):
        from orchestrator import worker_capabilities,final_quality
        from orchestrator.document_authoring import PPTX
        self.grant['workers'].append(worker_capabilities.entry(dict(type='gemini-code',model='fixture',runtime='a'*64)))
        self.grant['operations']=['pptx.create','pptx.render'];self.grant['deliverables']={'deck':PPTX,'preview':'application/pdf'}
        self.create();self.rt.tick('job')
        previews=intents.ports(dict(outputs=intents.native_outputs('pptx.render')))
        tasks=[worker('content',media='application/json'),native('build','pptx.create',[dict(task='content',output='result')]),native('preview','pptx.render',[dict(task='build',output='result')]),dict(id='review',objective='Inspect actual final documents and every page',review_of='preview')]
        self.invoke([('job_delegate',dict(request_id='visual',tasks=tasks))])
        spec=self.rt.spec(self.rt.task('job','review'))
        self.assertTrue({'files.binary','images.view'}<=set(spec['worker']['requires']))
        self.assertEqual({i['output'] for i in spec['inputs'] if i.get('from_task')=='preview'}, {v['path'] for v in previews.values()})
        self.assertIn('final_quality',spec)
        with self.assertRaisesRegex(ValueError,'structured dimension'):
            final_quality.validate_report(dict(decision='accept',checks=[dict(criterion=int(n),passed=True,evidence='Looks fine') for n in spec['final_quality']['checks']]),spec)

    def test_compilation_receipt_and_dispatch_are_atomic_on_interruption(self):
        self.create();self.rt.tick('job');session=self.session();original=session.rt.event
        args=dict(request_id='retry',tasks=[worker('write')])
        def interrupted(run,task,attempt,kind,data):
            if kind=='owner_intents_compiled':raise RuntimeError('interrupted commit')
            return original(run,task,attempt,kind,data)
        with patch.object(session.rt,'event',side_effect=interrupted),self.assertRaisesRegex(RuntimeError,'interrupted'):
            session.delegate(args)
        self.assertEqual(len(self.rt.status('job')['tasks']),1);self.assertEqual(len(self.factory.calls),1)
        receipt=session.delegate(args)
        self.assertEqual(len(self.factory.calls),1)
        self.rt.close();self.rt=Runtime(self.root,self.factory)
        self.rt.tick('job');self.assertEqual(self.rt.task('job','write')['attempts'],1)
        events=[json.loads(r[0]) for r in self.rt.db.execute("SELECT data FROM production_events WHERE run='job' AND kind='owner_intents_compiled'")]
        self.assertEqual(len(events),1);self.assertEqual(events[0]['intent'],args)

    def test_compacted_owner_delegates_without_reconstructing_worker_schema(self):
        from orchestrator.working_context import MAX_TEXT_BYTES,size
        ref=self.source('Synthetic evidence interval. '*9000)
        self.grant['operations']=['pptx.create'];self.create(inputs=[ref]);self.rt.tick('job')
        steps=[('job_catalog',dict(operations=['pptx.create']))]
        steps += [('job_read',dict(artifact=ref['artifact'],offset=i*24000,limit=24000)) for i in range(7)]
        steps += [('job_delegate',dict(request_id='after-reads',tasks=[worker('write',[dict(artifact=ref['artifact'])])]))]
        result,script=self.invoke(steps)
        self.assertEqual(result['outcome'],'yielded');self.assertEqual(len(script.requests),9)
        self.assertIn('Historical exchanges are preserved separately',c.encoded(script.requests[-1]))
        self.assertLessEqual(size(script.requests[-1]['contents']),MAX_TEXT_BYTES)
        self.assertIn('pptx.create',c.encoded(script.requests[-1]))
        self.rt.tick('job');self.assertEqual(self.rt.task('job','write')['attempts'],1)

    def test_saved_layout_reuses_exact_authorized_asset_aliases_across_versions(self):
        from tests.test_browser_screenshots import png
        from orchestrator import step_runner
        from orchestrator.document_authoring import PPTX
        from pptx import Presentation
        p=self.root/'photo.png';p.write_bytes(png());image=self.rt.register(p,'Exact synthetic pixels')
        source=dict(artifact=image,path='reference/photo.png',purpose='Original pixels',authority='Fixture pixels',media_type='image/png',visual_reference=True)
        self.grant['operations']=['pptx.create'];self.grant['deliverables']={'deck':PPTX}
        self.create(inputs=[source]);self.rt.tick('job')
        self.invoke([('job_delegate',dict(request_id='layout',tasks=[worker('content',[dict(artifact=image)],media='application/json'),native('build','pptx.create',[dict(task='content',output='result')])]))])
        def fragments(f):
            path=next(i['path'] for i in f['inputs'] if i['artifact']==image)
            page=dict(kind='page',slot=1,layout='image_grid',content=dict(title='Exact asset',items=[dict(image=path,label='Fixture',caption='Saved pixels',credit='Synthetic')],per_page=4,fit='contain'),notes='')
            return [('operation_submit',dict(request_key='page',revision=0,records=[page])),('operation_submit',dict(request_key='end',revision=1,records=[dict(kind='end',count=1)]))]
        self.run_worker('content',fragments)
        layout=self.rt.output('job','content','outputs/result.json');saved=Path(layout['blob']).read_bytes()
        self.create('next',inputs=[{**source,'path':'renamed/input.png'},dict(artifact=layout['id'],path='delivery/layout.json',purpose='Exact saved layout',authority='Unaccepted prior candidate',media_type='application/json')]);self.rt.tick('next')
        self.invoke([('job_delegate',dict(request_id='reuse',tasks=[native('rebuild','pptx.create',[dict(artifact=layout['id'])])]))],'next')
        self.rt.tick('next');s=self.factory.sessions[self.rt.task('next','rebuild')['latest']]
        self.assertIn('reference/photo.png',[i['path'] for i in s['frozen']['inputs']])
        control=Path(s['session']['control']);control.mkdir(parents=True,exist_ok=True)
        self.assertEqual(step_runner.execute(s['frozen'],control)['outcome'],'completed')
        deck=Presentation(str(Path(s['frozen']['workspace'])/s['frozen']['outputs'][0]['path']))
        self.assertEqual([s.image.blob for slide in deck.slides for s in slide.shapes if s.shape_type==13],[p.read_bytes()])
        self.assertEqual(Path(layout['blob']).read_bytes(),saved)

    def test_provider_variant_staged_authoring_supplies_executable_shared_schema(self):
        from orchestrator import parameter_contracts, web_sources
        for capability in ('web.sources.prepared','web.sources.prepared.codex'):
            context=json.loads(intents.authoring_instruction(capability).split('conferred.\n',1)[1])
            schema=context['preparation_document_schema']
            prepared=dict(version=1,targets=[dict(id='subject-1',queries=['Synthetic chair dimensions'],required_options=1)])
            parameter_contracts.validate(prepared,schema,'prepared')
            web_sources.validate_prepared(prepared)
            with self.assertRaises(ValueError):
                parameter_contracts.validate(dict(prepared,execution_ready=False),schema,'prepared')
            self.assertEqual(context['preparation_inputs'][0]['target'],'web.sources.prepared')

    def test_existing_task_names_reject_before_compilation_without_mutation(self):
        self.create();self.rt.tick('job');session=self.session()
        session.delegate(dict(request_id='first',tasks=[worker('saved')]))
        before=copy.deepcopy(self.rt.status('job'));calls=list(self.factory.calls)
        # This also contains an invalid operation. Identity conflicts must be
        # disclosed together before semantic compilation consumes corrections.
        with self.assertRaisesRegex(ValueError,'Task IDs already exist: owner, saved'):
            session.delegate(dict(request_id='conflict',tasks=[native('saved','not.registered'),worker('owner')]))
        self.assertEqual(self.rt.status('job'),before);self.assertEqual(self.factory.calls,calls)
        retry=session.delegate(dict(request_id='first',tasks=[worker('saved')]))
        self.assertIn('saved',retry['queued'])

    def test_native_preparation_aliases_bind_one_exact_named_port_across_hosts(self):
        from orchestrator import native_apps, producer_binding, execution
        self.grant.update(delegation_version=4,operations=list(native_apps.SCRIPT_OPERATIONS))
        checks=self.source('{"fixture":"checks"}','checks.json','application/json')
        alias=self.source('{"fixture":"checks"}','input.json','application/json')
        conflict=self.source('{"fixture":"different"}','other.json','application/json')
        script=self.source('pass\n','script.txt','text/plain')
        ruby=self.source('# fixture\n','ruby.txt','text/plain')
        scene=self.source('synthetic native bytes','scene.blend','application/x-blender')
        self.create(inputs=[checks,alias,conflict,script,ruby,scene]);self.rt.tick('job');session=self.session()
        def ref(x):return dict(artifact=x['artifact'])
        before=copy.deepcopy(self.rt.status('job'));calls=list(self.factory.calls)
        for cap in native_apps.SCRIPT_OPERATIONS:
            with self.subTest(cap=cap):
                item=native('build',cap,[ref(script if cap!='sketchup.run_ruby' else ruby),ref(checks)],
                    dict(scene_sha256=ref(scene) if cap=='blender.run_python' else None,
                         script_sha256=ref(ruby if cap=='sketchup.run_ruby' else script),checks_sha256=ref(checks),permissions='unrestricted_host'))
                item['data']=ref(alias);args=dict(request_id='native',tasks=[item]);original=copy.deepcopy(args)
                batch=intents.compile(session,args)
                producer_binding.compile(dict(backend=fixture.BACKEND,tasks=batch))
                batch[0].update(role='procedure',criteria=execution.REGISTRY[cap]['criteria'])
                c.assignment(batch[0])
                selected=[i for i in batch[0]['inputs'] if i['media_type']=='application/json']
                self.assertEqual([i['artifact'] for i in selected],[checks['artifact']])
                self.assertEqual(batch[0]['execution']['parameters']['checks_sha256'],self.rt.artifact(checks['artifact'])['sha256'])
                self.assertEqual(args,original)
                selected_script=next(i for i in batch[0]['inputs'] if i['artifact']==(ruby if cap=='sketchup.run_ruby' else script)['artifact'])
                self.assertEqual(selected_script['media_type'],native_apps.profile(cap).script_media)
                wrong=copy.deepcopy(args);wrong['tasks'][0]['parameters']['script_sha256']=ref(checks)
                with self.assertRaisesRegex(ValueError,'script_sha256 requires'):intents.compile(session,wrong)
                self.assertEqual(execution.adapter(cap).handoffs['ports'][0]['reference_parameter'],'checks_sha256')
                bad=copy.deepcopy(args);bad['tasks'][0]['data']=ref(conflict)
                with self.assertRaisesRegex(ValueError,'conflicts with named preparation'):intents.compile(session,bad)
        self.assertEqual(self.rt.status('job'),before);self.assertEqual(self.factory.calls,calls)

    def test_completed_named_hash_ports_bind_saved_outputs_after_restart_without_dispatch(self):
        from orchestrator import producer_binding,execution
        self.grant.update(delegation_version=4,operations=['rhino.run_python'])
        self.create();self.rt.tick('job')
        prepare=dict(id='prepare',objective='Save exact script and checks',outputs=[output('script','text/plain'),output('checks','application/json'),output('input','application/json')])
        refs={key:dict(task='prepare',output=name) for key,name in [('script_sha256','script'),('checks_sha256','checks')]}
        consumer=native('native','rhino.run_python',[dict(task='prepare',output='input')],dict(**refs,scene_sha256=None,permissions='unrestricted_host'))
        with self.assertRaisesRegex(ValueError,'existing artifact'):
            intents.compile(self.session(),dict(request_id='early',tasks=[prepare,consumer]))
        self.invoke([('job_delegate',dict(request_id='prepare',tasks=[prepare]))]);self.rt.tick('job')
        aid=self.rt.task('job','prepare')['latest'];self.factory.finish(aid)
        ws=self.factory.sessions[aid]['workspace']
        (ws/'outputs/script.txt').write_text('pass\n')
        for name in ('checks.json','input.json'):(ws/'outputs'/name).write_text('{}\n')
        self.rt.tick('job')
        from orchestrator.owner_compilation import label
        review=self.rt.task('job',label('prepare','review'))
        if review and review['latest']:
            self.factory.finish(review['latest'],decision='accept');self.rt.tick('job')
        self.assertEqual(self.rt.task('job','prepare')['status'],'completed')
        original=self.rt.spec(self.rt.task('job','prepare'));calls=list(self.factory.calls)
        self.rt.close();self.rt=Runtime(self.root,self.factory)
        batch=intents.compile(self.session(),dict(request_id='execute',tasks=[consumer]))
        producer_binding.compile(dict(backend=fixture.BACKEND,tasks=batch))
        batch[0].update(role='procedure',criteria=execution.REGISTRY['rhino.run_python']['criteria']);c.assignment(batch[0])
        self.assertTrue(all('artifact' in i and 'from_task' not in i for i in batch[0]['inputs']))
        for key,name in [('script_sha256','script.txt'),('checks_sha256','checks.json')]:
            artifact=self.rt.output('job','prepare','outputs/'+name)
            self.assertEqual(batch[0]['execution']['parameters'][key],artifact['sha256'])
        self.assertEqual(len(batch[0]['inputs']),2)
        self.assertEqual(next(i for i in batch[0]['inputs'] if i['path'].endswith('script.txt'))['media_type'],'text/x-python')
        self.assertEqual(next(o for o in original['outputs'] if o['path'].endswith('script.txt'))['media_type'],'text/plain')
        self.assertEqual(self.rt.spec(self.rt.task('job','prepare')),original)
        self.assertEqual(self.factory.calls,calls)
