"""Exact asset closure, real PPTX composition and generated evidence review admission."""
import copy,hashlib,io,json,unittest
from unittest.mock import patch
from orchestrator import handoff_closure as h,execution,document_compose,pptx_document,artifact_roles,worker_capabilities,operation_adapters
from task_relay import record_planning
from tests import test_producer_binding as fixtures
from tests.test_automatic_composition import pixel

class Tests(unittest.TestCase):
    def test_same_named_section_assets_survive_composition_without_historical_images(self):
        def inp(a,p,m):return dict(artifact=a,path=p,media_type=m,purpose='Exact fixture',authority='Fixture')
        old=inp('old','preview.png','image/png');left=inp('left','preview.png','image/png');right=inp('right','preview.png','image/png')
        first=inp('first','first.json','application/json');second=inp('second','second.json','application/json')
        producers={'first':dict(id='first',inputs=[left]),'second':dict(id='second',inputs=[right]),
            'right':dict(id='corrected',inputs=[old],companions=[inp('native','model.dat','application/octet-stream')])}
        build=dict(id='compose',execution=dict(capability='document.compose'),inputs=[first,second],outputs=[])
        render=dict(id='render',execution=dict(capability='pptx.create'),inputs=[dict(from_task='compose',output='result.json',path='result.json',media_type='application/json',purpose='Composed data',authority='Fixture')],outputs=[])
        review=dict(id='review',inputs=[dict(from_task='render',output='deck.pptx',path='deck.pptx',purpose='Rendered deck',authority='Fixture')],outputs=[])
        grant=h.capture(['document.compose','pptx.create']);tasks=[build,render,review]
        h.bind(tasks,lambda _:None,producers.get,
            preserve_paths=lambda t:bool(t.get('execution')) and execution.adapter(t['execution']['capability']).asset_binding=='path',
            inherit=lambda t,i:not t.get('execution') or i['media_type'] in grant.get(t['execution']['capability'],[]),
            inherit_ancestry=lambda t,v,d:h.inherits_assets(t,v,d,grant))
        self.assertEqual({i.get('artifact') for i in build['inputs']},{'first','second','left','right'})
        self.assertEqual({i.get('artifact') for i in render['inputs']},{None,'left','right'})
        self.assertTrue({'old','native','left','right'}<={i.get('artifact') for i in review['inputs']})
        self.assertEqual(len({i['path'] for i in build['inputs']}),4)
        images={'left':pixel('blue'),'right':pixel('red')};blobs={}
        for name,asset in [('first','left'),('second','right')]:
            doc=dict(version=1,title=name,slides=[dict(notes='Exact '+asset,elements=[dict(type='image',path='preview.png',x=1,y=1,w=3,h=3)])],asset_sha256={'preview.png':hashlib.sha256(images[asset]).hexdigest()})
            blobs[name+'.json']=json.dumps(doc).encode()
        blobs.update({i['path']:images[i['artifact']] for i in build['inputs'] if i['artifact'] in images})
        frozen=dict(inputs=[dict(i,sha256=hashlib.sha256(blobs[i['path']]).hexdigest()) for i in build['inputs']],execution=dict(parameters=dict(document_paths=['first.json','second.json'],image_paths=[],image_titles=[],title='Composed')))
        raw,_=document_compose.build(frozen,blobs.__getitem__)
        deck,_=pptx_document.create(json.loads(raw),{i['path']:images[i['artifact']] for i in render['inputs'] if i.get('artifact') in images})
        from pptx import Presentation
        reopened=Presentation(io.BytesIO(deck))
        self.assertEqual([hashlib.sha256(s.shapes[0].image.blob).hexdigest() for s in reopened.slides],[hashlib.sha256(images[x]).hexdigest() for x in ('left','right')])
        # Path-bound adapters still reject genuine direct dependency ambiguity.
        direct=dict(id='direct',execution=dict(capability='pptx.create'),inputs=[first,second],outputs=[])
        with self.assertRaisesRegex(ValueError,'path conflicts'):
            h.bind([direct],lambda _:None,producers.get,preserve_paths=lambda _:True,
                inherit=lambda t,i:i['media_type']=='image/png',inherit_ancestry=lambda t,v,d:h.inherits_assets(t,v,d,grant))

    def test_generated_record_review_binds_readers_before_worker_resolution(self):
        inputs=[dict(artifact=a,path=p,media_type=m,purpose='Saved evidence',authority='Synthetic') for a,p,m in
            [('inventory','inventory.json','application/json'),('sources','sources.json','application/json'),('photos','photos.zip','application/zip'),('room','room.png','image/png')]]
        producer=fixtures.writer('producer',inputs=copy.deepcopy(inputs),records=dict(family=dict(id='collection.options',version=1),
            inputs=dict(inventory_path=dict(source='inventory'),source_pack_path=dict(source='sources')),
            parameters=dict(field_definitions={'name':'text'},identity_field='name',options_per_item=1,count_policy='target',max_records=2)))
        producer['outputs'].append(dict(path='data.md',purpose='Summary',media_type='text/markdown'))
        reviewer=fixtures.writer('review',output='review.json',inputs=copy.deepcopy(inputs)+[fixtures.ref('producer','data.json')],review_of='producer')
        reviewer['outputs'].append(dict(path='review.md',purpose='Review',media_type='text/markdown'))
        backend=dict(type='codex-agent',model='fixture',reasoning='low',budget_semantics='cli-exchanges-v1')
        plan=dict(tasks=[producer,reviewer],backend=backend)
        # Hooks inspect only native representations; this fixture has ordinary image/ZIP inputs.
        with patch.object(operation_adapters,'resolve',return_value=lambda *args:{}):
            record_planning.bind(plan,workers=[worker_capabilities.entry(backend)],evidence_binder=lambda t:artifact_roles.bind(t,None,[]))
        for task in plan['tasks']:
            self.assertEqual(task['worker']['requires'],['files.text','images.view'])
            self.assertEqual(task['worker']['backend']['type'],'codex-agent')
            self.assertEqual(task['artifact_roles']['inputs']['photos.zip']['inspection']['tool'],'image_view')
            self.assertNotIn('python',task['tools'])
            self.assertTrue({'inventory','sources','photos','room'}<={i.get('artifact') for i in task['inputs']})
        denied=copy.deepcopy(reviewer);denied.pop('tools');denied['worker']=dict(requires=['files.text','images.view'])
        with self.assertRaisesRegex(ValueError,'No eligible'):
            worker_capabilities.resolve(denied,[],backend)
