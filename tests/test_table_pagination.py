import io,unittest
from orchestrator import slide_templates,pptx_document,table_layout

class Tests(unittest.TestCase):
    def document(self,rows):
        return dict(version=2,title='Synthetic procurement matrix',style='clean_minimal_v1',slides=[dict(layout='table',content=dict(title='Candidate comparison',rows=rows,caption='Evidence and current availability remain subject to review'))])

    def test_long_cells_paginate_without_clipping_or_losing_data(self):
        from pptx import Presentation
        header=['Option','Product and material','Price','Dimensions','Availability and delivery','Comparison']
        rows=[header]+[[f'Option {n}',f'Product {n}: '+('Textured warm upholstery with supported details. '*2),'USD 1200','62 x 36 x 33 inches','Availability unverified; delivery must be checked with the retailer.','Warm terracotta woven texture with a low profile. '*2] for n in range(1,8)]
        doc=self.document(rows);expanded=slide_templates.expand(doc)
        self.assertGreater(len(expanded['slides']),1)
        tables=[next(e for e in s['elements'] if e['type']=='table') for s in expanded['slides']]
        self.assertEqual([r for t in tables for r in t['rows'][1:]],rows[1:])
        for t in tables:
            self.assertEqual(t['rows'][0],header);self.assertGreaterEqual(t['font_size'],14)
            self.assertLessEqual(t['y']+t['h'],6.31)
            self.assertAlmostEqual(sum(t['row_heights']),t['h'])
        namespace={'__name__':'frozen_validator'}
        exec(pptx_document.validator_source(),namespace)
        self.assertEqual(namespace['validate'](doc),pptx_document.validate(doc))
        raw=pptx_document.create(doc,{})
        if isinstance(raw,tuple):raw=raw[0]
        deck=Presentation(io.BytesIO(raw))
        for s,t in zip(deck.slides,tables):
            shape=next(x for x in s.shapes if x.has_table)
            self.assertLessEqual(shape.top+shape.height,int(6.31*914400))
            self.assertEqual([[c.text for c in r.cells] for r in shape.table.rows],t['rows'])
            self.assertTrue(all(p.font.size.pt==t['font_size'] for r in shape.table.rows for c in r.cells for p in c.text_frame.paragraphs))

    def test_dense_columns_become_readable_details_with_every_value(self):
        rows=[[f'Field {n}' for n in range(10)],[('WIDE 詳細 '*25)+str(n) for n in range(10)]]
        tables=list(table_layout.pages(rows));self.assertGreater(len(tables),1)
        cells=[r for t in tables for r in t['rows'][1:]]
        self.assertEqual(cells,[[label,value] for label,value in zip(*rows)])
        self.assertTrue(all(sum(t['row_heights'])<=4.8 for t in tables))

    def test_rectangular_bounded_data_and_oversized_single_cell_fail_explicitly(self):
        for rows in [[['A','B'],['C']],[['A','B'],['X'*4000,'Y'*4000]]]:
            with self.assertRaises(ValueError):list(table_layout.pages(rows))
        bad=slide_templates.expand(self.document([['A'],['B']]))
        table=next(x for x in bad['slides'][0]['elements'] if x['type']=='table');table['row_heights'][0]+=10
        with self.assertRaisesRegex(ValueError,'row heights'):pptx_document.validate(bad)
