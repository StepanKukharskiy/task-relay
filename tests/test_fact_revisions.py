"""Controlled O14 proof: exact source facts, conservative impact, native revision."""

from contextlib import closing
import json
import hashlib
from pathlib import Path
import sqlite3
import tempfile
import unittest
from types import SimpleNamespace
import uuid

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import PatternFill

from orchestrator.runtime import Runtime
from task_relay.bridge import State
from task_relay import (fact_revisions as facts, impact_handoff, production_control as pc,
                        reviewed_links, revision_review, workflow_files)


CASE = json.loads((Path(__file__).parent / 'fixtures/o14_acceptance.json').read_text())


def workbook(path, sheet, headers, rows):
    book = Workbook()
    ws = book.active
    ws.title = sheet
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws['B2'].fill = PatternFill('solid', fgColor='FFF2CC')
    book.save(path)
    book.close()


class Tests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name).resolve()
        self.state = State(self.root / 'data/state.sqlite')
        self.rt = Runtime(pc.root(self.state), connection=self.state.db)
        self.job = 'job-' + 'a' * 24
        self.state.db.execute('''INSERT INTO relay_pipelines
            (id,request_id,request,title,spec,channel,provider,model,status,created)
            VALUES (?,?,?,?,?,?,?,?,?,?)''',
            (self.job, 1, 'Revise only the affected catalog material.', 'O14 fixture',
             '{}', 'telegram', 'fixture', 'fixture', 'active', 1))
        self.state.db.commit()
        spec = CASE['supplier']
        workbook(self.root / 'supplier.xlsx', spec['sheet'], spec['columns'], spec['before'])
        workbook(self.root / 'replacement.xlsx', spec['sheet'], spec['columns'], spec['replacement'])
        spec = CASE['catalog']
        workbook(self.root / 'catalog.xlsx', spec['sheet'], spec['columns'], spec['rows'])
        self.old = self.rt.register(self.root / 'supplier.xlsx', 'Exact supplier version')
        self.new = self.rt.register(self.root / 'replacement.xlsx', 'Explicit replacement version')
        self.catalog = self.rt.register(self.root / 'catalog.xlsx', 'Reviewed catalog baseline')
        self.rt.close()

    def tearDown(self):
        self.state.db.close()
        self.temp.cleanup()

    def bind(self, key, source_cell, output_cell, value, review_state='reviewed'):
        return facts.record_binding(self.state.db, job=self.job, entity_key=key,
            predicate='material', value=value, source_artifact=self.old,
            source_sheet='Supplier', source_key_column='A', source_cell=source_cell,
            evidence='Reviewed supplier row for exact product key.',
            output_artifact=self.catalog, output_sheet='Catalog', output_cell=output_cell,
            reviewer='fixture-reviewer', review_state=review_state)

    def freeze_coverage(self):
        for row in CASE['reviewed_bindings']:
            self.bind(row['key'], row['source_cell'], row['output_cell'], row['value'])
        for row in CASE['coverage']:
            facts.record_coverage(self.state.db, job=self.job, output_artifact=self.catalog,
                output_sheet='Catalog', output_cell=row['cell'], state=row['state'],
                note='Exact reviewed supplier fact.' if row['state'] == 'complete'
                     else 'No source evidence for this catalog row.')

    def plan(self):
        return facts.plan_impact(self.state.db, job=self.job, old_source=self.old,
                                 replacement_source=self.new, baseline_artifact=self.catalog)

    def test_frozen_external_handoff_admits_exact_agent_candidate_once(self):
        self.freeze_coverage()
        handoff = impact_handoff.record_xlsx_change(self.state.db, job=self.job,
            request_key='material-change-1', exact_request='Revise only the affected material.',
            actor='user', old_source=self.old, replacement_source=self.new,
            baseline_artifact=self.catalog)
        frozen = impact_handoff.inspect(self.state.db, handoff)
        self.assertEqual(frozen['status'], 'current')
        self.assertEqual(impact_handoff.record_xlsx_change(self.state.db, job=self.job,
            request_key='material-change-1', exact_request='Revise only the affected material.',
            actor='user', old_source=self.old, replacement_source=self.new,
            baseline_artifact=self.catalog), handoff)
        with self.assertRaisesRegex(ValueError, 'different frozen plan'):
            impact_handoff.record_xlsx_change(self.state.db, job=self.job,
                request_key='material-change-1', exact_request='Changed request',
                actor='user', old_source=self.old, replacement_source=self.new,
                baseline_artifact=self.catalog)
        candidate_path = self.root / 'external.xlsx'
        book = load_workbook(self.root / 'catalog.xlsx')
        book['Catalog']['B2'] = 'stainless steel'
        book.save(candidate_path)
        book.close()
        aid = facts.submit_external_xlsx(self.rt, handoff_id=handoff,
            submission_key='external-1', candidate_file=candidate_path,
            submitted_by='independent-producer')
        self.assertEqual(facts.submit_external_xlsx(self.rt, handoff_id=handoff,
            submission_key='external-1', candidate_file=candidate_path,
            submitted_by='independent-producer'), aid)
        self.assertEqual(self.state.db.execute(
            'SELECT COUNT(*) FROM relay_fact_external_submissions').fetchone()[0], 1)
        self.assertEqual(facts._verified_candidate(self.state.db, aid)[2]['changed_cells'],
                         ['Catalog!B2'])
        view = workflow_files.sync(self.state, self.job)
        with sqlite3.connect(view / '.relay/job.sqlite') as db:
            self.assertEqual(db.execute('SELECT COUNT(*) FROM fact_external_submissions').fetchone()[0], 1)
        book = load_workbook(candidate_path)
        book['Catalog']['B4'] = 'bronze'
        book.save(candidate_path)
        book.close()
        with self.assertRaisesRegex(ValueError, 'unsupported value'):
            facts.submit_external_xlsx(self.rt, handoff_id=handoff,
                submission_key='external-2', candidate_file=candidate_path,
                submitted_by='independent-producer')
        self.assertEqual(self.state.db.execute(
            'SELECT COUNT(*) FROM relay_fact_external_submissions').fetchone()[0], 1)

    def test_change_preserves_independent_cell_formula_style_and_versions(self):
        self.freeze_coverage()
        before = self.state.db.total_changes
        plan = self.plan()
        self.assertEqual(self.state.db.total_changes, before)
        for status in ('affected', 'unaffected', 'unknown'):
            self.assertEqual([x['location'] for x in plan[status]],
                             CASE['expected_impact'][status])
        self.assertIn('Planning only', plan['authorization'])
        candidate = facts.revise_xlsx(self.rt, plan=plan,
                                      patches={'Catalog!B2': 'stainless steel'},
                                      reviewer='fixture-reviewer')
        self.assertNotEqual(candidate, self.catalog)
        self.assertEqual(self.state.db.execute('SELECT count(*) FROM relay_fact_revisions').fetchone()[0], 1)
        with Path(self.rt.artifact(self.catalog)['blob']).open('rb') as stream:
            original = load_workbook(stream)
        with Path(self.rt.artifact(candidate)['blob']).open('rb') as stream:
            revised = load_workbook(stream)
        for location, value in CASE['expected_revision'].items():
            sheet, cell = location.split('!')
            self.assertEqual(revised[sheet][cell].value, value)
        self.assertEqual(revised['Catalog']['B2']._style, original['Catalog']['B2']._style)
        self.assertEqual(revised['Catalog']['C2'].value, original['Catalog']['C2'].value)
        self.assertEqual(original['Catalog']['B2'].value, 'steel')
        original.close();revised.close()
        root = workflow_files.sync(self.state, self.job)
        marker = json.loads((root / '.relay-workflow.json').read_text())
        manifest = json.loads((root / 'manifest.json').read_text())
        candidate_copy = next(a['copy_path'] for a in manifest['artifacts'] if a['id'] == candidate)
        self.assertIn('/drafts/', candidate_copy)
        self.assertFalse(next(a['selected'] for a in manifest['artifacts'] if a['id'] == candidate))
        with closing(sqlite3.connect(root / '.relay/job.sqlite')) as view:
            self.assertEqual(view.execute("SELECT value FROM meta WHERE key='snapshot'").fetchone()[0], marker['snapshot'])
            self.assertEqual(view.execute('SELECT count(*) FROM fact_bindings').fetchone()[0], 2)
            self.assertEqual(view.execute('SELECT state FROM coverage WHERE output_cell=?', ('B4',)).fetchone()[0], 'incomplete')
            self.assertEqual(view.execute('SELECT candidate_artifact FROM revisions').fetchone()[0], candidate)
        snapshot = root / '.relay/snapshots' / marker['snapshot'] / 'job.sqlite'
        self.assertEqual(snapshot.read_bytes(), (root / '.relay/job.sqlite').read_bytes())
        exported = json.loads((root / '.relay/snapshots' / marker['snapshot'] / 'job-state.json').read_text())
        self.assertEqual(len(exported['facts']['bindings']), 2)
        self.assertEqual(len(exported['facts']['revisions']), 1)
        facts.review_candidate(self.state.db, candidate_artifact=candidate, decision='accept',
                               reviewer='independent-reviewer', note='Verified the candidate against the frozen case.')
        facts.select_candidate(self.state.db, candidate_artifact=candidate,
                               selected_by='fixture-user', receipt='Synthetic exact candidate selection.')
        selected_root = workflow_files.sync(self.state, self.job)
        selected_manifest = json.loads((selected_root / 'manifest.json').read_text())
        selected_copy = next(a for a in selected_manifest['artifacts'] if a['id'] == candidate)
        self.assertTrue(selected_copy['selected'])
        self.assertIn('/selected/', selected_copy['copy_path'])
        with closing(sqlite3.connect(selected_root / '.relay/job.sqlite')) as view:
            self.assertEqual(view.execute('SELECT decision FROM candidate_reviews').fetchone()[0], 'accept')
            self.assertEqual(view.execute('SELECT candidate_artifact FROM selections').fetchone()[0], candidate)
        self.assertEqual(len(json.loads((selected_root / '.relay/snapshots' /
            json.loads((selected_root / '.relay-workflow.json').read_text())['snapshot'] /
            'job-state.json').read_text())['facts']['selections']), 1)
        newer = facts.revise_xlsx(self.rt, plan=self.plan(),
                                  patches={'Catalog!B2': 'stainless steel'},
                                  reviewer='fixture-reviewer')
        facts.review_candidate(self.state.db, candidate_artifact=newer, decision='accept',
                               reviewer='independent-reviewer', note='Verified the new exact version.')
        facts.select_candidate(self.state.db, candidate_artifact=newer,
                               selected_by='fixture-user', receipt='Synthetic later candidate selection.')
        latest = json.loads((workflow_files.sync(self.state, self.job) / 'manifest.json').read_text())
        by_id = {a['id']: a for a in latest['artifacts']}
        self.assertFalse(by_id[candidate]['selected'])
        self.assertTrue(by_id[newer]['selected'])

    def test_shared_impact_is_saved_and_cannot_claim_incomplete_cell_unaffected(self):
        self.freeze_coverage()
        plan = self.plan()
        original = {key: value for key, value in plan.items()
                    if key not in ('digest', 'reviewed_impact')}
        self.assertEqual(plan['digest'], hashlib.sha256(json.dumps(
            original, sort_keys=True).encode()).hexdigest())
        shared = plan['reviewed_impact']
        self.assertEqual(shared['schema'], reviewed_links.IMPACT_SCHEMA)
        self.assertEqual([(item['output_location'], item['status'])
                          for item in shared['outcomes']],
                         [('Catalog!B2', 'affected'), ('Catalog!B3', 'unaffected'),
                          ('Catalog!B4', 'unknown')])
        self.assertEqual(shared['baseline_sha256'], self.rt.artifact(self.catalog)['sha256'])
        forged = {**plan, 'unaffected': [*plan['unaffected'],
                  {'location': 'Catalog!B4', 'reason': 'Forged reuse.'}], 'unknown': []}
        with self.assertRaisesRegex(ValueError, 'complete reviewed'):
            reviewed_links.impact_record(reviewed_links.records(self.state.db, self.job),
                kind='material_cell', plan=forged, source_sha256=shared['old_source_sha256'],
                replacement_sha256=shared['replacement_sha256'])
        candidate = facts.revise_xlsx(self.rt, plan=plan,
            patches={'Catalog!B2': 'stainless steel'}, reviewer='fixture-reviewer')
        root = workflow_files.sync(self.state, self.job)
        with closing(sqlite3.connect(root / '.relay/job.sqlite')) as view:
            self.assertEqual(view.execute('SELECT count(*) FROM reviewed_links').fetchone()[0], 2)
            self.assertEqual(view.execute('SELECT state FROM reviewed_coverage WHERE output_location=?',
                                          ('Catalog!B4',)).fetchone()[0], 'incomplete')
            saved = view.execute('SELECT record FROM reviewed_impacts WHERE candidate_artifact=?',
                                 (candidate,)).fetchone()[0]
            self.assertEqual(json.loads(saved), shared)
        exported = json.loads((root / '.relay/snapshots' /
            json.loads((root / '.relay-workflow.json').read_text())['snapshot'] /
            'job-state.json').read_text())
        self.assertEqual(exported['reviewed_impacts'][0]['record'], shared)

    def test_selection_requires_independent_acceptance_and_exact_receipt(self):
        self.freeze_coverage()
        candidate = facts.revise_xlsx(self.rt, plan=self.plan(),
                                      patches={'Catalog!B2': 'stainless steel'},
                                      reviewer='fixture-reviewer')
        with self.assertRaisesRegex(ValueError, 'independently accepted'):
            facts.select_candidate(self.state.db, candidate_artifact=candidate,
                                   selected_by='fixture-user', receipt='Synthetic exact choice.')
        with self.assertRaisesRegex(ValueError, 'independent'):
            facts.review_candidate(self.state.db, candidate_artifact=candidate,
                                   decision='accept', reviewer='fixture-reviewer', note='Self review.')
        self.assertEqual(self.state.db.execute('SELECT count(*) FROM relay_fact_candidate_reviews').fetchone()[0], 0)
        facts.review_candidate(self.state.db, candidate_artifact=candidate,
                               decision='revise', reviewer='independent-reviewer', note='Request another version.')
        with self.assertRaisesRegex(ValueError, 'independently accepted'):
            facts.select_candidate(self.state.db, candidate_artifact=candidate,
                                   selected_by='fixture-user', receipt='Synthetic exact choice.')
        self.assertEqual(self.state.db.execute('SELECT count(*) FROM relay_fact_selections').fetchone()[0], 0)

    def test_generic_review_surface_also_handles_material_cell_candidate(self):
        self.freeze_coverage()
        candidate = facts.revise_xlsx(self.rt, plan=self.plan(),
                                      patches={'Catalog!B2': 'stainless steel'},
                                      reviewer='fixture-reviewer')
        paths = SimpleNamespace(state=self.root/'data/state.sqlite',
                                generated=self.root/'data/generated')
        view = revision_review.detail(candidate, paths)
        self.assertEqual(view['kind'], 'material_cell')
        self.assertTrue(view['verified'])
        self.assertEqual([item['location'] for item in view['plan']['affected']],
                         ['Catalog!B2'])
        self.assertEqual(view['evidence_links'][0]['source_location'], 'Supplier!B2')
        self.assertEqual(view['evidence_links'][0]['output_location'], 'Catalog!B2')
        with self.assertRaisesRegex(ValueError, 'independent'):
            revision_review.decide(candidate_artifact=candidate, verb='accept',
                actor='fixture-reviewer', note='Self review.',
                expected_sha256=view['candidate_sha256'],
                expected_plan_digest=view['plan']['digest'],
                request_id=str(uuid.uuid4()), paths=paths)
        self.assertEqual(self.state.db.execute('SELECT count(*) FROM relay_revision_action_receipts').fetchone()[0], 0)
        accepted = revision_review.decide(candidate_artifact=candidate, verb='accept',
            actor='independent-reviewer', note='Checked the exact material change.',
            expected_sha256=view['candidate_sha256'],
            expected_plan_digest=view['plan']['digest'],
            request_id=str(uuid.uuid4()), paths=paths)
        self.assertTrue(accepted['detail']['can_select'])

    def test_missing_or_duplicate_replacement_key_is_unknown(self):
        self.freeze_coverage()
        for rows in ([['XYZ-2040', 'aluminium']],
                     [['ABC-9182', 'steel'], ['ABC-9182', 'stainless steel'], ['XYZ-2040', 'aluminium']]):
            workbook(self.root / 'ambiguous.xlsx', 'Supplier', ['product_key', 'material'], rows)
            candidate = self.rt.register(self.root / 'ambiguous.xlsx', 'Ambiguous replacement')
            plan = facts.plan_impact(self.state.db, job=self.job, old_source=self.old,
                replacement_source=candidate, baseline_artifact=self.catalog)
            self.assertIn('Catalog!B2', [x['location'] for x in plan['unknown']])
            self.assertEqual([x['location'] for x in plan['unaffected']], ['Catalog!B3'])

    def test_reordered_source_columns_are_unknown(self):
        self.freeze_coverage()
        workbook(self.root / 'reordered.xlsx', 'Supplier', ['product_key', 'note', 'material'],
                 [['ABC-9182', 'unrelated', 'stainless steel'],
                  ['XYZ-2040', 'unrelated', 'aluminium']])
        reordered = self.rt.register(self.root / 'reordered.xlsx', 'Reordered supplier')
        plan = facts.plan_impact(self.state.db, job=self.job, old_source=self.old,
            replacement_source=reordered, baseline_artifact=self.catalog)
        self.assertEqual([x['location'] for x in plan['unknown']], ['Catalog!B2', 'Catalog!B3', 'Catalog!B4'])

    def test_unreviewed_or_conflicting_source_cannot_claim_complete_coverage(self):
        self.bind('ABC-9182', 'B2', 'B2', 'steel', review_state='pending')
        with self.assertRaisesRegex(ValueError, 'reviewed'):
            facts.record_coverage(self.state.db, job=self.job, output_artifact=self.catalog,
                output_sheet='Catalog', output_cell='B2', state='complete', note='Review pending.')
        self.assertEqual(self.state.db.execute('SELECT count(*) FROM relay_fact_coverage').fetchone()[0], 0)

    def test_replacement_conflicting_with_independent_reviewed_source_is_unknown(self):
        self.bind('ABC-9182', 'B2', 'B2', 'steel')
        other = self.rt.register(self.root / 'supplier.xlsx', 'Independent reviewed supplier copy')
        facts.record_binding(self.state.db, job=self.job, entity_key='ABC-9182',
            predicate='material', value='steel', source_artifact=other,
            source_sheet='Supplier', source_key_column='A', source_cell='B2',
            evidence='Separate reviewed source for the same product.',
            output_artifact=self.catalog, output_sheet='Catalog', output_cell='B2',
            reviewer='fixture-reviewer', review_state='reviewed')
        facts.record_coverage(self.state.db, job=self.job, output_artifact=self.catalog,
            output_sheet='Catalog', output_cell='B2', state='complete', note='Two reviewed sources.')
        plan = self.plan()
        self.assertEqual([x['location'] for x in plan['unknown']], ['Catalog!B2'])
        self.assertIn('conflicts', plan['unknown'][0]['reason'])

    def test_unsupported_chart_workbook_is_rejected(self):
        book = Workbook()
        ws = book.active
        ws.append(['Key', 'Value'])
        ws.append(['ABC-9182', 3])
        chart = BarChart()
        chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=2), titles_from_data=True)
        ws.add_chart(chart, 'D2')
        path = self.root / 'with-chart.xlsx'
        book.save(path)
        book.close()
        with self.assertRaisesRegex(ValueError, 'Unsupported XLSX'):
            facts._plain_xlsx(path)

    def test_stale_file_and_unapproved_patch_do_not_create_candidate(self):
        self.freeze_coverage()
        plan = self.plan()
        before = self.state.db.execute('SELECT count(*) FROM production_artifacts').fetchone()[0]
        with self.assertRaisesRegex(ValueError, 'exactly match'):
            facts.revise_xlsx(self.rt, plan=plan,
                              patches={'Catalog!B3': 'stainless steel'}, reviewer='reviewer')
        self.assertEqual(self.state.db.execute('SELECT count(*) FROM production_artifacts').fetchone()[0], before)
        blob = Path(self.rt.artifact(self.old)['blob'])
        blob.chmod(0o600)
        blob.write_bytes(b'corrupted')
        with self.assertRaisesRegex(ValueError, 'missing or changed|hash mismatch'):
            facts.revise_xlsx(self.rt, plan=plan,
                              patches={'Catalog!B2': 'stainless steel'}, reviewer='reviewer')
        self.assertEqual(self.state.db.execute('SELECT count(*) FROM relay_fact_revisions').fetchone()[0], 0)

    def test_restart_keeps_plan_and_does_not_accept_projection_edits(self):
        self.freeze_coverage()
        before = self.plan()
        root = workflow_files.sync(self.state, self.job)
        self.state.db.close()
        self.state = State(self.root / 'data/state.sqlite')
        self.rt = Runtime(pc.root(self.state), connection=self.state.db)
        self.assertEqual(before, self.plan())
        projection = root / '.relay/job.sqlite'
        projection.chmod(0o600)
        with closing(sqlite3.connect(projection)) as db:
            db.execute("UPDATE meta SET value='forged' WHERE key='authority'")
            db.commit()
        with self.assertRaisesRegex(ValueError, 'modified'):
            workflow_files.sync(self.state, self.job)
        self.assertEqual(self.plan(), before)
        self.assertEqual(self.state.db.execute('SELECT count(*) FROM relay_fact_bindings').fetchone()[0], 2)


if __name__ == '__main__':
    unittest.main()
