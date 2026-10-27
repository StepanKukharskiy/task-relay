"""Registered in-assignment work: scripted models, real local adapters, no paid calls."""
import copy, hashlib, io, json, unittest
from pathlib import Path
from unittest.mock import patch
from orchestrator import contracts as c, executors, job_owner, worker_capabilities, worker_tools
from orchestrator.document_authoring import PPTX
from orchestrator.gemini_worker import execute
from orchestrator.workers import atomic
from tests import test_owner_compilation as fixtures
from tests.test_operation_session import Script, CONFIG
from orchestrator.owner_compilation import label

class Tests(unittest.TestCase):
    create=fixtures.Tests.create
    source=fixtures.Tests.source
    session=fixtures.Tests.session
    invoke=fixtures.Tests.invoke
    def setUp(self):
        fixtures.Tests.setUp(self)
        self.grant.update(operations=['pptx.create','pptx.render','text.bundle'],deliverables={'deck':PPTX,'pdf':'application/pdf'})
        self.grant['task_limits'].update(seconds=600,tool_calls=24,provider_requests=24,output_bytes=20000000)
        self.grant['workers'].append(worker_capabilities.entry(dict(type='gemini-code',model='fixture',runtime='a'*64)))

    def admit(self,cap='pptx.create',inputs=None):
        self.create(inputs=inputs or []);self.rt.tick('job')
        intent=dict(id='make',objective='Create a readable evidence deck, inspect and correct it.',inputs=[dict(artifact=i['artifact']) for i in inputs or []],
            work=dict(operation=cap,mode='create' if cap=='pptx.create' else 'execute',parameters={},**(dict(content_parameters=dict(max_records=1)) if cap=='pptx.create' else {})))
        self.invoke([('job_delegate',dict(request_id='work',tasks=[intent]))])
        self.rt.tick('job');task=self.rt.task('job','make')
        if task['latest'] is None:raise AssertionError(dict(task))
        state=self.factory.sessions[task['latest']];f=state['frozen'];control=Path(state['session']['control']);control.mkdir(parents=True,exist_ok=True)
        atomic(control/'launch.json',dict(credential_fingerprint=executors.fingerprint(CONFIG,f['backend'])))
        return f,control,state

    def report(self,f):
        return ('finish',dict(report_json=c.encoded(dict(summary='Built, inspected and corrected the exact deck.',decision='delivered',instruction='',findings=[],
            checks={f'c{i}':dict(passed=True,evidence='Rendered exact revised deck and inspected its pages') for i in range(1,len(f['criteria'])+1)}))))

    def test_real_build_inspect_correct_in_one_assignment(self):
        from pptx import Presentation
        from PIL import Image,ImageDraw
        image_path=self.root/'fixture-item.png'
        picture=Image.new('RGB',(240,160),'#eee3d1');draw=ImageDraw.Draw(picture)
        draw.rounded_rectangle((20,40,220,135),radius=18,fill='#ae553c');picture.save(image_path)
        image_id=self.rt.register(image_path,'Synthetic image fixture')
        image_input=dict(artifact=image_id,path='prior/item.png',purpose='Exact synthetic item image',authority='Synthetic fixture, no real product claims',media_type='image/png')
        f,control,state=self.admit(inputs=[image_input])
        self.assertEqual(f['backend']['type'],'gemini-agent')
        self.assertEqual(f['limits']['provider_requests'],24)
        self.assertEqual({t['id'] for t in self.rt.status('job')['tasks']},{'owner','make',label('make','review')})
        page=dict(kind='page',slot=1,layout='text',content=dict(title='Wrong heading',body='Verified fixture evidence.'),notes='Synthetic plumbing test')
        corrected=copy.deepcopy(page);corrected['layout']='content_cards'
        corrected['content']=dict(title='Corrected heading',per_page=1,items=[dict(label='Synthetic item',body='Verified fixture evidence.',image='prior/item.png',link='https://example.org/item',credit='Synthetic test image')])
        steps=[('operation_submit',dict(request_key='page',revision=0,records=[page])),
            ('operation_submit',dict(request_key='end',revision=1,records=[dict(kind='end',count=1)])),
            ('work_run',dict(step='build',request_key='build-v1')),('work_run',dict(step='inspect',request_key='inspect-v1')),
            ('image_view',dict(path='outputs/work-inspect/slide-previews.zip',members=['slide-001.png'])),
            ('operation_amend',dict(request_key='correct-page',revision=2,expected_record_sha256=c.digest(page),record=corrected)),
            self.report(f), # old build/render must be rejected after amendment
            ('work_run',dict(step='build',request_key='build-v2')),('work_run',dict(step='inspect',request_key='inspect-v2')),
            self.report(f), # new render has not yet been observed
            ('image_view',dict(path='outputs/work-inspect/slide-previews.zip',members=['slide-001.png'])),self.report(f)]
        script=Script(steps)
        save=worker_tools.Session.save
        def interrupt_after_correction(session,contents,history,requests,tools,**options):
            save(session,contents,history,requests,tools,**options)
            if requests==6 and not options.get('pending'):raise InterruptedError('Controlled restart after committed correction')
        try:
            with patch.object(worker_tools.Session,'save',interrupt_after_correction):
                with self.assertRaisesRegex(InterruptedError,'Controlled restart'):
                    execute(f,control,script,lambda:(CONFIG,f['backend']))
            saved=json.loads((control/'work-session.json').read_text())
            self.assertEqual((saved['requests'],saved['tools']),(6,6))
            original_started=saved['started']
            execute(f,control,script,lambda:(CONFIG,f['backend']))
            self.assertEqual(json.loads((control/'work-clock.json').read_text())['started'],original_started)
        except BaseException:
            for p in sorted(control.glob('tool-*.json')):print(p.name,p.read_text()[:1800])
            raise
        self.assertEqual(len(script.requests),12)
        self.assertTrue(any('inlineData' in p for m in script.requests[5]['contents'] for p in m['parts']))
        deck=Presentation(str(Path(f['workspace'])/'outputs/work-build/result.pptx'))
        self.assertIn('Corrected heading','\n'.join(s.text for s in deck.slides[0].shapes if s.has_text_frame))
        self.assertTrue(any(shape.shape_type==13 for shape in deck.slides[0].shapes))
        self.assertTrue(any(rel.target_ref=='https://example.org/item' for rel in deck.slides[0].part.rels.values()))
        self.assertEqual(len(list((control/'work-operations').glob('*/outcome.json'))),4)
        self.assertTrue((Path(f['workspace'])/'outputs/work-inspect/preview.pdf').is_file())
        state['status']=dict(status='finished',exit_code=0,external_outcome='no_pending_response',pending_requests=[])
        self.rt.tick('job')
        self.assertEqual(self.rt.task('job','make')['status'],'awaiting_review')

        # A later assignment can inspect this loop's committed PDF and deck via
        # their matching pages, just as it can inspect standalone render results.
        from orchestrator import artifact_roles,pptx_render
        inputs=[]
        for output in f['outputs']:
            artifact=self.rt.output('job','make',output['path'])
            inputs.append(dict(artifact=artifact['id'],path=output['path'],media_type=output['media_type'],purpose='Exact prior work',authority='Unaccepted fixture'))
        follow=dict(id='follow',objective='Revise exact prior candidate',instruction='Inspect prior work before revising',inputs=inputs)
        worker_tools.configure(follow,dict(operation='pptx.create',mode='create',parameters={},content_parameters=dict(max_records=1)),job_owner.original(self.rt,'job')['job_owner'])
        artifact_roles.bind(follow,self.rt,[]);artifact_roles.verify(self.rt,follow)
        views=pptx_render.evidence_views(self.rt,inputs)
        self.assertEqual(views['outputs/work-inspect/preview.pdf']['path'],'outputs/work-inspect/slide-previews.zip')
        self.assertEqual(views['outputs/work-build/result.pptx']['proof']['candidate_sha256'],self.rt.output('job','make','outputs/work-build/result.pptx')['sha256'])
        self.assertEqual(artifact_roles.operation_inputs(follow,'build'),[])

        # Each native version is real and retained, and the PDF matches v2.
        from pypdf import PdfReader
        pdf=PdfReader(str(Path(f['workspace'])/'outputs/work-inspect/preview.pdf'))
        self.assertIn('Corrected heading',pdf.pages[0].extract_text())
        v1=control/'work-operations/0001/workspace/outputs/result.pptx'
        self.assertIn('Wrong heading','\n'.join(s.text for s in Presentation(str(v1)).slides[0].shapes if s.has_text_frame))
        for path in control.glob('tool-*.json'):
            result=json.loads(path.read_text())['result']
            if path.name=='tool-07-00.json':self.assertIn('stale',result['error'])
            elif path.name=='tool-10-00.json':self.assertIn('Inspect every',result['error'])
            else:self.assertNotIn('error',result,path.read_text())
        self.review_and_complete(f,control)
        import os,shutil
        if os.environ.get('RELAY_WORKER_PROOF_OUTPUT'):
            destination=Path(os.environ['RELAY_WORKER_PROOF_OUTPUT']);destination.mkdir(parents=True,exist_ok=True)
            for version,build,inspect in [('v1','0001','0002'),('v2','0003','0004')]:
                shutil.copyfile(control/f'work-operations/{build}/workspace/outputs/result.pptx',destination/(version+'.pptx'))
                shutil.copyfile(control/f'work-operations/{inspect}/workspace/delivery/preview.pdf',destination/(version+'.pdf'))
            shutil.copytree(control,destination/'worker-journal',dirs_exist_ok=True)
            atomic(destination/'execution-record.json',dict(controlled=True,live_provider_calls=0,scripted_worker_requests=len(script.requests),worker_assignment=f['assignment_id'],
                worker_limits=f['limits'],source_task_count=3,registered_calls=4,worker_restart_after_request=6,completed_job_status=self.rt.status('job')['status'],
                checks=['real PPTX/PDF v1 and v2','actual slide pixels supplied','missing image repaired','hyperlink preserved','stale build rejected','unobserved v2 rejected','independent exact-version review','original request/tool/byte/deadline allowance retained']))

    def review_and_complete(self,producer,producer_control):
        rid=label('make','review');state=self.factory.sessions[self.rt.task('job',rid)['latest']]
        f=state['frozen'];control=Path(state['session']['control']);control.mkdir(parents=True,exist_ok=True)
        atomic(control/'launch.json',dict(credential_fingerprint=executors.fingerprint(CONFIG,f['backend'])))
        preview=next(i for i in f['inputs'] if i['path'].endswith('slide-previews.zip'))
        from orchestrator.final_quality import DIMENSIONS
        checks={f'c{i}':dict(passed=True,evidence='Inspected synthetic final candidate') for i in range(1,len(f['criteria'])+1)}
        for number,binding in f['final_quality']['checks'].items():
            checked=[dict(path=p,location='Slide/page 1 of exact final version') for p in dict.fromkeys([binding['candidate'],*binding['inspection_inputs']])]
            checks['c'+number]['evidence']={dim:dict(verdict='satisfied',reason='Scripted judgment of synthetic candidate',inspected=checked) for dim in DIMENSIONS}
        report=dict(summary='Independently inspected exact final pages',decision='accept',instruction='',findings=[],checks=checks)
        script=Script([('image_view',dict(path=preview['path'],members=['slide-001.png'])),('finish',dict(report_json=c.encoded(report)))])
        with patch('task_relay.code_runtime.available',return_value=dict(tools=[])):
            execute(f,control,script,lambda:(CONFIG,f['backend']))
        state['status']=dict(status='finished',exit_code=0,external_outcome='no_pending_response',pending_requests=[])
        self.rt.tick('job')
        self.assertEqual(self.rt.task('job','make')['status'],'completed')
        deck=self.rt.output('job','make','outputs/work-build/result.pptx');pdf=self.rt.output('job','make','outputs/work-inspect/preview.pdf')
        attempt=self.rt.db.execute('SELECT * FROM production_attempts WHERE id=?',(f['assignment_id'],)).fetchone()
        self.assertTrue(job_owner.inspected_render(self.rt,attempt,f,deck))
        self.assertFalse(job_owner.inspected_render(self.rt,attempt,f,{**deck,'sha256':'0'*64}))
        self.assertFalse(job_owner.inspected_render(self.rt,attempt,f,{**deck,'attempt':'older-attempt'}))
        self.invoke([('job_complete',dict(deliverables_json=c.encoded(dict(deck=deck['id'],pdf=pdf['id'])),summary='Reviewed exact corrected deck and matching PDF'))])
        self.rt.tick('job');self.assertEqual(self.rt.status('job')['status'],'completed')

    def text_work(self):
        self.grant['deliverables']={'report':'text/plain'}
        return self.admit('text.bundle',[self.source('Exact synthetic input, retained verbatim.')])

    def test_blocked_worker_preserves_verified_candidates_without_delivery_or_reexecution(self):
        f,control,_=self.text_work()
        name,args=self.report(f);report=json.loads(args['report_json'])
        report.update(decision='blocked',summary='Candidate needs semantic correction',instruction='Correct the saved candidate')
        for check in report['checks'].values():check['passed']=False
        execute(f,control,Script([('work_run',dict(step='build',request_key='candidate')),
            ('finish',dict(report_json=c.encoded(report)))]),lambda:(CONFIG,f['backend']))
        self.assertFalse((control/'work-delivery.json').exists())
        with self.assertRaises(FileNotFoundError):worker_tools.verify_delivery(f,control)
        with patch.object(worker_tools,'execute_local',side_effect=AssertionError('Verification cannot execute')):
            proof=worker_tools.verify_delivery(f,control,candidate=True)
        self.assertFalse(proof['acceptance']);self.assertTrue(proof['candidate_only'])
        self.assertEqual(set(proof['steps']),{'build'})
        output=next(iter(f['work_tools']['steps']['build']['outputs'].values()))['path']
        (Path(f['workspace'])/output).write_text('Changed draft')
        with self.assertRaisesRegex(ValueError,'committed version'):
            worker_tools.verify_delivery(f,control,candidate=True)

    def test_restart_after_native_commit_reuses_response_and_result_without_dispatch(self):
        f,control,_=self.text_work();first=Script([('work_run',dict(step='build',request_key='one'))])
        with patch.object(worker_tools.Session,'publish',side_effect=KeyboardInterrupt('after commit')):
            with self.assertRaises(KeyboardInterrupt):execute(f,control,first,lambda:(CONFIG,f['backend']))
        folder=control/'work-operations/0001';before=(folder/'outcome.json').read_bytes()
        follow=Script([self.report(f)])
        with patch.object(worker_tools,'execute_local',side_effect=AssertionError('Confirmed operation must not replay')):
            execute(f,control,follow,lambda:(CONFIG,f['backend']))
        self.assertEqual(len(first.requests),1);self.assertEqual(len(follow.requests),1)
        self.assertEqual((folder/'outcome.json').read_bytes(),before)
        result=json.loads((control/'agent-result.json').read_text());self.assertEqual((result['requests'],result['tool_calls']),(2,2))
        self.assertEqual(len(list((control/'work-operations').iterdir())),1)
        worker_tools.verify_delivery(f,control)

    def test_unknown_native_dispatch_is_held_in_worker_and_factory(self):
        f,control,_=self.text_work();script=Script([('work_run',dict(step='build',request_key='unknown'))])
        with patch.object(worker_tools,'execute_local',side_effect=KeyboardInterrupt('lost native result')):
            with self.assertRaises(KeyboardInterrupt):execute(f,control,script,lambda:(CONFIG,f['backend']))
        follow=Script([])
        with self.assertRaisesRegex(RuntimeError,'uncertain'):
            execute(f,control,follow,lambda:(CONFIG,f['backend']))
        self.assertEqual(follow.requests,[])
        from orchestrator.adapters import GeminiFactory,CodexFactory
        with patch.object(CodexFactory,'inspect',return_value=dict(status='finished',exit_code=1)):
            receipt=GeminiFactory().inspect(dict(id=f['assignment_id'],control=str(control),backend=f['backend']))
        self.assertEqual(receipt['status'],'uncertain');self.assertIn('work-operations/0001',receipt['pending_requests'])

    def test_lost_generation_does_not_repeat_on_restart(self):
        f,control,_=self.text_work();script=Script([])
        with self.assertRaises(StopIteration):execute(f,control,script,lambda:(CONFIG,f['backend']))
        with self.assertRaisesRegex(RuntimeError,'unresolved'):
            execute(f,control,script,lambda:(CONFIG,f['backend']))
        self.assertEqual(len(script.requests),1)

    def test_existing_cumulative_request_tool_and_byte_allowance_survives_restart(self):
        self.grant['task_limits'].update(tool_calls=3,provider_requests=3,output_bytes=750)
        f,control,_=self.text_work()
        # Stop between complete turns, then spend the remaining two requests.
        save=worker_tools.Session.save
        def stop(session,*args,**kw):
            save(session,*args,**kw)
            if args[2]==1 and not kw.get('pending'):raise InterruptedError('checkpoint')
        with patch.object(worker_tools.Session,'save',stop):
            with self.assertRaises(InterruptedError):execute(f,control,Script([('work_run',dict(step='build',request_key='first'))]),lambda:(CONFIG,f['backend']))
        saved=json.loads((control/'work-session.json').read_text());self.assertEqual((saved['requests'],saved['tools']),(1,1))
        # An identical native request is reused and still consumes one tool call.
        script=Script([('work_run',dict(step='build',request_key='first')),self.report(f)])
        execute(f,control,script,lambda:(CONFIG,f['backend']))
        result=json.loads((control/'agent-result.json').read_text());self.assertEqual((result['requests'],result['tool_calls']),(3,3))
        self.assertEqual(len(list((control/'work-operations').glob('*/outcome.json'))),1)
        journal=json.loads((control/'work-operations/0001/outcome.json').read_text())
        self.assertLess(sum(o['bytes'] for o in journal['outputs'].values()),f['limits']['output_bytes'])

    def test_deadline_and_failed_native_bytes_are_not_reset(self):
        f,control,_=self.text_work();files=worker_tools.files(f,control);session=worker_tools.Session(files,control)
        def fail(sub,folder,*args):
            p=Path(sub['workspace'])/sub['outputs'][0]['path'];p.parent.mkdir(parents=True,exist_ok=True);p.write_bytes(b'x'*100)
            raise ValueError('Controlled confirmed failure')
        with patch.object(worker_tools,'execute_local',side_effect=fail):
            result=session.run(dict(step='build',request_key='failed'))
        self.assertEqual(result['effect'],'confirmed');self.assertEqual(session.bytes_used(),100)
        resumed=worker_tools.Session(worker_tools.files(f,control),control)
        self.assertEqual(resumed.bytes_used(),100)
        with self.assertRaisesRegex(ValueError,'already failed'):resumed.run(dict(step='build',request_key='failed'))
        with patch('orchestrator.worker_tools.time.time',return_value=session.started+f['limits']['seconds']+1),patch.object(worker_tools,'execute_local',side_effect=AssertionError('Expired work cannot run')):
            with self.assertRaisesRegex(ValueError,'deadline'):resumed.run(dict(step='build',request_key='new'))
        # Simulate a lower remaining byte allowance without adding authority.
        with patch.object(resumed,'bytes_used',return_value=f['limits']['output_bytes']):
            with self.assertRaisesRegex(ValueError,'allowance'):resumed.run(dict(step='build',request_key='another'))

    def test_ungranted_tools_custom_code_and_native_output_overwrites_fail(self):
        f,control,_=self.text_work()
        files=worker_tools.files(f,control)
        with self.assertRaisesRegex(ValueError,'code owned'):files.call('file_write',dict(path=f['outputs'][0]['path'],text='Forged native output'))
        with self.assertRaisesRegex(ValueError,'not granted'):files.call('python_run',dict(code='pass',seconds=1))
        script=Script([('work_run',dict(step='ungranted',request_key='bad')),('work_run',dict(step='build',request_key='ok')),self.report(f)])
        execute(f,control,script,lambda:(CONFIG,f['backend']))
        first=json.loads((control/'tool-01-00.json').read_text());self.assertIn('error',first['result'])
        self.assertEqual(len(list((control/'work-operations').glob('*/intent.json'))),1)
        for request in script.requests:
            self.assertNotIn('python_run',[d['name'] for d in request['tools'][0]['functionDeclarations']])

    def test_old_grant_cannot_acquire_worker_tools_or_extra_operations(self):
        self.create();self.rt.tick('job');session=self.session()
        task=dict(id='make',objective='Do work',work=dict(operation='pptx.create',parameters={}))
        session.grant.pop('worker_tools_version')
        from orchestrator.owner_intents import compile
        with self.assertRaisesRegex(ValueError,'outside this grant'):compile(session,dict(request_id='old',tasks=[task]))
        session.grant['worker_tools_version']=1;session.grant['operations'].pop('pptx.render')
        with self.assertRaisesRegex(ValueError,'inspection companion'):compile(session,dict(request_id='missing',tasks=[task]))
        task['work']['operation']='gemini.text'
        with self.assertRaisesRegex(ValueError,'outside this job grant'):compile(session,dict(request_id='paid',tasks=[task]))
        self.assertEqual(len(self.rt.status('job')['tasks']),1)

    def test_generation_retry_allowance_is_cumulative_after_checkpoint_restart(self):
        from task_relay import gemini
        self.grant['generation_recovery']=dict(version=1,max_retries=1)
        f,control,_=self.text_work();script=Script([('work_run',dict(step='build',request_key='one'))])
        seen=[]
        class Transport:
            def request(inner,endpoint,payload,**kw):
                seen.append(copy.deepcopy(payload))
                if len(seen)==1:raise gemini.ProviderError('timeout',uncertain=True,detail={'failure_kind':'timeout'})
                return script.request(endpoint,payload,**kw)
        save=worker_tools.Session.save
        def stop(session,*args,**kw):
            save(session,*args,**kw)
            if args[2]==2 and not kw.get('pending'):raise KeyboardInterrupt('checkpoint after one paid retry')
        with patch.object(worker_tools.Session,'save',stop):
            with self.assertRaises(KeyboardInterrupt):execute(f,control,Transport(),lambda:(CONFIG,f['backend']))
        class Timeout:
            def request(inner,*args,**kw):raise gemini.ProviderError('timeout',uncertain=True,detail={'failure_kind':'timeout'})
        with self.assertRaises(gemini.ProviderError):execute(f,control,Timeout(),lambda:(CONFIG,f['backend']))
        self.assertEqual(len(list(control.glob('api-*.request.json'))),3)
        self.assertEqual(len(list(control.glob('api-*.recovery.json'))),1)

    def test_received_invalid_envelope_corrects_inside_worker_without_owner_round_trip(self):
        f,control,_=self.text_work()
        invalid=dict(candidates=[dict(finishReason='STOP',content=dict(role='model',parts=[dict(text='No executable tool call')]))])
        script=Script([invalid,('work_run',dict(step='build',request_key='valid')),self.report(f)])
        result=execute(f,control,script,lambda:(CONFIG,f['backend']))
        self.assertEqual(result['decision'],'delivered');self.assertEqual(len(script.requests),3)
        self.assertEqual(json.loads((control/'api-01.validation.json').read_text())['effect'],'none')
        self.assertEqual(len(self.factory.calls),2) # owner and one productive worker; reviewer not dispatched yet

    def test_changed_input_and_altered_delivery_are_rejected_without_native_replay(self):
        f,control,_=self.text_work();script=Script([('work_run',dict(step='build',request_key='first')),self.report(f)])
        execute(f,control,script,lambda:(CONFIG,f['backend']))
        target=Path(f['workspace'])/f['outputs'][0]['path'];original=target.read_bytes();target.write_bytes(b'Altered output')
        with self.assertRaisesRegex(ValueError,'committed version'):worker_tools.verify_delivery(f,control)
        target.write_bytes(original)
        files=worker_tools.files(f,control);source=Path(f['workspace'])/f['inputs'][0]['path'];source.chmod(0o600);source.write_text('Changed input')
        with patch.object(worker_tools,'execute_local',side_effect=AssertionError('Changed evidence must not dispatch')):
            session=worker_tools.Session(files,control)
            with self.assertRaisesRegex(ValueError,'Frozen work input changed'):session.run(dict(step='build',request_key='new'))

    def test_real_native_versions_share_original_output_byte_ceiling(self):
        self.grant['task_limits']['output_bytes']=500
        f,control,_=self.text_work();session=worker_tools.Session(worker_tools.files(f,control),control)
        first=session.run(dict(step='build',request_key='v1'))
        self.assertNotIn('error',first);spent=session.bytes_used();self.assertGreater(spent,100)
        session=worker_tools.Session(worker_tools.files(f,control),control)
        results=[session.run(dict(step='build',request_key='v'+str(n))) for n in range(2,5)]
        self.assertTrue(any('error' in r for r in results))
        self.assertLessEqual(session.bytes_used(),500)
        for folder,intent,receipt in session.receipts:
            frozen=json.loads((folder/'workspace/.relay/ASSIGNMENT.json').read_text())
            prior=sum(sum(d['bytes'] for d in r.get('outputs',{}).values()) for p,_,r in session.receipts if p.name<folder.name)
            self.assertLessEqual(frozen['limits']['output_bytes'],500-prior)

    def test_compiler_defaults_fit_provider_ceiling_and_do_not_expose_paid_tools(self):
        self.grant['task_limits'].update(provider_requests=12,tool_calls=24)
        self.create();self.rt.tick('job');session=self.session();before=copy.deepcopy(session.grant)
        task=dict(id='make',objective='Build and inspect',work=dict(operation='pptx.create',mode='create'))
        session.delegate(dict(request_id='bounded-default',tasks=[task]))
        spec=self.rt.spec(self.rt.task('job','make'))
        # Three batched fragment submissions hold eight pages and their seal;
        # the remaining nine calls are still reserved for work and inspection.
        self.assertEqual(spec['operation_contract']['parameters']['max_records'],8)
        self.assertEqual(spec['limits']['provider_requests'],12);self.assertEqual(session.grant,before)
        from orchestrator.connector_protocol import capture
        from orchestrator.owner_intents import compile
        session.grant['operations']['gemini.text']=capture('gemini.text')
        session.grant['worker_tools_version']=1 # Captured legacy grants remain local-only.
        task['work']=dict(operation='gemini.text',parameters=dict(model='fixture',max_output_tokens=128))
        with self.assertRaisesRegex(ValueError,'local worker-tool interface'):compile(session,dict(request_id='no-paid-tool',tasks=[task]))
