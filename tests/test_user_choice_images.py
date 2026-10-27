"""User-choice sourcing through existing compiler, worker, native and export paths."""
import copy,json,unittest
from pathlib import Path
from unittest.mock import patch
from orchestrator import job_owner
from tests.test_correction_handoff_lifecycle import Fixtures
from tests.test_image_sources import photo,bundle,subject,page as source_page
from tests.test_owner_intents import worker


class UserChoice(Fixtures,unittest.TestCase):
    schema='furniture'
    def setUp(self):
        runtime=patch("task_relay.code_runtime.available",return_value={"tools":{}});runtime.start();self.addCleanup(runtime.stop)
        create=job_owner.create
        def admission(rt,run,request,backend,inputs,grant,**kwargs):
            grant.update(execution_policy=dict(version=2,review='user',public_read=False),
                         deliverables=getattr(self,'deliverables',dict(report='text/plain',photos='application/zip')),producer_binding_version=1)
            grant['operations'] += ['images.pptx','images.gallery','text.bundle']
            if getattr(self,'named_nodes',False):
                from orchestrator import node_spec
                grant['node_contract']=node_spec.capture(dict(schema_version=3,id='request',operation='model',objective=request,inputs={'evidence'+str(n):{'artifact':v['artifact']} for n,v in enumerate(inputs)},input_definitions={'evidence'+str(n):dict(type='artifact',label='Evidence',required=True,media_type=v['media_type']) for n,v in enumerate(inputs)},outputs=[dict(name=k,media_type=v,purpose='Requested file') for k,v in grant['deliverables'].items()]))
            return create(rt,run,request,backend,inputs,grant,**kwargs)
        with patch.object(job_owner,'create',side_effect=admission):Fixtures.setUp(self)

    def submit(self,key,tasks):
        s=self.session();result=s.delegate(dict(request_id=key,tasks=tasks));self.rt.tick('job');return result

    def test_typed_inventory_photos_advice_restart_and_all_candidate_export_without_review(self):
        tasks=[dict(id='item',objective='Identify observed subject; user assesses fit',content=dict(family='collection.subjects',sources=dict(source_path={'artifact':self.inv['artifact']}),parameters=dict(max_records=1,fields=[]))),
            dict(id='photos',objective='Acquire all observed candidates',operation='images.products',parameters=dict(inventory_path={'task':'item','output':'result'},catalog_path=None,source_pack_path={'artifact':self.pack['artifact']},count_policy='target')),
            worker('advice',[{'task':'photos','output':'result'}])]
        self.assertEqual(set(self.submit('flow',tasks)['queued']),{'item','photos','advice'})
        label=self.h.inventory['items'][0]['label']
        self.run_worker('item',lambda f:[('operation_submit',dict(request_key='item',revision=0,records=[dict(kind='item',slot=1,label=label,attributes={},quote=label,bbox=None,notes='Not independently reviewed')])),('operation_submit',dict(request_key='seal',revision=1,records=[dict(kind='end',count=1,notes='One requested subject')]))])
        with patch('orchestrator.image_sources.download',return_value=(photo(),'image/jpeg','https://images.example/product.jpg')):self.native('photos')
        photos=self.artifact('photos','outputs/result.zip');before=Path(photos['blob']).read_bytes()
        self.run_worker('advice',lambda f:[('file_write',dict(path=f['outputs'][0]['path'],text='Advisory only: prefer another style. All acquired candidates remain available. Buying details unknown. Target three not achieved.'))])
        self.rt.close()
        from orchestrator.runtime import Runtime
        self.rt=Runtime(self.root,self.factory,connection=self.state.db)
        result,script=self.invoke([('node_result',dict(outputs=dict(report={'task':'advice','output':'result'},photos={'task':'photos','output':'result'}),summary='Candidates ready for user choice; target shortfall disclosed.'))])
        self.rt.tick('job');self.assertEqual(self.rt.status('job')['status'],'completed')
        delivery=json.loads(self.rt.db.execute("SELECT data FROM production_events WHERE kind='owner_delivery_verified'").fetchone()[0])
        self.assertFalse(delivery['user_acceptance'])
        self.assertEqual({v['review_status'] for v in delivery['deliverables'].values()},{'unreviewed'})
        self.assertFalse(any(self.rt.spec(t).get('review_of') for t in self.rt.status('job')['tasks']))
        # Delivery copies only declared artifacts. No hidden conversion runs.
        from task_relay import workflow_files,result_handoff
        with patch.object(workflow_files,'folder_path',return_value=self.home/'export'), patch('task_relay.image_gallery.presentation',side_effect=AssertionError('Hidden PPTX conversion')), patch('task_relay.image_gallery.render',side_effect=AssertionError('Hidden gallery conversion')):
            folder=result_handoff.sync(self.state,'job')
        self.assertFalse(list(folder.rglob('*.pptx')))
        self.assertFalse(list(folder.rglob('*.html')))
        self.assertEqual(Path(photos['blob']).read_bytes(),before)
        self.assertEqual(self.rt.db.execute('SELECT count(*) FROM production_decisions').fetchone()[0],0)
        with patch.object(workflow_files,'folder_path',return_value=self.home/'automatic-export'):
            result_handoff.tick(self.state);result_handoff.tick(self.state)
        rows=self.state.db.execute("SELECT text FROM outbox WHERE id='production:job:files-ready'").fetchall()
        self.assertEqual(len(rows),1)
        self.assertIn('files ready for your choice',rows[0][0]);self.assertNotIn('selected results',rows[0][0])
        # Migrating an installation after this event must not publish a backlog.
        with self.state.db:
            self.state.db.execute("DELETE FROM outbox WHERE id='production:job:files-ready'")
            self.state.db.execute('DROP TABLE result_handoff_epoch')
            result_handoff.initialize(self.state.db)
        with patch.object(workflow_files,'folder_path',return_value=self.home/'automatic-export'):
            result_handoff.tick(self.state)
        self.assertIsNone(self.state.db.execute("SELECT 1 FROM outbox WHERE id='production:job:files-ready'").fetchone())

    def test_explicit_review_still_blocks_delivery(self):
        self.submit('explicit',[worker('advice'),dict(id='inspect',objective='Explicit review requested',review_of='advice'),dict(id='photos',objective='Acquire candidates',operation='images.products',parameters=dict(inventory_path={'artifact':self.inv['artifact']},catalog_path=None,source_pack_path={'artifact':self.pack['artifact']},count_policy='target'))])
        with patch('orchestrator.image_sources.download',return_value=(photo(),'image/jpeg','https://images.example/product.jpg')):self.native('photos')
        self.run_worker('advice',lambda f:[('file_write',dict(path=f['outputs'][0]['path'],text='Candidate'))])
        s=self.session()
        # Explicit review is retained and cannot be replaced by the user policy.
        self.assertEqual(self.rt.spec(self.rt.task('job','inspect'))['review_of'],'advice')
        self.assertNotEqual(self.rt.task('job','inspect')['status'],'completed')
        with self.assertRaisesRegex(ValueError,'review|active|running|unfinished|settling|finished'):
            s.complete(dict(deliverables_json=json.dumps(dict(report=self.artifact('advice','outputs/result.txt')['id'],photos=self.artifact('photos','outputs/result.zip')['id'])),summary='Attempt premature delivery'))

    def test_old_policy_and_intrinsic_review_contracts_remain_strict(self):
        from orchestrator.owner_compilation import final_companions
        task=dict(id='photos',objective='Photos',inputs=[],outputs=[dict(path='outputs/result.zip',media_type='application/zip',purpose='Photos')],dependencies=[],execution={'capability':'images.products'})
        before=copy.deepcopy(self.grant);before['execution_policy']['review']='owner'
        tasks=[copy.deepcopy(task)];final_companions(tasks,before);self.assertEqual(len(tasks),2)
        self.assertFalse(job_owner.user_choice(self.grant,dict(operation_contract=dict(id='collection.options',version=1))))
        self.assertFalse(job_owner.user_choice(self.grant,dict(execution=dict(capability='pptx.create'))))
        self.assertFalse(job_owner.user_choice(before,task))
        from orchestrator import document_authoring
        old=document_authoring.capture(['pptx.create'])
        old['formats'][document_authoring.PPTX]['operations']=['pptx.create','pptx.edit']
        document_authoring.validate(old,['pptx.create','images.pptx'])
        candidate=dict(execution=dict(capability='images.pptx'),outputs=[dict(path='outputs/result.pptx',media_type=document_authoring.PPTX)])
        with self.assertRaisesRegex(ValueError,'Use registered document operations'):
            document_authoring.check(candidate,old,['pptx.create','images.pptx'])
        bad=copy.deepcopy(self.grant);bad['execution_policy']['version']=1
        with self.assertRaises(ValueError):job_owner.validate_execution_policy(bad)

class ExportNodes(Fixtures,unittest.TestCase):
    schema='furniture'
    named_nodes=True
    setUp=UserChoice.setUp
    submit=UserChoice.submit
    deliverables=dict(deck='application/vnd.openxmlformats-officedocument.presentationml.presentation',gallery='text/html')

    def prepare_exports(self):
        # Public node_run uses the same operation definitions exposed to the UI.
        nodes=json.loads(Path('examples/image-candidates.workflow.json').read_text())['tasks'][-2:]
        for n in nodes:n['inputs']['recommendations']=dict(task='advice',output='result')
        tasks=[dict(schema_version=3,id='photos',objective='Acquire candidates',operation='images.products',inputs=dict(inventory_path={'artifact':self.inv['artifact']},catalog_path={'value':None},source_pack_path={'artifact':self.pack['artifact']},count_policy={'value':'target'})),
               dict(schema_version=3,id='advice',operation='model',objective='Compare all candidates',inputs=dict(photos=dict(task='photos',output='result')),input_definitions=dict(photos=dict(type='artifact',label='Images',required=True,media_type='application/zip')),outputs=[dict(name='result',media_type='text/plain',purpose='Advisory report')]),*nodes]
        result=self.session().call('node_run',dict(tasks=tasks));self.rt.tick('job')
        self.assertEqual(set(result['queued']),{'photos','advice','pptx','gallery'})
        with patch('orchestrator.image_sources.download',return_value=(photo(),'image/jpeg','https://images.example/product.jpg')):self.native('photos')
        self.run_worker('advice',lambda f:[('file_write',dict(path=f['outputs'][0]['path'],text='Advisory: keep every candidate. Buying details unknown. <script>unsafe</script>'))])
        self.assertIn(self.rt.task('job','pptx')['status'],('launching','running'))
        self.assertIn(self.rt.task('job','gallery')['status'],('launching','running'))

    def test_two_export_nodes_restart_exact_delivery_without_hidden_conversion(self):
        self.prepare_exports()
        self.native('pptx');self.native('gallery')
        deck=self.artifact('pptx','outputs/result.pptx');page=self.artifact('gallery','outputs/result.html')
        self.assertIn('&lt;script&gt;unsafe&lt;/script&gt;',Path(page['blob']).read_text())
        originals={a['id']:Path(a['blob']).read_bytes() for a in (deck,page)}
        self.rt.close()
        from orchestrator.runtime import Runtime
        self.rt=Runtime(self.root,self.factory,connection=self.state.db)
        self.invoke([('node_result',dict(outputs=dict(deck={'task':'pptx','output':'result'},gallery={'task':'gallery','output':'result'}),summary='All candidates, user chooses.'))])
        self.rt.tick('job');self.assertEqual(self.rt.status('job')['status'],'completed')
        from task_relay import result_handoff,workflow_files
        with patch.object(workflow_files,'folder_path',return_value=self.home/'export'), patch('task_relay.image_gallery.presentation',side_effect=AssertionError('Hidden conversion')), patch('task_relay.image_gallery.render',side_effect=AssertionError('Hidden conversion')):
            result_handoff.tick(self.state);result_handoff.tick(self.state)
        for a in (deck,page):
            copies=list((self.home/'export').rglob('*'+Path(a['path']).suffix))
            self.assertTrue(any(p.read_bytes()==originals[a['id']] for p in copies))
        self.assertFalse(any(self.rt.spec(t).get('review_of') for t in self.rt.status('job')['tasks']))
        self.assertEqual(self.rt.db.execute('SELECT count(*) FROM production_decisions').fetchone()[0],0)
        text=self.state.db.execute("SELECT text FROM outbox WHERE id='production:job:files-ready'").fetchone()[0]
        self.assertIn('.pptx',text);self.assertIn('.html',text)

    def failed_pptx(self):
        self.prepare_exports()
        from orchestrator import step_runner,operation_lifecycle
        session=self.factory.sessions[self.rt.task('job','pptx')['latest']]
        control=Path(session['session']['control']);control.mkdir(parents=True,exist_ok=True)
        with patch('task_relay.image_gallery.presentation',side_effect=ValueError('PPTX slide limit exceeded')):
            with self.assertRaisesRegex(ValueError,'slide limit'):
                step_runner.execute(session['frozen'],control)
        result=operation_lifecycle.terminal(session['frozen'],control,error='PPTX slide limit exceeded')
        self.assertEqual(result['effect'],'confirmed')
        session['status']=dict(status='finished',exit_code=1,external_outcome='no_pending_response',pending_requests=[])
        self.rt.tick('job');self.native('gallery')
        page=self.artifact('gallery','outputs/result.html');original=Path(page['blob']).read_bytes()
        self.assertNotEqual(self.rt.task('job','pptx')['status'],'completed')
        return session,control,page,original

    def test_failed_pptx_leaves_gallery_available_and_fresh_export_reuses_inputs(self):
        from orchestrator import step_runner
        session,control,page,original=self.failed_pptx()
        # A fresh explicit node is an authorized offline continuation, never a
        # replay of the dispatched failed operation or the successful sibling.
        self.submit('retry-export',[dict(id='pptx_again',objective='Retry confirmed local export',operation='images.pptx',parameters=dict(photos=dict(task='photos',output='result'),recommendations=dict(task='advice',output='result')))])
        self.native('pptx_again')
        self.assertEqual(Path(page['blob']).read_bytes(),original)
        self.assertEqual(len([s for s in self.factory.sessions.values() if s['frozen']['id']=='gallery']),1)
        self.assertEqual(len([s for s in self.factory.sessions.values() if s['frozen']['id']=='photos']),1)
        with self.assertRaisesRegex(ValueError,'replay'):
            step_runner.execute(session['frozen'],control)

    def test_terminal_pptx_failure_delivers_gallery_as_exact_partial(self):
        _,_,page,original=self.failed_pptx()
        from orchestrator.runtime_outcomes import Hold
        with self.assertRaisesRegex(Hold,'PPTX export failed'):
            self.invoke([('node_stop',dict(reason='PPTX export failed; gallery remains available.'))])
        # The fake process adapter reports the confirmed local owner exit.
        self.factory.sessions[self.rt.task('job','owner')['latest']]['status']=dict(status='finished',exit_code=1,external_outcome='no_pending_response',pending_requests=[])
        self.rt.tick('job')
        status=self.rt.status('job')
        self.assertEqual(status['status'],'blocked')
        delivery=status['delivery']
        self.assertEqual(delivery['missing_formats'],{'deck':self.deliverables['deck']})
        self.assertEqual([(a['artifact'],a['review_status']) for a in delivery['candidates']],[(page['id'],'unreviewed')])
        from task_relay import result_handoff,workflow_files
        with patch.object(workflow_files,'folder_path',return_value=self.home/'partial'), patch('task_relay.image_gallery.presentation',side_effect=AssertionError('Hidden conversion')):
            folder=result_handoff.sync(self.state,'job')
        self.assertTrue(any(p.read_bytes()==original for p in folder.rglob('*.html')))
        self.assertFalse(list(folder.rglob('*.pptx')))

    def test_named_export_inputs_reject_wrong_media_before_dispatch(self):
        from orchestrator import node_spec
        node=json.loads(Path('examples/image-candidates.workflow.json').read_text())['tasks'][-1]
        node['inputs']={k:dict(artifact='wrong') for k in node['inputs']}
        with self.assertRaisesRegex(ValueError,'photos requires'):
            node_spec.inspect_graph([node],artifact=lambda ident:dict(media_type='text/plain'))
        before=self.rt.db.execute('SELECT count(*) FROM production_tasks').fetchone()[0]
        with self.assertRaisesRegex(ValueError,'photos requires'):
            self.session().delegate(dict(request_id='wrong',tasks=[dict(id='bad',objective='Invalid export',operation='images.gallery',parameters=dict(photos={'artifact':self.inv['artifact']},recommendations={'artifact':self.inv['artifact']}))]))
        self.assertEqual(before,self.rt.db.execute('SELECT count(*) FROM production_tasks').fetchone()[0])

class DeliveryBytes(Fixtures,unittest.TestCase):
    schema='furniture'
    setUp=UserChoice.setUp
    submit=UserChoice.submit
    deliverables=dict(report='text/plain',second='text/plain')

    def test_large_worker_references_keep_job_cap_and_owner_writes_keep_control_cap(self):
        self.submit('large',[worker('left'),worker('right'),worker('small'),
            dict(id='large',objective='Combine exact text',operation='text.bundle',inputs=[dict(task='left',output='result'),dict(task='right',output='result')],parameters={}),
            dict(id='review_large',objective='Review exact combined text',review_of='large')])
        for ident,char in [('left','x'),('right','y')]:
            self.run_worker(ident,lambda f:[('file_write',dict(path=f['outputs'][0]['path'],text=char*20000))]+[('file_append',dict(path=f['outputs'][0]['path'],text=char*20000,expected_bytes=n*20000)) for n in range(1,6)])
        self.run_worker('small',lambda f:[('file_write',dict(path=f['outputs'][0]['path'],text='Short report'))])
        self.native('large')
        self.run_worker('review_large',lambda f:[('file_read',dict(path=i['path'],offset=n,limit=24000)) for i in f['inputs'] for n in range(0,(Path(f['workspace'])/i['path']).stat().st_size,24000)]+[('file_write',dict(path=f['outputs'][0]['path'],text='Exact synthetic joined text accepted'))],decision='accept')
        artifact=self.artifact('large','outputs/result.txt')
        size=Path(artifact['blob']).stat().st_size
        self.assertGreater(size,200000)
        refs=dict(report=artifact['id'],second=self.artifact('small','outputs/result.txt')['id'])
        s=self.session()
        with self.assertRaisesRegex(ValueError,'Original output limit'):
            s.call('job_write',dict(deliverable='report',text='y'*200001))
        original=s.grant['task_limits']['output_bytes']
        s.grant['task_limits']['output_bytes']=size
        with self.assertRaisesRegex(ValueError,'Combined outputs'):
            s.complete(dict(deliverables_json=json.dumps(refs),summary='Over delivery cap'))
        s.grant['task_limits']['output_bytes']=original
        result=s.complete(dict(deliverables_json=json.dumps(refs),summary='Exact existing worker outputs'))
        self.assertFalse(result['user_acceptance'])

    def test_multiple_owner_outputs_still_obey_combined_control_limit(self):
        s=self.session()
        refs={key:s.call('job_write',dict(deliverable=key,text='x'*110000))['artifact'] for key in self.deliverables}
        with self.assertRaisesRegex(ValueError,'Owner-authored outputs'):
            s.complete(dict(deliverables_json=json.dumps(refs),summary='Over owner control cap'))

class Gallery(unittest.TestCase):
    def test_saved_seven_node_workflow_has_no_review_and_exposes_all_photos(self):
        from orchestrator import node_spec
        w=json.loads(Path('examples/image-candidates.workflow.json').read_text())
        view=node_spec.resolve_workflow(w,artifact=lambda ident:{"media_type":"image/png"})
        self.assertFalse(view['ready'])
        self.assertEqual(len(w['tasks']),7)
        self.assertFalse(any(t['operation']=='review' for t in w['tasks']))
        self.assertEqual({t['operation'] for t in w['tasks'][-2:]},{'images.pptx','images.gallery'})
        for t in w['tasks'][-2:]:
            self.assertEqual(t['inputs']['photos'],{'task':'photos','output':'result'})
            self.assertEqual(t['inputs']['recommendations'],{'task':'recommendations','output':'report'})

    def test_all_candidates_escape_source_text_and_reject_changed_image(self):
        import io,zipfile
        from task_relay.image_gallery import render
        from orchestrator.image_sources import collect
        def fetch(url,limit,deadline):
            if 'api.php' in url:return json.dumps({'query':{'pages':{'1':source_page()}}}).encode(),'application/json',url
            return photo(),'image/jpeg',url
        raw,_=collect([subject('a'),subject('b'),subject('c')],fetch=fetch)
        page=render(raw).decode();self.assertEqual(page.count('<article data-candidate='),3)
        out=io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(raw)) as z,zipfile.ZipFile(out,'w') as changed:
            for name in z.namelist():changed.writestr(name,b'wrong' if name.endswith('.jpg') else z.read(name))
        with self.assertRaises(ValueError):render(out.getvalue())

    def test_pptx_keeps_all_candidates_exact_images_and_full_advice(self):
        import io
        from pptx import Presentation
        from task_relay.image_gallery import presentation
        from orchestrator.image_sources import collect
        def fetch(url,limit,deadline):
            if 'api.php' in url:return json.dumps({'query':{'pages':{'1':source_page()}}}).encode(),'application/json',url
            return photo(),'image/jpeg',url
        raw,_=collect([subject('a'),subject('b'),subject('c')],fetch=fetch)
        advice='I prefer candidate b; c is lower-ranked. Unknown buying details. '*150
        deck,receipt=presentation(raw,[('advice',advice)])
        slides=Presentation(io.BytesIO(deck)).slides
        pictures=[s for slide in slides for s in slide.shapes if s.shape_type==13]
        self.assertEqual(len(pictures),3)
        self.assertTrue(all(s.image.blob==photo() for s in pictures))
        text=''.join(s.text for slide in slides for s in slide.shapes if s.has_text_frame)
        self.assertNotIn('c is lower-ranked.',text)
        marker='Advisory text / unavailable records (continued in following slide notes if needed)\n'
        retained=''.join(slide.notes_slide.notes_text_frame.text.split(marker,1)[1] for slide in slides if marker in slide.notes_slide.notes_text_frame.text)
        self.assertIn(advice,retained)
        self.assertLessEqual(len(slides),2)
        self.assertEqual(receipt['candidates'],3)
        self.assertFalse(receipt['user_acceptance'])
        self.assertTrue(all('Source:' in slide.notes_slide.notes_text_frame.text for slide in slides if any(s.shape_type==13 for s in slide.shapes)))

    def test_named_cards_keep_long_labels_in_notes_and_report_after_gallery(self):
        import io,zipfile,hashlib
        from pptx import Presentation
        from task_relay import image_gallery,workflow_files
        from orchestrator.image_sources import collect
        def fetch(url,limit,deadline):
            if 'api.php' in url:return json.dumps({'query':{'pages':{'1':source_page()}}}).encode(),'application/json',url
            return photo(),'image/jpeg',url
        raw,_=collect([subject('a'),subject('b'),subject('c'),subject('d')],fetch=fetch)
        report='# Green wooden stool alternatives\n\nThe user chooses. Unknown buying details. <script>bad</script>'
        deck,receipt=image_gallery.presentation(raw,[('Advice',report)])
        view=Presentation(io.BytesIO(deck));self.assertEqual(len(view.slides),2)
        self.assertEqual(view.core_properties.title,'Green wooden stool alternatives')
        self.assertEqual([sum(s.shape_type==13 for s in slide.shapes) for slide in view.slides],[3,1])
        html=image_gallery.render(raw,[('Advice',report)]).decode()
        self.assertIn('<title>Green wooden stool alternatives</title>',html)
        self.assertLess(html.index('class="grid"'),html.index('<summary>Advice</summary>'))
        self.assertIn('&lt;script&gt;bad&lt;/script&gt;',html)
        import tempfile
        with tempfile.TemporaryDirectory() as folder:
            for extension,data in [('pptx',deck),('html',html.encode())]:
                p=Path(folder)/('result.'+extension);p.write_bytes(data)
                a=dict(path=p.name,blob=str(p),sha256=hashlib.sha256(data).hexdigest())
                self.assertEqual(workflow_files.display_name(a),'Green-wooden-stool-alternatives.'+extension)
                a['sha256']='0'*64
                self.assertEqual(workflow_files.display_name(a),p.name)
