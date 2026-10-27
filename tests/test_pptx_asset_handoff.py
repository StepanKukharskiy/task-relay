"""Exact pixel bindings survive node-local aliases; no provider or host calls."""
import copy
import hashlib
import io
import json
from pathlib import Path
import tempfile
import unittest

from orchestrator import layout_contract, pptx_document
from tests.test_automatic_composition import pixel


class Tests(unittest.TestCase):
    def fixture(self):
        raw=pixel('blue');path='inputs/producer/preview.png'
        doc=dict(version=1,title='Exact image',slides=[dict(notes='Retained notes',elements=[
            dict(type='image',path=path,x=1,y=1,w=6,h=4)])],
            asset_sha256={path:hashlib.sha256(raw).hexdigest()})
        return doc,raw,path

    def test_alias_uses_exact_pixels_without_mutating_saved_content(self):
        doc,raw,path=self.fixture();before=copy.deepcopy(doc)
        images={'inputs/consumer/preview.png':raw};original=images.copy()
        result,assets=pptx_document.prepare(doc,images)
        self.assertEqual(result,doc);self.assertEqual(doc,before)
        self.assertEqual(images,original);self.assertEqual(assets[path],raw)

    def test_changed_same_path_rejects_even_if_correct_bytes_exist_elsewhere(self):
        doc,raw,path=self.fixture()
        with self.assertRaisesRegex(ValueError,'image version changed'):
            pptx_document.prepare(doc,{path:pixel('red'),'elsewhere.png':raw})

    def test_missing_or_unbound_images_do_not_gain_authority_from_filenames(self):
        doc,raw,path=self.fixture()
        with self.assertRaisesRegex(ValueError,'image version changed'):
            pptx_document.prepare(doc,{'consumer/preview.png':pixel('red')})
        doc.pop('asset_sha256')
        with self.assertRaisesRegex(ValueError,'exact declared'):
            pptx_document.prepare(doc,{'consumer/preview.png':raw})
        with self.assertRaisesRegex(ValueError,'exact declared'):
            pptx_document.prepare(doc,{})

    def test_malformed_hash_and_traversal_reject(self):
        doc,raw,path=self.fixture()
        for binding in ({path:[]},{path:'not-a-hash'},[]):
            bad=copy.deepcopy(doc);bad['asset_sha256']=binding
            with self.assertRaises(ValueError):pptx_document.prepare(bad,{'new.png':raw})
        doc['slides'][0]['elements'][0]['path']='../outside.png'
        doc['asset_sha256']={'../outside.png':hashlib.sha256(raw).hexdigest()}
        with self.assertRaises(ValueError):pptx_document.prepare(doc,{'new.png':raw})

    def test_layout_projection_preflight_create_and_native_reopen(self):
        # Same boundary as a model-produced registered layout feeding a new node.
        from types import SimpleNamespace
        from pptx import Presentation
        raw=pixel('blue');old='inputs/author/preview.png';new='inputs/export/preview.png'
        params=dict(renderer='pptx.create',title='Composed handoff',style='clean_minimal_v1',max_records=1)
        records=[dict(kind='page',slot=1,layout='image_text',content=dict(title='Built form',image=old,caption='Exact native preview',body='Geometry evidence'),notes='Original citation'),dict(kind='end',count=1)]
        store=SimpleNamespace(contract={'parameters':params},context=lambda:{'images':{old:raw}})
        doc=layout_contract.projection(store,records);saved=json.dumps(doc).encode()
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);spec=root/'slides.json';pic=root/'preview.png'
            spec.write_bytes(saved);pic.write_bytes(raw)
            inputs=[dict(path='spec.json',media_type='application/json',blob=spec),dict(path=new,media_type='image/png',blob=pic)]
            prepared,assets=pptx_document.preflight_inputs(inputs)
            result,receipt=pptx_document.create(doc,{new:raw})
            self.assertTrue(receipt['reopened']);self.assertEqual(spec.read_bytes(),saved)
            deck=Presentation(io.BytesIO(result))
            pictures=[s for s in deck.slides[0].shapes if s.shape_type==13]
            self.assertEqual([s.image.blob for s in pictures],[raw])
            self.assertEqual(prepared,pptx_document.expand(doc));self.assertEqual(assets[old],raw)


if __name__=='__main__':unittest.main()
