"""Bounded source-name revision pilot for the selected motorcycle workbook.

The saved catalog description and URL are historical evidence. A changed source
name invalidates a translation claim for review; this module never invents a new
Russian name or repeats web research.
"""

from contextlib import closing
import hashlib
import io
import json
from pathlib import Path
import tempfile
import time
import uuid
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from orchestrator.storage import transaction
from . import fact_revisions as facts


SHEET = 'Пилот 10 позиций'
SUPPORTED = 'подтверждено'
UNRESOLVED = 'требует проверки'
SOURCE_KINDS = {'yamaha_txt', 'hd_xlsx'}
MAIN_NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PACKAGE_NS = 'http://schemas.openxmlformats.org/package/2006/relationships'


def _surgical_candidate(source, destination, patches, warning_style_id):
    """Edit named cells and their hyperlink relationship without rewriting other ZIP parts."""
    sheet_part = 'xl/worksheets/sheet1.xml'
    rel_part = 'xl/worksheets/_rels/sheet1.xml.rels'
    ET.register_namespace('', MAIN_NS)
    ET.register_namespace('r', REL_NS)
    with ZipFile(source) as original:
        sheet = ET.fromstring(original.read(sheet_part))
        cells = {cell.get('r'): cell for cell in sheet.findall('.//{' + MAIN_NS + '}c')}
        clear_links = set()
        for location, value in patches.items():
            sheet_name, address = location.rsplit('!', 1)
            if sheet_name != SHEET or address not in cells:
                raise ValueError('Candidate patch targets an unknown pilot cell.')
            cell = cells[address]
            for child in list(cell):
                cell.remove(child)
            if value is None:
                cell.attrib.pop('t', None)
                if address.startswith('G'):
                    clear_links.add(address)
            elif isinstance(value, int):
                cell.attrib.pop('t', None)
                ET.SubElement(cell, '{' + MAIN_NS + '}v').text = str(value)
            elif isinstance(value, str):
                cell.set('t', 'inlineStr')
                inline = ET.SubElement(cell, '{' + MAIN_NS + '}is')
                text = ET.SubElement(inline, '{' + MAIN_NS + '}t')
                text.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
                text.text = value
            else:
                raise ValueError('Unsupported candidate value type.')
            if address.startswith('H'):
                cell.set('s', str(warning_style_id))
        removed_ids = set()
        hyperlinks = sheet.find('{' + MAIN_NS + '}hyperlinks')
        if clear_links and hyperlinks is not None:
            for link in list(hyperlinks):
                if link.get('ref') in clear_links:
                    removed_ids.add(link.get('{' + REL_NS + '}id'))
                    hyperlinks.remove(link)
        replacements = {sheet_part: ET.tostring(sheet, encoding='utf-8', xml_declaration=True)}
        if removed_ids:
            relationships = ET.fromstring(original.read(rel_part))
            for relation in list(relationships):
                if relation.get('Id') in removed_ids:
                    relationships.remove(relation)
            replacements[rel_part] = ET.tostring(relationships, encoding='utf-8', xml_declaration=True)
        with ZipFile(destination, 'w') as candidate:
            for info in original.infolist():
                candidate.writestr(info, replacements.get(info.filename, original.read(info.filename)))
    with ZipFile(source) as old_zip, ZipFile(destination) as new_zip:
        for name in old_zip.namelist():
            if name not in replacements and old_zip.read(name) != new_zip.read(name):
                raise ValueError('Candidate changed an unrelated workbook part.')


def initialize(db):
    db.execute('''CREATE TABLE IF NOT EXISTS relay_translation_links(
        id TEXT PRIMARY KEY, job TEXT NOT NULL, source_artifact TEXT NOT NULL,
        source_sha256 TEXT NOT NULL, source_kind TEXT NOT NULL, source_row INTEGER NOT NULL,
        part_number TEXT NOT NULL, source_name TEXT NOT NULL, source_record_sha256 TEXT NOT NULL,
        output_artifact TEXT NOT NULL, output_sha256 TEXT NOT NULL, output_row INTEGER NOT NULL,
        russian_name TEXT, status TEXT NOT NULL, evidence_url TEXT,
        evidence_title TEXT, evidence_locator TEXT NOT NULL,
        review_report_artifact TEXT NOT NULL, review_report_sha256 TEXT NOT NULL,
        review_receipt TEXT NOT NULL, created REAL NOT NULL,
        UNIQUE(job,output_artifact,output_row))''')
    db.execute('''CREATE TABLE IF NOT EXISTS relay_translation_revisions(
        id TEXT PRIMARY KEY, job TEXT NOT NULL, baseline_artifact TEXT NOT NULL,
        candidate_artifact TEXT NOT NULL UNIQUE, old_source TEXT NOT NULL,
        replacement_source TEXT NOT NULL, plan_digest TEXT NOT NULL,
        patches TEXT NOT NULL, created REAL NOT NULL)''')
    db.execute('''CREATE TABLE IF NOT EXISTS relay_translation_reviews(
        id TEXT PRIMARY KEY, job TEXT NOT NULL, candidate_artifact TEXT NOT NULL UNIQUE,
        candidate_sha256 TEXT NOT NULL, decision TEXT NOT NULL, reviewer TEXT NOT NULL,
        note TEXT NOT NULL, checks TEXT NOT NULL, created REAL NOT NULL)''')
    db.execute('''CREATE TABLE IF NOT EXISTS relay_translation_selections(
        id TEXT PRIMARY KEY, job TEXT NOT NULL, candidate_artifact TEXT NOT NULL UNIQUE,
        candidate_sha256 TEXT NOT NULL, baseline_artifact TEXT NOT NULL,
        selected_by TEXT NOT NULL, receipt TEXT NOT NULL, created REAL NOT NULL)''')


def _artifact(db, ident, suffix):
    row = db.execute('SELECT * FROM production_artifacts WHERE id=?', (ident,)).fetchone()
    if row is None:
        raise ValueError('Unknown registered artifact.')
    row = dict(row)
    path = Path(row['blob'])
    if not row['path'].lower().endswith(suffix) or path.is_symlink() or not path.is_file():
        raise ValueError('Registered artifact format or file is unavailable.')
    if path.stat().st_size != row['bytes'] or row['bytes'] > 30_000_000 or hashlib.sha256(path.read_bytes()).hexdigest() != row['sha256']:
        raise ValueError('Registered artifact bytes changed or exceed the pilot limit.')
    return row


def _source_rows(artifact, kind):
    if kind == 'yamaha_txt':
        lines = Path(artifact['blob']).read_text(encoding='utf-8-sig').splitlines()
        for number, line in enumerate(lines, 1):
            fields = line.split(';')
            if len(fields) != 14:
                raise ValueError('Yamaha source row does not have 14 fields.')
            yield {'row': number, 'key': fields[1], 'name': fields[4], 'record': line}
    elif kind == 'hd_xlsx':
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(Path(artifact['blob']).read_bytes()),
                             read_only=True, data_only=False)
        try:
            if 'Лист1' not in book:
                raise ValueError('Harley source sheet is unavailable.')
            sheet = book['Лист1']
            if tuple(sheet.cell(1, col).value for col in (1, 2, 3)) != (
                    'Артикул детали', 'Наименование детали', 'Стоимость детали'):
                raise ValueError('Harley source columns changed.')
            for number, (key, name, price) in enumerate(
                    sheet.iter_rows(min_row=2, min_col=1, max_col=3, values_only=True), 2):
                record = {'A': str(key) if key is not None else '',
                          'B': name, 'C': str(price) if price is not None else ''}
                yield {'row': number, 'key': record['A'], 'name': name,
                       'record': json.dumps(record, ensure_ascii=False, sort_keys=True)}
        finally:
            book.close()
    else:
        raise ValueError('Unsupported source format for this pilot.')


def _record_digest(record):
    return hashlib.sha256(record.encode('utf-8')).hexdigest()


def _pilot_row(book, number):
    if SHEET not in book or number < 2 or number > book[SHEET].max_row:
        raise ValueError('Pilot output row is unavailable.')
    sheet = book[SHEET]
    return {column: sheet[f'{column}{number}'].value for column in
            ('A', 'B', 'C', 'D', 'F', 'G', 'H', 'K', 'O', 'P', 'Q', 'R', 'T')}


def _source_label(filename, kind, row):
    if kind == 'yamaha_txt':
        return filename + '; строка ' + str(row)
    if kind == 'hd_xlsx':
        return filename + '; Лист1; строка ' + str(row)
    raise ValueError('Unsupported source format for this pilot.')


def record_link(db, *, job, source_artifact, source_kind, output_artifact,
                output_row, review_report_artifact, review_receipt):
    """Record one exact row mapping from previously selected reviewed outputs."""
    if (source_kind not in SOURCE_KINDS or not isinstance(review_receipt, str)
            or not review_receipt.strip() or type(output_row) is not int):
        raise ValueError('Explicit source format, output row and selection receipt are required.')
    if not db.execute('SELECT 1 FROM relay_pipelines WHERE id=?', (job,)).fetchone():
        raise ValueError('Unknown job.')
    source = _artifact(db, source_artifact, '.txt' if source_kind == 'yamaha_txt' else '.xlsx')
    output = _artifact(db, output_artifact, '.xlsx')
    review = _artifact(db, review_report_artifact, '.md')
    with closing(facts._plain_xlsx(output['blob'])) as book:
        row = _pilot_row(book, output_row)
    brand = 'Yamaha' if source_kind == 'yamaha_txt' else 'Harley-Davidson'
    if row['B'] != brand or type(row['A']) is not int:
        raise ValueError('Pilot brand or source row mapping differs.')
    matches = [entry for entry in _source_rows(source, source_kind) if entry['row'] == row['A']]
    if len(matches) != 1 or matches[0]['key'] != str(row['C']) or matches[0]['name'] != row['D']:
        raise ValueError('Pilot part number or English name differs from exact source row.')
    found = matches[0]
    if row['K'] != _source_label(Path(source['path']).name, source_kind, found['row']):
        raise ValueError('Pilot source locator differs from the exact registered source row.')
    if source_kind == 'yamaha_txt':
        if row['P'] != found['record']:
            raise ValueError('Pilot full Yamaha source record differs.')
    elif json.loads(row['P']) != json.loads(found['record']):
        raise ValueError('Pilot full Harley source record differs.')
    if row['H'] == SUPPORTED:
        if not all(isinstance(row[c], str) and row[c].strip() for c in ('F','G','Q','T')):
            raise ValueError('Supported translation lacks its saved name or evidence locator.')
    elif row['H'] == UNRESOLVED:
        if row['F'] is not None or row['G'] is not None:
            raise ValueError('Unresolved translation must remain blank.')
    else:
        raise ValueError('Unexpected pilot review state.')
    ident = uuid.uuid4().hex
    with transaction(db):
        db.execute('''INSERT INTO relay_translation_links VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',
            (ident, job, source_artifact, source['sha256'], source_kind, found['row'],
             found['key'], found['name'], _record_digest(found['record']),
             output_artifact, output['sha256'], output_row, row['F'], row['H'],
             row['G'], row['Q'], row['T'] or '', review_report_artifact,
             review['sha256'], review_receipt, time.time()))
    return ident


def records(db, job):
    def saved(table):
        if not db.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name=?",
                          (table,)).fetchone():
            return []
        return [dict(r) for r in db.execute('SELECT * FROM ' + table +
                                            ' WHERE job=? ORDER BY created,id', (job,))]
    return {'links': [dict(r) for r in db.execute(
        'SELECT * FROM relay_translation_links WHERE job=? ORDER BY output_row', (job,))],
        'revisions': [dict(r) for r in db.execute(
        'SELECT * FROM relay_translation_revisions WHERE job=? ORDER BY created,id', (job,))],
        'reviews': saved('relay_translation_reviews'),
        'selections': saved('relay_translation_selections')}


def plan_impact(db, *, job, old_source, replacement_source, baseline_artifact):
    """Read-only translation impact, bounded to exact selected pilot rows."""
    source_row = db.execute('SELECT path FROM production_artifacts WHERE id=?',
                            (old_source,)).fetchone()
    if source_row is None:
        raise ValueError('Unknown registered source artifact.')
    old = _artifact(db, old_source,
                    '.txt' if source_row['path'].lower().endswith('.txt') else '.xlsx')
    replacement = _artifact(db, replacement_source,
        '.txt' if old['path'].lower().endswith('.txt') else '.xlsx')
    baseline = _artifact(db, baseline_artifact, '.xlsx')
    if old_source == replacement_source:
        raise ValueError('A distinct source version is required.')
    links = [r for r in records(db, job)['links'] if r['output_artifact'] == baseline_artifact]
    if not links:
        raise ValueError('No reviewed pilot row links for this output.')
    relevant = [r for r in links if r['source_artifact'] == old_source]
    if not relevant:
        raise ValueError('The replaced source is not linked to this pilot.')
    replacement_kind = relevant[0]['source_kind']
    new_by_key = {}
    for entry in _source_rows(replacement, replacement_kind):
        new_by_key.setdefault(entry['key'], []).append(entry)
    affected, unaffected, unknown = [], [], []
    for link in links:
        location = f'{SHEET}!F{link["output_row"]}'
        item = {'part_number': link['part_number'], 'translation_cell': location,
                'pilot_row': link['output_row']}
        try:
            source = _artifact(db, link['source_artifact'],
                '.txt' if link['source_kind'] == 'yamaha_txt' else '.xlsx')
            report = _artifact(db, link['review_report_artifact'], '.md')
            if (source['sha256'] != link['source_sha256'] or
                    baseline['sha256'] != link['output_sha256'] or
                    report['sha256'] != link['review_report_sha256']):
                raise ValueError('Recorded source, output or review report version differs.')
            original = [entry for entry in _source_rows(source, link['source_kind'])
                        if entry['row'] == link['source_row']]
            if (len(original) != 1 or original[0]['key'] != link['part_number'] or
                    original[0]['name'] != link['source_name'] or
                    _record_digest(original[0]['record']) != link['source_record_sha256']):
                raise ValueError('Exact historical source row differs.')
            if link['source_artifact'] == old_source:
                current = new_by_key.get(link['part_number'], [])
                if len(current) != 1:
                    raise ValueError('Replacement part number is missing or duplicated.')
                if not isinstance(current[0]['name'], str) or not current[0]['name'].strip():
                    raise ValueError('Replacement English name is missing.')
                if current[0]['name'] != link['source_name']:
                    affected.append({**item, 'before': link['source_name'],
                        'after': current[0]['name'], 'new_source_row': current[0]['row'],
                        'new_source_record': current[0]['record'],
                        'reason': 'English source name changed; Russian name requires review.',
                        'saved_evidence_url': link['evidence_url'],
                        'saved_catalog_title': link['evidence_title']})
                    continue
                if _record_digest(current[0]['record']) != link['source_record_sha256']:
                    raise ValueError('Replacement source row changed outside the English name.')
                item.update(new_source_row=current[0]['row'])
            if link['status'] == SUPPORTED:
                unaffected.append({**item, 'reason': 'Exact source name and selected catalog evidence are unchanged.'})
            else:
                unknown.append({**item, 'reason': 'Original pilot left the Russian name unverified.'})
        except (ValueError, OSError) as exc:
            unknown.append({**item, 'reason': str(exc)})
    result = {'schema': 'task-relay.translation-impact', 'version': 2,
              'job': job, 'old_source': old_source, 'replacement_source': replacement_source,
              'replacement_filename': Path(replacement['path']).name,
              'replacement_kind': replacement_kind,
              'baseline_artifact': baseline_artifact, 'old_sha256': old['sha256'],
              'replacement_sha256': replacement['sha256'], 'baseline_sha256': baseline['sha256'],
              'affected': affected, 'unaffected': unaffected, 'unknown': unknown,
              'coverage': 'Only ten explicitly linked selected pilot rows; saved catalog evidence is historical.',
              'authorization': 'Planning only; no translation, review acceptance or dispatch.'}
    result['digest'] = hashlib.sha256(json.dumps(result, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
    from . import reviewed_links
    result['reviewed_impact'] = reviewed_links.impact_record(
        reviewed_links.records(db, job), kind='parts_translation', plan=result,
        source_sha256=old['sha256'], replacement_sha256=replacement['sha256'])
    return result


def _candidate_patches(db, plan, sheet):
    patches = {}
    for item in (*plan['unaffected'], *plan['unknown']):
        if 'new_source_row' not in item:
            continue
        row = item['pilot_row']
        if sheet[f'C{row}'].value != item['part_number']:
            raise ValueError('Pilot row identity changed before source rebinding.')
        source_row_location = f'{SHEET}!A{row}'
        if sheet[f'A{row}'].value != item['new_source_row']:
            patches[source_row_location] = item['new_source_row']
        location = f'{SHEET}!K{row}'
        label = _source_label(plan['replacement_filename'], plan['replacement_kind'],
                              item['new_source_row'])
        if sheet[f'K{row}'].value != label:
            patches[location] = label
    for item in plan['affected']:
        row = item['pilot_row']
        if sheet[f'C{row}'].value != item['part_number']:
            raise ValueError('Pilot row identity changed before candidate review.')
        new_name = item['after']
        note = ('Source English name changed from ' + repr(item['before']) +
                ' to ' + repr(new_name) + '; saved catalog title ' +
                repr(item['saved_catalog_title']) + ' requires review for this exact part.')
        changes = {'A': item['new_source_row'], 'D': new_name, 'F': None, 'G': None,
                   'H': UNRESOLVED, 'O': note, 'P': item['new_source_record'],
                   'K': _source_label(plan['replacement_filename'], plan['replacement_kind'],
                                      item['new_source_row'])}
        for column, value in changes.items():
            cell = f'{column}{row}'
            if sheet[cell].value != value:
                patches[f'{SHEET}!{cell}'] = value
    return patches


def _checked_cells(before_book, after_book, plan, patches):
    before, after = facts._cell_state(before_book), facts._cell_state(after_book)
    sheet = before_book[SHEET]
    warning_cells = [sheet[f'H{row}'] for row in range(2, sheet.max_row + 1)
                     if sheet[f'H{row}'].value == UNRESOLVED]
    if not warning_cells:
        raise ValueError('Pilot has no reviewed unresolved status style.')
    warning_style = (warning_cells[0]._style if warning_cells[0]._style
                     and any(warning_cells[0]._style) else None)
    def preserved(key, old):
        location = key[0] + '!' + key[1]
        if location not in patches:
            return after[key] == old
        expected_link = None if key[1].startswith('G') else old[3]
        expected_style = warning_style if key[1].startswith('H') else old[2]
        return (after[key][0] == patches[location] and after[key][2] == expected_style
                and after[key][3] == expected_link and after[key][1] != 'f')
    mismatches = [key for key, old in before.items()
                  if key not in after or not preserved(key, old)]
    if before.keys() != after.keys() or mismatches:
        raise ValueError('Translation candidate changed an unrelated value, formula or style: '
                         + repr(mismatches[:10]))
    return {'changed_cells': sorted(patches),
            'untouched_cells_verified': len(before) - len(patches)}


def create_candidate(rt, *, plan):
    """Clear changed translations for review; preserve other pilot cells exactly."""
    fresh = plan_impact(rt.db, job=plan['job'], old_source=plan['old_source'],
        replacement_source=plan['replacement_source'], baseline_artifact=plan['baseline_artifact'])
    if fresh != plan or not plan['affected']:
        raise ValueError('Translation impact plan is stale or has no changed source name.')
    baseline = _artifact(rt.db, plan['baseline_artifact'], '.xlsx')
    with closing(facts._plain_xlsx(baseline['blob'])) as book:
        sheet = book[SHEET]
        warning_cells = [sheet[f'H{row}'] for row in range(2, sheet.max_row + 1)
                         if sheet[f'H{row}'].value == UNRESOLVED]
        if not warning_cells:
            raise ValueError('Pilot has no reviewed unresolved status style.')
        warning_style_id = warning_cells[0].style_id
        patches = _candidate_patches(rt.db, plan, sheet)
        with tempfile.TemporaryDirectory(dir=rt.root) as folder:
            path = Path(folder) / 'pilot_translation_revision.xlsx'
            _surgical_candidate(baseline['blob'], path, patches, warning_style_id)
            with closing(facts._plain_xlsx(path)) as checked:
                _checked_cells(book, checked, plan, patches)
            with transaction(rt.db):
                if plan_impact(rt.db, job=plan['job'], old_source=plan['old_source'],
                    replacement_source=plan['replacement_source'],
                    baseline_artifact=plan['baseline_artifact']) != plan:
                    raise ValueError('Translation impact changed before candidate registration.')
                candidate = rt.register(path, 'Unreviewed motorcycle translation revision candidate',
                    task='translation_revision', path='delivery/pilot_translation_revision.xlsx')
                rt.db.execute('''INSERT INTO relay_translation_revisions VALUES (?,?,?,?,?,?,?,?,?)''',
                    (uuid.uuid4().hex, plan['job'], plan['baseline_artifact'], candidate,
                     plan['old_source'], plan['replacement_source'], plan['digest'],
                     json.dumps(patches, sort_keys=True, ensure_ascii=False), time.time()))
                from . import reviewed_links
                reviewed_links.save_impact(rt.db, candidate_artifact=candidate, plan=plan)
    return candidate


def verified_candidate(db, candidate_artifact):
    """Recheck the exact registered candidate before showing or deciding it."""
    revision = db.execute('SELECT * FROM relay_translation_revisions WHERE candidate_artifact=?',
                          (candidate_artifact,)).fetchone()
    if revision is None:
        raise ValueError('Unknown translation revision candidate.')
    revision = dict(revision)
    candidate = _artifact(db, candidate_artifact, '.xlsx')
    baseline = _artifact(db, revision['baseline_artifact'], '.xlsx')
    plan = plan_impact(db, job=revision['job'], old_source=revision['old_source'],
        replacement_source=revision['replacement_source'],
        baseline_artifact=revision['baseline_artifact'])
    if plan['digest'] != revision['plan_digest']:
        raise ValueError('Candidate plan is stale or predates source-locator rebinding; create a fresh revision.')
    from . import reviewed_links
    reviewed_links.verify_saved_impact(db, candidate_artifact=candidate_artifact, plan=plan)
    patches = json.loads(revision['patches'])
    with closing(facts._plain_xlsx(baseline['blob'])) as before_book, \
            closing(facts._plain_xlsx(candidate['blob'])) as after_book:
        if patches != _candidate_patches(db, plan, before_book[SHEET]):
            raise ValueError('Candidate patches differ from the current impact plan.')
        checks = _checked_cells(before_book, after_book, plan, patches)
    with ZipFile(baseline['blob']) as old_zip, ZipFile(candidate['blob']) as new_zip:
        if old_zip.namelist() != new_zip.namelist() or any(
                old_zip.read(name) != new_zip.read(name)
                for name in old_zip.namelist()
                if name not in ('xl/worksheets/sheet1.xml',
                                'xl/worksheets/_rels/sheet1.xml.rels')):
            raise ValueError('Candidate changed an unrelated workbook part.')
    return revision, candidate, plan, {
        **checks, 'baseline_sha256': baseline['sha256'],
        'candidate_sha256': candidate['sha256'], 'plan_digest': plan['digest']}


def review_candidate(db, *, candidate_artifact, decision, reviewer, note):
    """Record a human review of the exact candidate; no selection is inferred."""
    if decision not in ('accept', 'revise') or not all(
            isinstance(value, str) and value.strip() for value in (reviewer, note)):
        raise ValueError('An explicit reviewer, accept/revise decision and note are required.')
    revision, candidate, _, checks = verified_candidate(db, candidate_artifact)
    ident = uuid.uuid4().hex
    with transaction(db):
        if db.execute('SELECT 1 FROM relay_translation_reviews WHERE candidate_artifact=?',
                      (candidate_artifact,)).fetchone():
            raise ValueError('Candidate already reviewed; create a new version to revise it.')
        db.execute('''INSERT INTO relay_translation_reviews VALUES (?,?,?,?,?,?,?,?,?)''',
                   (ident, revision['job'], candidate_artifact, candidate['sha256'],
                    decision, reviewer.strip(), note.strip(),
                    json.dumps(checks, sort_keys=True), time.time()))
    return ident


def select_candidate(db, *, candidate_artifact, selected_by, receipt):
    """Record an explicit exact-version selection after acceptance; never dispatch."""
    if not all(isinstance(value, str) and value.strip()
               for value in (selected_by, receipt)):
        raise ValueError('Selection needs the actor and exact action receipt.')
    ident = uuid.uuid4().hex
    with transaction(db):
        revision, candidate, plan, _ = verified_candidate(db, candidate_artifact)
        review = db.execute('SELECT * FROM relay_translation_reviews WHERE candidate_artifact=?',
                            (candidate_artifact,)).fetchone()
        if (review is None or review['decision'] != 'accept'
                or review['candidate_sha256'] != candidate['sha256']):
            raise ValueError('Select only the exact accepted candidate.')
        if db.execute('SELECT 1 FROM relay_translation_selections WHERE candidate_artifact=?',
                      (candidate_artifact,)).fetchone():
            raise ValueError('This exact candidate was already selected.')
        prior_links = [link for link in records(db, revision['job'])['links']
                       if link['output_artifact'] == revision['baseline_artifact']]
        outcomes = {item['pilot_row']: (status, item)
                    for status in ('affected', 'unaffected', 'unknown')
                    for item in plan[status]}
        if (len(prior_links) != len(outcomes)
                or {link['output_row'] for link in prior_links} != set(outcomes)):
            raise ValueError('Selected candidate lacks complete reviewed row lineage.')
        for link in prior_links:
            _, item = outcomes[link['output_row']]
            if (link['source_artifact'] == revision['old_source']
                    and 'new_source_row' not in item
                    and link['status'] == SUPPORTED):
                raise ValueError('A supported translation has an unresolved replacement source row.')
        db.execute('''INSERT INTO relay_translation_selections VALUES (?,?,?,?,?,?,?,?)''',
                   (ident, revision['job'], candidate_artifact, candidate['sha256'],
                    revision['baseline_artifact'], selected_by.strip(),
                    receipt.strip(), time.time()))
        for link in prior_links:
            status, item = outcomes[link['output_row']]
            moved = (link['source_artifact'] == revision['old_source']
                     and 'new_source_row' in item)
            source_artifact = (revision['replacement_source'] if moved
                               else link['source_artifact'])
            lineage = json.dumps({
                'schema': 'task-relay.translation-link-continuation', 'version': 1,
                'selection_id': ident, 'parent_link_id': link['id'],
                'candidate_sha256': candidate['sha256'],
                'impact_status': status,
                'source_binding': 'replacement' if moved else 'historical',
            }, sort_keys=True)
            record_link(db, job=revision['job'], source_artifact=source_artifact,
                source_kind=link['source_kind'], output_artifact=candidate_artifact,
                output_row=link['output_row'],
                review_report_artifact=link['review_report_artifact'],
                review_receipt=lineage)
    return ident
