"""Reviewed XLSX fact bindings and conservative, read-only revision plans.

This is the bounded O14 material-cell pilot. It never infers a binding from a
declared workflow input and never starts a worker from an impact report.
"""

from collections import defaultdict
from contextlib import closing
import hashlib
import json
from pathlib import Path
import re
import tempfile
import time
import uuid
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from orchestrator.storage import transaction


SCHEMA = 'task-relay.reviewed-facts'
VERSION = 1
CELL = re.compile(r'[A-Z]{1,3}[1-9][0-9]{0,6}\Z')
COLUMN = re.compile(r'[A-Z]{1,3}\Z')
ALLOWED_PARTS = re.compile(
    r'(?:\[Content_Types\]\.xml|_rels/\.rels|docProps/(?:app|core)\.xml|'
    r'xl/(?:workbook\.xml|styles\.xml|sharedStrings\.xml|theme/theme1\.xml|'
    r'_rels/workbook\.xml\.rels|worksheets/sheet[1-9][0-9]*\.xml|'
    r'worksheets/_rels/sheet[1-9][0-9]*\.xml\.rels))\Z')


def initialize(db):
    db.execute('''CREATE TABLE IF NOT EXISTS relay_fact_bindings(
        id TEXT PRIMARY KEY, job TEXT NOT NULL, entity_type TEXT NOT NULL,
        entity_key TEXT NOT NULL, predicate TEXT NOT NULL, value TEXT NOT NULL,
        source_artifact TEXT NOT NULL, source_sha256 TEXT NOT NULL,
        source_sheet TEXT NOT NULL, source_key_column TEXT NOT NULL,
        source_cell TEXT NOT NULL, evidence TEXT NOT NULL,
        review_state TEXT NOT NULL, reviewer TEXT NOT NULL, review_note TEXT NOT NULL,
        output_artifact TEXT NOT NULL, output_sha256 TEXT NOT NULL,
        output_sheet TEXT NOT NULL, output_cell TEXT NOT NULL,
        check_kind TEXT NOT NULL, created REAL NOT NULL)''')
    db.execute('''CREATE TABLE IF NOT EXISTS relay_fact_coverage(
        job TEXT NOT NULL, output_artifact TEXT NOT NULL, output_sheet TEXT NOT NULL,
        output_cell TEXT NOT NULL, state TEXT NOT NULL, note TEXT NOT NULL,
        PRIMARY KEY(job,output_artifact,output_sheet,output_cell))''')
    db.execute('''CREATE TABLE IF NOT EXISTS relay_fact_revisions(
        id TEXT PRIMARY KEY, job TEXT NOT NULL, baseline_artifact TEXT NOT NULL,
        candidate_artifact TEXT NOT NULL UNIQUE, old_source TEXT NOT NULL,
        replacement_source TEXT NOT NULL, plan_digest TEXT NOT NULL,
        patches TEXT NOT NULL, reviewer TEXT NOT NULL, created REAL NOT NULL)''')
    db.execute('''CREATE TABLE IF NOT EXISTS relay_fact_candidate_reviews(
        id TEXT PRIMARY KEY, job TEXT NOT NULL, candidate_artifact TEXT NOT NULL UNIQUE,
        candidate_sha256 TEXT NOT NULL, decision TEXT NOT NULL, reviewer TEXT NOT NULL,
        note TEXT NOT NULL, checks TEXT NOT NULL, created REAL NOT NULL)''')
    db.execute('''CREATE TABLE IF NOT EXISTS relay_fact_selections(
        id TEXT PRIMARY KEY, job TEXT NOT NULL, candidate_artifact TEXT NOT NULL UNIQUE,
        candidate_sha256 TEXT NOT NULL, baseline_artifact TEXT NOT NULL,
        selected_by TEXT NOT NULL, receipt TEXT NOT NULL, created REAL NOT NULL)''')
    db.execute('''CREATE TABLE IF NOT EXISTS relay_fact_external_submissions(
        id TEXT PRIMARY KEY, job TEXT NOT NULL, submission_key TEXT NOT NULL,
        handoff_id TEXT NOT NULL, candidate_artifact TEXT NOT NULL UNIQUE,
        candidate_sha256 TEXT NOT NULL, checks TEXT NOT NULL,
        submitted_by TEXT NOT NULL, created REAL NOT NULL,
        UNIQUE(job,submission_key))''')


def _artifact(db, ident):
    row = db.execute('SELECT * FROM production_artifacts WHERE id=?', (ident,)).fetchone()
    if row is None:
        raise ValueError('Unknown registered artifact: ' + str(ident))
    row = dict(row)
    path = Path(row['blob'])
    if not path.is_file() or path.is_symlink() or path.stat().st_size != row['bytes']:
        raise ValueError('Registered artifact is missing or changed: ' + ident)
    if row['bytes'] > 30_000_000:
        raise ValueError('XLSX exceeds the bounded fact revision pilot.')
    if hashlib.sha256(path.read_bytes()).hexdigest() != row['sha256']:
        raise ValueError('Registered artifact hash mismatch: ' + ident)
    if path.suffix.lower() != '.xlsx' and not row['path'].lower().endswith('.xlsx'):
        raise ValueError('The fact revision pilot accepts XLSX artifacts only.')
    return row


def _plain_xlsx(path):
    with ZipFile(path) as archive:
        names = archive.namelist()
        if (len(names) != len(set(names)) or
                sum(info.file_size for info in archive.infolist()) > 50_000_000 or
                any(not ALLOWED_PARTS.fullmatch(name) for name in names)):
            raise ValueError('Unsupported XLSX feature or archive part; preserve this workbook without editing.')
        for name in names:
            if name.startswith('xl/worksheets/_rels/'):
                root = ET.fromstring(archive.read(name))
                for relation in root:
                    if (relation.get('Type') !=
                            'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink'
                            or relation.get('TargetMode') != 'External'
                            or not relation.get('Target', '').startswith('https://')):
                        raise ValueError('Unsupported XLSX relationship; preserve this workbook without editing.')
    from openpyxl import load_workbook
    # Registered blobs are named "content"; pass a stream so openpyxl checks the
    # archive itself rather than inferring the format from that storage name.
    with Path(path).open('rb') as stream:
        book = load_workbook(stream, data_only=False, read_only=False, keep_links=False)
    try:
        for sheet in book:
            if (sheet._charts or sheet._images or sheet.tables or sheet.merged_cells.ranges
                    or sheet.conditional_formatting or sheet.data_validations.dataValidation
                    or sheet._pivots or sheet.max_row * sheet.max_column > 10000):
                raise ValueError('Unsupported XLSX feature; preserve this workbook without editing.')
        if book.defined_names or len(book.sheetnames) > 10:
            raise ValueError('Unsupported XLSX feature; preserve this workbook without editing.')
    except BaseException:
        book.close()
        raise
    return book


def _location(sheet, cell):
    if not isinstance(sheet, str) or not sheet or len(sheet) > 31 or not CELL.fullmatch(cell):
        raise ValueError('An exact XLSX sheet and cell are required.')
    return sheet + '!' + cell


def _value(book, sheet, cell):
    if sheet not in book:
        raise ValueError('Recorded sheet is missing: ' + sheet)
    return book[sheet][cell].value


def _matching_rows(book, sheet, key_column, key):
    if sheet not in book or not COLUMN.fullmatch(key_column):
        raise ValueError('Recorded product-key location is invalid.')
    ws = book[sheet]
    return [row for row in range(2, ws.max_row + 1) if ws[f'{key_column}{row}'].value == key]


def record_binding(db, *, job, entity_key, predicate, value, source_artifact,
                   source_sheet, source_key_column, source_cell, evidence,
                   output_artifact, output_sheet, output_cell, reviewer,
                   review_state, review_note='', entity_type='product'):
    """Save an explicit source fact and exact output binding in one transaction."""
    _location(source_sheet, source_cell)
    _location(output_sheet, output_cell)
    if (not all(isinstance(v, str) and v.strip() for v in
                (job, entity_type, entity_key, predicate, value, evidence, reviewer))
            or review_state not in ('reviewed', 'pending')
            or not COLUMN.fullmatch(source_key_column)):
        raise ValueError('Fact, evidence, product key and explicit review state are required.')
    if db.execute('SELECT 1 FROM relay_pipelines WHERE id=?', (job,)).fetchone() is None:
        raise ValueError('Unknown job.')
    source, output = _artifact(db, source_artifact), _artifact(db, output_artifact)
    with closing(_plain_xlsx(source['blob'])) as sb, closing(_plain_xlsx(output['blob'])) as ob:
        rows = _matching_rows(sb, source_sheet, source_key_column, entity_key)
        if rows != [sb[source_sheet][source_cell].row]:
            raise ValueError('Source product key is missing, duplicated or does not match the source cell.')
        if _value(sb, source_sheet, source_cell) != value:
            raise ValueError('Recorded fact differs from the exact source cell.')
        if _value(ob, output_sheet, output_cell) != value:
            raise ValueError('Output check failed: exact cell differs from the reviewed fact.')
    ident = uuid.uuid4().hex
    with transaction(db):
        db.execute('''INSERT INTO relay_fact_bindings VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',
                   (ident, job, entity_type, entity_key, predicate, json.dumps(value),
                    source_artifact, source['sha256'], source_sheet, source_key_column,
                    source_cell, evidence, review_state, reviewer, review_note,
                    output_artifact, output['sha256'], output_sheet, output_cell,
                    'exact_value', time.time()))
    return ident


def record_coverage(db, *, job, output_artifact, output_sheet, output_cell, state, note):
    """Declare the review boundary; complete requires a consistent reviewed link."""
    _location(output_sheet, output_cell)
    if state not in ('complete', 'incomplete') or not isinstance(note, str) or not note.strip():
        raise ValueError('Coverage requires complete/incomplete and an explanation.')
    _artifact(db, output_artifact)
    bindings = [dict(r) for r in db.execute('''SELECT * FROM relay_fact_bindings WHERE
        job=? AND output_artifact=? AND output_sheet=? AND output_cell=?''',
        (job, output_artifact, output_sheet, output_cell))]
    if state == 'complete' and (not bindings or any(b['review_state'] != 'reviewed' for b in bindings)
                                or len({b['value'] for b in bindings}) != 1):
        raise ValueError('Complete coverage requires consistent reviewed bindings.')
    with transaction(db):
        db.execute('''INSERT INTO relay_fact_coverage VALUES (?,?,?,?,?,?)
            ON CONFLICT(job,output_artifact,output_sheet,output_cell)
            DO UPDATE SET state=excluded.state,note=excluded.note''',
            (job, output_artifact, output_sheet, output_cell, state, note))


def records(db, job):
    """Read exact committed records for a job projection; no inference."""
    return {
        'bindings': [dict(r) for r in db.execute(
            'SELECT * FROM relay_fact_bindings WHERE job=? ORDER BY created,id', (job,))],
        'coverage': [dict(r) for r in db.execute(
            'SELECT * FROM relay_fact_coverage WHERE job=? ORDER BY output_artifact,output_sheet,output_cell', (job,))],
        'revisions': [dict(r) for r in db.execute(
            'SELECT * FROM relay_fact_revisions WHERE job=? ORDER BY created,id', (job,))],
        'candidate_reviews': [dict(r) for r in db.execute(
            'SELECT * FROM relay_fact_candidate_reviews WHERE job=? ORDER BY created,id', (job,))],
        'selections': [dict(r) for r in db.execute(
            'SELECT * FROM relay_fact_selections WHERE job=? ORDER BY rowid', (job,))],
        'external_submissions': [dict(r) for r in db.execute(
            'SELECT * FROM relay_fact_external_submissions WHERE job=? ORDER BY created,id', (job,))],
    }


def plan_impact(db, *, job, old_source, replacement_source, baseline_artifact):
    """Classify only declared output locations. This read never revises a file."""
    old = _artifact(db, old_source)
    replacement = _artifact(db, replacement_source)
    baseline = _artifact(db, baseline_artifact)
    if old_source == replacement_source:
        raise ValueError('Replacement must be a distinct registered artifact version.')
    rows = records(db, job)
    coverage = [c for c in rows['coverage'] if c['output_artifact'] == baseline_artifact]
    if not coverage:
        raise ValueError('No declared output coverage for this job and artifact.')
    bindings = defaultdict(list)
    for b in rows['bindings']:
        if b['output_artifact'] == baseline_artifact:
            bindings[(b['output_sheet'], b['output_cell'])].append(b)
    affected, unaffected, unknown = [], [], []
    with closing(_plain_xlsx(replacement['blob'])) as new_book, closing(_plain_xlsx(old['blob'])) as old_book:
        for c in coverage:
            name = _location(c['output_sheet'], c['output_cell'])
            related = bindings[(c['output_sheet'], c['output_cell'])]
            item = {'location': name, 'output_artifact': baseline_artifact}
            if c['state'] != 'complete' or not related:
                unknown.append({**item, 'reason': 'Incomplete reviewed coverage: ' + c['note']})
                continue
            if any(b['review_state'] != 'reviewed' for b in related):
                unknown.append({**item, 'reason': 'Unreviewed binding.'})
                continue
            if len({b['value'] for b in related}) != 1:
                unknown.append({**item, 'reason': 'Conflicting reviewed facts.'})
                continue
            changes = []
            gap = None
            for b in related:
                try:
                    source = _artifact(db, b['source_artifact'])
                    if source['sha256'] != b['source_sha256'] or baseline['sha256'] != b['output_sha256']:
                        raise ValueError('Recorded artifact hash changed.')
                    if b['source_artifact'] == old_source:
                        column = re.match(r'[A-Z]+', b['source_cell']).group()
                        if (_value(old_book, b['source_sheet'], f'{column}1') !=
                                _value(new_book, b['source_sheet'], f'{column}1') or
                                _value(old_book, b['source_sheet'], f'{b["source_key_column"]}1') !=
                                _value(new_book, b['source_sheet'], f'{b["source_key_column"]}1')):
                            raise ValueError('Replacement source columns changed; fact location is ambiguous.')
                        matches = _matching_rows(new_book, b['source_sheet'], b['source_key_column'], b['entity_key'])
                        if len(matches) != 1:
                            raise ValueError('Replacement product key is missing or ambiguous.')
                        new_value = _value(new_book, b['source_sheet'], f'{column}{matches[0]}')
                        if not isinstance(new_value, str) or not new_value.strip():
                            raise ValueError('Replacement fact is missing or unsupported.')
                        changes.append((json.loads(b['value']), new_value))
                    else:
                        with closing(_plain_xlsx(source['blob'])) as other:
                            if _value(other, b['source_sheet'], b['source_cell']) != json.loads(b['value']):
                                raise ValueError('Independent source no longer matches the reviewed fact.')
                except (ValueError, OSError) as exc:
                    gap = str(exc)
                    break
            if gap:
                unknown.append({**item, 'reason': gap})
            elif len({json.dumps(new) for _, new in changes}) > 1:
                unknown.append({**item, 'reason': 'Conflicting replacement facts.'})
            elif changes and any(b['source_artifact'] != old_source
                                 and json.loads(b['value']) != changes[0][1] for b in related):
                unknown.append({**item, 'reason': 'Replacement conflicts with another reviewed source.'})
            elif changes and any(before != after for before, after in changes):
                affected.append({**item, 'before': changes[0][0], 'after': changes[0][1],
                                 'reason': 'Reviewed source fact changed.'})
            else:
                unaffected.append({**item, 'reason': 'Complete reviewed bindings remain valid.'})
    result = {'schema': 'task-relay.fact-impact', 'version': 1, 'job': job,
              'old_source': old_source, 'replacement_source': replacement_source,
              'baseline_artifact': baseline_artifact,
              'source_sha256': old['sha256'], 'replacement_sha256': replacement['sha256'],
              'baseline_sha256': baseline['sha256'], 'affected': affected,
              'unaffected': unaffected, 'unknown': unknown,
              'coverage_scope': 'Declared output cells only; all other workbook locations are unassessed.',
              'authorization': 'Planning only; no rebuild, selection or dispatch.'}
    result['digest'] = hashlib.sha256(json.dumps(result, sort_keys=True).encode()).hexdigest()
    from . import reviewed_links
    result['reviewed_impact'] = reviewed_links.impact_record(
        reviewed_links.records(db, job), kind='material_cell', plan=result,
        source_sha256=old['sha256'], replacement_sha256=replacement['sha256'])
    return result


def _cell_state(book):
    return {(ws.title, cell.coordinate): (cell.value, cell.data_type,
            cell._style if cell._style and any(cell._style) else None,
            cell.hyperlink.target if cell.hyperlink else None)
            for ws in book for row in ws for cell in row}


def _external_checks(db, handoff_id, candidate_path):
    from . import impact_handoff
    handoff = impact_handoff.inspect(db, handoff_id)
    if handoff['kind'] != impact_handoff.XLSX_FACT_CHANGE or handoff['status'] != 'current':
        raise ValueError('External XLSX candidate needs a current frozen XLSX handoff.')
    plan = handoff['plan']
    baseline = _artifact(db, plan['baseline_artifact'])
    candidate_path = Path(candidate_path)
    if candidate_path.is_symlink() or not candidate_path.is_file():
        raise ValueError('External XLSX candidate must be a regular file.')
    if candidate_path.stat().st_size > 30_000_000:
        raise ValueError('External XLSX candidate exceeds the bounded pilot.')
    expected = {item['location']: item['after'] for item in plan['affected']}
    if not expected:
        raise ValueError('Frozen XLSX handoff has no affected location to revise.')
    with closing(_plain_xlsx(baseline['blob'])) as old_book, closing(_plain_xlsx(candidate_path)) as new_book:
        if (old_book.sheetnames != new_book.sheetnames
                or any(old_book[name].sheet_state != new_book[name].sheet_state
                       for name in old_book.sheetnames)):
            raise ValueError('External XLSX candidate changed sheet identity or visibility.')
        before, after = _cell_state(old_book), _cell_state(new_book)
        if before.keys() != after.keys():
            raise ValueError('External XLSX candidate changed workbook cell locations.')
        changed = []
        for key, old in before.items():
            name = key[0] + '!' + key[1]
            new = after[key]
            target = expected.get(name, old[0])
            if new != (target, old[1], old[2], old[3]):
                raise ValueError('External XLSX candidate changed an unsupported value, formula, style or link: ' + name)
            if new != old:
                changed.append(name)
        if set(changed) != set(expected):
            raise ValueError('External XLSX candidate omitted an affected location.')
    with ZipFile(baseline['blob']) as old_zip, ZipFile(candidate_path) as new_zip:
        old_members = {name: old_zip.read(name) for name in old_zip.namelist()}
        new_members = {name: new_zip.read(name) for name in new_zip.namelist()}
    for member in ('xl/styles.xml', 'xl/theme/theme1.xml'):
        if old_members.get(member) != new_members.get(member):
            raise ValueError('External XLSX candidate changed shared styles or theme.')
    return handoff, {'schema': 'task-relay.external-xlsx-checks', 'version': 1,
                     'handoff_id': handoff_id, 'handoff_digest': handoff['plan_digest'],
                     'baseline_sha256': baseline['sha256'],
                     'candidate_sha256': hashlib.sha256(candidate_path.read_bytes()).hexdigest(),
                     'changed_cells': sorted(changed),
                     'untouched_cells_verified': len(before) - len(changed),
                     'package_members_changed': sorted(
                         name for name in old_members.keys() | new_members.keys()
                         if old_members.get(name) != new_members.get(name)),
                     'package_members_preserved': sum(
                         old_members.get(name) == new_members.get(name)
                         for name in old_members.keys() & new_members.keys())}


def submit_external_xlsx(rt, *, handoff_id, submission_key, candidate_file, submitted_by):
    """Admit an agent's exact workbook only after frozen impact and preservation checks."""
    if not all(isinstance(value, str) and value.strip() for value in
               (handoff_id, submission_key, submitted_by)) or len(submission_key) > 200:
        raise ValueError('External XLSX submission needs an exact handoff, key and actor.')
    from orchestrator.runtime import safe_file
    path = Path(candidate_file).absolute()
    safe_file(Path(path.anchor), str(path.relative_to(path.anchor)))
    handoff, checks = _external_checks(rt.db, handoff_id, path)
    with transaction(rt.db):
        current, locked = _external_checks(rt.db, handoff_id, path)
        if current['plan_digest'] != handoff['plan_digest'] or locked != checks:
            raise ValueError('External XLSX candidate changed before registration.')
        prior = rt.db.execute('''SELECT * FROM relay_fact_external_submissions
            WHERE job=? AND submission_key=?''', (handoff['job'], submission_key)).fetchone()
        if prior:
            if (prior['handoff_id'] != handoff_id or prior['candidate_sha256'] != checks['candidate_sha256']
                    or prior['checks'] != json.dumps(checks, sort_keys=True)
                    or prior['submitted_by'] != submitted_by):
                raise ValueError('External XLSX submission key belongs to another candidate.')
            _verified_candidate(rt.db, prior['candidate_artifact'])
            from .execution_capture import notify
            notify(rt.db, 'xlsx_candidate', prior['id'])
            return prior['candidate_artifact']
        aid = rt.register(path, 'Agent-submitted XLSX revision',
                          task='fact_revision', path='delivery/revised_catalog.xlsx')
        if _artifact(rt.db, aid)['sha256'] != checks['candidate_sha256']:
            raise ValueError('External XLSX changed during registration.')
        plan = handoff['plan']
        rt.db.execute('''INSERT INTO relay_fact_revisions VALUES (?,?,?,?,?,?,?,?,?,?)''',
            (uuid.uuid4().hex, handoff['job'], plan['baseline_artifact'], aid,
             plan['old_source'], plan['replacement_source'], plan['digest'],
             json.dumps({item['location']: item['after'] for item in plan['affected']}, sort_keys=True),
             submitted_by, time.time()))
        from . import reviewed_links
        reviewed_links.save_impact(rt.db, candidate_artifact=aid, plan=plan)
        rt.db.execute('''INSERT INTO relay_fact_external_submissions VALUES (?,?,?,?,?,?,?,?,?)''',
            (uuid.uuid4().hex, handoff['job'], submission_key, handoff_id, aid,
             checks['candidate_sha256'], json.dumps(checks, sort_keys=True),
             submitted_by, time.time()))
        from .execution_capture import notify
        notify(rt.db, 'xlsx_candidate', rt.db.execute('SELECT id FROM relay_fact_external_submissions WHERE candidate_artifact=?', (aid,)).fetchone()[0])
        return aid


def revise_xlsx(rt, *, plan, patches, reviewer):
    """Register a distinct candidate for explicit reviewed patches in a simple XLSX."""
    from . import reviewed_links
    if not isinstance(reviewer, str) or not reviewer.strip():
        raise ValueError('An explicit reviewer is required.')
    fresh = plan_impact(rt.db, job=plan['job'], old_source=plan['old_source'],
                        replacement_source=plan['replacement_source'],
                        baseline_artifact=plan['baseline_artifact'])
    if fresh != plan:
        raise ValueError('Impact plan is stale; inspect a fresh plan before revision.')
    expected = {a['location']: a['after'] for a in plan['affected']}
    if expected != patches or not expected:
        raise ValueError('Patches must exactly match reviewed affected cells and replacement facts.')
    baseline = _artifact(rt.db, plan['baseline_artifact'])
    with closing(_plain_xlsx(baseline['blob'])) as book:
        before = _cell_state(book)
        for name, value in patches.items():
            sheet, cell = name.rsplit('!', 1)
            target = book[sheet][cell]
            if target.data_type == 'f':
                raise ValueError('Formula cells are outside the direct-value revision pilot.')
            if target.hyperlink:
                raise ValueError('Linked cells are outside the direct-value revision pilot.')
            target.value = value
        with tempfile.TemporaryDirectory(dir=rt.root) as folder:
            candidate = Path(folder) / 'candidate.xlsx'
            book.save(candidate)
            with closing(_plain_xlsx(candidate)) as checked:
                after = _cell_state(checked)
                if before.keys() != after.keys() or any(
                    after[key] != (patches.get(key[0] + '!' + key[1], old[0]), old[1], old[2], old[3])
                    for key, old in before.items()):
                    raise ValueError('XLSX candidate changed an untouched value, formula or style.')
            with transaction(rt.db):
                locked_plan = plan_impact(rt.db, job=plan['job'], old_source=plan['old_source'],
                    replacement_source=plan['replacement_source'],
                    baseline_artifact=plan['baseline_artifact'])
                if locked_plan != plan:
                    raise ValueError('Impact plan changed before candidate registration.')
                aid = rt.register(candidate, 'O14 XLSX revision candidate',
                                  task='fact_revision', path='delivery/revised_catalog.xlsx')
                rt.db.execute('''INSERT INTO relay_fact_revisions VALUES (?,?,?,?,?,?,?,?,?,?)''',
                              (uuid.uuid4().hex, plan['job'], plan['baseline_artifact'], aid,
                               plan['old_source'], plan['replacement_source'], plan['digest'],
                               json.dumps(patches, sort_keys=True), reviewer, time.time()))
                reviewed_links.save_impact(rt.db, candidate_artifact=aid, plan=plan)
    return aid


def _verified_candidate(db, candidate_artifact):
    revision = db.execute('SELECT * FROM relay_fact_revisions WHERE candidate_artifact=?',
                          (candidate_artifact,)).fetchone()
    if revision is None:
        raise ValueError('Unknown XLSX revision candidate.')
    revision = dict(revision)
    candidate = _artifact(db, candidate_artifact)
    baseline = _artifact(db, revision['baseline_artifact'])
    plan = plan_impact(db, job=revision['job'], old_source=revision['old_source'],
        replacement_source=revision['replacement_source'],
        baseline_artifact=revision['baseline_artifact'])
    if plan['digest'] != revision['plan_digest']:
        raise ValueError('Candidate plan is stale; review a fresh revision.')
    from . import reviewed_links
    reviewed_links.verify_saved_impact(db, candidate_artifact=candidate_artifact, plan=plan)
    patches = json.loads(revision['patches'])
    if patches != {a['location']: a['after'] for a in plan['affected']}:
        raise ValueError('Candidate patches differ from the reviewed impact plan.')
    with closing(_plain_xlsx(baseline['blob'])) as before_book, closing(_plain_xlsx(candidate['blob'])) as after_book:
        before, after = _cell_state(before_book), _cell_state(after_book)
        if before.keys() != after.keys() or any(
                after[key] != (patches.get(key[0] + '!' + key[1], old[0]), old[1], old[2], old[3])
                for key, old in before.items()):
            raise ValueError('Candidate changed an unapproved cell, formula or style.')
    return revision, candidate, {'baseline_sha256': baseline['sha256'],
        'candidate_sha256': candidate['sha256'], 'plan_digest': plan['digest'],
        'changed_cells': sorted(patches), 'untouched_cells_verified': len(before) - len(patches)}


def review_candidate(db, *, candidate_artifact, decision, reviewer, note):
    """Record an independent accept/revise decision for the exact candidate."""
    if decision not in ('accept', 'revise') or not all(
            isinstance(v, str) and v.strip() for v in (reviewer, note)):
        raise ValueError('An explicit reviewer, accept/revise decision and note are required.')
    revision, candidate, checks = _verified_candidate(db, candidate_artifact)
    if reviewer == revision['reviewer']:
        raise ValueError('Candidate review must be independent of its patch reviewer.')
    ident = uuid.uuid4().hex
    with transaction(db):
        if db.execute('SELECT 1 FROM relay_fact_candidate_reviews WHERE candidate_artifact=?',
                      (candidate_artifact,)).fetchone():
            raise ValueError('Candidate already has an independent review; create a new version to revise it.')
        db.execute('''INSERT INTO relay_fact_candidate_reviews VALUES (?,?,?,?,?,?,?,?,?)''',
                   (ident, revision['job'], candidate_artifact, candidate['sha256'],
                    decision, reviewer, note, json.dumps(checks, sort_keys=True), time.time()))
    return ident


def select_candidate(db, *, candidate_artifact, selected_by, receipt):
    """Record an exact user choice after independent acceptance; never dispatch."""
    if not all(isinstance(v, str) and v.strip() for v in (selected_by, receipt)):
        raise ValueError('Selection needs the actor and exact user-action receipt.')
    revision, candidate, _ = _verified_candidate(db, candidate_artifact)
    review = db.execute('SELECT * FROM relay_fact_candidate_reviews WHERE candidate_artifact=?',
                        (candidate_artifact,)).fetchone()
    if review is None or review['decision'] != 'accept' or review['candidate_sha256'] != candidate['sha256']:
        raise ValueError('Select only the exact independently accepted candidate.')
    ident = uuid.uuid4().hex
    with transaction(db):
        if db.execute('SELECT 1 FROM relay_fact_selections WHERE candidate_artifact=?',
                      (candidate_artifact,)).fetchone():
            raise ValueError('This exact candidate was already selected.')
        db.execute('''INSERT INTO relay_fact_selections VALUES (?,?,?,?,?,?,?,?)''',
                   (ident, revision['job'], candidate_artifact, candidate['sha256'],
                    revision['baseline_artifact'], selected_by, receipt, time.time()))
    return ident
