"""Check and enable the bundled native Python toolchain, without installing tools."""
import json
from pathlib import Path
import tempfile
import time

from orchestrator.contracts import digest
from orchestrator.workers import atomic
from .relay_paths import PATHS
from . import native_code_host

PROBE = '''import io,json,os,pathlib,socket
root=pathlib.Path(os.environ['RELAY_OUTPUTS'])
checks={}
try: pathlib.Path(os.environ['RELAY_INPUTS']).parent.joinpath('outside.txt').read_text(); checks['outside_read_denied']=False
except PermissionError: checks['outside_read_denied']=True
try: socket.socket().connect(('127.0.0.1',9)); checks['network_denied']=False
except PermissionError: checks['network_denied']=True
except OSError: checks['network_denied']=False
try:
    child=os.fork()
    if child==0:os._exit(0)
    os.waitpid(child,0);checks['fork_denied']=False
except PermissionError: checks['fork_denied']=True
tools={}
for module in ('docx','pptx','pypdf','reportlab','openpyxl','PIL'):
    try:
        imported=__import__(module);tools[module]={'available':True,'version':getattr(imported,'__version__','unknown'),'checked':'import'}
    except ImportError:tools[module]={'available':False}
if tools['docx']['available']:
    import docx
    b=io.BytesIO();d=docx.Document();d.add_paragraph('Relay fixture');d.save(b);b.seek(0)
    tools['docx']['checked']='save/reopen' if docx.Document(b).paragraphs[0].text=='Relay fixture' else 'failed'
if tools['pptx']['available']:
    from pptx import Presentation
    b=io.BytesIO();d=Presentation();s=d.slides.add_slide(d.slide_layouts[6]);s.shapes.add_textbox(0,0,1000000,1000000).text='Relay fixture';d.save(b);b.seek(0)
    tools['pptx']['checked']='save/reopen' if Presentation(b).slides[0].shapes[0].text=='Relay fixture' else 'failed'
if tools['openpyxl']['available']:
    import openpyxl
    b=io.BytesIO();d=openpyxl.Workbook();d.active['A1']=7;d.save(b);b.seek(0)
    tools['openpyxl']['checked']='save/reopen' if openpyxl.load_workbook(b).active['A1'].value==7 else 'failed'
if tools['pypdf']['available']:
    import pypdf
    b=io.BytesIO();d=pypdf.PdfWriter();d.add_blank_page(100,100);d.write(b);b.seek(0)
    tools['pypdf']['checked']='save/reopen' if len(pypdf.PdfReader(b).pages)==1 else 'failed'
if tools['reportlab']['available']:
    from reportlab.pdfgen import canvas
    b=io.BytesIO();d=canvas.Canvas(b);d.drawString(10,10,'Relay fixture');d.save()
    tools['reportlab']['checked']='save' if b.getvalue().startswith(b'%PDF-') else 'failed'
if tools['PIL']['available']:
    from PIL import Image
    b=io.BytesIO();Image.new('RGB',(2,2),'white').save(b,format='PNG');b.seek(0)
    tools['PIL']['checked']='save/reopen' if Image.open(b).size==(2,2) else 'failed'
(root/'tools.json').write_text(json.dumps({'isolation':checks,'tools':tools}))
'''


def path():return PATHS.data/'code-runtime.json'


def read():
    try:
        value=json.loads(path().read_text())
        if not isinstance(value,dict) or value.get('version')!=1 or value.get('id')!=digest(value.get('runtime')) or not isinstance(value.get('tools'),dict) or type(value.get('enabled')) is not bool:raise ValueError('Invalid code runtime receipt.')
        return value
    except FileNotFoundError:return None


def status():
    try:
        value=read();runtime=native_code_host.identity()
        current=bool(value and value['runtime']==runtime)
        return {'available':True,'enabled':value.get('enabled',False) if value else True,'ready':bool(current and value.get('enabled')),
                'tools':value.get('tools',{}) if value else {},'checked_at':value.get('checked_at') if value else None,
                'detail':('Runtime changed; verification is required before use. ' if value and not current else 'Local verification runs before first use. ' if not value else '')+'Bundled Python and document libraries; no extra installation. Code works on assigned files without network or subprocess access.'}
    except (ValueError,OSError,KeyError) as exc:return {'available':False,'enabled':False,'error':str(exc)}


def configure(enabled):
    from .app_access import settings_lock
    with settings_lock(path().parent):return _configure(enabled)


def _configure(enabled):
    if type(enabled) is not bool:raise ValueError('Choose whether to enable native code workers.')
    if not enabled:
        value=read()
        # Remember an explicit Off even before this runtime has been checked.
        path().parent.mkdir(parents=True,exist_ok=True)
        atomic(path(),{**value,'enabled':False} if value else {'version':1,'id':digest(None),'enabled':False,'runtime':None,'tools':{},'checked_at':None})
        return status()
    try:old=read()
    except (ValueError,OSError):old=None
    if old:atomic(path(),{**old,'enabled':False})
    runtime=native_code_host.identity()
    path().parent.mkdir(parents=True,exist_ok=True)
    atomic(path(),{'version':1,'id':digest(runtime),'enabled':False,'runtime':runtime,'tools':{},'checked_at':None})
    with tempfile.TemporaryDirectory(prefix='code-check-',dir=path().parent) as temporary:
        folder=Path(temporary);(folder/'inputs').mkdir();(folder/'outside.txt').write_text('must remain inaccessible')
        result=native_code_host.run(runtime,folder,PROBE,30,1000000)
        if result['returncode']:raise ValueError('Bundled Python isolation/tool check failed. Code workers remain off.')
        value=json.loads((folder/'outputs/tools.json').read_text())
        if value['isolation']!={'outside_read_denied':True,'network_denied':True,'fork_denied':True}:raise ValueError('Native isolation did not enforce every required boundary; code workers remain off.')
        if any(t.get('checked')=='failed' for t in value['tools'].values()):raise ValueError('Document tool qualification failed; code workers remain off.')
    atomic(path(),{'version':1,'id':digest(runtime),'enabled':True,'runtime':runtime,'tools':value['tools'],'checked_at':time.time()})
    return status()


def available(ident=None):
    try:
        value=read()
        if value is None:
            if ident is not None:raise ValueError('Native runtime receipt is missing; prepare a fresh plan before execution.')
            native_code_host.identity()  # Unsupported hosts do not create setup state.
            from .app_access import settings_lock
            with settings_lock(path().parent):
                if read() is None:_configure(True)
                value=read()
        if not value or not value.get('enabled'):raise ValueError('Enable Code and document tools in Relay Settings first.')
        if (ident is not None and value['id']!=ident) or value['runtime']!=native_code_host.identity():raise ValueError('Native runtime changed; check Code and document tools again.')
        return value
    except (OSError,KeyError,TypeError) as exc:raise ValueError('Native runtime unavailable; check Code and document tools again.') from exc
