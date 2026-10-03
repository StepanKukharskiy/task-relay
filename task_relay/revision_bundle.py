"""Versioned PPTX and JSON companion revisions against one frozen handoff.

The two companion documents are checked by exact JSON patches. This bounded
adapter also checks that the plant schedule, photo card, image provenance and
notes agree with the admitted native PPTX candidate. Submission never selects.
"""

import hashlib
import json
from pathlib import Path
import re
import time
import uuid

from orchestrator import pptx_edit
from orchestrator.storage import transaction
from . import agent_candidate, impact_handoff, native_links


SCHEMA = 'task-relay.revision-bundle'
VERSION = 1
ROLES = {'slides': 'slides.json', 'photo_manifest': 'selected-photo-manifest.json'}


def _json(value):
    return json.dumps(value, sort_keys=True, ensure_ascii=False, separators=(',', ':'))


def _digest(value):
    return hashlib.sha256(_json(value).encode('utf-8')).hexdigest()


def _document(raw):
    if len(raw) > 10_000_000:
        raise ValueError('Companion JSON exceeds 10 MB.')
    return native_links._json(raw)


def _parts(pointer):
    if not isinstance(pointer, str) or not pointer.startswith('/') or len(pointer) > 500:
        raise ValueError('Companion patch needs a bounded JSON pointer.')
    result = []
    for part in pointer[1:].split('/'):
        if '~' in part and re.search(r'~(?![01])', part):
            raise ValueError('Invalid JSON pointer escape.')
        result.append(part.replace('~1', '/').replace('~0', '~'))
    return result


def _parent(document, pointer):
    parts = _parts(pointer)
    parent = document
    for part in parts[:-1]:
        if isinstance(parent, list):
            if not part.isascii() or not part.isdecimal() or (len(part) > 1 and part[0] == '0'):
                raise ValueError('Invalid JSON array index.')
            index = int(part)
            if index >= len(parent):
                raise ValueError('JSON pointer is outside the baseline.')
            parent = parent[index]
        elif isinstance(parent, dict):
            if part not in parent:
                raise ValueError('JSON pointer is missing from the baseline.')
            parent = parent[part]
        else:
            raise ValueError('JSON pointer traverses a scalar.')
    return parent, parts[-1]


def apply_patches(raw, patches):
    """Return a revised document; reject stale, duplicate or broad undeclared edits."""
    if not isinstance(patches, list) or not 1 <= len(patches) <= 100:
        raise ValueError('Companion plan needs 1–100 exact patches.')
    document = _document(raw)
    seen = set()
    for patch in patches:
        if (not isinstance(patch, dict) or patch.get('op') not in ('replace', 'add')
                or set(patch) != ({'op', 'path', 'old', 'new'} if patch.get('op') == 'replace'
                                  else {'op', 'path', 'new'})):
            raise ValueError('Invalid companion patch.')
        pointer = patch['path']
        if pointer in seen:
            raise ValueError('Duplicate companion patch location.')
        seen.add(pointer)
        parent, key = _parent(document, pointer)
        if isinstance(parent, list):
            if not key.isascii() or not key.isdecimal() or (len(key) > 1 and key[0] == '0'):
                raise ValueError('Invalid JSON array index.')
            index = int(key)
            if index >= len(parent) or patch['op'] != 'replace':
                raise ValueError('JSON array patch needs an existing element.')
            if parent[index] != patch['old'] or patch['old'] == patch['new']:
                raise ValueError('Companion baseline value changed or patch is empty.')
            parent[index] = patch['new']
        elif isinstance(parent, dict):
            if patch['op'] == 'add':
                if key in parent:
                    raise ValueError('Companion added field already exists.')
            elif key not in parent or parent[key] != patch['old'] or patch['old'] == patch['new']:
                raise ValueError('Companion baseline value changed or patch is empty.')
            parent[key] = patch['new']
        else:
            raise ValueError('JSON patch parent is not a container.')
    return document


def initialize(db):
    db.execute('''CREATE TABLE IF NOT EXISTS relay_revision_bundle_plans(
        id TEXT PRIMARY KEY, job TEXT NOT NULL, request_key TEXT NOT NULL,
        handoff_id TEXT NOT NULL, exact_request TEXT NOT NULL,
        plan_digest TEXT NOT NULL, plan TEXT NOT NULL, actor TEXT NOT NULL,
        created REAL NOT NULL, UNIQUE(job,request_key))''')
    db.execute('''CREATE TABLE IF NOT EXISTS relay_revision_bundles(
        id TEXT PRIMARY KEY, job TEXT NOT NULL, submission_key TEXT NOT NULL,
        plan_id TEXT NOT NULL, pptx_candidate TEXT NOT NULL,
        set_digest TEXT NOT NULL, checks TEXT NOT NULL,
        submitted_by TEXT NOT NULL, created REAL NOT NULL,
        UNIQUE(job,submission_key))''')
    db.execute('''CREATE TABLE IF NOT EXISTS relay_revision_bundle_files(
        bundle_id TEXT NOT NULL, job TEXT NOT NULL, role TEXT NOT NULL,
        baseline_artifact TEXT NOT NULL, baseline_sha256 TEXT NOT NULL,
        candidate_artifact TEXT NOT NULL UNIQUE, candidate_sha256 TEXT NOT NULL,
        PRIMARY KEY(bundle_id,role))''')
    db.execute('''CREATE TABLE IF NOT EXISTS relay_revision_bundle_reviews(
        id TEXT PRIMARY KEY, bundle_id TEXT NOT NULL UNIQUE, job TEXT NOT NULL,
        set_digest TEXT NOT NULL, decision TEXT NOT NULL, reviewer TEXT NOT NULL,
        exact_feedback TEXT NOT NULL, note TEXT NOT NULL, created REAL NOT NULL)''')
    db.execute('''CREATE TABLE IF NOT EXISTS relay_revision_bundle_selections(
        bundle_id TEXT PRIMARY KEY, job TEXT NOT NULL, set_digest TEXT NOT NULL,
        selected_by TEXT NOT NULL, receipt TEXT NOT NULL, created REAL NOT NULL)''')


def _validated_plan(db, handoff_id, companions, cross_checks=None):
    handoff = impact_handoff.inspect(db, handoff_id)
    if (handoff['status'] != 'current'
            or handoff['kind'] != impact_handoff.NATIVE_REPLACEMENT):
        raise ValueError('Bundle requires a current native replacement handoff.')
    if (not isinstance(companions, list) or len(companions) != 2
            or {item.get('role') for item in companions if isinstance(item, dict)} != set(ROLES)):
        raise ValueError('Bundle needs the slides and photo-manifest companions.')
    frozen = []
    for item in sorted(companions, key=lambda value: value['role']):
        if set(item) != {'role', 'baseline_artifact', 'patches'}:
            raise ValueError('Companion declaration has unsupported fields.')
        source = native_links._artifact(db, item['baseline_artifact'])
        if source['path'].split('/')[-1] != ROLES[item['role']]:
            raise ValueError('Companion baseline has the wrong registered role.')
        revised = apply_patches(Path(source['blob']).read_bytes(), item['patches'])
        frozen.append({'role': item['role'], 'baseline_artifact': source['id'],
                       'baseline_sha256': source['sha256'], 'patches': item['patches'],
                       'expected_digest': _digest(revised)})
    if cross_checks is not None:
        if (not isinstance(cross_checks, list) or not 1 <= len(cross_checks) <= 100
                or any(not isinstance(check, dict) or check.get('kind') not in
                       ('table_cell', 'text_run', 'notes_run', 'notes_block',
                        'picture', 'forbid_terms') for check in cross_checks)):
            raise ValueError('Bundle needs bounded typed cross-file checks.')
        if not {'picture', 'forbid_terms'}.issubset(
                {check['kind'] for check in cross_checks}):
            raise ValueError('Bundle needs picture identity and stale-term checks.')
    plan = {'schema': SCHEMA + '.plan', 'version': VERSION,
                     'handoff_id': handoff_id, 'handoff_digest': handoff['plan_digest'],
                     'job': handoff['job'], 'entity_key': handoff['inputs']['entity_key'],
                     'image_artifact': handoff['inputs']['replacement_artifact'],
                     'image_sha256': handoff['plan']['replacement_sha256'],
                     'companions': frozen}
    if cross_checks is not None:
        plan['cross_checks'] = cross_checks
    return handoff, plan


def record_plan(db, *, handoff_id, request_key, exact_request, actor, companions,
                cross_checks=None):
    if (not all(isinstance(value, str) and value.strip()
                for value in (request_key, exact_request, actor))
            or len(request_key) > 200 or len(exact_request) > 100_000):
        raise ValueError('Bundle plan needs exact bounded request and actor.')
    with transaction(db):
        handoff, plan = _validated_plan(db, handoff_id, companions, cross_checks)
        if exact_request != handoff['exact_request']:
            raise ValueError('Bundle exact request differs from the frozen handoff.')
        digest = _digest(plan)
        prior = db.execute('''SELECT * FROM relay_revision_bundle_plans
            WHERE job=? AND request_key=?''', (handoff['job'], request_key)).fetchone()
        if prior:
            if (prior['handoff_id'] != handoff_id or prior['exact_request'] != exact_request
                    or prior['plan_digest'] != digest or prior['plan'] != _json(plan)
                    or prior['actor'] != actor):
                raise ValueError('Bundle request key belongs to another plan.')
            return prior['id']
        ident = uuid.uuid4().hex
        db.execute('INSERT INTO relay_revision_bundle_plans VALUES (?,?,?,?,?,?,?,?,?)',
                   (ident, handoff['job'], request_key, handoff_id, exact_request,
                    digest, _json(plan), actor, time.time()))
        return ident


def inspect_plan(db, plan_id):
    row = db.execute('SELECT * FROM relay_revision_bundle_plans WHERE id=?',
                     (plan_id,)).fetchone()
    if row is None:
        raise ValueError('Unknown bundle plan.')
    row = dict(row)
    try:
        saved = native_links._json(row['plan'])
        if _digest(saved) != row['plan_digest']:
            raise ValueError('Saved bundle plan digest changed.')
        handoff, current = _validated_plan(db, row['handoff_id'],
            [{key: item[key] for key in ('role', 'baseline_artifact', 'patches')}
             for item in saved['companions']], saved.get('cross_checks'))
        if current != saved or handoff['job'] != row['job']:
            raise ValueError('Bundle source evidence changed.')
    except (ValueError, TypeError, KeyError, OSError) as exc:
        return {**row, 'status': 'stale', 'reason': str(exc)}
    return {**row, 'plan': saved, 'status': 'current',
            'reason': 'Frozen companion patches match current registered evidence.'}


def _couplings(db, pptx_candidate, documents, plan):
    slides = documents['slides']['slides']
    photos = documents['photo_manifest']['photos']
    if len(slides) != 45:
        raise ValueError('Companion deck slide count changed.')
    image = native_links._artifact(db, plan['image_artifact'])
    from PIL import Image
    import io
    with Image.open(io.BytesIO(Path(image['blob']).read_bytes())) as source_image:
        dimensions = [source_image.width, source_image.height]
    rose = [photo for photo in photos if photo.get('id') == 'rose-unverified']
    if len(rose) != 1:
        raise ValueError('Photo manifest needs one provisional rose entry.')
    rose = rose[0]
    if (rose.get('sha256') != plan['image_sha256']
            or rose.get('bytes') != image['bytes']
            or rose.get('dimensions') != dimensions
            or rose.get('media_type') != 'image/png'
            or rose.get('status') != 'user_provided_unverified'
            or any(rose.get(key) is not None for key in
                   ('author', 'license', 'license_url', 'source_url', 'download_url'))
            or rose.get('staged_path') != 'inputs/rose.png'):
        raise ValueError('Photo entry does not match the pinned image and unknown rights.')
    photo_slide = slides[38]['elements']
    if photo_slide[8].get('path') != rose['staged_path']:
        raise ValueError('Slide image path differs from the pinned photo entry.')
    raw = Path(pptx_candidate['blob']).read_bytes()
    listing = pptx_edit.inspect(raw)['slides']
    shapes18 = {shape['shape_id']: shape for shape in listing[17]['shapes']}
    cells = {(cell['row'], cell['column']): cell['text_runs']
             for cell in shapes18[6]['table_cells']}
    schedule = slides[17]['elements']
    for column in range(4):
        if schedule[4]['rows'][3][column] != '\n'.join(cells[(3, column)]):
            raise ValueError('Shrub schedule differs between JSON and PPTX.')
    if (schedule[2]['text'] != shapes18[4]['text_runs'][0]
            or slides[17]['notes'] != listing[17]['notes_text_runs'][0]):
        raise ValueError('Shrub source footer or note differs between JSON and PPTX.')
    shapes39 = {shape['shape_id']: shape for shape in listing[38]['shapes']}
    for index, shape_id in ((6, 8), (7, 9), (9, 11), (10, 12)):
        if photo_slide[index].get('text') != shapes39[shape_id]['text_runs'][0]:
            raise ValueError('Photo card differs between JSON and PPTX.')
    if shapes39[10]['image_sha256'] != plan['image_sha256']:
        raise ValueError('Native photo differs from the pinned image.')
    from pptx import Presentation
    import io
    native = Presentation(io.BytesIO(raw))
    picture = next(shape for shape in native.slides[38].shapes if shape.shape_id == 10)
    for key, value in (('x', picture.left), ('y', picture.top),
                       ('w', picture.width), ('h', picture.height)):
        if abs(photo_slide[8].get(key, -1) * 914400 - value) > 1:
            raise ValueError('Companion image geometry differs from the native picture.')
    if slides[38].get('notes', '').splitlines() != listing[38]['notes_text_runs'][1:7]:
        raise ValueError('Companion photo notes differ from the native attribution.')
    if any('oleander' in _json(document).casefold() or 'nerium' in _json(document).casefold()
           for document in documents.values()):
        raise ValueError('Withdrawn subject remains in revised companions.')
    return {'schedule_cells': 4, 'photo_card_texts': 4,
            'source_footer_and_note': 2, 'photo_notes': 6,
            'pinned_picture_and_manifest': True, 'native_picture_geometry': True}


def _value(document, pointer):
    parent, key = _parent(document, pointer)
    if isinstance(parent, list):
        if not key.isascii() or not key.isdecimal() or int(key) >= len(parent):
            raise ValueError('Cross-file JSON pointer is outside an array.')
        return parent[int(key)]
    if not isinstance(parent, dict) or key not in parent:
        raise ValueError('Cross-file JSON pointer is missing.')
    return parent[key]


def _generic_couplings(db, pptx_candidate, documents, plan):
    checks = plan['cross_checks']
    listing = pptx_edit.inspect(Path(pptx_candidate['blob']).read_bytes())['slides']
    native = None
    passed = []
    keys = {
        'table_cell': {'kind', 'role', 'pointer', 'slide', 'shape_id', 'row', 'column'},
        'text_run': {'kind', 'role', 'pointer', 'slide', 'shape_id'},
        'notes_run': {'kind', 'role', 'pointer', 'slide', 'index'},
        'notes_block': {'kind', 'role', 'pointer', 'slide', 'start_index', 'count'},
        'picture': {'kind', 'role', 'pointer', 'slide', 'shape_id',
                    'manifest_role', 'manifest_pointer', 'unknown_fields', 'status'},
        'forbid_terms': {'kind', 'roles', 'terms'},
    }
    for check in checks:
        kind = check['kind']
        if set(check) != keys[kind]:
            raise ValueError('Cross-file check has unsupported fields.')
        if kind == 'forbid_terms':
            if (not isinstance(check['roles'], list) or not check['roles']
                    or not isinstance(check['terms'], list) or not check['terms']):
                raise ValueError('Forbidden-term check needs roles and terms.')
            for role in check['roles']:
                if role not in documents:
                    raise ValueError('Forbidden-term check names an unknown role.')
                content = _json(documents[role]).casefold()
                if any(not isinstance(term, str) or not term
                       or term.casefold() in content for term in check['terms']):
                    raise ValueError('Withdrawn term remains in a companion.')
            passed.append(kind)
            continue
        role = check['role']
        if role not in documents or type(check['slide']) is not int or not 1 <= check['slide'] <= len(listing):
            raise ValueError('Cross-file check targets an unknown role or slide.')
        source = _value(documents[role], check['pointer'])
        slide = listing[check['slide']-1]
        if kind in ('table_cell', 'text_run', 'picture'):
            if type(check['shape_id']) is not int:
                raise ValueError('Cross-file shape ID must be an integer.')
            shapes = [item for item in slide['shapes'] if item['shape_id'] == check['shape_id']]
            if len(shapes) != 1:
                raise ValueError('Cross-file native shape is missing.')
            shape = shapes[0]
        if kind == 'table_cell':
            if type(check['row']) is not int or type(check['column']) is not int:
                raise ValueError('Cross-file table coordinates must be integers.')
            cells = [item for item in shape.get('table_cells', [])
                     if (item['row'], item['column']) == (check['row'], check['column'])]
            if len(cells) != 1 or source != '\n'.join(cells[0]['text_runs']):
                raise ValueError('Companion table cell differs from the native slide.')
        elif kind == 'text_run':
            if not isinstance(source, str) or shape['text_runs'] != [source]:
                raise ValueError('Companion text differs from the native slide.')
        elif kind in ('notes_run', 'notes_block'):
            index = check['index'] if kind == 'notes_run' else check['start_index']
            if type(index) is not int or not 0 <= index < len(slide['notes_text_runs']):
                raise ValueError('Cross-file notes index is invalid.')
            if kind == 'notes_run':
                if not isinstance(source, str) or slide['notes_text_runs'][index] != source:
                    raise ValueError('Companion note differs from native speaker notes.')
            else:
                count = check['count']
                if (type(count) is not int or not 1 <= count <= 100
                        or index + count > len(slide['notes_text_runs'])
                        or not isinstance(source, str)
                        or source.splitlines() != slide['notes_text_runs'][index:index+count]):
                    raise ValueError('Companion notes block differs from native speaker notes.')
        else:
            if (not isinstance(source, dict) or check['manifest_role'] not in documents
                    or not isinstance(check['unknown_fields'], list)
                    or not check['unknown_fields']):
                raise ValueError('Cross-file picture declaration is invalid.')
            entry = _value(documents[check['manifest_role']], check['manifest_pointer'])
            image = native_links._artifact(db, plan['image_artifact'])
            from PIL import Image
            import io
            with Image.open(io.BytesIO(Path(image['blob']).read_bytes())) as bitmap:
                dimensions = [bitmap.width, bitmap.height]
            if (not isinstance(entry, dict) or entry.get('sha256') != plan['image_sha256']
                    or shape.get('image_sha256') != plan['image_sha256']
                    or entry.get('bytes') != image['bytes']
                    or entry.get('dimensions') != dimensions
                    or entry.get('media_type') != 'image/png'
                    or entry.get('status') != check['status']
                    or source.get('path') != entry.get('staged_path')
                    or any(entry.get(field) is not None for field in check['unknown_fields'])):
                raise ValueError('Companion photo or rights differ from the pinned image.')
            if native is None:
                from pptx import Presentation
                native = Presentation(io.BytesIO(Path(pptx_candidate['blob']).read_bytes()))
            picture = next((item for item in native.slides[check['slide']-1].shapes
                            if item.shape_id == check['shape_id']), None)
            if picture is None:
                raise ValueError('Native picture geometry is missing.')
            for field, coordinate in (('x', picture.left), ('y', picture.top),
                                      ('w', picture.width), ('h', picture.height)):
                value = source.get(field)
                if type(value) not in (int, float) or abs(value * 914400 - coordinate) > 1:
                    raise ValueError('Companion picture geometry differs from native geometry.')
        passed.append(kind)
    return {'declared': len(checks), 'passed': len(passed),
            'kinds': sorted(set(passed)), 'image_sha256': plan['image_sha256']}


def _check(db, plan_id, pptx_candidate_artifact, files):
    frozen = inspect_plan(db, plan_id)
    if frozen['status'] != 'current':
        raise ValueError('Bundle plan is stale: ' + frozen['reason'])
    plan = frozen['plan']
    row, pptx_candidate, handoff, _ = agent_candidate.verified(db, pptx_candidate_artifact)
    if (row['job'] != frozen['job'] or row['handoff_id'] != plan['handoff_id']
            or handoff['plan_digest'] != plan['handoff_digest']):
        raise ValueError('PPTX candidate is not from this frozen bundle plan.')
    if set(files) != set(ROLES):
        raise ValueError('Bundle requires exactly two declared companion files.')
    documents, members = {}, []
    for item in plan['companions']:
        role = item['role']
        baseline = native_links._artifact(db, item['baseline_artifact'])
        expected = apply_patches(Path(baseline['blob']).read_bytes(), item['patches'])
        if _digest(expected) != item['expected_digest']:
            raise ValueError('Frozen companion patch result changed.')
        candidate = _document(files[role])
        if candidate != expected:
            raise ValueError('Companion contains an undeclared or missing JSON change: ' + role)
        documents[role] = candidate
        members.append({'role': role, 'baseline_artifact': baseline['id'],
                        'baseline_sha256': baseline['sha256'],
                        'candidate_sha256': hashlib.sha256(files[role]).hexdigest(),
                        'patch_count': len(item['patches'])})
    couplings = (_generic_couplings(db, pptx_candidate, documents, plan)
                 if 'cross_checks' in plan else
                 _couplings(db, pptx_candidate, documents, plan))
    digest = _digest({'plan_digest': frozen['plan_digest'],
                      'pptx_artifact': pptx_candidate_artifact,
                      'pptx_sha256': pptx_candidate['sha256'],
                      'image_artifact': plan['image_artifact'],
                      'image_sha256': plan['image_sha256'],
                      'members': members})
    return frozen, members, {'schema': SCHEMA + '.checks', 'version': VERSION,
        'set_digest': digest, 'plan_digest': frozen['plan_digest'],
        'pptx_sha256': pptx_candidate['sha256'], 'companions': members,
        'couplings': couplings, 'review': 'pending', 'selection': 'none'}


def submit(rt, *, plan_id, pptx_candidate_artifact, submission_key,
           companion_files, submitted_by):
    if (not all(isinstance(value, str) and value.strip()
                for value in (submission_key, submitted_by)) or len(submission_key) > 200):
        raise ValueError('Bundle submission needs key and agent identity.')
    from orchestrator.runtime import safe_file
    raw = {}
    for role, source in companion_files.items():
        path = Path(source).absolute()
        safe_file(Path(path.anchor), str(path.relative_to(path.anchor)))
        raw[role] = path.read_bytes()
    frozen, members, checks = _check(rt.db, plan_id, pptx_candidate_artifact, raw)
    with transaction(rt.db):
        second, _, current_checks = _check(rt.db, plan_id, pptx_candidate_artifact, raw)
        if second['plan_digest'] != frozen['plan_digest'] or current_checks != checks:
            raise ValueError('Bundle changed before registration.')
        prior = rt.db.execute('''SELECT * FROM relay_revision_bundles
            WHERE job=? AND submission_key=?''', (frozen['job'], submission_key)).fetchone()
        if prior:
            if (prior['plan_id'] != plan_id or prior['pptx_candidate'] != pptx_candidate_artifact
                    or prior['set_digest'] != checks['set_digest']
                    or prior['checks'] != _json(checks)
                    or prior['submitted_by'] != submitted_by):
                raise ValueError('Bundle submission key belongs to another candidate set.')
            from .execution_capture import notify
            notify(rt.db, 'revision_bundle', prior['id'])
            return prior['id']
        ident = uuid.uuid4().hex
        registered = {}
        for role, source in companion_files.items():
            aid = rt.register(source, 'Agent-submitted revision companion',
                task='revision_bundle', path='delivery/' + ROLES[role])
            if rt.artifact(aid)['sha256'] != next(item['candidate_sha256']
                    for item in members if item['role'] == role):
                raise ValueError('Companion changed during registration.')
            registered[role] = aid
        rt.db.execute('INSERT INTO relay_revision_bundles VALUES (?,?,?,?,?,?,?,?,?)',
            (ident, frozen['job'], submission_key, plan_id, pptx_candidate_artifact,
             checks['set_digest'], _json(checks), submitted_by, time.time()))
        for member in members:
            rt.db.execute('INSERT INTO relay_revision_bundle_files VALUES (?,?,?,?,?,?,?)',
                (ident, frozen['job'], member['role'], member['baseline_artifact'],
                 member['baseline_sha256'], registered[member['role']],
                 member['candidate_sha256']))
        from .execution_capture import notify
        notify(rt.db, 'revision_bundle', ident)
    return ident


def verified(db, bundle_id):
    row = db.execute('SELECT * FROM relay_revision_bundles WHERE id=?',
                     (bundle_id,)).fetchone()
    if row is None:
        raise ValueError('Unknown revision bundle.')
    row = dict(row)
    members = [dict(item) for item in db.execute('''SELECT * FROM relay_revision_bundle_files
        WHERE bundle_id=? ORDER BY role''', (bundle_id,))]
    files = {item['role']: Path(native_links._artifact(db, item['candidate_artifact'])['blob']).read_bytes()
             for item in members}
    frozen, checked_members, checks = _check(db, row['plan_id'],
                                              row['pptx_candidate'], files)
    if (row['job'] != frozen['job'] or row['set_digest'] != checks['set_digest']
            or row['checks'] != _json(checks) or len(members) != len(checked_members)
            or any(item['candidate_sha256'] != hashlib.sha256(files[item['role']]).hexdigest()
                   for item in members)):
        raise ValueError('Saved bundle version or checks changed.')
    return row, checks


def review(db, *, bundle_id, decision, reviewer, exact_feedback, note,
           expected_set_digest):
    """Record an independent exact-version decision without selecting the set."""
    if (decision not in ('accept', 'revise')
            or not all(isinstance(value, str) and value.strip() for value in
                       (reviewer, exact_feedback, note, expected_set_digest))
            or len(exact_feedback) > 100_000 or len(note) > 10_000):
        raise ValueError('Bundle review needs an exact response, actor and reason.')
    with transaction(db):
        bundle, _ = verified(db, bundle_id)
        if (bundle['set_digest'] != expected_set_digest
                or reviewer == bundle['submitted_by']):
            raise ValueError('Bundle review needs an independent actor and exact set digest.')
        if db.execute('SELECT 1 FROM relay_revision_bundle_reviews WHERE bundle_id=?',
                      (bundle_id,)).fetchone():
            raise ValueError('Bundle already reviewed.')
        ident = uuid.uuid4().hex
        db.execute('INSERT INTO relay_revision_bundle_reviews VALUES (?,?,?,?,?,?,?,?,?)',
                   (ident, bundle_id, bundle['job'], bundle['set_digest'], decision,
                    reviewer, exact_feedback, note, time.time()))
        return ident


def select(db, *, bundle_id, selected_by, receipt, expected_set_digest):
    """Select an accepted set and its native candidate in one transaction."""
    if not all(isinstance(value, str) and value.strip() for value in
               (selected_by, receipt, expected_set_digest)):
        raise ValueError('Bundle selection needs actor, receipt and exact digest.')
    with transaction(db):
        bundle, _ = verified(db, bundle_id)
        if bundle['set_digest'] != expected_set_digest:
            raise ValueError('Bundle set changed before selection.')
        approval = db.execute('''SELECT * FROM relay_revision_bundle_reviews
            WHERE bundle_id=?''', (bundle_id,)).fetchone()
        if (approval is None or approval['decision'] != 'accept'
                or approval['set_digest'] != expected_set_digest):
            raise ValueError('An exact independent bundle acceptance is required.')
        if db.execute('SELECT 1 FROM relay_revision_bundle_selections WHERE bundle_id=?',
                      (bundle_id,)).fetchone():
            raise ValueError('Bundle already selected.')
        plan = db.execute('SELECT handoff_id FROM relay_revision_bundle_plans WHERE id=?',
                          (bundle['plan_id'],)).fetchone()
        if plan is None:
            raise ValueError('Bundle frozen plan is missing.')
        other = db.execute('''SELECT 1 FROM relay_revision_bundle_selections s
            JOIN relay_revision_bundles b ON b.id=s.bundle_id
            JOIN relay_revision_bundle_plans p ON p.id=b.plan_id
            WHERE p.handoff_id=? AND s.bundle_id<>?''',
            (plan['handoff_id'], bundle_id)).fetchone()
        if other:
            raise ValueError('Another revision set for this handoff is already selected.')
        candidate, native_artifact, handoff, _ = agent_candidate.verified(
            db, bundle['pptx_candidate'])
        native_review = db.execute('''SELECT * FROM relay_agent_candidate_reviews
            WHERE candidate_artifact=?''', (bundle['pptx_candidate'],)).fetchone()
        if (native_review is None or native_review['decision'] != 'accept'
                or native_review['candidate_sha256'] != native_artifact['sha256']
                or native_review['plan_digest'] != handoff['plan_digest']):
            raise ValueError('The exact native candidate needs independent acceptance.')
        agent_candidate.select(db, candidate_artifact=bundle['pptx_candidate'],
            selected_by=selected_by, receipt=receipt,
            expected_sha256=native_artifact['sha256'],
            expected_plan_digest=handoff['plan_digest'])
        db.execute('INSERT INTO relay_revision_bundle_selections VALUES (?,?,?,?,?,?)',
                   (bundle_id, bundle['job'], expected_set_digest, selected_by,
                    receipt, time.time()))
    return bundle_id


def records(db, job):
    return {'plans': [dict(row) for row in db.execute(
                'SELECT * FROM relay_revision_bundle_plans WHERE job=? ORDER BY created,id', (job,))],
            'bundles': [dict(row) for row in db.execute(
                'SELECT * FROM relay_revision_bundles WHERE job=? ORDER BY created,id', (job,))],
            'files': [dict(row) for row in db.execute(
                'SELECT * FROM relay_revision_bundle_files WHERE job=? ORDER BY bundle_id,role', (job,))],
            'reviews': [dict(row) for row in db.execute(
                'SELECT * FROM relay_revision_bundle_reviews WHERE job=? ORDER BY created,id', (job,))],
            'selections': [dict(row) for row in db.execute(
                'SELECT * FROM relay_revision_bundle_selections WHERE job=? ORDER BY created,bundle_id',
                (job,))]}


def status(db, bundle_id):
    if db.execute('SELECT 1 FROM relay_revision_bundle_selections WHERE bundle_id=?',
                  (bundle_id,)).fetchone():
        return 'selected'
    review = db.execute('SELECT decision FROM relay_revision_bundle_reviews WHERE bundle_id=?',
                        (bundle_id,)).fetchone()
    return ({'accept': 'accepted', 'revise': 'revision_requested'}[review['decision']]
            if review else 'verified_pending_review')


def main(argv=None):
    """Desktop-independent freeze, submit and inspect commands for another agent."""
    import argparse
    import sqlite3
    from .bridge import State
    from . import production_control, workflow_files
    from orchestrator.runtime import Runtime, safe_file

    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest='action', required=True)
    freeze = sub.add_parser('plan')
    freeze.add_argument('--database', required=True)
    freeze.add_argument('--handoff-id', required=True)
    freeze.add_argument('--request-key', required=True)
    freeze.add_argument('--request-file', required=True)
    freeze.add_argument('--spec', required=True)
    freeze.add_argument('--actor', required=True)
    send = sub.add_parser('submit')
    send.add_argument('--database', required=True)
    send.add_argument('--plan-id', required=True)
    send.add_argument('--pptx-candidate', required=True)
    send.add_argument('--submission-key', required=True)
    send.add_argument('--files-root', required=True)
    send.add_argument('--actor', required=True)
    check = sub.add_parser('verify')
    check.add_argument('--database', required=True)
    check.add_argument('--bundle-id', required=True)
    args = parser.parse_args(argv)
    supplied = Path(args.database).expanduser()
    if not supplied.is_file() or supplied.is_symlink():
        parser.error('An existing authoritative Relay database is required.')
    database = supplied.resolve()
    if args.action == 'verify':
        with sqlite3.connect(database.as_uri() + '?mode=ro', uri=True) as db:
            db.row_factory = sqlite3.Row
            row, checks = verified(db, args.bundle_id)
            result = {'bundle_id': row['id'], 'job': row['job'],
                      'set_digest': row['set_digest'], 'checks_at_submission': checks,
                      'status': status(db, row['id'])}
    else:
        state = State(database)
        try:
            if args.action == 'plan':
                for named in (args.request_file, args.spec):
                    path = Path(named).expanduser().absolute()
                    safe_file(Path(path.anchor), str(path.relative_to(path.anchor)))
                spec = native_links._json(Path(args.spec).read_text(encoding='utf-8'))
                if not isinstance(spec, dict) or set(spec) != {'companions', 'cross_checks'}:
                    raise ValueError('Bundle specification needs companions and cross_checks.')
                ident = record_plan(state.db, handoff_id=args.handoff_id,
                    request_key=args.request_key,
                    exact_request=Path(args.request_file).read_text(encoding='utf-8'),
                    actor=args.actor, companions=spec['companions'],
                    cross_checks=spec['cross_checks'])
                row = inspect_plan(state.db, ident)
                job = row['job']
                result = {'plan_id': ident, 'plan_digest': row['plan_digest'],
                          'status': row['status'], 'plan': row['plan']}
            else:
                root = Path(args.files_root).expanduser().absolute()
                files = {role: root / filename for role, filename in ROLES.items()}
                ident = submit(Runtime(production_control.root(state), connection=state.db),
                    plan_id=args.plan_id, pptx_candidate_artifact=args.pptx_candidate,
                    submission_key=args.submission_key, companion_files=files,
                    submitted_by=args.actor)
                row, checks = verified(state.db, ident)
                job = row['job']
                result = {'bundle_id': ident, 'set_digest': row['set_digest'],
                          'status': status(state.db, ident),
                          'checks_at_submission': checks}
            owner = state.db.execute('SELECT channel FROM relay_pipelines WHERE id=?',
                                     (job,)).fetchone()
            if owner is None:
                raise ValueError('Bundle job is missing.')
            state.channel = owner['channel']
            view = workflow_files.sync(state, job)
            result['job_view'] = str(view/'.relay/job.sqlite')
        finally:
            state.db.close()
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
