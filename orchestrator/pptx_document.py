"""Bounded data-only PPTX authoring. No application launch or network access."""
import io
import json
import math
import re
from .contracts import relative
from .handoff_contracts import PPTX_SLIDES, PPTX_ELEMENTS

MIME = 'application/vnd.openxmlformats-officedocument.presentationml.presentation'
VERSION = '1.0.2'
TITLE_MAX = 255
DESCRIPTION = '''Version 1 slide specification: {version:1,title:string,slides:[
{background?:"FFFFFF",notes?:string,elements:[...]}],size?:[width,height],font?:"Arial"}.
Dimensions and x,y,w,h are inches; default size is [13.333333,7.5].
Each element has type, x,y,w,h and optional name. Text, shape and image elements
may include hyperlink: an absolute HTTPS purchase/source URL without credentials.
The hyperlink is retained on the native object and checked after reopening. Supported elements:
text: {text:string,font_size?:24,color?:"172B4D",bold?:false,align?:"left"|"center"|"right"}.
shape: {shape:"rectangle"|"ellipse",fill?:"E8EEF5",color?:"172B4D",text?:string,font_size?:24,bold?:false,align?:"left"|"center"|"right"}.
image: {path:exact staged PNG/JPEG input path,fit?:"contain"|"cover",source_region?:[left,top,right,bottom]}.
source_region uses normalized original-image coordinates with positive area; original pixels stay embedded.
Default contain preserves the whole image; explicit cover center-crops to fill its box.
Sourced image bundles from images.collect are also accepted: use
the exact ZIP input path + "/" + the image member path in manifest.json. For
example, input assets/photos.zip with member images/oak.jpg uses
assets/photos.zip/images/oak.jpg (no angle brackets). Full image credits are retained in
slide notes. The planner should include visible captions/credits and review identity.
table: {rows:[[string,...],...],font_size?:18,color?:"172B4D"}; first row is the header.
chart: {chart:"column"|"bar"|"line",categories:[string,...],series:[{name:string,values:[number,...]}]}.
Use native objects, readable text and uncluttered layouts. All boxes must fit the slide.
Font sizes are 8–96 points. Text frames wrap without automatic resizing, with
0.06 inch left/right and 0.03 inch top/bottom margins. Tables divide their box
equally among rows and columns. Text-fit estimates are not rendered measurements;
report assumptions and distinguish confirmed errors from visual review notes.
Limits: 1–50 slides, 1–50 elements/slide, 1000 elements/deck, 4000 characters/text,
20x10 table cells (300 characters each), 50 chart categories and 8 series.
No arbitrary code, remote image loading, native PPTX import, animations, or font installation.
Creation reopens and checks native objects, text, table/chart data and image bytes.
These are structural checks, not visual layout approval or Keynote qualification.
'''
from .slide_templates import DESCRIPTION as TEMPLATE_DESCRIPTION, expand
DESCRIPTION += '\n'+TEMPLATE_DESCRIPTION


def available():
    try:
        import pptx
        if pptx.__version__ != VERSION:
            raise ValueError('PPTX creation requires python-pptx==' + VERSION)
    except ImportError as exc:
        raise ValueError('PPTX dependency missing. Install task-relay[presentations] in the worker runtime.') from exc


def _object(value, required, optional=()):
    if not isinstance(value, dict) or not set(required) <= set(value) or set(value) - set(required) - set(optional):
        raise ValueError('Invalid slide specification fields; see pptx.create slide_schema.')


def _text(value, maximum, empty=True):
    if not isinstance(value, str) or len(value) > maximum or (not empty and not value.strip()):
        raise ValueError('Missing or oversized presentation text.')
    # Avoid OOXML escaping/substitution changing exact source text on reopen.
    if any(ord(ch) < 32 and ch not in '\n\t' or 0xD800 <= ord(ch) <= 0xDFFF or ord(ch) in (0xFFFE, 0xFFFF) for ch in value):
        raise ValueError('Unsupported control character in presentation text.')


def _number(value, low, high):
    if type(value) not in (int, float) or not math.isfinite(value) or not low <= value <= high:
        raise ValueError('Presentation numbers must be finite and within bounds.')


def _color(value):
    if not isinstance(value, str) or not re.fullmatch('[0-9A-Fa-f]{6}', value):
        raise ValueError('Use six-digit RGB colors.')


def _hyperlink(value):
    from urllib.parse import urlsplit
    _text(value, 4096, False)
    if any(ch.isspace() for ch in value) or any(ord(ch) < 32 for ch in value):
        raise ValueError('Hyperlink must be an absolute HTTPS URL without credentials.')
    try:
        parsed = urlsplit(value)
        valid = parsed.scheme == 'https' and parsed.hostname and not parsed.username and not parsed.password and parsed.port in (None, 443)
    except ValueError:
        valid = False
    if not valid:
        raise ValueError('Hyperlink must be an absolute HTTPS URL without credentials.')


def load(text):
    def unique(pairs):
        result = {}
        for key, value in pairs:
            if key in result:
                raise ValueError('Duplicate slide specification key: ' + key)
            result[key] = value
        return result
    return json.loads(text, object_pairs_hook=unique)


def validate(document, image_paths=()):
    document = expand(document)
    _object(document, ('version', 'title', 'slides'), ('size', 'font', 'asset_sha256'))
    if type(document['version']) is not int or document['version'] != 1:
        raise ValueError('Unsupported slide specification version.')
    _text(document['title'], TITLE_MAX, False)
    _text(document.get('font', 'Arial'), 100, False)
    bindings=document.get('asset_sha256',{})
    if not isinstance(bindings,dict) or any(not isinstance(p,str) or not isinstance(h,str) or not re.fullmatch('[a-f0-9]{64}',h) for p,h in bindings.items()):
        raise ValueError('Invalid document asset identities')
    size = document.get('size', [13.333333, 7.5])
    if not isinstance(size, list) or len(size) != 2:
        raise ValueError('Slide size must contain width and height.')
    for n in size:
        _number(n, 1, 40)
    slides = document['slides']
    if not isinstance(slides, list) or not 1 <= len(slides) <= PPTX_SLIDES:
        raise ValueError('Expected 1–50 slides.')
    count = 0
    for slide in slides:
        _object(slide, ('elements',), ('background', 'notes'))
        _color(slide.get('background', 'FFFFFF'))
        _text(slide.get('notes', ''), 10000)
        if not isinstance(slide['elements'], list) or not 1 <= len(slide['elements']) <= 50:
            raise ValueError('Expected 1–50 elements per slide.')
        count += len(slide['elements'])
        if count > PPTX_ELEMENTS:
            raise ValueError('Too many presentation elements.')
        for item in slide['elements']:
            if not isinstance(item, dict) or not isinstance(item.get('type'), str):
                raise ValueError('Missing presentation element type.')
            kind = item['type']
            fields = {
                'text': (('text',), ('font_size', 'color', 'bold', 'align', 'font')),
                'shape': (('shape',), ('fill', 'text', 'font_size', 'color', 'bold', 'align', 'font')),
                'image': (('path',), ('fit','source_region')),
                'table': (('rows',), ('font_size', 'color', 'font', 'row_heights')),
                'chart': (('chart', 'categories', 'series'), ()),
            }
            if kind not in fields:
                raise ValueError('Unsupported presentation element type.')
            required, optional = fields[kind]
            _object(item, ('type', 'x', 'y', 'w', 'h', *required), ('name', *optional, *(('hyperlink',) if kind in ('text', 'shape', 'image') else ())))
            if 'hyperlink' in item:
                _hyperlink(item['hyperlink'])
            for key in ('x', 'y', 'w', 'h'):
                _number(item[key], 0 if key in ('x', 'y') else 0.01, 40)
            if item['x'] + item['w'] > size[0] + 1e-6 or item['y'] + item['h'] > size[1] + 1e-6:
                raise ValueError('Presentation element extends beyond slide bounds.')
            _text(item.get('name', ''), 100)
            if 'row_heights' in item:
                heights=item['row_heights']
                if not isinstance(heights,list) or len(heights)!=len(item['rows']):raise ValueError('Table row heights must match rows')
                for height in heights:_number(height,.01,40)
                if abs(sum(heights)-item['h'])>1e-6:raise ValueError('Table row heights exceed its box')
            if 'font_size' in item:
                _number(item['font_size'], 8, 96)
            if 'font' in item:
                _text(item['font'],100,False)
            if 'source_region' in item:
                region=item['source_region']
                if not isinstance(region,list) or len(region)!=4:raise ValueError('Image source region needs four normalized coordinates')
                for value in region:_number(value,0,1)
                if region[0]>=region[2] or region[1]>=region[3]:raise ValueError('Image source region must have positive area')
            if 'fit' in item and item['fit'] not in ('contain','cover'):
                raise ValueError('Image fit must be contain or cover.')
            for key in ('color', 'fill'):
                if key in item:
                    _color(item[key])
            if 'bold' in item and type(item['bold']) is not bool:
                raise ValueError('bold must be boolean.')
            if 'align' in item and item['align'] not in ('left', 'center', 'right'):
                raise ValueError('Unsupported text alignment.')
            if kind in ('text', 'shape'):
                _text(item.get('text', ''), 4000)
            if kind == 'shape' and item['shape'] not in ('rectangle', 'ellipse'):
                raise ValueError('Unsupported native shape.')
            if kind == 'image':
                path = relative(item['path'])
                if path not in image_paths:
                    raise ValueError('Image must name an exact declared PNG/JPEG input: ' + path)
            if kind == 'table':
                rows = item['rows']
                if not isinstance(rows, list) or not 1 <= len(rows) <= 20 or not isinstance(rows[0], list) or not 1 <= len(rows[0]) <= 10:
                    raise ValueError('Table must contain 1–20 rows and 1–10 columns.')
                for row in rows:
                    if not isinstance(row, list) or len(row) != len(rows[0]):
                        raise ValueError('Table rows must have the same width.')
                    for cell in row:
                        _text(cell, 300)
            if kind == 'chart':
                if item['chart'] not in ('column', 'bar', 'line'):
                    raise ValueError('Unsupported native chart.')
                categories, series = item['categories'], item['series']
                if not isinstance(categories, list) or not 1 <= len(categories) <= 50 or not isinstance(series, list) or not 1 <= len(series) <= 8:
                    raise ValueError('Invalid chart dimensions.')
                for category in categories:
                    _text(category, 100, False)
                    if category.startswith('='):
                        raise ValueError('Chart labels must be literal text, not spreadsheet formulas.')
                for entry in series:
                    _object(entry, ('name', 'values'))
                    _text(entry['name'], 100, False)
                    if entry['name'].startswith('='):
                        raise ValueError('Chart labels must be literal text, not spreadsheet formulas.')
                    if not isinstance(entry['values'], list) or len(entry['values']) != len(categories):
                        raise ValueError('Chart series must match categories.')
                    for n in entry['values']:
                        _number(n, -1e12, 1e12)
    return document


def validator_source():
    """Freeze the actual stdlib schema checker for isolated draft/review workers."""
    import inspect
    from .slide_templates import CATALOG, STYLES, COLLECTION_LAYOUTS, grid_pages, card_pages, table_pages
    functions=(relative,_object,_text,_number,_color,_hyperlink,load,validate)
    compiler=inspect.getsource(grid_pages).replace('    from .pptx_document import _object, _text\n','')+'\n'+inspect.getsource(card_pages).replace('    from .pptx_document import _object, _text, _hyperlink\n','')+'\n'+inspect.getsource(expand).replace('    from .pptx_document import _object, _text, _color, _number, TITLE_MAX\n','').replace('    from .table_layout import paginate_text\n','')
    from . import table_layout
    compiler=inspect.getsource(table_layout).replace('import math','').replace('import re','').replace('import unicodedata','')+'\n'+inspect.getsource(table_pages).replace('    from .pptx_document import _object,_text\n','').replace('    from .table_layout import pages\n','')+'\n'+compiler
    return ('import json, math, re, sys, copy, unicodedata\nfrom pathlib import PurePosixPath\n\n'+
            'TITLE_MAX = '+repr(TITLE_MAX)+'\nPPTX_SLIDES = '+repr(PPTX_SLIDES)+'\nPPTX_ELEMENTS = '+repr(PPTX_ELEMENTS)+'\n'+
            'CATALOG = '+repr(CATALOG)+'\nSTYLES = '+repr(STYLES)+'\nCOLLECTION_LAYOUTS = '+repr(COLLECTION_LAYOUTS)+'\n\n'+compiler+'\n\n'+
            '\n\n'.join(inspect.getsource(f) for f in functions)+
            '\nif __name__ == "__main__":\n'
            '    with open(sys.argv[1], encoding="utf-8") as source:\n'
            '        validate(load(source.read()), sys.argv[2:])\n'
            '    print("RELAY_PPTX_CONTRACT_VALID")\n')


def _register_notes_master(deck):
    """Register the notes relationship omitted by python-pptx 1.0.2.

    Keynote requires this presentation-level reference even though python-pptx
    can reopen the notes through slide relationships without it (upstream #1051).
    Keep the standard CT_Presentation element order and the existing notes bytes.
    """
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.oxml.ns import qn
    relations = [r for r in deck.part.rels.values() if r.reltype == RT.NOTES_MASTER]
    if not relations:
        return
    if len(relations) != 1:
        raise ValueError('PPTX must have one notes master.')
    root = deck.part._element
    lists = root.findall(qn('p:notesMasterIdLst'))
    if not lists:
        listing = OxmlElement('p:notesMasterIdLst')
        masters = root.find(qn('p:sldMasterIdLst'))
        root.insert(root.index(masters) + 1 if masters is not None else 0, listing)
        entry = OxmlElement('p:notesMasterId')
        entry.set(qn('r:id'), relations[0].rId)
        listing.append(entry)
    _inspect_notes_master(deck)


def _inspect_notes_master(deck):
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.oxml.ns import qn
    relations = [r for r in deck.part.rels.values() if r.reltype == RT.NOTES_MASTER]
    lists = deck.part._element.findall(qn('p:notesMasterIdLst'))
    if not relations and not lists:
        return
    if (len(relations) != 1 or len(lists) != 1 or len(lists[0]) != 1
            or lists[0][0].tag != qn('p:notesMasterId')
            or lists[0][0].get(qn('r:id')) != relations[0].rId
            or relations[0].is_external):
        raise ValueError('PPTX reopen failed: notes master registration is missing or invalid.')
    root = deck.part._element
    if any(root.index(e) < root.index(lists[0]) for e in root
           if e.tag not in (qn('p:sldMasterIdLst'), qn('p:notesMasterIdLst'))):
        raise ValueError('PPTX reopen failed: notes master registration is out of order.')


def prepare(document, images=None, bundles=None):
    """Validate the same expanded inputs used by creation, without producing a deck."""
    from .slide_templates import expand
    document = expand(document)
    images = dict(images or {})
    if bundles:
        from .image_sources import presentation_inputs
        document,images=presentation_inputs(document,images,bundles)
    import hashlib
    bindings=document.get('asset_sha256',{})
    # Validate the authored envelope before resolving code-owned staged aliases.
    # A saved pixel identity can follow the same bytes into a different node;
    # names alone never authorize a substitute or an undeclared filesystem read.
    validate(document, set(images) | (set(bindings) if isinstance(bindings,dict) else set()))
    by_hash={hashlib.sha256(raw).hexdigest():raw for raw in images.values()}
    for path,digest in bindings.items():
        if path not in images and digest in by_hash:
            images[path]=by_hash[digest]
        if path not in images or hashlib.sha256(images[path]).hexdigest()!=digest:
            raise ValueError('Saved document image version changed: '+path)
    validate(document, images)
    from PIL import Image
    for raw in images.values():
        with Image.open(io.BytesIO(raw)) as image:
            if image.format not in ('PNG', 'JPEG') or image.width * image.height > 40000000:
                raise ValueError('Only bounded PNG/JPEG images are supported.')
            image.verify()
    return document, images


def preflight_inputs(inputs):
    manifests=[i for i in inputs if i.get('media_type')=='application/json']
    if len(manifests)!=1:
        raise ValueError('PPTX creation requires exactly one JSON slide specification.')
    images={i['path']:i['blob'].read_bytes() for i in inputs if i.get('media_type') in ('image/png','image/jpeg')}
    bundles={i['path']:i['blob'].read_bytes() for i in inputs if i.get('media_type')=='application/zip'}
    return prepare(load(manifests[0]['blob'].read_text()),images,bundles)


def image_region(item,width,height):
    left,top,right,bottom=item.get('source_region',[0,0,1,1])
    if item.get('fit','contain')=='cover':
        ratio=width*(right-left)/(height*(bottom-top));target=item['w']/item['h']
        if ratio>target:
            inset=(right-left)*(1-target/ratio)/2;left+=inset;right-=inset
        else:
            inset=(bottom-top)*(1-ratio/target)/2;top+=inset;bottom-=inset
    return left,top,right,bottom


def create(document, images=None, max_bytes=50000000, bundles=None):
    """Return checked PPTX bytes and structural evidence, before any output write."""
    available()
    document, images = prepare(document, images, bundles)
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
    from pptx.util import Inches, Pt
    from PIL import Image

    dimensions = {}
    for path, raw in images.items():
        with Image.open(io.BytesIO(raw)) as image:
            if image.format not in ('PNG', 'JPEG') or image.width * image.height > 40000000:
                raise ValueError('Only bounded PNG/JPEG images are supported.')
            image.verify()
            dimensions[path] = image.size
    deck = Presentation()
    width, height = document.get('size', [13.333333, 7.5])
    deck.slide_width, deck.slide_height = Inches(width), Inches(height)
    deck.core_properties.title = document['title']
    deck.core_properties.author = 'Task Relay'

    def text(frame, value, item, default_size=24):
        frame.text = value
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.margin_left = frame.margin_right = Inches(0.06)
        frame.margin_top = frame.margin_bottom = Inches(0.03)
        for paragraph in frame.paragraphs:
            # Preserve literal text while making supplied HTTPS citations usable.
            # No URL discovery, rewriting or inference of source support occurs.
            literal=paragraph.text
            matches=list(re.finditer(r'https://[^\s<>]+',literal))
            if matches:
                paragraph.clear();position=0
                for match in matches:
                    url=match.group().rstrip('.,;')
                    if match.start()>position:paragraph.add_run().text=literal[position:match.start()]
                    run=paragraph.add_run();run.text=url
                    try:_hyperlink(url)
                    except ValueError:pass
                    else:run.hyperlink.address=url
                    position=match.start()+len(url)
                if position<len(literal):paragraph.add_run().text=literal[position:]
            paragraph.font.name = item.get('font', document.get('font', 'Arial'))
            paragraph.font.size = Pt(item.get('font_size', default_size))
            paragraph.font.bold = item.get('bold', False)
            paragraph.font.color.rgb = RGBColor.from_string(item.get('color', '172B4D'))
            paragraph.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}[item.get('align', 'left')]

    for spec in document['slides']:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(spec.get('background', 'FFFFFF'))
        if 'notes' in spec:
            slide.notes_slide.notes_text_frame.text = spec['notes']
        for index, item in enumerate(spec['elements']):
            box = [Inches(item[k]) for k in ('x', 'y', 'w', 'h')]
            kind = item['type']
            if kind == 'text':
                shape = slide.shapes.add_textbox(*box)
                text(shape.text_frame, item['text'], item)
            elif kind == 'shape':
                shape = slide.shapes.add_shape({'rectangle': MSO_SHAPE.RECTANGLE, 'ellipse': MSO_SHAPE.OVAL}[item['shape']], *box)
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor.from_string(item.get('fill', 'E8EEF5'))
                shape.line.fill.background()
                text(shape.text_frame, item.get('text', ''), item)
                if item['shape']=='ellipse' and item.get('text','').isdigit():
                    # The ellipse's inscribed text area is narrower than its
                    # bounding box. Preserve short numeric diagram markers.
                    shape.text_frame.margin_left=shape.text_frame.margin_right=0
                    shape.text_frame.word_wrap=False
            elif kind == 'image':
                iw, ih = dimensions[item['path']]
                region=image_region(item,iw,ih);left,top,right,bottom=region
                if item.get('fit','contain')=='cover':
                    shape=slide.shapes.add_picture(io.BytesIO(images[item['path']]),*box)
                else:
                    iw*=right-left;ih*=bottom-top
                    scale=min(box[2]/iw,box[3]/ih);w,h=round(iw*scale),round(ih*scale)
                    shape=slide.shapes.add_picture(io.BytesIO(images[item['path']]),round(box[0]+(box[2]-w)/2),round(box[1]+(box[3]-h)/2),w,h)
                shape.crop_left=left;shape.crop_top=top;shape.crop_right=1-right;shape.crop_bottom=1-bottom
            elif kind == 'table':
                rows = item['rows']
                shape = slide.shapes.add_table(len(rows), len(rows[0]), *box)
                if 'row_heights' in item:
                    for r,height in enumerate(item['row_heights']):shape.table.rows[r].height=Inches(height)
                for r, row in enumerate(rows):
                    for col, value in enumerate(row):
                        cell = shape.table.cell(r, col)
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = RGBColor.from_string('E8EEF5' if r == 0 else 'FFFFFF')
                        text(cell.text_frame, value, {**item, 'bold': r == 0}, 18)
            else:
                data = CategoryChartData()
                data.categories = item['categories']
                for series in item['series']:
                    data.add_series(series['name'], series['values'])
                shape = slide.shapes.add_chart({'column': XL_CHART_TYPE.COLUMN_CLUSTERED, 'bar': XL_CHART_TYPE.BAR_CLUSTERED, 'line': XL_CHART_TYPE.LINE}[item['chart']], *box, data)
                shape.chart.has_legend = len(item['series']) > 1
                shape.chart.font.name = document.get('font', 'Arial')
                shape.chart.font.size = Pt(16)
            if 'hyperlink' in item:
                shape.click_action.hyperlink.address = item['hyperlink']
            shape.name = item.get('name') or f'{kind}-{index + 1}'
    _register_notes_master(deck)
    buffer = io.BytesIO()
    deck.save(buffer)
    raw = buffer.getvalue()
    if len(raw) > max_bytes:
        raise ValueError('PPTX output exceeds its byte limit.')
    evidence = inspect(raw, document, images)
    return raw, evidence


def inspect(raw, document, images):
    """Reopen the actual saved bytes and compare semantic native content."""
    from pptx import Presentation
    from pptx.util import Inches
    deck = Presentation(io.BytesIO(raw))
    _inspect_notes_master(deck)
    if len(deck.slides) != len(document['slides']) or deck.core_properties.title != document['title']:
        raise ValueError('PPTX reopen failed: slide count or title changed.')
    size = document.get('size', [13.333333, 7.5])
    if (deck.slide_width, deck.slide_height) != tuple(Inches(n) for n in size):
        raise ValueError('PPTX reopen failed: slide dimensions changed.')
    counts = dict(text=0, shape=0, table=0, chart=0, image=0)
    for slide, spec in zip(deck.slides, document['slides']):
        if len(slide.shapes) != len(spec['elements']):
            raise ValueError('PPTX reopen failed: native objects changed.')
        if 'notes' in spec and slide.notes_slide.notes_text_frame.text != spec['notes']:
            raise ValueError('PPTX reopen failed: notes changed.')
        for shape, item in zip(slide.shapes, spec['elements']):
            kind = item['type']
            counts[kind] += 1
            if kind in ('text', 'shape', 'image') and shape.click_action.hyperlink.address != item.get('hyperlink'):
                raise ValueError('PPTX reopen failed: purchase/source hyperlink changed.')
            if kind in ('text', 'shape') and (not shape.has_text_frame or shape.text != item.get('text', '')):
                raise ValueError('PPTX reopen failed: editable text changed.')
            if kind in ('text','shape'):
                expected=[]
                for match in re.finditer(r'https://[^\s<>]+',item.get('text','')):
                    url=match.group().rstrip('.,;')
                    try:_hyperlink(url)
                    except ValueError:continue
                    expected.append((url,url))
                actual=[(run.text,run.hyperlink.address) for paragraph in shape.text_frame.paragraphs
                        for run in paragraph.runs if run.hyperlink.address]
                if actual!=expected:raise ValueError('PPTX reopen failed: literal source hyperlink changed.')
            if kind == 'table':
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows] if shape.has_table else None
                if rows != item['rows']:
                    raise ValueError('PPTX reopen failed: editable table data changed.')
                if 'row_heights' in item:
                    if any(abs(row.height-Inches(height))>2 for row,height in zip(shape.table.rows,item['row_heights'])) or shape.top+shape.height>deck.slide_height:
                        raise ValueError('PPTX reopen failed: table row geometry changed or clipped')
                    if any(p.font.size is None or abs(p.font.size.pt-item['font_size'])>.01 for row in shape.table.rows for cell in row.cells for p in cell.text_frame.paragraphs):
                        raise ValueError('PPTX reopen failed: table typography changed')
            if kind == 'chart':
                if not shape.has_chart or list(shape.chart.plots[0].categories) == []:
                    raise ValueError('PPTX reopen failed: missing chart.')
                series = [{'name': s.name, 'values': list(s.values)} for s in shape.chart.series]
                labels = [category.label for category in shape.chart.plots[0].categories]
                matches = len(series) == len(item['series']) and all(
                    actual['name'] == expected['name'] and len(actual['values']) == len(expected['values']) and all(
                        math.isclose(a, b, rel_tol=1e-14, abs_tol=0) for a, b in zip(actual['values'], expected['values']))
                    for actual, expected in zip(series, item['series']))
                if labels != item['categories'] or not matches:
                    raise ValueError('PPTX reopen failed: editable chart data changed.')
            if kind == 'image':
                if shape.image.blob!=images[item['path']]:raise ValueError('PPTX reopen failed: image content changed.')
                from PIL import Image
                with Image.open(io.BytesIO(images[item['path']])) as original:region=image_region(item,*original.size)
                actual=(shape.crop_left,shape.crop_top,1-shape.crop_right,1-shape.crop_bottom)
                if any(abs(a-b)>0.000011 for a,b in zip(actual,region)):raise ValueError('PPTX reopen failed: image source region changed.')
    return {'slide_count': len(deck.slides), 'native_objects': counts, 'reopened': True,
            'visual_review': 'not_performed', 'keynote_import': 'not_tested', 'python_pptx': VERSION,
            'chart_numeric_relative_tolerance': 1e-14}
