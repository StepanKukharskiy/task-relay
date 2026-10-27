#!/usr/bin/env python3
"""Opt-in native composed qualification; scripted agents, no model API calls.

Runs real registered Blender, Rhino, public Commons and PPTX operations. Retains
the isolated runtime and original journals even on failure. Never installs code.
"""
import argparse
import copy
import json
from pathlib import Path
import sys
import time
import traceback
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

RHINO_SCRIPT = '''import json
from pathlib import Path
import Rhino
root = next(p for p in Path(__file__).resolve().parents if (p / '.relay/ASSIGNMENT.json').exists())
assignment = json.loads((root / '.relay/ASSIGNMENT.json').read_text())
sources = [root / i['path'] for i in assignment['inputs'] if i['path'].endswith('.txt')]
assert len(sources) == 1, 'One exact Blender inventory bundle required'
inventory = json.loads(sources[0].read_text().split('\\n', 1)[1])
meshes = [o for o in inventory['objects'] if o['type'] == 'MESH']
assert len(meshes) == 1 and meshes[0]['geometry']['vertices'] == 8
obj = meshes[0]
assert all(abs(v) < 1e-6 for v in obj['rotation_euler'])
x, y, z = obj['location']; dx, dy, dz = obj['dimensions']
box = Rhino.Geometry.BoundingBox(Rhino.Geometry.Point3d(x-dx/2,y-dy/2,z-dz/2), Rhino.Geometry.Point3d(x+dx/2,y+dy/2,z+dz/2))
attributes = Rhino.DocObjects.ObjectAttributes(); attributes.Name = 'TransferredBox'
doc.Objects.AddBrep(Rhino.Geometry.Brep.CreateFromBox(box), attributes)
for name, w, h, offset in [('Front',dx,dz,3),('Side',dy,dz,6),('Top',dx,dy,9)]:
    points = [Rhino.Geometry.Point3d(offset+a,b,0) for a,b in [(0,0),(w,0),(w,h),(0,h),(0,0)]]
    attributes = Rhino.DocObjects.ObjectAttributes(); attributes.Name = name
    doc.Objects.AddPolyline(points, attributes)
'''


def run(folder):
    from orchestrator import contracts as c, executors, job_owner, worker_tools, work_authorization
    from orchestrator.agent_loop import execute
    from orchestrator.runtime import Runtime, file_hash
    from orchestrator.runtime_outcomes import Hold
    from orchestrator.workers import atomic
    from orchestrator.owner_compilation import label
    from orchestrator.owner_intents import ports
    from orchestrator.step_runner import execute as native
    from orchestrator.execution_audit import export
    from tests.test_worker_tools import Tests
    from tests.test_orchestrator import FakeFactory
    from tests.test_operation_session import Script, CONFIG
    from tests.test_blender_operations import scene_data
    from task_relay.host_apps import blender, rhino
    from orchestrator.document_authoring import PPTX

    test=Tests();test.setUp()
    test.rt.close()
    test.root=folder/'runtime';test.factory=FakeFactory();test.rt=Runtime(test.root,test.factory)
    test.grant.update(max_tasks=20, operations=['blender.scene','blender.inspect','text.bundle','rhino.run_python','images.collect','pptx.create','pptx.render'])
    test.grant['task_limits'].update(seconds=1800,tool_calls=100,provider_requests=24,output_bytes=100000000)
    summary=dict(started=time.time(),scope='Actual installed CAD and public image HTTP; scripted agent judgments; source runtime, no deployment',paid_model_calls=0,passed=False,applications=[blender(),rhino()],stages=[])

    def saved(ident):
        test.rt.tick('job');row=test.rt.task('job',ident)
        if not row['latest']:raise ValueError('No dispatch for '+ident+': '+str(row))
        state=test.factory.sessions[row['latest']];f=state['frozen'];control=Path(state['session']['control']);control.mkdir(parents=True,exist_ok=True)
        atomic(control/'launch.json',dict(credential_fingerprint=executors.fingerprint(CONFIG,f['backend'])))
        return f,control,state

    def finish_state(state):
        state['status']=dict(status='finished',exit_code=0,external_outcome='no_pending_response',pending_requests=[])
        test.rt.tick('job')

    def report(f,decision='delivered'):
        from orchestrator.final_quality import DIMENSIONS
        checks={f'c{i}':dict(passed=True,evidence='Controlled qualification: exact native receipts and artifact checks; semantic judgment scripted') for i in range(1,len(f['criteria'])+1)}
        for n,binding in f.get('final_quality',{}).get('checks',{}).items():
            inspected=[dict(path=p,location='Exact qualified candidate') for p in dict.fromkeys([binding['candidate'],*binding['inspection_inputs']])]
            checks['c'+n]['evidence']={d:dict(verdict='satisfied',reason='Scripted qualification judgment; not independent model quality evidence',inspected=inspected) for d in DIMENSIONS}
        return ('finish',dict(report_json=c.encoded(dict(summary='Controlled qualification',decision=decision,instruction='',findings=[],checks=checks))))

    def drive(ident,steps,approve=False,restart_after=None):
        f,control,state=saved(ident);script=Script(steps(f)+[report(f)])
        save=worker_tools.Session.save
        interrupted=False
        def checkpoint(session,contents,history,requests,tools,**kw):
            nonlocal interrupted
            save(session,contents,history,requests,tools,**kw)
            if requests==restart_after and not kw.get('pending') and not interrupted:
                interrupted=True;raise InterruptedError('Qualification restart after committed correction')
        while True:
            try:
                with patch.object(worker_tools.Session,'save',checkpoint):execute(f,control,script,lambda:(CONFIG,f['backend']))
                break
            except InterruptedError:
                original=(control/'work-clock.json').read_bytes();test.rt.close();test.rt=Runtime(test.root,test.factory)
                assert (control/'work-clock.json').read_bytes()==original
                summary['stages'].append(dict(task=ident,restart='after committed correction',requests=len(script.requests)))
            except Hold as exc:
                proposals=[p for p in control.glob('work-operations/*/host-proposal.json') if not p.with_name('host-authorization.json').exists()]
                if not approve or not proposals:raise
                for p in proposals:
                    with test.rt.transaction():work_authorization.authorize(test.rt,f['assignment_id'],p,dict(source='User explicitly requested actual Blender and Rhino composed qualification; fixed script in qualify_native_composed.py',script_sha256=file_hash(Path(f['workspace'])/next(i['path'] for i in f['inputs'] if i['media_type']=='text/x-python'))))
                summary['stages'].append(dict(task=ident,approval=str(exc),proposals=[str(p) for p in proposals]))
        finish_state(state)
        summary['stages'].append(dict(task=ident,requests=len(script.requests),status=test.rt.task('job',ident)['status']))
        print(ident+': '+test.rt.task('job',ident)['status'],flush=True)
        return f,control

    def review(ident):
        rid=label(ident,'review')
        if not test.rt.db.execute('SELECT 1 FROM production_tasks WHERE run=? AND id=?',('job',rid)).fetchone():return
        f,control,state=saved(rid);steps=[]
        for i in f['inputs']:
            if i['media_type'].startswith('image/'):
                steps.append(('image_view',dict(path=i['path'],members=[])))
            elif i['path'].endswith('slide-previews.zip'):
                steps.append(('image_view',dict(path=i['path'],members=['slide-001.png'])))
        script=Script(steps+[report(f,'accept')])
        # The test grant's code reviewer is a scripted peer, not an installed
        # code-agent runtime. No code tools are called by this review fixture.
        with patch('task_relay.code_runtime.available',return_value=dict(tools=[])):
            execute(f,control,script,lambda:(CONFIG,f['backend']))
        finish_state(state)
        assert test.rt.task('job',ident)['status']=='completed'

    try:
        data=scene_data();data['objects'][0]['rotation']=[0,0,0];data['resolution']=[256,256]
        scene=test.source(json.dumps(data),'scene.json','application/json')
        code=test.source(RHINO_SCRIPT,'model.py','text/x-python')
        checks=dict(mode='create',units='Meters',changed_objects=[],allow_additions=True,expected_object_count=4,expected_dimensions={'TransferredBox':[1,1,2],'Front':[1,2,0],'Side':[1,2,0],'Top':[1,1,0]},preview={'resolution':[320,240]})
        check=test.source(json.dumps(checks),'checks.json','application/json')
        test.create(request='Qualify actual Blender to Rhino projections and sourced-image PPTX with correction and restart. No paid model calls.',inputs=[scene,code,check]);test.rt.tick('job')
        # Named ports are derived from the same compiled task contract used by models.
        first=dict(id='blender',objective='Create, reopen, render and inspect the exact fixture cube',inputs=[dict(artifact=scene['artifact'])],work=dict(operation='blender.scene',steps=[dict(id='inventory',operation='blender.inspect',from_steps=['build'])]))
        test.invoke([('job_delegate',dict(request_id='native-start',tasks=[first]))]);test.rt.tick('job')
        drive('blender',lambda f:[('work_run',dict(step='build',request_key='scene')),('image_view',dict(path='outputs/work-build/preview.png',members=[])),('work_run',dict(step='inventory',request_key='inventory'))]);review('blender')
        spec=test.rt.spec(test.rt.task('job','blender'));port=next(k for k,v in ports(spec).items() if v['path'].endswith('/inspection.json'))
        tasks=[dict(id='handoff',objective='Preserve exact Blender inventory as a text evidence bundle',operation='text.bundle',inputs=[dict(task='blender',output=port)]),
            dict(id='rhino',objective='Use exact Blender dimensions to create a Rhino model and three projection outlines',inputs=[dict(task='handoff',output='result')],work=dict(operation='rhino.run_python',parameters=dict(scene_sha256=None,script_sha256=dict(artifact=code['artifact']),checks_sha256=dict(artifact=check['artifact']),permissions='unrestricted_host'))),
            dict(id='sources',objective='Collect a public Eiffel Tower photo without any model call',operation='images.collect',parameters=dict(subjects=[dict(id='chair',label='Eiffel Tower',query='Eiffel Tower')]))]
        test.invoke([('job_delegate',dict(request_id='native-handoff',tasks=tasks))]);test.rt.tick('job')
        for ident in ('handoff','sources'):
            f,control,state=saved(ident);result=native(f,control);assert result['outcome']=='completed',result;finish_state(state);summary['stages'].append(dict(task=ident,result=result));print(ident+': completed',flush=True)
        rf,rc=drive('rhino',lambda f:[('work_run',dict(step='build',request_key='native')),('image_view',dict(path='outputs/work-build/preview.png',members=[]))],approve=True);review('rhino')
        from PIL import Image
        with Image.open(Path(rf['workspace'])/'outputs/work-build/preview.png') as preview:
            # The simple fixture has no background geometry: clipped black
            # object edges must never touch the capture boundary.
            rgb=preview.convert('RGB');w,h=rgb.size
            edge=[rgb.getpixel((x,y)) for x in range(w) for y in (0,h-1)]+[rgb.getpixel((x,y)) for y in range(h) for x in (0,w-1)]
            assert all(min(pixel)>60 for pixel in edge),'Native preview clips fixture geometry'
        refs=[dict(task='sources',output='result')]
        for ident in ('blender','rhino'):
            p=next(k for k,v in ports(test.rt.spec(test.rt.task('job',ident))).items() if v['path']=='outputs/work-build/preview.png');refs.append(dict(task=ident,output=p))
        test.invoke([('job_delegate',dict(request_id='deck',tasks=[dict(id='make',objective='Make a sourced image deck with actual CAD previews; correct and restart',inputs=refs,work=dict(operation='pptx.create',content_parameters=dict(max_records=1)))]))]);test.rt.tick('job')
        page={}
        def deck_steps(f):
            import zipfile
            bundle=next(i for i in f['inputs'] if i['media_type']=='application/zip')
            with zipfile.ZipFile(Path(f['workspace'])/bundle['path']) as z:member=next(n for n in z.namelist() if n.startswith('images/'))
            images=[bundle['path']+'/'+member]+[i['path'] for i in f['inputs'] if i['media_type']=='image/png']
            page.update(kind='page',slot=1,layout='content_cards',content=dict(title='Initial heading',per_page=3,items=[dict(label=title,body='Qualification fixture; no reconstruction fidelity claim',image=p,link=None,credit='Source details retained in notes and audit') for title,p in zip(['Sourced landmark','Blender cube','Rhino projections'],images)]),notes='Actual source credit retained in the image bundle; CAD geometry is a fixed fixture.')
            revised=copy.deepcopy(page);revised['content']['title']='Sourced image and native CAD qualification'
            return [('image_view',dict(path=bundle['path'],members=[member]))]*3+[
                ('operation_submit',dict(request_key='page',revision=0,records=[page])),('operation_submit',dict(request_key='seal',revision=1,records=[dict(kind='end',count=1)])),
                ('work_run',dict(step='build',request_key='v1')),('work_run',dict(step='inspect',request_key='render1')),('image_view',dict(path='outputs/work-inspect/slide-previews.zip',members=['slide-001.png'])),
                ('operation_amend',dict(request_key='correct',revision=2,expected_record_sha256=c.digest(page),record=revised)),
                report(f), # must reject stale v1 after amendment
                ('work_run',dict(step='build',request_key='v2')),('work_run',dict(step='inspect',request_key='render2')),('image_view',dict(path='outputs/work-inspect/slide-previews.zip',members=['slide-001.png']))]
        f,control=drive('make',deck_steps,restart_after=9);review('make')
        assert 'stale' in (control/'tool-10-00.json').read_text()
        from pptx import Presentation
        from pypdf import PdfReader
        deck=Path(f['workspace'])/'outputs/work-build/result.pptx';pdf=Path(f['workspace'])/'outputs/work-inspect/preview.pdf'
        doc=Presentation(deck);assert sum(s.shape_type==13 for s in doc.slides[0].shapes)==3
        assert 'Sourced image and native CAD qualification' in PdfReader(pdf).pages[0].extract_text()
        outputs={key:test.rt.output('job','make',str(p.relative_to(f['workspace'])))['id'] for key,p in [('deck',deck),('pdf',pdf)]}
        test.invoke([('job_complete',dict(deliverables_json=c.encoded(outputs),summary='Controlled native composed qualification completed; model judgments scripted'))]);test.rt.tick('job')
        assert test.rt.status('job')['status']=='completed'
        summary.update(passed=True,deliverables={key:str(p) for key,p in [('deck',deck),('pdf',pdf)]})
    except BaseException:
        summary['error']=traceback.format_exc();print(summary['error'],flush=True)
    finally:
        summary['finished']=time.time();atomic(folder/'summary.json',summary)
        database=Path(test.rt.db.execute('PRAGMA database_list').fetchone()[2])
        if test.rt.db.execute("SELECT 1 FROM production_runs WHERE id='job'").fetchone():export(database,test.root/'workers','job',folder/'audit')
        test.doCleanups()
    return summary['passed']


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--host',action='store_true',required=True)
    parser.add_argument('--output',type=Path,required=True)
    args=parser.parse_args();folder=args.output.resolve();folder.mkdir(parents=True,exist_ok=False)
    (folder/'qualification-source.py').write_bytes(Path(__file__).read_bytes())
    (folder/'command.json').write_text(json.dumps(dict(command=[sys.executable,*sys.argv]),indent=2))
    sys.exit(0 if run(folder) else 1)
